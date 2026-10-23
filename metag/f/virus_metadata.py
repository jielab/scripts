#!/usr/bin/env python3
"""Fetch/freeze ICTV VMR, map to SAME-build NCBI taxonomy, query offline.

Sequence mol-type is deliberately never used as virus genome type evidence.
"""
from __future__ import annotations
import argparse
from collections import defaultdict
from datetime import datetime, timezone
from functools import lru_cache
import csv
import fcntl
import gzip
import hashlib
import http.client
import io
import json
from pathlib import Path
import re
import sqlite3
import sys
import tempfile
import time
import urllib.parse
import urllib.request
import zipfile

from taxonomy import Taxonomy, atomic_json, prepare_sqlite, sha256, taxonomy_identity

SCHEMA = 'metag-virus-v2'
RULE_VERSION = 'genome-normalization-and-inheritance-v2.1'
VMR_URL = 'https://ictv.global/vmr/current'
SOURCE_FILE = 'virus_metadata.source.json'
MAIN_FIELDS = ['taxid','taxon_rank','ncbi_name','species_taxid','ncbi_species_name',
    'ictv_species_name','genome_class','genome_type_detail','genome_composition_raw',
    'baltimore_group','entity_type','evidence_scope','evidence_method','source_ref',
    'source_version','source_accession','mapping_method','conflict_status','notes',
    'resolution_status','evidence_taxid']
ACCESSION = re.compile(r'(?<![A-Za-z0-9_])(?:[A-Z]{1,2}_?[0-9]{5,9})(?:\.[0-9]+)?(?![A-Za-z0-9_.])')


def utcnow():
    return datetime.now(timezone.utc).isoformat()


def write_table(path, fields, rows):
    with Path(path).open('w', newline='') as f:
        w = csv.DictWriter(f, fields, delimiter='\t', extrasaction='ignore')
        w.writeheader(); w.writerows(rows)


def read_table(path):
    with Path(path).open(newline='', encoding='utf-8-sig') as f:
        return list(csv.DictReader(f, delimiter='\t'))


def normkey(value):
    return re.sub(r'\s+', '', str(value or '').strip().lower().replace('−','-').replace('–','-'))


# A whitelist, not a substring preference. Unknown future labels remain unknown.
NORMAL_TYPES = {}
for detail in ('dsDNA','ssDNA','ssDNA(+)','ssDNA(-)','ssDNA(+/-)','dsDNA-RT',
               'dsRNA','ssRNA','ssRNA(+)','ssRNA(-)','ssRNA(+/-)','ssRNA-RT'):
    NORMAL_TYPES[normkey(detail)] = ('DNA' if 'DNA' in detail else 'RNA', detail)
for alias, detail in {'DNA reverse-transcribing':'dsDNA-RT', 'RNA reverse-transcribing':'ssRNA-RT',
                     'ambisense RNA':'ssRNA(+/-)', 'ssRNA(±)':'ssRNA(+/-)',
                     'ssDNA(±)':'ssDNA(+/-)', 'ssRNA(+) / ssRNA(-)':'ssRNA(+/-)',
                     'ssDNA(+) / ssDNA(-)':'ssDNA(+/-)'}.items():
    NORMAL_TYPES[normkey(alias)] = NORMAL_TYPES[normkey(detail)]
NORMAL_TYPES.update({'dna': ('DNA','unknown'), 'rna': ('RNA','unknown')})


def normalize_genome_type(raw):
    key = normkey(raw)
    if key in NORMAL_TYPES:
        c, d = NORMAL_TYPES[key]
        return c, d, 'recognized'
    if key in ('', 'unknown', 'na', 'n/a', 'notknown', 'unassigned'):
        return 'unknown', 'unknown', 'missing'
    # Legal combinations must consist entirely of recognized labels. +/- inside
    # parentheses is protected by only allowing separators before another type.
    pieces = re.split(r'(?:;|,|/|\+|and|&)(?=(?:ss|ds)?(?:dna|rna))', key)
    if len(pieces) > 1 and all(p in NORMAL_TYPES for p in pieces):
        kinds = {NORMAL_TYPES[p][0] for p in pieces}
        if len(kinds) == 1:
            return next(iter(kinds)), ';'.join(sorted({NORMAL_TYPES[p][1] for p in pieces})), 'recognized'
        return 'unknown', 'unknown', 'mixed'
    return 'unknown', 'unknown', 'unrecognized'


def accessions(value):
    return sorted(set(ACCESSION.findall(str(value or ''))))


def headerkey(value):
    return re.sub(r'[^a-z0-9]+','',str(value or '').lower())


def vmr_records(path):
    """Locate headers by names on ANY worksheet/row; never rely on column letters."""
    path = Path(path)
    book = None
    if path.suffix.lower() == '.xlsx':
        try:
            from openpyxl import load_workbook
        except ImportError as e:
            raise RuntimeError('VMR XLSX requires openpyxl; install environment.yml') from e
        book = load_workbook(path, read_only=True, data_only=True)
        sheets = ((s.title, s.iter_rows(values_only=True)) for s in book.worksheets)
    else:
        sheets = [('TSV', csv.reader(io.StringIO(path.read_text(encoding='utf-8-sig')), delimiter='\t'))]
    found = False
    try:
        for sheet, rows in sheets:
            columns = None
            for lineno, values in enumerate(rows, 1):
                if columns is None:
                    keys = [headerkey(v) for v in values]
                    species = next((i for i,k in enumerate(keys) if k in ('species','ictvspeciesname','speciesname')), None)
                    genome = next((i for i,k in enumerate(keys) if k in ('genome','genomecomposition','genomecompositionraw')), None)
                    acc = [i for i,k in enumerate(keys) if ('accession' in k and ('genbank' in k or 'refseq' in k))
                           or k in ('accession','accessions','sourceaccession')]
                    if species is not None and genome is not None and acc:
                        columns = (species, genome, acc); found = True
                    continue
                si, gi, ai = columns
                cell = lambda i: str(values[i] or '').strip() if i < len(values) else ''
                if not cell(si):
                    continue
                yield dict(ictv_species_name=cell(si), genome_composition_raw=cell(gi),
                           accessions=accessions(';'.join(cell(i) for i in ai)),
                           accession_raw=';'.join(cell(i) for i in ai),
                           source_row=f'{sheet}:{lineno}')
        if not found:
            raise RuntimeError('VMR headers missing: species, genome composition, GenBank/RefSeq accession')
    finally:
        if book:
            book.close()


def request_bytes(url, data=None):
    last = None
    for attempt in range(4):
        try:
            req = urllib.request.Request(url, data=data, headers={'User-Agent':'metag-virus/2.0'})
            with urllib.request.urlopen(req, timeout=60) as response:
                return response.read(), response.geturl(), dict(response.headers)
        except (OSError, ValueError, http.client.HTTPException) as e:
            last = e
            if attempt < 3:
                time.sleep(2 ** attempt)
    raise RuntimeError(f'Failed retrieval after retries: {url}: {last}')


def fetch(args):
    args.out_dir.mkdir(parents=True, exist_ok=True)
    with (args.out_dir / '.fetch.lock').open('a') as lock:
        try:
            fcntl.flock(lock, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError as error:
            raise RuntimeError(f'Another ICTV fetch holds the lock: {args.out_dir}') from error
        existing = [p for p in args.out_dir.iterdir() if p.name != '.fetch.lock']
        if existing:
            books = [p for p in existing if p.suffix.lower() == '.xlsx' and p.is_file()]
            if len(books) != 1 or set(existing) != {books[0], books[0].with_suffix('.source.json')}:
                raise RuntimeError(f'Incomplete/unrecognized ICTV directory: {args.out_dir}; '
                                   'preserved for inspection. Use a new directory.')
            path = books[0]
            source = json.loads(path.with_suffix('.source.json').read_text())
            if source.get('source_url') != args.url or source.get('sha256') != sha256(path):
                raise RuntimeError(f'ICTV source URL/checksum mismatch: {args.out_dir}; use a new directory.')
            count = sum(1 for _ in vmr_records(path))
            if not count or source.get('rows') != count:
                raise RuntimeError(f'ICTV row count mismatch: {path}')
            print(path)
            return
        # A failed download leaves a retryable empty directory (plus the lock).
        data, final, headers = request_bytes(args.url)
        if not zipfile.is_zipfile(io.BytesIO(data)):
            raise RuntimeError('ICTV endpoint did not return XLSX; retry the same command.')
        name = Path(urllib.parse.unquote(urllib.parse.urlparse(final).path)).name
        if not name.lower().endswith('.xlsx'):
            match = re.search(r'filename="?([^";]+)', headers.get('Content-Disposition',''))
            name = Path(match[1]).name if match else 'VMR.xlsx'
        with tempfile.TemporaryDirectory(prefix='.metag-vmr-', dir=args.out_dir.parent) as temporary:
            staged = Path(temporary) / name
            staged.write_bytes(data)
            count = sum(1 for _ in vmr_records(staged))
            if not count:
                raise RuntimeError('ICTV XLSX contains no species rows; retry the same command.')
            receipt = staged.with_suffix('.source.json')
            atomic_json(receipt, dict(source_url=args.url, resolved_url=final,
                        source_version=name, retrieved_utc=utcnow(), sha256=sha256(staged), rows=count))
            path = args.out_dir / name
            staged.replace(path)
            receipt.replace(path.with_suffix('.source.json'))
        print(path)


def parse_esummary(data, requested):
    """NCBI nuccore ESummary JSON: result.uids, accessionversion, taxid."""
    obj = json.loads(data)
    result = obj.get('result', {})
    if not isinstance(result.get('uids'), list):
        raise RuntimeError('Unexpected NCBI ESummary schema/error')
    expected = set(requested)
    bases = {a.split('.')[0] for a in requested if '.' not in a}
    out = []
    for uid in result['uids']:
        record = result.get(str(uid), {})
        acc = record.get('accessionversion', '')
        if not ACCESSION.fullmatch(acc):
            raise RuntimeError('NCBI record lacks validated accessionversion')
        if acc not in expected and acc.split('.')[0] not in bases:
            raise RuntimeError(f'Unexpected accession from NCBI: {acc}')
        # Some valid nuccore records (including patents) have no assigned taxid.
        # Keep their raw response and let cache_accessions list them as unmatched.
        if record.get('taxid') in (None, '', 0, '0'):
            continue
        if not str(record['taxid']).isdigit():
            raise RuntimeError(f'Invalid NCBI taxid for {acc}')
        out.append({'accession': acc, 'taxid': int(record['taxid'])})
    return out


def cache_accessions(args):
    wanted = sorted({a for r in vmr_records(args.vmr) for a in r['accessions']})
    if args.accessions:
        wanted = sorted({a for line in args.accessions.read_text().splitlines() for a in accessions(line)})
    args.cache_dir.mkdir(parents=True, exist_ok=True)
    rows = []
    for i in range(0, len(wanted), 100):
        batch = wanted[i:i+100]
        key = hashlib.sha256('\n'.join(batch).encode()).hexdigest()
        cache = args.cache_dir / f'{key}.json'
        if not cache.exists():
            # No credentials; no large all-organism mapping downloads.
            query = urllib.parse.urlencode(dict(db='nuccore', id=','.join(batch), retmode='json'))
            url = 'https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esummary.fcgi?' + query
            data, _, _ = request_bytes(url)
            parse_esummary(data, batch)
            cache.write_bytes(data)
            atomic_json(cache.with_suffix('.source.json'), dict(url=url, retrieved_utc=utcnow(), sha256=sha256(cache)))
            time.sleep(0.4)  # <3 requests/s, including no API key
        receipt = json.loads(cache.with_suffix('.source.json').read_text())
        if receipt['sha256'] != sha256(cache):
            raise RuntimeError(f'Cached ESummary checksum mismatch: {cache}')
        rows.extend(parse_esummary(cache.read_bytes(), batch))
    dest = args.cache_dir / 'accession_taxid.tsv'
    write_table(dest, ['accession','taxid'], rows)
    obtained = {r['accession'] for r in rows} | {r['accession'].split('.')[0] for r in rows}
    write_table(args.cache_dir / 'unmatched.tsv', ['accession'], ({'accession':a} for a in wanted if a not in obtained))
    atomic_json(args.cache_dir / 'cache.source.json', dict(schema=SCHEMA, requested=len(wanted),
                returned=len(rows), mapping_sha256=sha256(dest), created_utc=utcnow(),
                note='Cross-reference only; never use sequence mol-type as genome class.'))
    print(dest)


def load_xref(path, wanted):
    """Stream huge local maps; retain only VMR accessions. Never choose first match."""
    out = defaultdict(set)
    if path is None:
        return out
    bases = {a.split('.')[0] for a in wanted}
    opener = gzip.open if str(path).endswith('.gz') else open
    with opener(path, 'rt', newline='') as f:
        first = f.readline().rstrip('\r\n').split('\t')
        headers = [headerkey(x) for x in first]
        ti = headers.index('taxid') if 'taxid' in headers else None
        ai = next((headers.index(k) for k in ('accessionversion','accession','seqid') if k in headers), None)
        def add(parts, accidx, taxidx):
            if max(accidx,taxidx) >= len(parts):
                raise RuntimeError(f'Malformed accession-taxid row in {path}')
            for acc in accessions(parts[accidx]):
                if acc in wanted or acc.split('.')[0] in bases:
                    if not parts[taxidx].isdigit():
                        raise RuntimeError(f'Invalid taxid in {path}')
                    out[acc].add(int(parts[taxidx]))
                    out[acc.split('.')[0]].add(int(parts[taxidx]))
        if ti is None or ai is None:
            # k2 2.17.1 preliminary maps carry comment/source columns through
            # kraken2-build; the first two columns still identify seqid/taxid.
            if len(first) < 2 or not first[1].isdigit():
                raise RuntimeError('Cross-reference requires accession + taxid columns or seqid2taxid.map')
            ai, ti = 0, 1
            add(first, ai, ti)
        for line in f:
            if line.strip():
                add(line.rstrip('\r\n').split('\t'), ai, ti)
    return out


class VirusGenomeResolver:
    def __init__(self, path, tax, verify=True):
        self.path, self.tax = Path(path).resolve(), tax
        self.con = sqlite3.connect(self.path.as_uri() + '?mode=ro&immutable=1', uri=True)
        meta = self.con.execute('SELECT value FROM metadata WHERE key="source"').fetchone()
        if meta is None:
            self.con.close(); raise RuntimeError('Virus map source metadata missing')
        self.source = json.loads(meta[0])
        if self.source.get('schema') != SCHEMA or self.source.get('rule_version') != RULE_VERSION:
            self.con.close(); raise RuntimeError('Incompatible virus mapping schema/rules')
        if self.source['taxonomy_identity'] != tax.identity:
            self.con.close(); raise RuntimeError('Virus mapping belongs to a different taxonomy snapshot')
        self.identity = {k:self.source[k] for k in ('schema','rule_version','taxonomy_identity','build_signature','inputs')}
        self.identity['mapping_schema'] = SCHEMA
        self.identity['mapping_sha256'] = sha256(self.path)
        self._verify = verify
        if verify:
            try:
                if self.source['inputs'].get('implementation') != sha256(Path(__file__)) or self.source['inputs'].get('taxonomy_implementation') != sha256(Path(__file__).with_name('taxonomy.py')):
                    raise RuntimeError('Mapping implementation changed; build a NEW mapping version')
                self.assert_unchanged()
            except BaseException:
                self.con.close()
                raise

    def assert_unchanged(self):
        if sha256(self.path) != self.identity['mapping_sha256']:
            raise RuntimeError('Virus mapping content changed during run')
        receipt = json.loads((self.path.parent / SOURCE_FILE).read_text())
        if receipt.get('mapping_sha256') != self.identity['mapping_sha256'] or any(
            receipt.get(k) != v for k,v in self.source.items()):
            raise RuntimeError('Virus mapping checksum/source signature invalid')

    @lru_cache(maxsize=50000)
    def evidence(self, tid):
        return tuple(json.loads(r[0]) for r in self.con.execute('SELECT payload FROM evidence WHERE taxid=? ORDER BY id', (tid,)))

    @lru_cache(maxsize=100000)
    def resolve(self, original):
        tid = self.tax.canonical_taxid(original)
        lineage = self.tax.lineage(tid)
        evidence = []
        for ancestor in lineage:
            for e in self.evidence(ancestor):
                if ancestor == tid or e['evidence_scope'] == 'subtree' or (
                    e['evidence_scope'] == 'species' and self.tax.node(ancestor)[1] == 'species'):
                    evidence.append(e)
        result = dict(genome_class='unknown', genome_type_detail='unknown', evidence_taxid='',
                      evidence_method='', resolution_status='missing', conflict_status='none',
                      source_ref='', source_version='', source_accession='', notes='', evidence_json='[]')
        if not evidence:
            if self.tax.node(tid)[2] in ('Viruses','Riboviria'):
                result['resolution_status'] = 'mixed'
            return result
        result['evidence_json'] = json.dumps(evidence, ensure_ascii=False, sort_keys=True)
        for key in ('source_ref','source_version','source_accession'):
            result[key] = ';'.join(sorted({str(e.get(key,'')) for e in evidence if e.get(key)}))
        # Only the closest applicable explicit review may resolve a contradiction.
        overrides = [e for e in evidence if e['evidence_method'] == 'reviewed_override']
        chosen = evidence
        # An ancestor override cannot silently erase more specific evidence.
        overrides = [e for e in overrides if lineage.index(e['taxid']) <= min(lineage.index(x['taxid']) for x in evidence)]
        if overrides:
            closest = min(lineage.index(e['taxid']) for e in overrides)
            chosen = [e for e in overrides if lineage.index(e['taxid']) == closest]
        classes = {e['genome_class'] for e in chosen if e['genome_class'] != 'unknown'}
        has_unknown = any(e['genome_class'] == 'unknown' for e in chosen)
        all_classes = {e['genome_class'] for e in evidence if e['genome_class'] != 'unknown'}
        result['evidence_taxid'] = chosen[0]['taxid']
        result['evidence_method'] = chosen[0]['evidence_method']
        if len(classes) > 1:
            result.update(resolution_status='conflict', conflict_status='unresolved')
        elif has_unknown:
            states = {e.get('normalization_status') for e in chosen}
            result['resolution_status'] = 'mixed' if 'mixed' in states else 'unrecognized' if 'unrecognized' in states else 'missing'
        elif classes:
            result.update(genome_class=next(iter(classes)),
                          genome_type_detail=';'.join(sorted({e['genome_type_detail'] for e in chosen})),
                          resolution_status='override' if overrides else 'direct' if chosen[0]['taxid']==tid else 'inherited',
                          conflict_status='resolved_by_override' if overrides and (len(all_classes)>1 or len(evidence)>len(chosen)) else 'none')
        return result

    def close(self):
        self.evidence.cache_clear(); self.resolve.cache_clear(); self.con.close()


def build_mapping(db, vmr, out, accession_taxid=None, rules=None, overrides=None, source_version=None):
    db, vmr, out = Path(db).resolve(), Path(vmr).resolve(), Path(out).resolve()
    if out.exists():
        raise RuntimeError('Mapping output exists: choose a NEW frozen version directory')
    tax = Taxonomy(prepare_sqlite(db))
    records = list(vmr_records(vmr))
    if not records:
        tax.close(); raise RuntimeError('VMR contains no species rows')
    if accession_taxid is None:
        candidates = [p for p in (db/'seqid2taxid.map', db/'taxonomy/seqid2taxid.map', db/'accession_taxid.tsv') if p.is_file()]
        if len(candidates)>1:
            tax.close(); raise RuntimeError('Multiple possible cross-references; specify --accession-taxid')
        accession_taxid = candidates[0] if candidates else None
    accession_taxid = Path(accession_taxid).resolve() if accession_taxid else None
    wanted = {a for r in records for a in r['accessions']}
    xref = load_xref(accession_taxid, wanted)
    inputs = {'vmr':sha256(vmr), 'accession_taxid':sha256(accession_taxid) if accession_taxid else None,
              'rules':sha256(Path(rules)) if rules else None, 'overrides':sha256(Path(overrides)) if overrides else None,
              'implementation':sha256(Path(__file__)), 'taxonomy_implementation':sha256(Path(__file__).with_name('taxonomy.py'))}
    vmr_source = vmr.with_suffix('.source.json')
    receipt = json.loads(vmr_source.read_text()) if vmr_source.exists() else {}
    if receipt and receipt.get('sha256') != inputs['vmr']:
        tax.close(); raise RuntimeError('VMR source checksum mismatch')
    version = source_version or receipt.get('source_version') or vmr.name
    source = dict(schema=SCHEMA, rule_version=RULE_VERSION, taxonomy_identity=tax.identity,
                  inputs=inputs, source_version=version, vmr_provenance=receipt,
                  vmr_source=str(vmr), xref_source=str(accession_taxid) if accession_taxid else None,
                  created_utc=utcnow())
    source['build_signature'] = hashlib.sha256(json.dumps({k:source[k] for k in
        ('schema','rule_version','taxonomy_identity','inputs','source_version','vmr_provenance')}, sort_keys=True).encode()).hexdigest()
    unmatched, unknown, evidence = [], [], []
    def add_evidence(tid, row, method, scope, source_ref, extra=None):
        kind, detail, status = normalize_genome_type(row['genome_composition_raw'])
        e = dict(row, taxid=tid, genome_class=kind, genome_type_detail=detail,
                 normalization_status=status, evidence_scope=scope, evidence_method=method,
                 source_ref=source_ref, source_version=version, entity_type='unknown', baltimore_group='unknown')
        if extra:
            e.update(extra)
        evidence.append(e)
        if status == 'unrecognized' and method != 'ICTV_VMR':
            unknown.append(e)
    try:
        for r in records:
            if normalize_genome_type(r['genome_composition_raw'])[2] == 'unrecognized':
                unknown.append(dict(r, taxid='', source_ref=version+'#'+r['source_row']))
            species, invalid = set(), []
            for acc in r['accessions']:
                # An explicit version must match that version; no silent version replacement.
                hits = xref.get(acc, set())
                if not hits:
                    unmatched.append(dict(r, accession=acc, reason='accession_not_in_cross_reference'))
                for hit in hits:
                    try:
                        can = tax.canonical_taxid(hit)
                        group, _, sp, _ = tax.base_route(can)
                        if group != 'virus' or not sp:
                            invalid.append(f'{acc}:{hit}:not_viral_species')
                        else:
                            species.add(sp)
                    except RuntimeError as e:
                        invalid.append(f'{acc}:{hit}:{e}')
            mapping = 'accession_taxid_species'
            if len(species)>1 or invalid:
                unmatched.append(dict(r, reason='one_to_many_or_invalid_taxonomy', details=json.dumps(sorted(species)+invalid)))
                continue
            if not species:
                # UNIQUE exact scientific/official synonym match, never common/fuzzy name.
                matches = {tax.canonical_taxid(x[0]) for x in tax.con.execute('SELECT taxid FROM names WHERE name=?', (r['ictv_species_name'],))}
                matches = {t for t in matches if tax.node(t)[1]=='species' and tax.base_route(t)[0]=='virus'}
                if len(matches) == 1:
                    species = matches; mapping = 'unique_exact_taxonomy_name'
                else:
                    unmatched.append(dict(r, reason='ambiguous_name' if len(matches)>1 else 'unmatched_name_and_accession'))
                    continue
            tid = next(iter(species))
            add_evidence(tid, r, 'ICTV_VMR', 'species', receipt.get('resolved_url', 'https://ictv.global/vmr')+'#'+r['source_row'],
                         dict(mapping_method=mapping, source_accession=';'.join(r['accessions'])))
        for path, override in ((rules,False),(overrides,True)):
            if path is None:
                continue
            for number, r in enumerate(read_table(path), 2):
                required = ('taxid','evidence_scope','genome_composition_raw','source_ref','source_version','reason','review_date')
                if any(not r.get(k) for k in required):
                    raise RuntimeError(f'{path}:{number}: rules/overrides require {required}')
                datetime.strptime(r['review_date'], '%Y-%m-%d')
                tid = tax.canonical_taxid(int(r['taxid']))
                if tax.base_route(tid)[0] != 'virus':
                    raise RuntimeError('Virus rule cannot move a nonviral taxon into Viruses')
                if r['evidence_scope'] not in ('taxon','species','subtree') or (
                    r['evidence_scope']=='species' and tax.node(tid)[1]!='species'):
                    raise RuntimeError('Invalid evidence_scope for taxon rank')
                kind = normalize_genome_type(r['genome_composition_raw'])[0]
                if tax.node(tid)[2] in ('Viruses','Riboviria') and kind != 'unknown':
                    raise RuntimeError('Known mixed virus root cannot receive uniform DNA/RNA rule')
                add_evidence(tid, r, 'reviewed_override' if override else 'authoritative_rule', r['evidence_scope'], r['source_ref'],
                             dict(source_version=r['source_version'], mapping_method='explicit_canonical_taxid',
                                  source_accession=r.get('source_accession',''), notes=r['reason']))
        if not evidence:
            raise RuntimeError('No usable viral evidence mapped; use cache-accessions or provide same-build cross-reference')
        out.mkdir(parents=True)
        path = out/'virus_genome_types.sqlite'
        con = sqlite3.connect(path)
        try:
            con.executescript('CREATE TABLE evidence(id INTEGER PRIMARY KEY,taxid INTEGER,payload TEXT); '
                              'CREATE INDEX evidence_taxid ON evidence(taxid); CREATE TABLE metadata(key TEXT PRIMARY KEY,value TEXT);')
            con.executemany('INSERT INTO evidence(taxid,payload) VALUES (?,?)', ((e['taxid'],json.dumps(e,sort_keys=True,ensure_ascii=False)) for e in evidence))
            con.execute('INSERT INTO metadata VALUES ("source",?)',(json.dumps(source,sort_keys=True),)); con.commit()
        finally:
            con.close()
        resolver = VirusGenomeResolver(path, tax, verify=False)
        rows, conflicts = [], []
        try:
            for tid in sorted({e['taxid'] for e in evidence}):
                ann = resolver.resolve(tid)
                es = resolver.evidence(tid)
                _,_,sp,spname = tax.base_route(tid)
                row = dict(taxid=tid, taxon_rank=tax.node(tid)[1], ncbi_name=tax.node(tid)[2],
                           species_taxid=sp, ncbi_species_name=spname, **ann)
                for key in MAIN_FIELDS:
                    if key not in row:
                        row[key] = ';'.join(sorted({str(e.get(key,'')) for e in es if e.get(key)}))
                rows.append(row)
                if ann['conflict_status'] != 'none':
                    conflicts.append(dict(taxid=tid, **ann))
        finally:
            resolver.close()
        if taxonomy_identity(db)!=tax.identity:
            raise RuntimeError('Taxonomy changed during mapping build')
        for key,p in (('vmr',vmr),('accession_taxid',accession_taxid),('rules',rules),('overrides',overrides)):
            if p and sha256(Path(p)) != inputs[key]:
                raise RuntimeError('Metadata input changed during mapping build')
        write_table(out/'virus_genome_types.tsv', MAIN_FIELDS, rows)
        write_table(out/'virus_metadata.evidence.tsv', ['taxid','evidence_json'],
                    ({'taxid':e['taxid'],'evidence_json':json.dumps(e,sort_keys=True,ensure_ascii=False)} for e in evidence))
        write_table(out/'virus_metadata.unmatched.tsv', ['ictv_species_name','source_row','accession','accession_raw','reason','details'], unmatched)
        write_table(out/'virus_metadata.unknown_labels.tsv', ['taxid','genome_composition_raw','source_ref'], unknown)
        write_table(out/'virus_metadata.conflicts.tsv', ['taxid','resolution_status','conflict_status','evidence_json'], conflicts)
        atomic_json(out/'virus_metadata.build.json', dict(status='complete', schema=SCHEMA, vmr_rows=len(records),
                    evidence_rows=len(evidence), annotated_taxa=len(rows), unmatched_records=len(unmatched),
                    unknown_labels=len(unknown), conflicts=len(conflicts), created_utc=utcnow()))
        atomic_json(out/SOURCE_FILE, dict(source, mapping_sha256=sha256(path)))
        return path
    finally:
        tax.close()


def parser():
    p = argparse.ArgumentParser(description=__doc__)
    sub = p.add_subparsers(dest='command',required=True)
    a=sub.add_parser('fetch',help='Download VMR; retry an empty directory or verify/reuse a completed version')
    a.add_argument('--out-dir',type=Path,required=True); a.add_argument('--url',default=VMR_URL); a.set_defaults(func=fetch)
    a=sub.add_parser('cache-accessions',help='Batch NCBI cross-references with resumable response cache')
    a.add_argument('--vmr',type=Path,required=True); a.add_argument('--cache-dir',type=Path,required=True)
    a.add_argument('--accessions',type=Path,help='Optional newline accession subset'); a.set_defaults(func=cache_accessions)
    a=sub.add_parser('build',help='Build a NEW immutable metadata version against same-build taxonomy')
    for key in ('db','vmr','out'):
        a.add_argument('--'+key,type=Path,required=True)
    for key in ('accession-taxid','rules','overrides'):
        a.add_argument('--'+key,type=Path)
    a.add_argument('--source-version')
    a.set_defaults(func=lambda a: print(build_mapping(a.db,a.vmr,a.out,a.accession_taxid,a.rules,a.overrides,a.source_version)))
    return p


if __name__=='__main__':
    try:
        args=parser().parse_args(); args.func(args)
    except (RuntimeError,OSError,ValueError,sqlite3.Error) as e:
        print(f'ERROR: {e}',file=sys.stderr); sys.exit(1)
