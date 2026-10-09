#!/usr/bin/env python3
"""PRSformer cohort/PGEN preparation and GRID result publication.

Individual exchanges and prepare.json belong to the temporary cache. Formal
individual outputs are RDS; every formal figure has a same-name XLSX workbook.
"""
from __future__ import annotations

import argparse
import gzip
import hashlib
import json
import math
import os
from pathlib import Path
import shutil
import sys
import tempfile
import time

import numpy as np
import pandas as pd


# 🚩 Table and identity handling
def log(message):
	print(message, flush = True)


def require(path):
	p = Path(path).expanduser()
	if not p.is_file():
		raise ValueError(f"Missing input: {p}")
	return p.resolve()


def optional(value):
	return None if value is None or str(value).strip().lower() in {"", "none"} else require(value)


def read_table(path):
	p = require(path)
	if p.suffix.lower() == ".rds":
		import pyreadr
		objects = pyreadr.read_r(str(p))
		if len(objects) != 1 or not isinstance(next(iter(objects.values())), pd.DataFrame):
			raise ValueError(f"Expected one data.frame/data.table in {p}")
		return next(iter(objects.values()))
	opener = gzip.open if p.suffix.lower() == ".gz" else open
	with opener(p, "rt") as stream:
		header = stream.readline()
	# A whitespace regex would collapse an empty TSV field and shift its row.
	separator = "\t" if "\t" in header else r"\s+"
	return pd.read_csv(p, sep = separator, dtype = str, keep_default_na = True)


def normalize_ids(values):
	s = pd.Series(values, copy = False).astype("string").str.strip()
	# R stores many UKB IDs as integral doubles. Preserve all other IDs literally.
	s = s.str.replace(r"^([+-]?\d+)\.0+$", r"\1", regex = True)
	if s.isna().any() or s.isin(["", "NA", "nan", "None"]).any():
		raise ValueError("Missing/empty individual ID")
	if s.str.contains(r"\s", regex = True).any():
		raise ValueError("Whitespace in individual ID")
	return s.astype(str).to_numpy()


def identified(table, label):
	cols = [c for c in ("eid", "IID", "#IID", "id") if c in table.columns]
	if not cols:
		raise ValueError(f"{label}: need eid or IID column")
	d = table.copy()
	d["eid"] = normalize_ids(d[cols[0]])
	if d.eid.duplicated().any():
		raise ValueError(f"{label}: duplicate individual IDs")
	return d


def read_id_list(path):
	p = optional(path)
	if p is None or p.stat().st_size == 0:
		return None if p is None else set()
	d = pd.read_csv(p, sep = r"\s+", header = None, dtype = str, comment = "#")
	if d.empty:
		return set()
	s = d.iloc[:, min(1, d.shape[1] - 1)]
	if str(s.iloc[0]).upper() in {"IID", "EID", "ID"}:
		s = s.iloc[1:]
	return set(normalize_ids(s))


def numeric(series, name):
	v = pd.to_numeric(series, errors = "coerce")
	bad = series.notna() & v.isna() & ~series.astype(str).isin(["", "NA", "NaN", "nan", "."])
	if bad.any():
		raise ValueError(f"{name}: nonnumeric values; encode categorical covariates explicitly")
	return v.where(np.isfinite(v), np.nan).astype(float)


def chromosomes(value):
	out = []
	for part in value.split(","):
		if "-" in part:
			a, b = map(int, part.split("-"))
			out.extend(range(a, b + 1))
		else:
			out.append(int(part))
	if not out or len(out) != len(set(out)) or any(c < 1 or c > 22 for c in out):
		raise ValueError("--chrs must contain unique autosomes 1..22")
	return sorted(out)


def genotype_files(directory, chrs):
	root = Path(directory).expanduser().resolve()
	out = []
	for c in chrs:
		prefix = root / f"chr{c}"
		if prefix.with_suffix(".pgen").is_file():
			pvar = prefix.with_suffix(".pvar")
			if not pvar.is_file():
				pvar = prefix.with_suffix(".pvar.zst")
			out.append((c, require(prefix.with_suffix(".pgen")), require(pvar), require(prefix.with_suffix(".psam")), "pgen"))
		else:
			out.append((c, require(prefix.with_suffix(".bed")), require(prefix.with_suffix(".bim")), require(prefix.with_suffix(".fam")), "bed"))
	return out


def sample_ids(sample_path, mode):
	if mode == "bed":
		d = pd.read_csv(sample_path, sep = r"\s+", header = None, dtype = str)
		if d.shape[1] < 2:
			raise ValueError(f"Invalid FAM: {sample_path}")
		s = normalize_ids(d.iloc[:, 1])
	else:
		with open(sample_path) as stream:
			skip = 0
			for line in stream:
				if line.startswith("##"):
					skip += 1
				else:
					break
		d = pd.read_csv(sample_path, sep = r"\s+", skiprows = skip, dtype = str)
		col = "IID" if "IID" in d else "#IID"
		if col not in d:
			raise ValueError(f"No IID in {sample_path}")
		s = normalize_ids(d[col])
	if len(set(s)) != len(s):
		raise ValueError(f"Duplicate IID (including across FIDs): {sample_path}")
	return s


def decode(value):
	return value.decode() if isinstance(value, bytes) else str(value)


def variant_filter(path):
	p = optional(path)
	if p is None:
		return None
	d = pd.read_csv(p, sep = r"\s+", dtype = str)
	lookup = {str(c).upper().lstrip("#"): c for c in d.columns}
	snp_col = next((lookup[x] for x in ("SNP", "RSID", "ID") if x in lookup), None)
	if snp_col is None:
		# Plain one-column, headerless ID files are also supported.
		d = pd.read_csv(p, sep = r"\s+", header = None, dtype = str)
		if d.shape[1] != 1:
			raise ValueError("--snp-list needs SNP/RSID/ID header, or one column without header")
		snp_col = 0
		lookup = {}
	if d[snp_col].isna().any() or d[snp_col].duplicated().any():
		raise ValueError("SNP list contains missing or duplicate IDs")
	c_col = next((lookup[x] for x in ("CHR", "CHROM") if x in lookup), None)
	p_col = next((lookup[x] for x in ("BP", "POS") if x in lookup), None)
	if (c_col is None) != (p_col is None):
		raise ValueError("SNP-list coordinate check needs both CHR and BP/POS")
	if c_col is None:
		return {s: None for s in d[snp_col]}
	return {s: (int(str(c).removeprefix("chr")), int(p)) for s, c, p in zip(d[snp_col], d[c_col], d[p_col])}


# 🚩 Cohort, labels and fixed splits
def make_cohort(args, files):
	traits = args.traits.split(",")
	covars = [] if args.covariates.lower() == "none" else args.covariates.split(",")
	log(f"COHORT: reading phenotype {args.pheno_file}")
	phenotype = read_table(args.pheno_file)
	# The UKB table contains many unrelated endpoints. Discard them before
	# identified()/merges copy the table, to bound cohort preparation memory.
	needed = {"eid", "IID", "#IID", "id", args.group_col, *covars}
	needed.update(getattr(args, trait + "_col") for trait in traits)
	if "t2dm" in traits and args.t2dm_col == "auto":
		needed.update(("t2dm.Yr2e", "t2dm.Yt2e"))
	phenotype = phenotype[[name for name in phenotype.columns if name in needed]]
	d = identified(phenotype, "Phenotype")
	del phenotype
	if args.group_col not in d:
		log(f"COHORT: reading ancestry {args.ancestry_file}")
		a = identified(read_table(args.ancestry_file), "Ancestry")
		if args.group_col not in a:
			raise ValueError(f"Missing ancestry column: {args.group_col}")
		d = d.merge(a[["eid", args.group_col]], on = "eid", how = "left", validate = "one_to_one")
	d["target"] = d[args.group_col].fillna("UNASSIGNED").replace("", "UNASSIGNED").astype(str)
	source_traits = {getattr(args, t + "_col") for t in traits}
	if "t2dm" in traits and args.t2dm_col == "auto":
		source_traits.update({"t2dm.Yr2e", "t2dm.Yt2e"})
	if set(covars) & (set(traits) | source_traits):
		raise ValueError("An outcome cannot be used as a covariate")
	for trait in traits:
		col = getattr(args, trait + "_col")
		if trait == "t2dm" and col == "auto":
			if not {"t2dm.Yr2e", "t2dm.Yt2e"}.issubset(d.columns):
				raise ValueError("T2DM auto needs t2dm.Yr2e/Yt2e; alternatively set --t2dm-col to an explicit baseline 0/1 column")
			pre = numeric(d["t2dm.Yr2e"], "t2dm.Yr2e")
			inc = numeric(d["t2dm.Yt2e"], "t2dm.Yt2e")
			d[trait] = np.where(pre == 1, 1., np.where((pre == 0) | inc.isin([0, 1]), 0., np.nan))
		else:
			if col not in d:
				raise ValueError(f"Missing trait column: {col}")
			d[trait] = numeric(d[col], col)
		if trait == "t2dm" and not d[trait].dropna().isin([0., 1.]).all():
			raise ValueError("T2DM must be binary 0/1; time-to-event outcomes need a different model")
	for col in covars:
		if col not in d:
			raise ValueError(f"Missing covariate: {col}")
		d[col] = numeric(d[col], col)
	keep, remove = read_id_list(args.keep), read_id_list(args.remove)
	d = d[~d.eid.str.startswith("-")]
	if keep is not None:
		d = d[d.eid.isin(keep)]
	if remove is not None:
		d = d[~d.eid.isin(remove)]
	d = d[d[covars].notna().all(axis = 1) & d[traits].notna().any(axis = 1)]
	common = set(d.eid)
	first_ids = None
	for _, _, _, sample, mode in files:
		ids = sample_ids(sample, mode)
		if first_ids is None:
			first_ids = ids
		common.intersection_update(ids)
	order = [eid for eid in first_ids if eid in common]
	if args.max_samples:
		# Explicit smoke-only sampling; preserve the genotype input order afterward.
		rng = np.random.default_rng(args.seed)
		chosen = set(rng.choice(order, min(args.max_samples, len(order)), replace = False))
		order = [eid for eid in order if eid in chosen]
	d = d.set_index("eid").loc[order].reset_index()
	if len(d) < 10:
		raise ValueError(f"Only {len(d)} complete-covariate, genotype-matched participants")
	if optional(args.split_file) is not None:
		s = identified(read_table(args.split_file), "Split file")
		if "split" not in s:
			raise ValueError("Split file needs eid,split")
		d = d.merge(s[["eid", "split"]], on = "eid", how = "left", validate = "one_to_one", sort = False)
		if not d["split"].isin(["train", "validation", "test"]).all():
			raise ValueError("Every retained individual must have split=train|validation|test")
	else:
		d["split"] = ""
		group_file = optional(args.split_group_file)
		if group_file is not None:
			g = identified(read_table(group_file), "Split group file")
			if "group" not in g:
				raise ValueError("Split group file needs eid,group (related-family/kinship component)")
			d = d.merge(g[["eid", "group"]], on = "eid", how = "left", validate = "one_to_one", sort = False)
			if d["group"].isna().any():
				raise ValueError("Missing split group for retained individual")
			from sklearn.model_selection import GroupShuffleSplit
			g1 = GroupShuffleSplit(n_splits = 1, train_size = args.train_fraction, random_state = args.seed)
			tr, rest = next(g1.split(d, groups = d["group"]))
			g2 = GroupShuffleSplit(n_splits = 1, train_size = args.validation_fraction / (1 - args.train_fraction), random_state = args.seed + 1)
			va, te = next(g2.split(d.iloc[rest], groups = d.iloc[rest]["group"]))
			d.loc[tr, "split"] = "train"
			d.loc[rest[va], "split"] = "validation"
			d.loc[rest[te], "split"] = "test"
		else:
			rng = np.random.default_rng(args.seed)
			strata = d.target + (":" + d.t2dm.fillna(-1).astype(str) if "t2dm" in traits else "")
			for indices in d.groupby(strata, sort = True).indices.values():
				ix = rng.permutation(indices)
				n = len(ix)
				nt = max(1, int(math.floor(n * args.train_fraction)))
				nv = int(math.floor(n * args.validation_fraction))
				if n >= 3:
					nt = min(nt, n - 2)
					nv = max(1, min(nv, n - nt - 1))
				d.loc[ix[:nt], "split"] = "train"
				d.loc[ix[nt:nt + nv], "split"] = "validation"
				d.loc[ix[nt + nv:], "split"] = "test"
			log("SPLIT: stratified by ancestry and T2DM; individual random split. Supply --split-group-file for related-family separation.")
	# Even when an explicit split file is supplied, verify family-group separation.
	if optional(args.split_group_file) is not None and "group" not in d:
		g = identified(read_table(args.split_group_file), "Split group file")
		if "group" not in g:
			raise ValueError("Split group file needs eid,group")
		d = d.merge(g[["eid", "group"]], on = "eid", how = "left", validate = "one_to_one", sort = False)
	if "group" in d:
		if d["group"].isna().any() or d.groupby("group")["split"].nunique().max() != 1:
			raise ValueError("Related-family group spans splits or contains missing values")
	for split in ("train", "validation", "test"):
		if not (d["split"] == split).any():
			raise ValueError(f"Empty {split} partition")
	for trait in traits:
		for split in ("train", "validation", "test"):
			y = d.loc[d["split"] == split, trait].dropna()
			if len(y) < 2 or (trait == "t2dm" and y.nunique() != 2):
				raise ValueError(f"{trait}/{split}: too few outcomes or only one binary class")
		log(f"LABEL {trait}: " + ", ".join(f"{s}={d.loc[d['split'] == s, trait].notna().sum():,}" for s in ("train", "validation", "test")))
	cols = ["eid", "target", "split"] + traits + covars + (["group"] if "group" in d else [])
	return d[cols]


# 🚩 Ordered biallelic variants and train-only genotype QC
def scan_variants(files, allowed, limit):
	import pgenlib
	records = []
	seen = set()
	coordinate_mismatches = 0
	for c, _, vp, _, _ in files:
		with pgenlib.PvarReader(os.fsencode(vp)) as pv:
			for j in range(pv.get_variant_ct()):
				snp = decode(pv.get_variant_id(j))
				if allowed is not None and snp not in allowed:
					continue
				if pv.get_allele_ct(j) != 2:
					continue
				chrom = decode(pv.get_variant_chrom(j)).removeprefix("chr")
				if chrom != str(c):
					raise ValueError(f"chr{c} input contains selected SNP on chromosome {chrom}")
				bp = int(pv.get_variant_pos(j))
				if allowed is not None and allowed[snp] is not None and allowed[snp] != (c, bp):
					coordinate_mismatches += 1
					continue
				ref, alt = decode(pv.get_allele_code(j, 0)), decode(pv.get_allele_code(j, 1))
				if snp in {".", "", "NA"} or not ref or not alt or ref == alt:
					continue
				if snp in seen:
					raise ValueError(f"Duplicate selected SNP ID: {snp}")
				seen.add(snp)
				records.append((c, bp, snp, ref, alt, j))
	if coordinate_mismatches:
		raise ValueError(f"{coordinate_mismatches} SNP-list/target CHR:BP mismatches; harmonize genome build first")
	d = pd.DataFrame(records, columns = ["CHR", "BP", "SNP", "REF", "ALT", "source_index"])
	d = d.sort_values(["CHR", "BP", "SNP"], kind = "stable").reset_index(drop = True)
	if limit and len(d) > limit:
		# Evenly spread an explicitly requested smoke subset across the genome.
		d = d.iloc[np.linspace(0, len(d) - 1, limit, dtype = int)].reset_index(drop = True)
	if len(d) == 0:
		raise ValueError("No selected biallelic autosomal variants")
	return d


def sha256(path):
	h = hashlib.sha256()
	with open(path, "rb") as stream:
		for block in iter(lambda: stream.read(1024 * 1024), b""):
			h.update(block)
	return h.hexdigest()


def identity(args, files):
	options = {k: v for k, v in vars(args).items() if k not in {"command", "func", "check", "replace", "cache_dir"}}
	paths = [args.pheno_file, args.ancestry_file, args.snp_list, args.keep, args.remove, args.split_file, args.split_group_file]
	paths += [str(p) for row in files for p in row[1:4]]
	stats = []
	for value in paths:
		# An ancestry file is optional only if its column is already in the phenotype.
		if value == args.ancestry_file and not Path(value).expanduser().is_file():
			stats.append({"path": value, "exists": False})
			continue
		p = optional(value)
		if p is not None:
			stat = p.stat()
			stats.append({"path": str(p), "size": stat.st_size, "mtime_ns": stat.st_mtime_ns})
	x = {"format_version": 1, "options": options, "inputs": stats}
	return hashlib.sha256(json.dumps(x, sort_keys = True).encode()).hexdigest(), x


def validate_cache(cache, signature):
	p = cache / "prepare.json"
	if not p.is_file():
		return False
	meta = json.loads(p.read_text())
	if meta.get("signature") != signature:
		return False
	for name in ("data.tsv.gz", "variants.tsv.gz"):
		if not (cache / name).is_file() or sha256(cache / name) != meta.get("sha256", {}).get(name):
			raise ValueError(f"Prepared cache is incomplete or changed: {cache / name}; use a new cache or --replace")
	a = np.load(cache / "genotypes.npy", mmap_mode = "r", allow_pickle = False)
	if list(a.shape) != meta["shape"] or a.dtype != np.float16:
		raise ValueError("Prepared genotype shape/type mismatch")
	return True


def prepare(args):
	files = genotype_files(args.dir_gen, chromosomes(args.chrs))
	sig, meta = identity(args, files)
	cache = Path(args.cache_dir).expanduser().resolve()
	if not args.replace and validate_cache(cache, sig):
		log(f"REUSE verified prepared cache: {cache}")
		return
	if (cache / "prepare.json").exists() and not args.replace:
		raise ValueError("Cache input/settings mismatch; choose another --cache-dir or explicitly --replace")
	d = make_cohort(args, files)
	v = scan_variants(files, variant_filter(args.snp_list), args.max_variants)
	n, m = len(d), len(v)
	peak_bytes = 4 * n * m + 512 * 1024 ** 2
	log(f"PREPARE: {n:,} individuals x {m:,} candidate SNPs; ALT dosage, missing=-1, float16")
	log(f"DISK: upper bound {peak_bytes / 1024 ** 3:.2f} GiB for two genotype layouts plus reserve; train-only MAF>={args.maf}, call rate>={args.call_rate}")
	if args.max_samples or args.max_variants:
		log("SMOKE SUBSET requested explicitly; results are not a full-cohort/full-genome analysis.")
	if args.check:
		log("CHECK OK: cohort, labels, splits and variant metadata checked. Genotype content QC occurs during prepare.")
		return
	cache.mkdir(parents = True, exist_ok = True)
	if shutil.disk_usage(cache).free < peak_bytes:
		raise ValueError("Insufficient cache disk space; use --cache-dir on a large local SSD")
	import pgenlib
	train = (d["split"] == "train").to_numpy()
	kept = []
	with tempfile.TemporaryDirectory(prefix = "prepare-", dir = cache) as td:
		work = Path(td)
		# Sequential writes in variant-major order avoid tiny strided disk writes.
		raw_path = work / "variant_major.bin"
		raw = np.memmap(raw_path, mode = "w+", dtype = np.float16, shape = (m, n))
		write_row = 0
		last_progress = time.monotonic()
		for c, gp, vp, sp, mode in files:
			vc = v[v.CHR == c]
			if vc.empty:
				continue
			ids = sample_ids(sp, mode)
			lookup = {eid: i for i, eid in enumerate(ids)}
			original = np.asarray([lookup[eid] for eid in d.eid], dtype = np.uint32)
			subset = np.sort(original)
			reorder = np.searchsorted(subset, original)
			with pgenlib.PvarReader(os.fsencode(vp)) as pv:
				kwargs = {"raw_sample_ct": len(ids), "sample_subset": subset, "pvar": pv}
				with pgenlib.PgenReader(os.fsencode(gp), **kwargs) as pg:
					for start in range(0, len(vc), args.chunk_variants):
						block = vc.iloc[start:start + args.chunk_variants]
						buf = np.empty((len(block), n), dtype = np.float32, order = "C")
						pg.read_dosages_list(np.asarray(block.source_index, dtype = np.uint32), buf, allele_idx = 1)
						buf = np.ascontiguousarray(buf[:, reorder])
						missing = (buf == -9) | ~np.isfinite(buf)
						if np.any((~missing) & ((buf < 0) | (buf > 2))):
							raise ValueError(f"chr{c}: dosage outside [0,2] with unrecognized missing encoding")
						tr = buf[:, train]
						ok = ~missing[:, train]
						calls = ok.sum(axis = 1)
						freq = np.divide(np.where(ok, tr, 0).sum(axis = 1, dtype = np.float64), 2 * calls, out = np.zeros(len(block)), where = calls > 0)
						maf = np.minimum(freq, 1 - freq)
						call_rate = calls / train.sum()
						use = (maf >= args.maf) & (call_rate >= args.call_rate) & (calls > 0) & (maf > 0)
						buf[missing] = -1
						if use.any():
							count = int(use.sum())
							raw[write_row:write_row + count] = buf[use]
							selected = block.loc[use].copy()
							selected["MAF_train"] = maf[use]
							selected["call_rate_train"] = call_rate[use]
							kept.append(selected)
							write_row += count
						if time.monotonic() - last_progress >= 60:
							log(f"GENOTYPES chr{c}: checked {min(start + args.chunk_variants, len(vc)):,}/{len(vc):,}; cumulative {write_row:,} SNPs retained")
							last_progress = time.monotonic()
			log(f"GENOTYPES chr{c}: cumulative {write_row:,} SNPs retained")
		if write_row == 0:
			raise ValueError("No SNP passed train-only MAF/call-rate QC")
		raw.flush()
		output = np.lib.format.open_memmap(work / "genotypes.npy", mode = "w+", dtype = np.float16, shape = (n, write_row))
		# Bound the transpose work buffer at approximately 64 MiB.
		batch = max(1, min(1024, (64 * 1024 ** 2) // max(2 * write_row, 1)))
		log(f"TRANSPOSE: writing {n:,} sample-major genotype rows")
		for start in range(0, n, batch):
			output[start:start + batch] = raw[:write_row, start:start + batch].T
			if time.monotonic() - last_progress >= 60:
				log(f"TRANSPOSE: {min(start + batch, n):,}/{n:,} rows")
				last_progress = time.monotonic()
		output.flush()
		del output, raw
		raw_path.unlink()
		final_v = pd.concat(kept, ignore_index = True).drop(columns = "source_index")
		d.to_csv(work / "data.tsv.gz", sep = "\t", index = False, na_rep = "NA")
		final_v.to_csv(work / "variants.tsv.gz", sep = "\t", index = False)
		meta.update({"signature": sig, "shape": [n, write_row], "dosage": "ALT, float16; missing=-1", "split_counts": d["split"].value_counts().to_dict(), "sha256": {name: sha256(work / name) for name in ("data.tsv.gz", "variants.tsv.gz")}})
		(work / "prepare.json").write_text(json.dumps(meta, indent = 2) + "\n")
		# The final manifest is installed last; interrupted preparation is never reusable.
		(cache / "prepare.json").unlink(missing_ok = True)
		for name in ("genotypes.npy", "data.tsv.gz", "variants.tsv.gz", "prepare.json"):
			os.replace(work / name, cache / name)
	log(f"PREPARE OK: {n:,} individuals x {write_row:,} SNPs -> {cache}")


# 🚩 Formal RDS and same-name PNG/XLSX result pairs
def write_rds(frame, path):
	import pyreadr
	p = Path(path)
	p.parent.mkdir(parents = True, exist_ok = True)
	data = frame.copy()
	for col in data:
		if data[col].dtype.name in {"string", "category"}:
			data[col] = data[col].astype(object)
	with tempfile.NamedTemporaryFile(prefix = p.stem + ".", suffix = ".rds", dir = p.parent, delete = False) as stream:
		tmp = Path(stream.name)
	try:
		pyreadr.write_rds(str(tmp), data, compress = "gzip")
		os.replace(tmp, p)
	finally:
		tmp.unlink(missing_ok = True)


def workbook(frame, path, title):
	from openpyxl import Workbook
	from openpyxl.styles import Font, PatternFill
	from openpyxl.utils import get_column_letter
	wb = Workbook()
	ws = wb.active
	ws.title = title[:31]
	ws.append(list(frame.columns))
	for row in frame.itertuples(index = False, name = None):
		ws.append([None if pd.isna(x) else x.item() if isinstance(x, np.generic) else x for x in row])
	for cell in ws[1]:
		cell.font = Font(bold = True, color = "FFFFFF")
		cell.fill = PatternFill("solid", fgColor = "234D65")
	ws.freeze_panes = "A2"
	ws.auto_filter.ref = ws.dimensions
	for i, col in enumerate(frame.columns, 1):
		ws.column_dimensions[get_column_letter(i)].width = min(34, max(13, len(str(col)) + 2))
	wb.save(path)


def attach_covariate_values(scores, cohort, covariates):
	"""Retain original outcome-model covariates for direct benchmark verification."""
	names = [] if covariates.strip().lower() in ("", "none") else [c.strip() for c in covariates.split(",")]
	if len(names) != len(set(names)) or any(name not in cohort for name in names):
		raise ValueError("Published covariates must be unique and present in the prepared cohort")
	if not names:
		return scores.copy()
	source = cohort[["eid"] + names].copy()
	for name in names:
		source[name] = numeric(source[name], name)
		if source[name].isna().any():
			raise ValueError(f"Prepared model covariate {name} is missing or nonfinite")
	source = source.rename(columns={name: f"covariate.{name}" for name in names})
	out = scores.merge(source, on="eid", how="left", validate="one_to_one", indicator=True, sort=False)
	if not out["_merge"].eq("both").all():
		raise ValueError("A published prediction has no matching original covariates")
	return out.drop(columns="_merge")


def canonical_metrics(metrics):
	"""Upgrade old metric names by their original formulas; never refit a model."""
	metrics = metrics.copy()
	metrics.loc[metrics.metric == "prediction_R2", "metric"] = "residual_correlation_R2"
	keys = ["trait", "target", "split", "metric"]
	if metrics.duplicated(keys).any():
		raise ValueError("Duplicate metric names after explicit legacy-label migration")
	derived = metrics[metrics.metric == "full_R2"].copy()
	derived["metric"] = "predictive_R2"
	present = set(map(tuple, metrics[keys].to_numpy()))
	derived = derived[[tuple(row) not in present for row in derived[keys].to_numpy()]]
	return pd.concat([metrics, derived], ignore_index=True)


def publish(args):
	run = Path(args.run_dir).expanduser().resolve()
	cache = Path(args.cache_dir).expanduser().resolve()
	for name in ("model.pt", "metrics.tsv", "history.tsv", "test_predictions.tsv.gz"):
		require(run / name)
	meta = json.loads(require(cache / "prepare.json").read_text())
	if not validate_cache(cache, meta["signature"]):
		raise ValueError("Cannot validate cache before publication")
	evaluation = json.loads(require(run / "evaluation.json").read_text())
	if evaluation.get("format_version") != 1 or evaluation.get("preparation_signature") != meta["signature"]:
		raise ValueError("Evaluation was generated from a different prepared cache")
	for name, key in (("data.tsv.gz", "data_sha256"), ("variants.tsv.gz", "variants_sha256")):
		if evaluation["inputs"].get(key) != sha256(cache / name):
			raise ValueError(f"Evaluation input mismatch: {name}")
	gen_stat = (cache / "genotypes.npy").stat()
	gen_ident = evaluation["inputs"]["genotypes"]
	if gen_ident.get("size") != gen_stat.st_size or gen_ident.get("mtime_ns") != gen_stat.st_mtime_ns or list(gen_ident.get("shape", [])) != meta["shape"]:
		raise ValueError("Evaluation genotype file was replaced or modified")
	for name in ("metrics.tsv", "history.tsv", "test_predictions.tsv.gz"):
		if evaluation["outputs_sha256"].get(name) != sha256(run / name):
			raise ValueError(f"Evaluation output was changed after completion: {name}")
	model_stat = (run / "model.pt").stat()
	if evaluation["model"].get("size") != model_stat.st_size or evaluation["model"].get("mtime_ns") != model_stat.st_mtime_ns:
		raise ValueError("Evaluation checkpoint was replaced or modified")
	pred = read_table(run / "test_predictions.tsv.gz")
	pred["eid"] = normalize_ids(pred.eid)
	needed = {"target", "split", "trait", "y", "baseline", "prediction", "genetic_score"}
	if not needed.issubset(pred.columns) or not (pred["split"] == "test").all():
		raise ValueError("Only held-out test predictions can be published")
	if pred.duplicated(["eid", "trait"]).any():
		raise ValueError("Duplicate test prediction ID/trait")
	cohort = identified(read_table(cache / "data.tsv.gz"), "Prepared cohort")
	test_ids = set(cohort.loc[cohort["split"] == "test", "eid"])
	if not set(pred.eid).issubset(test_ids):
		raise ValueError("Prediction includes a participant outside prepared test split")
	traits = args.traits.split(",")
	if set(pred.trait) != set(traits):
		raise ValueError("Requested traits do not match prediction output")
	for trait in traits:
		if set(pred.loc[pred.trait == trait, "eid"]) != test_ids:
			raise ValueError(f"Incomplete test predictions: {trait}")
	metrics = canonical_metrics(pd.read_csv(run / "metrics.tsv", sep = "\t"))
	history = pd.read_csv(run / "history.tsv", sep = "\t")
	if not (metrics["split"] == "test").all():
		raise ValueError("Performance table contains non-test rows")
	root = Path(args.output_root).expanduser().resolve() / "prsformer"
	score_root = Path(args.score_dir).expanduser().resolve()
	destinations = [root / x for x in ("3.prsformer.pt", "3.prsformer.split.rds", "3.prsformer.performance.png", "3.prsformer.performance.xlsx", "3.prsformer.training.png", "3.prsformer.training.xlsx")]
	destinations += [score_root / t / "3.prsformer.scores.rds" for t in traits]
	if not args.replace and any(p.exists() for p in destinations):
		raise ValueError("Formal PRSformer outputs already exist; select another output root/score dir or explicitly --replace")
	root.mkdir(parents = True, exist_ok = True)
	import matplotlib
	matplotlib.use("Agg")
	import matplotlib.pyplot as plt
	plt.rcParams.update({"font.size": 10, "axes.spines.top": False, "axes.spines.right": False, "savefig.dpi": 160})
	with tempfile.TemporaryDirectory(prefix = "publish-", dir = root) as td:
		stage = Path(td)
		# Primary metric above, all remaining exact metrics in the table below.
		fig, axs = plt.subplots(2, len(traits), figsize = (6.4 * len(traits), 11), gridspec_kw = {"height_ratios": [1, 2.2]}, squeeze = False)
		for j, trait in enumerate(traits):
			metric = "AUC" if trait == "t2dm" else "predictive_R2"
			sub = metrics[(metrics.trait == trait) & (metrics.metric == metric)]
			ax = axs[0, j]
			ax.bar(sub.target.astype(str), sub.value, color = "#347F9C")
			ax.set_title(f"{trait}: {metric}", loc = "left")
			ax.tick_params(axis = "x", labelrotation = 30)
			ax.axhline(0, color = "#777777", linewidth = .5)
			detail = metrics[metrics.trait == trait].pivot(index = "metric", columns = "target", values = "value")
			counts = metrics[metrics.trait == trait].groupby("target").n.first()
			detail.loc["n"] = counts
			labels = [[f"{x:.5g}" if pd.notna(x) else "NA" for x in row] for row in detail.to_numpy()]
			ax = axs[1, j]
			ax.axis("off")
			table = ax.table(cellText = labels, rowLabels = list(detail.index), colLabels = list(detail.columns), cellLoc = "right", loc = "upper center", bbox = [.28, .04, .72, .92])
			table.auto_set_font_size(False)
			table.set_fontsize(7)
		fig.suptitle("PRSformer: fixed held-out test set; point estimates", y = 1)
		fig.tight_layout()
		fig.savefig(stage / "3.prsformer.performance.png", bbox_inches = "tight")
		plt.close(fig)
		workbook(metrics, stage / "3.prsformer.performance.xlsx", "performance")
		fig, axs = plt.subplots(1, len(traits), figsize = (5 * len(traits), 3.6), squeeze = False)
		for ax, trait in zip(axs[0], traits):
			for split, color in (("train", "#347F9C"), ("validation", "#D18443")):
				sub = history[(history.trait == trait) & (history["split"] == split)]
				ax.plot(sub.epoch, sub.loss, label = split, color = color)
			ax.set(title = trait, xlabel = "Epoch", ylabel = "Standardized residual MSE" if trait != "t2dm" else "Binary cross entropy")
			ax.legend(frameon = False)
		fig.tight_layout()
		fig.savefig(stage / "3.prsformer.training.png", bbox_inches = "tight")
		plt.close(fig)
		workbook(history, stage / "3.prsformer.training.xlsx", "training")
		shutil.copyfile(run / "model.pt", stage / "3.prsformer.pt")
		write_rds(cohort[[c for c in ("eid", "target", "split", "group") if c in cohort]], stage / "3.prsformer.split.rds")
		for trait in traits:
			x = pred[pred.trait == trait].drop(columns = "trait").copy()
			for col in ("y", "baseline", "prediction", "genetic_score"):
				x[col] = pd.to_numeric(x[col], errors = "raise")
			x = x.rename(columns = {"y": "outcome", "genetic_score": "prsformer"})
			x = attach_covariate_values(x, cohort, meta["options"]["covariates"])
			x["prsformer_scale"] = "log_odds_increment" if trait == "t2dm" else "outcome_units"
			x["covariate_names"] = meta["options"]["covariates"]
			x["endpoint_definition"] = (
				"baseline_t2dm_from_Yr2e_Yt2e" if trait == "t2dm" and meta["options"]["t2dm_col"] == "auto"
				else meta["options"][trait + "_col"]
			)
			x["preparation_signature"] = meta["signature"]
			write_rds(x, stage / f"{trait}.scores.rds")
		for path in stage.iterdir():
			if path.name.endswith(".scores.rds"):
				trait = path.name.removesuffix(".scores.rds")
				dest = score_root / trait / "3.prsformer.scores.rds"
			else:
				dest = root / path.name
			dest.parent.mkdir(parents = True, exist_ok = True)
			# Cross-filesystem safe, per-file atomic publication.
			with tempfile.NamedTemporaryFile(prefix = dest.stem + ".", dir = dest.parent, delete = False) as stream:
				tmp = Path(stream.name)
			try:
				shutil.copyfile(path, tmp)
				os.replace(tmp, dest)
			finally:
				tmp.unlink(missing_ok = True)
	log(f"PUBLISH OK: model and aggregate figures/workbooks at {root}; test-only RDS under {score_root}/<trait>/")


# 🚩 Command line
def parser():
	p = argparse.ArgumentParser(description = __doc__)
	sub = p.add_subparsers(dest = "command", required = True)
	a = sub.add_parser("prepare", help = "Align individual data, fix splits, apply train-only QC and write dosage memmap")
	a.add_argument("--dir-gen", default = "/mnt/f/gen/ukb/37/hap")
	a.add_argument("--pheno-file", default = "/mnt/d/data/ukb/phe/Rdata/all.rds")
	a.add_argument("--ancestry-file", default = "/mnt/d/data/ukb/pca_proj/ukb.ancestry.auto.tsv.gz")
	a.add_argument("--group-col", default = "genetic_ancestry")
	a.add_argument("--traits", default = "height,ldl,t2dm")
	a.add_argument("--covariates", default = "age,sex,PC1,PC2")
	a.add_argument("--height-col", default = "height")
	a.add_argument("--ldl-col", default = "ldl")
	a.add_argument("--t2dm-col", default = "auto")
	a.add_argument("--snp-list", default = "/mnt/f/refLD/csx/snpinfo_mult_1kg_hm3")
	a.add_argument("--chrs", default = "1-22")
	a.add_argument("--keep", default = "none")
	a.add_argument("--remove", default = "/mnt/d/files/ukb.exclude.id")
	a.add_argument("--split-file", default = "none")
	a.add_argument("--split-group-file", default = "none")
	a.add_argument("--train-fraction", type = float, default = .6)
	a.add_argument("--validation-fraction", type = float, default = .2)
	a.add_argument("--seed", type = int, default = 20260904)
	a.add_argument("--maf", type = float, default = .01)
	a.add_argument("--call-rate", type = float, default = .98)
	a.add_argument("--chunk-variants", type = int, default = 32)
	a.add_argument("--max-samples", type = int, default = 0)
	a.add_argument("--max-variants", type = int, default = 0)
	a.add_argument("--cache-dir", required = True)
	a.add_argument("--check", action = "store_true")
	a.add_argument("--replace", action = "store_true")
	a.set_defaults(func = prepare)
	b = sub.add_parser("publish", help = "Publish held-out RDS predictions, reusable model and aggregate figure/workbook pairs")
	b.add_argument("--run-dir", required = True)
	b.add_argument("--cache-dir", required = True)
	b.add_argument("--output-root", default = "/mnt/d/analysis/grid")
	b.add_argument("--score-dir", default = "/mnt/d/data/ukb/pgs")
	b.add_argument("--traits", default = "height,ldl,t2dm")
	b.add_argument("--replace", action = "store_true")
	b.set_defaults(func = publish)
	return p


def main():
	a = parser().parse_args()
	a.traits = ",".join(x.strip() for x in a.traits.split(","))
	traits = a.traits.split(",")
	if not set(traits).issubset({"height", "ldl", "t2dm"}) or len(traits) != len(set(traits)):
		raise ValueError("Choose unique traits from height,ldl,t2dm")
	if a.command == "prepare":
		a.covariates = ",".join(x.strip() for x in a.covariates.split(","))
		if len(set(a.covariates.split(","))) != len(a.covariates.split(",")):
			raise ValueError("Covariate names must not repeat")
		if not (0 < a.train_fraction < 1 and 0 < a.validation_fraction < 1 - a.train_fraction):
			raise ValueError("Need positive train, validation and test fractions")
		if not (0 <= a.maf < .5 and 0 < a.call_rate <= 1 and a.chunk_variants > 0 and a.max_samples >= 0 and a.max_variants >= 0):
			raise ValueError("Invalid MAF/call-rate/chunk/subset setting")
	a.func(a)


if __name__ == "__main__":
	try:
		main()
	except (ValueError, FileNotFoundError, ImportError) as exc:
		raise SystemExit(f"ERROR: {exc}")
