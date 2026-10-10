#!/usr/bin/env python3
"""PRSformer: cohort preparation, official-model training, prediction and reports.

Public entry: grid.sh prsformer. Temporary individual data stay in the cache;
formal predictions are RDS and aggregate reports are paired XLSX/PNG files.
PRSformer was developed by 23andMe, Inc. This adapter imports the user's
checkout without redistributing upstream code; see the upstream LICENSE.txt.
Training uses masked height/ldl/baseline-t2dm tasks and training-only transforms.
"""

from __future__ import annotations
import argparse
import ast
import contextlib
import hashlib
import importlib
import json
import math
import os
from pathlib import Path
import random
import shutil
import subprocess
import sys
import time
import types
import numpy as np
import pandas as pd
from scipy.optimize import minimize
from scipy.special import expit
from scipy.stats import rankdata
import gzip
import tempfile

sys.dont_write_bytecode = True
SUPPORTED_TRAITS = {"height": "continuous", "ldl": "continuous", "t2dm": "binary"}

def log(message):
	print(message, flush = True)


def require(path):
	p = Path(path).expanduser()
	if not p.is_file():
		raise ValueError(f"Missing input: {p}")
	return p.resolve()


def optional(value):
	return None if value is None or str(value).strip().lower() in {"", "none"} else require(value)


def read_table(path):
	p = require(path)
	if p.suffix.lower() == ".rds":
		import pyreadr
		objects = pyreadr.read_r(str(p))
		if len(objects) != 1 or not isinstance(next(iter(objects.values())), pd.DataFrame):
			raise ValueError(f"Expected one data.frame/data.table in {p}")
		return next(iter(objects.values()))
	opener = gzip.open if p.suffix.lower() == ".gz" else open
	with opener(p, "rt") as stream:
		header = stream.readline()
	# A whitespace regex would collapse an empty TSV field and shift its row.
	separator = "\t" if "\t" in header else r"\s+"
	return pd.read_csv(p, sep = separator, dtype = str, keep_default_na = True)


def normalize_ids(values):
	s = pd.Series(values, copy = False).astype("string").str.strip()
	# R stores many UKB IDs as integral doubles. Preserve all other IDs literally.
	s = s.str.replace(r"^([+-]?\d+)\.0+$", r"\1", regex = True)
	if s.isna().any() or s.isin(["", "NA", "nan", "None"]).any():
		raise ValueError("Missing/empty individual ID")
	if s.str.contains(r"\s", regex = True).any():
		raise ValueError("Whitespace in individual ID")
	return s.astype(str).to_numpy()


def identified(table, label):
	cols = [c for c in ("eid", "IID", "#IID", "id") if c in table.columns]
	if not cols:
		raise ValueError(f"{label}: need eid or IID column")
	d = table.copy()
	d["eid"] = normalize_ids(d[cols[0]])
	if d.eid.duplicated().any():
		raise ValueError(f"{label}: duplicate individual IDs")
	return d


def read_id_list(path):
	p = optional(path)
	if p is None or p.stat().st_size == 0:
		return None if p is None else set()
	d = pd.read_csv(p, sep = r"\s+", header = None, dtype = str, comment = "#")
	if d.empty:
		return set()
	s = d.iloc[:, min(1, d.shape[1] - 1)]
	if str(s.iloc[0]).upper() in {"IID", "EID", "ID"}:
		s = s.iloc[1:]
	return set(normalize_ids(s))


def numeric(series, name):
	v = pd.to_numeric(series, errors = "coerce")
	bad = series.notna() & v.isna() & ~series.astype(str).isin(["", "NA", "NaN", "nan", "."])
	if bad.any():
		raise ValueError(f"{name}: nonnumeric values; encode categorical covariates explicitly")
	return v.where(np.isfinite(v), np.nan).astype(float)


def chromosomes(value):
	out = []
	for part in value.split(","):
		if "-" in part:
			a, b = map(int, part.split("-"))
			out.extend(range(a, b + 1))
		else:
			out.append(int(part))
	if not out or len(out) != len(set(out)) or any(c < 1 or c > 22 for c in out):
		raise ValueError("--chrs must contain unique autosomes 1..22")
	return sorted(out)


def genotype_files(directory, chrs):
	root = Path(directory).expanduser().resolve()
	out = []
	for c in chrs:
		prefix = root / f"chr{c}"
		if prefix.with_suffix(".pgen").is_file():
			pvar = prefix.with_suffix(".pvar")
			if not pvar.is_file():
				pvar = prefix.with_suffix(".pvar.zst")
			out.append((c, require(prefix.with_suffix(".pgen")), require(pvar), require(prefix.with_suffix(".psam")), "pgen"))
		else:
			out.append((c, require(prefix.with_suffix(".bed")), require(prefix.with_suffix(".bim")), require(prefix.with_suffix(".fam")), "bed"))
	return out


def sample_ids(sample_path, mode):
	if mode == "bed":
		d = pd.read_csv(sample_path, sep = r"\s+", header = None, dtype = str)
		if d.shape[1] < 2:
			raise ValueError(f"Invalid FAM: {sample_path}")
		s = normalize_ids(d.iloc[:, 1])
	else:
		with open(sample_path) as stream:
			skip = 0
			for line in stream:
				if line.startswith("##"):
					skip += 1
				else:
					break
		d = pd.read_csv(sample_path, sep = r"\s+", skiprows = skip, dtype = str)
		col = "IID" if "IID" in d else "#IID"
		if col not in d:
			raise ValueError(f"No IID in {sample_path}")
		s = normalize_ids(d[col])
	if len(set(s)) != len(s):
		raise ValueError(f"Duplicate IID (including across FIDs): {sample_path}")
	return s


def decode(value):
	return value.decode() if isinstance(value, bytes) else str(value)


def variant_filter(path):
	p = optional(path)
	if p is None:
		return None
	d = pd.read_csv(p, sep = r"\s+", dtype = str)
	lookup = {str(c).upper().lstrip("#"): c for c in d.columns}
	snp_col = next((lookup[x] for x in ("SNP", "RSID", "ID") if x in lookup), None)
	if snp_col is None:
		# Plain one-column, headerless ID files are also supported.
		d = pd.read_csv(p, sep = r"\s+", header = None, dtype = str)
		if d.shape[1] != 1:
			raise ValueError("--snp-list needs SNP/RSID/ID header, or one column without header")
		snp_col = 0
		lookup = {}
	if d[snp_col].isna().any() or d[snp_col].duplicated().any():
		raise ValueError("SNP list contains missing or duplicate IDs")
	c_col = next((lookup[x] for x in ("CHR", "CHROM") if x in lookup), None)
	p_col = next((lookup[x] for x in ("BP", "POS") if x in lookup), None)
	if (c_col is None) != (p_col is None):
		raise ValueError("SNP-list coordinate check needs both CHR and BP/POS")
	if c_col is None:
		return {s: None for s in d[snp_col]}
	return {s: (int(str(c).removeprefix("chr")), int(p)) for s, c, p in zip(d[snp_col], d[c_col], d[p_col])}


def make_cohort(args, files):
	traits = args.traits.split(",")
	covars = [] if args.covariates.lower() == "none" else args.covariates.split(",")
	log(f"COHORT: reading phenotype {args.pheno_file}")
	phenotype = read_table(args.pheno_file)
	# The UKB table contains many unrelated endpoints. Discard them before
	# identified()/merges copy the table, to bound cohort preparation memory.
	needed = {"eid", "IID", "#IID", "id", args.group_col, *covars}
	needed.update(getattr(args, trait + "_col") for trait in traits)
	if "t2dm" in traits and args.t2dm_col == "auto":
		needed.update(("t2dm.Yr2e", "t2dm.Yt2e"))
	phenotype = phenotype[[name for name in phenotype.columns if name in needed]]
	d = identified(phenotype, "Phenotype")
	del phenotype
	if args.group_col not in d:
		log(f"COHORT: reading ancestry {args.ancestry_file}")
		a = identified(read_table(args.ancestry_file), "Ancestry")
		if args.group_col not in a:
			raise ValueError(f"Missing ancestry column: {args.group_col}")
		d = d.merge(a[["eid", args.group_col]], on = "eid", how = "left", validate = "one_to_one")
	d["target"] = d[args.group_col].fillna("UNASSIGNED").replace("", "UNASSIGNED").astype(str)
	source_traits = {getattr(args, t + "_col") for t in traits}
	if "t2dm" in traits and args.t2dm_col == "auto":
		source_traits.update({"t2dm.Yr2e", "t2dm.Yt2e"})
	if set(covars) & (set(traits) | source_traits):
		raise ValueError("An outcome cannot be used as a covariate")
	for trait in traits:
		col = getattr(args, trait + "_col")
		if trait == "t2dm" and col == "auto":
			if not {"t2dm.Yr2e", "t2dm.Yt2e"}.issubset(d.columns):
				raise ValueError("T2DM auto needs t2dm.Yr2e/Yt2e; alternatively set --t2dm-col to an explicit baseline 0/1 column")
			pre = numeric(d["t2dm.Yr2e"], "t2dm.Yr2e")
			inc = numeric(d["t2dm.Yt2e"], "t2dm.Yt2e")
			d[trait] = np.where(pre == 1, 1., np.where((pre == 0) | inc.isin([0, 1]), 0., np.nan))
		else:
			if col not in d:
				raise ValueError(f"Missing trait column: {col}")
			d[trait] = numeric(d[col], col)
		if trait == "t2dm" and not d[trait].dropna().isin([0., 1.]).all():
			raise ValueError("T2DM must be binary 0/1; time-to-event outcomes need a different model")
	for col in covars:
		if col not in d:
			raise ValueError(f"Missing covariate: {col}")
		d[col] = numeric(d[col], col)
	keep, remove = read_id_list(args.keep), read_id_list(args.remove)
	d = d[~d.eid.str.startswith("-")]
	if keep is not None:
		d = d[d.eid.isin(keep)]
	if remove is not None:
		d = d[~d.eid.isin(remove)]
	d = d[d[covars].notna().all(axis = 1) & d[traits].notna().any(axis = 1)]
	common = set(d.eid)
	first_ids = None
	for _, _, _, sample, mode in files:
		ids = sample_ids(sample, mode)
		if first_ids is None:
			first_ids = ids
		common.intersection_update(ids)
	order = [eid for eid in first_ids if eid in common]
	if args.max_samples:
		# Explicit smoke-only sampling; preserve the genotype input order afterward.
		rng = np.random.default_rng(args.seed)
		chosen = set(rng.choice(order, min(args.max_samples, len(order)), replace = False))
		order = [eid for eid in order if eid in chosen]
	d = d.set_index("eid").loc[order].reset_index()
	if len(d) < 10:
		raise ValueError(f"Only {len(d)} complete-covariate, genotype-matched participants")
	if optional(args.split_file) is not None:
		s = identified(read_table(args.split_file), "Split file")
		if "split" not in s:
			raise ValueError("Split file needs eid,split")
		d = d.merge(s[["eid", "split"]], on = "eid", how = "left", validate = "one_to_one", sort = False)
		if not d["split"].isin(["train", "validation", "test"]).all():
			raise ValueError("Every retained individual must have split=train|validation|test")
	else:
		d["split"] = ""
		group_file = optional(args.split_group_file)
		if group_file is not None:
			g = identified(read_table(group_file), "Split group file")
			if "group" not in g:
				raise ValueError("Split group file needs eid,group (related-family/kinship component)")
			d = d.merge(g[["eid", "group"]], on = "eid", how = "left", validate = "one_to_one", sort = False)
			if d["group"].isna().any():
				raise ValueError("Missing split group for retained individual")
			from sklearn.model_selection import GroupShuffleSplit
			g1 = GroupShuffleSplit(n_splits = 1, train_size = args.train_fraction, random_state = args.seed)
			tr, rest = next(g1.split(d, groups = d["group"]))
			g2 = GroupShuffleSplit(n_splits = 1, train_size = args.validation_fraction / (1 - args.train_fraction), random_state = args.seed + 1)
			va, te = next(g2.split(d.iloc[rest], groups = d.iloc[rest]["group"]))
			d.loc[tr, "split"] = "train"
			d.loc[rest[va], "split"] = "validation"
			d.loc[rest[te], "split"] = "test"
		else:
			rng = np.random.default_rng(args.seed)
			strata = d.target + (":" + d.t2dm.fillna(-1).astype(str) if "t2dm" in traits else "")
			for indices in d.groupby(strata, sort = True).indices.values():
				ix = rng.permutation(indices)
				n = len(ix)
				nt = max(1, int(math.floor(n * args.train_fraction)))
				nv = int(math.floor(n * args.validation_fraction))
				if n >= 3:
					nt = min(nt, n - 2)
					nv = max(1, min(nv, n - nt - 1))
				d.loc[ix[:nt], "split"] = "train"
				d.loc[ix[nt:nt + nv], "split"] = "validation"
				d.loc[ix[nt + nv:], "split"] = "test"
			log("SPLIT: stratified by ancestry and T2DM; individual random split. Supply --split-group-file for related-family separation.")
	# Even when an explicit split file is supplied, verify family-group separation.
	if optional(args.split_group_file) is not None and "group" not in d:
		g = identified(read_table(args.split_group_file), "Split group file")
		if "group" not in g:
			raise ValueError("Split group file needs eid,group")
		d = d.merge(g[["eid", "group"]], on = "eid", how = "left", validate = "one_to_one", sort = False)
	if "group" in d:
		if d["group"].isna().any() or d.groupby("group")["split"].nunique().max() != 1:
			raise ValueError("Related-family group spans splits or contains missing values")
	for split in ("train", "validation", "test"):
		if not (d["split"] == split).any():
			raise ValueError(f"Empty {split} partition")
	for trait in traits:
		for split in ("train", "validation", "test"):
			y = d.loc[d["split"] == split, trait].dropna()
			if len(y) < 2 or (trait == "t2dm" and y.nunique() != 2):
				raise ValueError(f"{trait}/{split}: too few outcomes or only one binary class")
		log(f"LABEL {trait}: " + ", ".join(f"{s}={d.loc[d['split'] == s, trait].notna().sum():,}" for s in ("train", "validation", "test")))
	cols = ["eid", "target", "split"] + traits + covars + (["group"] if "group" in d else [])
	return d[cols]


def scan_variants(files, allowed, limit):
	import pgenlib
	records = []
	seen = set()
	coordinate_mismatches = 0
	for c, _, vp, _, _ in files:
		with pgenlib.PvarReader(os.fsencode(vp)) as pv:
			for j in range(pv.get_variant_ct()):
				snp = decode(pv.get_variant_id(j))
				if allowed is not None and snp not in allowed:
					continue
				if pv.get_allele_ct(j) != 2:
					continue
				chrom = decode(pv.get_variant_chrom(j)).removeprefix("chr")
				if chrom != str(c):
					raise ValueError(f"chr{c} input contains selected SNP on chromosome {chrom}")
				bp = int(pv.get_variant_pos(j))
				if allowed is not None and allowed[snp] is not None and allowed[snp] != (c, bp):
					coordinate_mismatches += 1
					continue
				ref, alt = decode(pv.get_allele_code(j, 0)), decode(pv.get_allele_code(j, 1))
				if snp in {".", "", "NA"} or not ref or not alt or ref == alt:
					continue
				if snp in seen:
					raise ValueError(f"Duplicate selected SNP ID: {snp}")
				seen.add(snp)
				records.append((c, bp, snp, ref, alt, j))
	if coordinate_mismatches:
		raise ValueError(f"{coordinate_mismatches} SNP-list/target CHR:BP mismatches; harmonize genome build first")
	d = pd.DataFrame(records, columns = ["CHR", "BP", "SNP", "REF", "ALT", "source_index"])
	d = d.sort_values(["CHR", "BP", "SNP"], kind = "stable").reset_index(drop = True)
	if limit and len(d) > limit:
		# Evenly spread an explicitly requested smoke subset across the genome.
		d = d.iloc[np.linspace(0, len(d) - 1, limit, dtype = int)].reset_index(drop = True)
	if len(d) == 0:
		raise ValueError("No selected biallelic autosomal variants")
	return d


def sha256(path):
	h = hashlib.sha256()
	with open(path, "rb") as stream:
		for block in iter(lambda: stream.read(1024 * 1024), b""):
			h.update(block)
	return h.hexdigest()


def identity(args, files):
	options = {k: v for k, v in vars(args).items() if k not in {"command", "func", "check", "replace", "cache_dir"}}
	paths = [args.pheno_file, args.ancestry_file, args.snp_list, args.keep, args.remove, args.split_file, args.split_group_file]
	paths += [str(p) for row in files for p in row[1:4]]
	stats = []
	for value in paths:
		# An ancestry file is optional only if its column is already in the phenotype.
		if value == args.ancestry_file and not Path(value).expanduser().is_file():
			stats.append({"path": value, "exists": False})
			continue
		p = optional(value)
		if p is not None:
			stat = p.stat()
			stats.append({"path": str(p), "size": stat.st_size, "mtime_ns": stat.st_mtime_ns})
	x = {"format_version": 1, "options": options, "inputs": stats}
	return hashlib.sha256(json.dumps(x, sort_keys = True).encode()).hexdigest(), x


def validate_cache(cache, signature):
	p = cache / "prepare.json"
	if not p.is_file():
		return False
	meta = json.loads(p.read_text())
	if meta.get("signature") != signature:
		return False
	for name in ("data.tsv.gz", "variants.tsv.gz"):
		if not (cache / name).is_file() or sha256(cache / name) != meta.get("sha256", {}).get(name):
			raise ValueError(f"Prepared cache is incomplete or changed: {cache / name}; use a new cache or --replace")
	a = np.load(cache / "genotypes.npy", mmap_mode = "r", allow_pickle = False)
	if list(a.shape) != meta["shape"] or a.dtype != np.float16:
		raise ValueError("Prepared genotype shape/type mismatch")
	return True


def prepare(args):
	files = genotype_files(args.dir_gen, chromosomes(args.chrs))
	sig, meta = identity(args, files)
	cache = Path(args.cache_dir).expanduser().resolve()
	if not args.replace and validate_cache(cache, sig):
		log(f"REUSE verified prepared cache: {cache}")
		return
	if (cache / "prepare.json").exists() and not args.replace:
		raise ValueError("Cache input/settings mismatch; choose another --cache-dir or explicitly --replace")
	d = make_cohort(args, files)
	v = scan_variants(files, variant_filter(args.snp_list), args.max_variants)
	n, m = len(d), len(v)
	peak_bytes = 4 * n * m + 512 * 1024 ** 2
	log(f"PREPARE: {n:,} individuals x {m:,} candidate SNPs; ALT dosage, missing=-1, float16")
	log(f"DISK: upper bound {peak_bytes / 1024 ** 3:.2f} GiB for two genotype layouts plus reserve; train-only MAF>={args.maf}, call rate>={args.call_rate}")
	if args.max_samples or args.max_variants:
		log("SMOKE SUBSET requested explicitly; results are not a full-cohort/full-genome analysis.")
	if args.check:
		log("CHECK OK: cohort, labels, splits and variant metadata checked. Genotype content QC occurs during prepare.")
		return
	cache.mkdir(parents = True, exist_ok = True)
	if shutil.disk_usage(cache).free < peak_bytes:
		raise ValueError("Insufficient cache disk space; use --cache-dir on a large local SSD")
	import pgenlib
	train = (d["split"] == "train").to_numpy()
	kept = []
	with tempfile.TemporaryDirectory(prefix = "prepare-", dir = cache) as td:
		work = Path(td)
		# Sequential writes in variant-major order avoid tiny strided disk writes.
		raw_path = work / "variant_major.bin"
		raw = np.memmap(raw_path, mode = "w+", dtype = np.float16, shape = (m, n))
		write_row = 0
		last_progress = time.monotonic()
		for c, gp, vp, sp, mode in files:
			vc = v[v.CHR == c]
			if vc.empty:
				continue
			ids = sample_ids(sp, mode)
			lookup = {eid: i for i, eid in enumerate(ids)}
			original = np.asarray([lookup[eid] for eid in d.eid], dtype = np.uint32)
			subset = np.sort(original)
			reorder = np.searchsorted(subset, original)
			with pgenlib.PvarReader(os.fsencode(vp)) as pv:
				kwargs = {"raw_sample_ct": len(ids), "sample_subset": subset, "pvar": pv}
				with pgenlib.PgenReader(os.fsencode(gp), **kwargs) as pg:
					for start in range(0, len(vc), args.chunk_variants):
						block = vc.iloc[start:start + args.chunk_variants]
						buf = np.empty((len(block), n), dtype = np.float32, order = "C")
						pg.read_dosages_list(np.asarray(block.source_index, dtype = np.uint32), buf, allele_idx = 1)
						buf = np.ascontiguousarray(buf[:, reorder])
						missing = (buf == -9) | ~np.isfinite(buf)
						if np.any((~missing) & ((buf < 0) | (buf > 2))):
							raise ValueError(f"chr{c}: dosage outside [0,2] with unrecognized missing encoding")
						tr = buf[:, train]
						ok = ~missing[:, train]
						calls = ok.sum(axis = 1)
						freq = np.divide(np.where(ok, tr, 0).sum(axis = 1, dtype = np.float64), 2 * calls, out = np.zeros(len(block)), where = calls > 0)
						maf = np.minimum(freq, 1 - freq)
						call_rate = calls / train.sum()
						use = (maf >= args.maf) & (call_rate >= args.call_rate) & (calls > 0) & (maf > 0)
						buf[missing] = -1
						if use.any():
							count = int(use.sum())
							raw[write_row:write_row + count] = buf[use]
							selected = block.loc[use].copy()
							selected["MAF_train"] = maf[use]
							selected["call_rate_train"] = call_rate[use]
							kept.append(selected)
							write_row += count
						if time.monotonic() - last_progress >= 60:
							log(f"GENOTYPES chr{c}: checked {min(start + args.chunk_variants, len(vc)):,}/{len(vc):,}; cumulative {write_row:,} SNPs retained")
							last_progress = time.monotonic()
			log(f"GENOTYPES chr{c}: cumulative {write_row:,} SNPs retained")
		if write_row == 0:
			raise ValueError("No SNP passed train-only MAF/call-rate QC")
		raw.flush()
		output = np.lib.format.open_memmap(work / "genotypes.npy", mode = "w+", dtype = np.float16, shape = (n, write_row))
		# Bound the transpose work buffer at approximately 64 MiB.
		batch = max(1, min(1024, (64 * 1024 ** 2) // max(2 * write_row, 1)))
		log(f"TRANSPOSE: writing {n:,} sample-major genotype rows")
		for start in range(0, n, batch):
			output[start:start + batch] = raw[:write_row, start:start + batch].T
			if time.monotonic() - last_progress >= 60:
				log(f"TRANSPOSE: {min(start + batch, n):,}/{n:,} rows")
				last_progress = time.monotonic()
		output.flush()
		del output, raw
		raw_path.unlink()
		final_v = pd.concat(kept, ignore_index = True).drop(columns = "source_index")
		d.to_csv(work / "data.tsv.gz", sep = "\t", index = False, na_rep = "NA")
		final_v.to_csv(work / "variants.tsv.gz", sep = "\t", index = False)
		meta.update({"signature": sig, "shape": [n, write_row], "dosage": "ALT, float16; missing=-1", "split_counts": d["split"].value_counts().to_dict(), "sha256": {name: sha256(work / name) for name in ("data.tsv.gz", "variants.tsv.gz")}})
		(work / "prepare.json").write_text(json.dumps(meta, indent = 2) + "\n")
		# The final manifest is installed last; interrupted preparation is never reusable.
		(cache / "prepare.json").unlink(missing_ok = True)
		for name in ("genotypes.npy", "data.tsv.gz", "variants.tsv.gz", "prepare.json"):
			os.replace(work / name, cache / name)
	log(f"PREPARE OK: {n:,} individuals x {write_row:,} SNPs -> {cache}")


def write_rds(frame, path):
	import pyreadr
	p = Path(path)
	p.parent.mkdir(parents = True, exist_ok = True)
	data = frame.copy()
	for col in data:
		if data[col].dtype.name in {"string", "category"}:
			data[col] = data[col].astype(object)
	with tempfile.NamedTemporaryFile(prefix = p.stem + ".", suffix = ".rds", dir = p.parent, delete = False) as stream:
		tmp = Path(stream.name)
	try:
		pyreadr.write_rds(str(tmp), data, compress = "gzip")
		os.replace(tmp, p)
	finally:
		tmp.unlink(missing_ok = True)


def workbook(frame, path, title):
	from openpyxl import Workbook
	from openpyxl.styles import Font, PatternFill
	from openpyxl.utils import get_column_letter
	wb = Workbook()
	ws = wb.active
	ws.title = title[:31]
	ws.append(list(frame.columns))
	for row in frame.itertuples(index = False, name = None):
		ws.append([None if pd.isna(x) else x.item() if isinstance(x, np.generic) else x for x in row])
	for cell in ws[1]:
		cell.font = Font(bold = True, color = "FFFFFF")
		cell.fill = PatternFill("solid", fgColor = "234D65")
	ws.freeze_panes = "A2"
	ws.auto_filter.ref = ws.dimensions
	for i, col in enumerate(frame.columns, 1):
		ws.column_dimensions[get_column_letter(i)].width = min(34, max(13, len(str(col)) + 2))
	wb.save(path)


def attach_covariate_values(scores, cohort, covariates):
	"""Retain original outcome-model covariates for direct benchmark verification."""
	names = [] if covariates.strip().lower() in ("", "none") else [c.strip() for c in covariates.split(",")]
	if len(names) != len(set(names)) or any(name not in cohort for name in names):
		raise ValueError("Published covariates must be unique and present in the prepared cohort")
	if not names:
		return scores.copy()
	source = cohort[["eid"] + names].copy()
	for name in names:
		source[name] = numeric(source[name], name)
		if source[name].isna().any():
			raise ValueError(f"Prepared model covariate {name} is missing or nonfinite")
	source = source.rename(columns={name: f"covariate.{name}" for name in names})
	out = scores.merge(source, on="eid", how="left", validate="one_to_one", indicator=True, sort=False)
	if not out["_merge"].eq("both").all():
		raise ValueError("A published prediction has no matching original covariates")
	return out.drop(columns="_merge")


def canonical_metrics(metrics):
	"""Upgrade old metric names by their original formulas; never refit a model."""
	metrics = metrics.copy()
	metrics.loc[metrics.metric == "prediction_R2", "metric"] = "residual_correlation_R2"
	keys = ["trait", "target", "split", "metric"]
	if metrics.duplicated(keys).any():
		raise ValueError("Duplicate metric names after explicit legacy-label migration")
	derived = metrics[metrics.metric == "full_R2"].copy()
	derived["metric"] = "predictive_R2"
	present = set(map(tuple, metrics[keys].to_numpy()))
	derived = derived[[tuple(row) not in present for row in derived[keys].to_numpy()]]
	return pd.concat([metrics, derived], ignore_index=True)


def publish(args):
	run = Path(args.run_dir).expanduser().resolve()
	cache = Path(args.cache_dir).expanduser().resolve()
	for name in ("model.pt", "metrics.tsv", "history.tsv", "test_predictions.tsv.gz"):
		require(run / name)
	meta = json.loads(require(cache / "prepare.json").read_text())
	if not validate_cache(cache, meta["signature"]):
		raise ValueError("Cannot validate cache before publication")
	evaluation = json.loads(require(run / "evaluation.json").read_text())
	if evaluation.get("format_version") != 1 or evaluation.get("preparation_signature") != meta["signature"]:
		raise ValueError("Evaluation was generated from a different prepared cache")
	for name, key in (("data.tsv.gz", "data_sha256"), ("variants.tsv.gz", "variants_sha256")):
		if evaluation["inputs"].get(key) != sha256(cache / name):
			raise ValueError(f"Evaluation input mismatch: {name}")
	gen_stat = (cache / "genotypes.npy").stat()
	gen_ident = evaluation["inputs"]["genotypes"]
	if gen_ident.get("size") != gen_stat.st_size or gen_ident.get("mtime_ns") != gen_stat.st_mtime_ns or list(gen_ident.get("shape", [])) != meta["shape"]:
		raise ValueError("Evaluation genotype file was replaced or modified")
	for name in ("metrics.tsv", "history.tsv", "test_predictions.tsv.gz"):
		if evaluation["outputs_sha256"].get(name) != sha256(run / name):
			raise ValueError(f"Evaluation output was changed after completion: {name}")
	model_stat = (run / "model.pt").stat()
	if evaluation["model"].get("size") != model_stat.st_size or evaluation["model"].get("mtime_ns") != model_stat.st_mtime_ns:
		raise ValueError("Evaluation checkpoint was replaced or modified")
	pred = read_table(run / "test_predictions.tsv.gz")
	pred["eid"] = normalize_ids(pred.eid)
	needed = {"target", "split", "trait", "y", "baseline", "prediction", "genetic_score"}
	if not needed.issubset(pred.columns) or not (pred["split"] == "test").all():
		raise ValueError("Only held-out test predictions can be published")
	if pred.duplicated(["eid", "trait"]).any():
		raise ValueError("Duplicate test prediction ID/trait")
	cohort = identified(read_table(cache / "data.tsv.gz"), "Prepared cohort")
	test_ids = set(cohort.loc[cohort["split"] == "test", "eid"])
	if not set(pred.eid).issubset(test_ids):
		raise ValueError("Prediction includes a participant outside prepared test split")
	traits = args.traits.split(",")
	if set(pred.trait) != set(traits):
		raise ValueError("Requested traits do not match prediction output")
	for trait in traits:
		if set(pred.loc[pred.trait == trait, "eid"]) != test_ids:
			raise ValueError(f"Incomplete test predictions: {trait}")
	metrics = canonical_metrics(pd.read_csv(run / "metrics.tsv", sep = "\t"))
	history = pd.read_csv(run / "history.tsv", sep = "\t")
	if not (metrics["split"] == "test").all():
		raise ValueError("Performance table contains non-test rows")
	root = Path(args.output_root).expanduser().resolve() / "prsformer"
	score_root = Path(args.score_dir).expanduser().resolve()
	destinations = [root / x for x in ("3.prsformer.pt", "3.prsformer.split.rds", "3.prsformer.performance.png", "3.prsformer.performance.xlsx", "3.prsformer.training.png", "3.prsformer.training.xlsx")]
	destinations += [score_root / t / "3.prsformer.scores.rds" for t in traits]
	if not args.replace and any(p.exists() for p in destinations):
		raise ValueError("Formal PRSformer outputs already exist; select another output root/score dir or explicitly --replace")
	root.mkdir(parents = True, exist_ok = True)
	import matplotlib
	matplotlib.use("Agg")
	import matplotlib.pyplot as plt
	plt.rcParams.update({"font.size": 10, "axes.spines.top": False, "axes.spines.right": False, "savefig.dpi": 160})
	with tempfile.TemporaryDirectory(prefix = "publish-", dir = root) as td:
		stage = Path(td)
		# Primary metric above, all remaining exact metrics in the table below.
		fig, axs = plt.subplots(2, len(traits), figsize = (6.4 * len(traits), 11), gridspec_kw = {"height_ratios": [1, 2.2]}, squeeze = False)
		for j, trait in enumerate(traits):
			metric = "AUC" if trait == "t2dm" else "predictive_R2"
			sub = metrics[(metrics.trait == trait) & (metrics.metric == metric)]
			ax = axs[0, j]
			ax.bar(sub.target.astype(str), sub.value, color = "#347F9C")
			ax.set_title(f"{trait}: {metric}", loc = "left")
			ax.tick_params(axis = "x", labelrotation = 30)
			ax.axhline(0, color = "#777777", linewidth = .5)
			detail = metrics[metrics.trait == trait].pivot(index = "metric", columns = "target", values = "value")
			counts = metrics[metrics.trait == trait].groupby("target").n.first()
			detail.loc["n"] = counts
			labels = [[f"{x:.5g}" if pd.notna(x) else "NA" for x in row] for row in detail.to_numpy()]
			ax = axs[1, j]
			ax.axis("off")
			table = ax.table(cellText = labels, rowLabels = list(detail.index), colLabels = list(detail.columns), cellLoc = "right", loc = "upper center", bbox = [.28, .04, .72, .92])
			table.auto_set_font_size(False)
			table.set_fontsize(7)
		fig.suptitle("PRSformer: fixed held-out test set; point estimates", y = 1)
		fig.tight_layout()
		fig.savefig(stage / "3.prsformer.performance.png", bbox_inches = "tight")
		plt.close(fig)
		workbook(metrics, stage / "3.prsformer.performance.xlsx", "performance")
		fig, axs = plt.subplots(1, len(traits), figsize = (5 * len(traits), 3.6), squeeze = False)
		for ax, trait in zip(axs[0], traits):
			for split, color in (("train", "#347F9C"), ("validation", "#D18443")):
				sub = history[(history.trait == trait) & (history["split"] == split)]
				ax.plot(sub.epoch, sub.loss, label = split, color = color)
			ax.set(title = trait, xlabel = "Epoch", ylabel = "Standardized residual MSE" if trait != "t2dm" else "Binary cross entropy")
			ax.legend(frameon = False)
		fig.tight_layout()
		fig.savefig(stage / "3.prsformer.training.png", bbox_inches = "tight")
		plt.close(fig)
		workbook(history, stage / "3.prsformer.training.xlsx", "training")
		shutil.copyfile(run / "model.pt", stage / "3.prsformer.pt")
		write_rds(cohort[[c for c in ("eid", "target", "split", "group") if c in cohort]], stage / "3.prsformer.split.rds")
		for trait in traits:
			x = pred[pred.trait == trait].drop(columns = "trait").copy()
			for col in ("y", "baseline", "prediction", "genetic_score"):
				x[col] = pd.to_numeric(x[col], errors = "raise")
			x = x.rename(columns = {"y": "outcome", "genetic_score": "prsformer"})
			x = attach_covariate_values(x, cohort, meta["options"]["covariates"])
			x["prsformer_scale"] = "log_odds_increment" if trait == "t2dm" else "outcome_units"
			x["covariate_names"] = meta["options"]["covariates"]
			x["endpoint_definition"] = (
				"baseline_t2dm_from_Yr2e_Yt2e" if trait == "t2dm" and meta["options"]["t2dm_col"] == "auto"
				else meta["options"][trait + "_col"]
			)
			x["preparation_signature"] = meta["signature"]
			write_rds(x, stage / f"{trait}.scores.rds")
		for path in stage.iterdir():
			if path.name.endswith(".scores.rds"):
				trait = path.name.removesuffix(".scores.rds")
				dest = score_root / trait / "3.prsformer.scores.rds"
			else:
				dest = root / path.name
			dest.parent.mkdir(parents = True, exist_ok = True)
			# Cross-filesystem safe, per-file atomic publication.
			with tempfile.NamedTemporaryFile(prefix = dest.stem + ".", dir = dest.parent, delete = False) as stream:
				tmp = Path(stream.name)
			try:
				shutil.copyfile(path, tmp)
				os.replace(tmp, dest)
			finally:
				tmp.unlink(missing_ok = True)
	log(f"PUBLISH OK: model and aggregate figures/workbooks at {root}; test-only RDS under {score_root}/<trait>/")


def csv_list(value):
	return [item.strip() for item in (value or "").split(",") if item.strip()]


def digest_json(value):
	return hashlib.sha256(json.dumps(value, sort_keys=True, allow_nan=False).encode()).hexdigest()


def file_sha256(path):
	hash_value = hashlib.sha256()
	with open(path, "rb") as handle:
		for block in iter(lambda: handle.read(8 * 1024 ** 2), b""):
			hash_value.update(block)
	return hash_value.hexdigest()


def validate_genotypes(values):
	if not np.isfinite(values).all() or not np.all((values == -1) | ((values >= 0) & (values <= 2))):
		raise ValueError("Genotypes must contain dosage 0..2 or the missing value -1; do not standardize SNP dosages.")


def identifier_hashes(values):
	"""Canonical UKB IDs; hashes serve as membership keys, not anonymization."""
	values = pd.Series(values, dtype="string").str.strip().str.replace(r"^([+-]?\d+)\.0+$", r"\1", regex=True)
	if values.isna().any() or values.eq("").any():
		raise ValueError("Missing individual/family identifier in split provenance.")
	return {hashlib.sha256(value.encode()).hexdigest() for value in values}


def validate_family_splits(data):
	for name in ("group", "family_id"):
		if name in data:
			identifier_hashes(data[name])
			canonical = data[name].astype("string").str.strip().str.replace(r"^([+-]?\d+)\.0+$", r"\1", regex=True)
			if data.groupby(canonical)["split"].nunique().max() > 1:
				raise ValueError(f"Related-family column {name} spans train/validation/test partitions.")


def split_registry(data):
	development = data[data["split"].isin(["train", "validation"])]
	return {
		"format_version": 1,
		"development_ids": sorted(identifier_hashes(development["eid"])),
		"development_families": {
			name: sorted(identifier_hashes(development[name]))
			for name in ("group", "family_id") if name in development
		},
	}


def validate_prediction_holdout(data, identity, checkpoint):
	"""A changed split label cannot turn a development participant into a test."""
	registry = checkpoint.get("split_registry")
	if registry is None:
		if identity.get("data_sha256") != checkpoint.get("training_identity", {}).get("data_sha256"):
			raise ValueError("Legacy checkpoint has no development-ID registry. Reuse its exact original cohort/split or retrain before predicting on a different cohort.")
		return
	if registry.get("format_version") != 1:
		raise ValueError("Unrecognized checkpoint split-registry version.")
	test = data[data["split"] == "test"]
	if identifier_hashes(test["eid"]) & set(registry["development_ids"]):
		raise ValueError("Requested test individuals were used for model training or validation; they are not held out.")
	for name, recorded in registry.get("development_families", {}).items():
		if name not in test:
			raise ValueError(f"Prediction needs the original related-family column {name} to verify independence.")
		if identifier_hashes(test[name]) & set(recorded):
			raise ValueError("Requested test families overlap model training/validation families.")


def read_inputs(args, traits, require_training=True):
	input_paths = [Path(args.data), Path(args.variants), Path(args.genotypes)]
	initial_stamps = [(path.stat().st_size, path.stat().st_mtime_ns, path.stat().st_ino) for path in input_paths]
	data = pd.read_csv(args.data, sep="\t", dtype={"eid": str, "target": str, "split": str, "group": str, "family_id": str})
	required = ["eid", "target", "split"] + traits
	missing = [column for column in required if column not in data]
	if missing:
		raise ValueError(f"Missing columns in --data: {missing}")
	if data["eid"].isna().any() or data["eid"].duplicated().any():
		raise ValueError("--data must have one unique nonmissing eid per genotype row.")
	if len(identifier_hashes(data["eid"])) != len(data):
		raise ValueError("--data contains duplicate individual IDs after normalizing integral UKB identifiers.")
	if data["target"].isna().any() or data["split"].isna().any():
		raise ValueError("Every sample needs a target label and an explicit split.")
	allowed_splits = {"train", "validation", "test"}
	if not set(data["split"]).issubset(allowed_splits):
		raise ValueError("split must be exactly train, validation, or test.")
	if require_training and set(data["split"]) != allowed_splits:
		raise ValueError("Training requires nonempty train, validation, and test partitions.")
	if not np.any(data["split"] == "test"):
		raise ValueError("No held-out test samples were provided.")
	validate_family_splits(data)
	genotypes = np.load(args.genotypes, mmap_mode="r", allow_pickle=False)
	if genotypes.ndim != 2 or genotypes.shape[0] != len(data):
		raise ValueError("The sample-major genotype matrix must have exactly one row per row of --data, in the same order.")
	if not genotypes.flags.c_contiguous:
		raise ValueError("--genotypes must be a C-contiguous sample-major .npy matrix.")
	if genotypes.dtype not in (np.dtype("int8"), np.dtype("float16"), np.dtype("float32")):
		raise ValueError("Supported genotype dtypes are int8, float16 and float32.")
	variants = pd.read_csv(args.variants, sep="\t", dtype=str)
	variant_columns = ["CHR", "BP", "SNP", "REF", "ALT"]
	if any(column not in variants for column in variant_columns):
		raise ValueError("--variants needs CHR, BP, SNP, REF, ALT columns in genotype-column order.")
	if len(variants) != genotypes.shape[1] or not len(variants):
		raise ValueError("Variant rows must exactly match the genotype columns.")
	if variants[variant_columns].isna().any().any() or variants["SNP"].duplicated().any():
		raise ValueError("Variant metadata must be complete with unique SNP IDs.")
	chromosomes = pd.to_numeric(variants["CHR"].str.replace(r"^chr", "", regex=True, case=False), errors="raise").to_numpy(float)
	positions = pd.to_numeric(variants["BP"], errors="raise").to_numpy(float)
	if not np.all(np.isfinite(chromosomes) & (chromosomes == np.floor(chromosomes)) & (chromosomes >= 1) & (chromosomes <= 22)):
		raise ValueError("PRSformer inputs must use autosomes 1..22.")
	if not np.all(np.isfinite(positions) & (positions == np.floor(positions)) & (positions > 0)):
		raise ValueError("Variant BP must be a positive integer.")
	ordered = (chromosomes[1:] > chromosomes[:-1]) | ((chromosomes[1:] == chromosomes[:-1]) & (positions[1:] >= positions[:-1]))
	if not ordered.all():
		raise ValueError("Input columns must be ordered by numeric chromosome and then base-pair position.")
	if not variants["REF"].str.upper().isin(list("ACGT")).all() or not variants["ALT"].str.upper().isin(list("ACGT")).all():
		raise ValueError("Only biallelic A/C/G/T SNPs are supported by this preparation workflow.")
	if (variants["REF"].str.upper() == variants["ALT"].str.upper()).any():
		raise ValueError("REF and ALT must differ.")
	labels = data[traits].apply(pd.to_numeric, errors="raise").to_numpy(dtype=np.float64)
	if np.isinf(labels).any():
		raise ValueError("Phenotypes must be finite or missing, never +/- infinity.")
	for column, trait in enumerate(traits):
		observed = labels[np.isfinite(labels[:, column]), column]
		if SUPPORTED_TRAITS[trait] == "binary" and not np.isin(observed, [0, 1]).all():
			raise ValueError("t2dm must be explicitly defined baseline case/control status coded 0/1/NA. Incident events are not substituted.")
		if require_training:
			for split in ("train", "validation"):
				values = labels[(data["split"].to_numpy() == split) & np.isfinite(labels[:, column]), column]
				if len(values) < 2:
					raise ValueError(f"{trait} needs at least two observed phenotypes in {split}.")
				if SUPPORTED_TRAITS[trait] == "binary" and len(np.unique(values)) < 2:
					raise ValueError(f"t2dm requires both cases and controls in {split}.")
	# This is an inexpensive format check; each actual batch is checked again.
	sample_rows = np.unique(np.linspace(0, len(data) - 1, min(24, len(data)), dtype=int))
	for row in sample_rows:
		validate_genotypes(genotypes[row])
	variant_identity = variants[variant_columns].to_csv(sep="\t", index=False)
	variant_digest = hashlib.sha256(variant_identity.encode()).hexdigest()
	data_digest = hashlib.sha256(data.to_csv(sep="\t", index=False).encode()).hexdigest()
	genotype_stat = Path(args.genotypes).stat()
	identity = {
		"data_sha256": data_digest, "variants_sha256": variant_digest,
		"data_file_sha256": file_sha256(args.data), "variants_file_sha256": file_sha256(args.variants),
		"genotypes_size": int(genotype_stat.st_size), "genotypes_mtime_ns": int(genotype_stat.st_mtime_ns),
		"shape": list(genotypes.shape), "genotypes_dtype": str(genotypes.dtype),
	}
	final_stamps = [(path.stat().st_size, path.stat().st_mtime_ns, path.stat().st_ino) for path in input_paths]
	if initial_stamps != final_stamps:
		raise RuntimeError("An input file changed while it was being read; retry with a stable prepared cache.")
	return data, labels, identity, variants[variant_columns]


def covariate_design(data, columns, training_rows=None, transform=None):
	missing = [column for column in columns if column not in data]
	if missing:
		raise ValueError(f"Covariate columns not found: {missing}")
	values = data[columns].apply(pd.to_numeric, errors="raise").to_numpy(dtype=np.float64)
	if np.isinf(values).any():
		raise ValueError("Covariates must be finite or missing, not infinity.")
	if transform is None:
		means, scales = [], []
		for column, name in enumerate(columns):
			observed = values[training_rows, column]
			observed = observed[np.isfinite(observed)]
			if not len(observed):
				raise ValueError(f"No observed training values for covariate {name}.")
			means.append(float(observed.mean()))
			scale = float(observed.std())
			scales.append(scale if scale > 1e-12 else 1.0)
		transform = {"columns": columns, "means": means, "scales": scales}
	means, scales = np.asarray(transform["means"]), np.asarray(transform["scales"])
	if columns:
		values = np.where(np.isfinite(values), values, means)
		values = (values - means) / scales
	design = np.column_stack([np.ones(len(data)), values])
	return design, transform


def fit_binary_baseline(design, phenotype, ridge=1e-4):
	# A small fixed penalty stabilizes covariate-only fits; the intercept is unpenalized.
	initial = np.zeros(design.shape[1])
	prevalence = float(phenotype.mean())
	initial[0] = math.log(prevalence / (1 - prevalence))

	def objective(coefficients):
		logits = design @ coefficients
		penalty = coefficients.copy()
		penalty[0] = 0
		loss = np.mean(np.logaddexp(0, logits) - phenotype * logits) + ridge * np.sum(penalty ** 2) / 2
		gradient = design.T @ (expit(logits) - phenotype) / len(phenotype) + ridge * penalty
		return float(loss), gradient

	result = minimize(objective, initial, jac=True, method="L-BFGS-B", options={"maxiter": 1000, "ftol": 1e-12, "gtol": 1e-7})
	if not result.success or not np.isfinite(result.x).all():
		raise ValueError(f"Covariate-only logistic regression did not converge: {result.message}")
	return result.x


def fit_baselines(data, labels, traits, columns):
	training_rows = np.flatnonzero(data["split"].to_numpy() == "train")
	design, transform = covariate_design(data, columns, training_rows=training_rows)
	baselines, baseline_logits = np.zeros_like(labels), np.zeros_like(labels)
	prepared_labels = np.full_like(labels, np.nan)
	task_parameters = {}
	for column, trait in enumerate(traits):
		fit_rows = training_rows[np.isfinite(labels[training_rows, column])]
		phenotype = labels[fit_rows, column]
		if SUPPORTED_TRAITS[trait] == "continuous":
			coefficients = np.linalg.lstsq(design[fit_rows], phenotype, rcond=None)[0]
			baseline = design @ coefficients
			scale = float(np.std(phenotype - baseline[fit_rows], ddof=0))
			if not np.isfinite(scale) or scale < 1e-12:
				raise ValueError(f"{trait} has no residual training variance after covariate adjustment.")
			baselines[:, column] = baseline
			prepared_labels[:, column] = (labels[:, column] - baseline) / scale
		else:
			coefficients = fit_binary_baseline(design[fit_rows], phenotype)
			baseline_logits[:, column] = design @ coefficients
			baselines[:, column] = expit(baseline_logits[:, column])
			prepared_labels[:, column] = labels[:, column]
			scale = 1.0
		task_parameters[trait] = {
			"type": SUPPORTED_TRAITS[trait], "coefficients": coefficients.tolist(),
			"residual_scale": scale, "n_training": int(len(fit_rows)),
			"logistic_ridge": 1e-4 if SUPPORTED_TRAITS[trait] == "binary" else None,
		}
	return {"covariates": transform, "tasks": task_parameters}, prepared_labels, baselines, baseline_logits


def apply_baselines(data, traits, preprocessing):
	transform = preprocessing["covariates"]
	design, _ = covariate_design(data, transform["columns"], transform=transform)
	baselines = np.zeros((len(data), len(traits)), dtype=np.float64)
	baseline_logits = np.zeros_like(baselines)
	for column, trait in enumerate(traits):
		linear_prediction = design @ np.asarray(preprocessing["tasks"][trait]["coefficients"])
		if SUPPORTED_TRAITS[trait] == "continuous":
			baselines[:, column] = linear_prediction
		else:
			baseline_logits[:, column] = linear_prediction
			baselines[:, column] = expit(linear_prediction)
	return baselines, baseline_logits


def neighborhood_runtime():
	"""Validate the two supported binary stacks; never select a dense fallback."""
	torch = import_torch()
	try:
		natten = importlib.import_module("natten")
	except ImportError as exc:
		raise RuntimeError("NATTEN is missing; run bash install_grid.sh.") from exc
	version = str(getattr(natten, "__version__", "")).split("+")[0]
	stack = (str(torch.__version__).split("+")[0], torch.version.cuda, version)
	if stack == ("2.6.0", "12.4", "0.17.5"):
		if not all(hasattr(natten, key) for key in ("NeighborhoodAttention1D", "use_fused_na", "is_fna_enabled")):
			raise RuntimeError("Legacy NATTEN lacks the upstream fused-attention API.")
		if torch.cuda.is_available() and torch.cuda.get_device_capability()[0] >= 10:
			raise RuntimeError("Blackwell GPUs require the modern CUDA 13.0 stack; run bash install_grid.sh.")
	elif stack == ("2.12.0", "13.0", "0.21.7"):
		if not getattr(natten, "HAS_LIBNATTEN", False) or not hasattr(natten, "na1d"):
			raise RuntimeError("NATTEN's compiled extension is missing; install the matching torch2120cu130 wheel.")
	else:
		raise RuntimeError(f"Unsupported torch/CUDA/NATTEN stack {stack}; run install_grid.sh for the supported environment.")
	return natten, version


def modern_natten_compatibility(natten):
	"""Map PRSformer's old constructor to 0.21.7 with explicit CUTLASS FNA.

	The qkv/proj modules and state-dict keys are inherited unchanged. Only the
	removed zero-dropout/bias options and fused-backend controls are translated.
	"""
	class NeighborhoodAttention1D(natten.NeighborhoodAttention1D):
		def __init__(self, dim, num_heads, kernel_size, dilation=1, is_causal=False,
				rel_pos_bias=False, qkv_bias=True, qk_scale=None, attn_drop=0., proj_drop=0.):
			if rel_pos_bias or attn_drop != 0.:
				raise ValueError("The modern adapter requires rel_pos_bias=False and attn_drop=0, as in official PRSformer.")
			super().__init__(embed_dim=dim, num_heads=num_heads, kernel_size=kernel_size,
				dilation=dilation, is_causal=is_causal, qkv_bias=qkv_bias,
				qk_scale=qk_scale, proj_drop=proj_drop)

		def forward(self, x):
			batch, length, channels = x.shape
			q, k, v = self.qkv(x).reshape(batch, length, 3, self.num_heads, self.head_dim).unbind(2)
			attended = natten.na1d(q, k, v, kernel_size=self.kernel_size,
				dilation=self.dilation, is_causal=self.is_causal, scale=self.scale,
				backend="cutlass-fna")
			return self.proj_drop(self.proj(attended.reshape(batch, length, channels)))

	def require_fused(enabled):
		if enabled is not True:
			raise ValueError("The PRSformer compatibility adapter always requires CUTLASS fused neighborhood attention.")

	return types.SimpleNamespace(NeighborhoodAttention1D=NeighborhoodAttention1D,
		use_fused_na=require_fused, is_fna_enabled=lambda: bool(natten.HAS_LIBNATTEN))


def load_official_model(upstream_dir, attention):
	upstream_dir = Path(upstream_dir).resolve()
	source_directory = upstream_dir / "src"
	paths = {name: source_directory / f"{name}.py" for name in ("utils", "modules", "model")}
	if not all(path.is_file() for path in paths.values()):
		raise ValueError("--upstream-dir must contain the official src/model.py, src/modules.py, and src/utils.py.")
	compatibility = []
	natten_adapter = None
	if attention == "neighborhood":
		natten, version = neighborhood_runtime()
		if version == "0.21.7":
			natten_adapter = modern_natten_compatibility(natten)
			compatibility.append("NATTEN 0.21.7: translate dim/zero dropout/relative-bias options; explicitly use cutlass-fna with unchanged qkv/proj parameters.")
	previous_modules = {name: sys.modules.get(name) for name in paths}
	loaded = {}
	try:
		for name, path in paths.items():
			tree = ast.parse(path.read_text(), filename=str(path))
			if name == "utils":
				lightning_imports = [node for node in tree.body if isinstance(node, ast.Import) and any(alias.name == "pytorch_lightning" for alias in node.names)]
				if lightning_imports:
					aliases = {alias.asname or alias.name for node in lightning_imports for alias in node.names if alias.name == "pytorch_lightning"}
					if any(isinstance(node, ast.Name) and isinstance(node.ctx, ast.Load) and node.id in aliases for node in ast.walk(tree)):
						raise RuntimeError("Upstream now uses pytorch_lightning actively; review this import adapter before using that revision.")
					for node in lightning_imports:
						node.names = [alias for alias in node.names if alias.name != "pytorch_lightning"]
					tree.body = [node for node in tree.body if not isinstance(node, ast.Import) or node.names]
					compatibility.append("Skipped the unused pytorch_lightning import in utils.py; all functions are unchanged.")
			if name == "modules" and (attention == "global" or natten_adapter is not None):
				# These names occur only inside the unused neighborhood-attention class.
				tree.body = [node for node in tree.body if not (
					(isinstance(node, ast.ImportFrom) and node.module == "natten") or
					(isinstance(node, ast.Import) and any(alias.name == "natten" for alias in node.names))
				)]
				if attention == "global":
					compatibility.append("Explicit small global-attention mode skips NATTEN imports and uses the upstream PyTorch attention branch.")
			module = types.ModuleType(f"grid_prsformer_official_{name}")
			if name == "modules" and natten_adapter is not None:
				module.__dict__.update(natten=natten_adapter,
					NeighborhoodAttention1D=natten_adapter.NeighborhoodAttention1D,
					is_fna_enabled=natten_adapter.is_fna_enabled)
			module.__file__ = str(path)
			module.__package__ = ""
			sys.modules[name] = module
			sys.modules[module.__name__] = module
			exec(compile(ast.fix_missing_locations(tree), str(path), "exec"), module.__dict__)
			loaded[name] = module
	finally:
		for name, original in previous_modules.items():
			if original is None:
				sys.modules.pop(name, None)
			else:
				sys.modules[name] = original
	model_class = getattr(loaded["model"], "g2p_transformer_ExplicitNaNDose2", None)
	if model_class is None:
		raise RuntimeError("The official g2p_transformer_ExplicitNaNDose2 model class was not found.")
	commit = subprocess.run(["git", "-C", str(upstream_dir), "rev-parse", "HEAD"], capture_output=True, text=True, check=False)
	metadata = {
		"repository": "https://github.com/23andMe/PRSformer",
		"class": "g2p_transformer_ExplicitNaNDose2",
		"commit": commit.stdout.strip() if commit.returncode == 0 else None,
		"source_sha256": {name: hashlib.sha256(path.read_bytes()).hexdigest() for name, path in paths.items()},
		"import_compatibility": compatibility,
	}
	return model_class, metadata


def architecture_from_args(args, length, tasks):
	dilation = [int(value) for value in csv_list(args.dilation)]
	if len(dilation) == 1:
		dilation *= args.layers
	if len(dilation) != args.layers or any(value < 1 for value in dilation):
		raise ValueError("--dilation needs one positive integer, or one positive integer per layer.")
	if min(length, tasks, args.embed_dim, args.heads, args.layers, args.ff_dim) < 1 or args.embed_dim % args.heads:
		raise ValueError("Model dimensions must be positive; --embed-dim must be divisible by --heads.")
	if args.attention == "global" and length > 4096:
		raise ValueError("Global attention is only an explicit small-data check (at most 4096 SNPs), not a genome-scale fallback.")
	if args.attention == "neighborhood" and (args.kernel_size < 2 or args.kernel_size * max(dilation) > length):
		raise ValueError("Neighborhood attention requires kernel_size >= 2 and kernel_size * max(dilation) <= the SNP count.")
	return {
		"seq_len": int(length), "embed_dim": args.embed_dim, "num_heads": args.heads,
		"dim_feedforward": args.ff_dim, "num_layers": args.layers, "num_covars": 0,
		"num_phenos": int(tasks), "kernel_size": args.kernel_size if args.attention == "neighborhood" else None,
		"dilation": dilation, "dlm_reprs": None, "weight_the_loss": False,
		"use_snp_annots": False, "snp_indices": None,
	}


def parameter_estimate(architecture):
	length, dimension = architecture["seq_len"], architecture["embed_dim"]
	feedforward, tasks = architecture["dim_feedforward"], architecture["num_phenos"]
	per_layer = 4 * dimension ** 2 + 2 * dimension * feedforward + 9 * dimension + feedforward
	return int(length * dimension * (2 + tasks) + architecture["num_layers"] * per_layer + 2 * dimension + tasks)


def import_torch():
	try:
		return importlib.import_module("torch")
	except ImportError as exc:
		raise RuntimeError("PyTorch is missing. Use the shared GRID environment installed by install_grid.sh.") from exc


def runtime(args, attention):
	torch = import_torch()
	torch.set_num_threads(args.threads)
	device = torch.device(args.device)
	if device.type not in ("cuda", "cpu"):
		raise ValueError("This adapter supports cuda or cpu devices.")
	if device.type == "cuda" and not torch.cuda.is_available():
		raise RuntimeError("CUDA is unavailable; genome-scale PRSformer training needs a compatible NVIDIA GPU.")
	if device.type == "cuda":
		if device.index is None:
			device = torch.device("cuda", torch.cuda.current_device())
		torch.cuda.set_device(device)
	if attention == "neighborhood" and device.type != "cuda":
		raise RuntimeError("The upstream fused neighborhood attention path requires CUDA. CPU is supported only for an explicit small global-attention check.")
	if args.amp == "bf16" and device.type == "cuda" and not torch.cuda.is_bf16_supported():
		raise RuntimeError("This GPU does not support BF16; select --amp fp16 or off.")
	if device.type == "cpu" and args.amp != "off":
		raise ValueError("For CPU small-data checks set --amp off explicitly.")
	dtype = {"fp16": torch.float16, "bf16": torch.bfloat16, "off": torch.float32}[args.amp]
	return torch, device, dtype


def checkpoint_layers(model, torch):
	from torch.utils.checkpoint import checkpoint
	for layer in model.transformer_encoder:
		original_forward = layer.forward

		def forward(value=None, *, x=None, _layer=layer, _forward=original_forward):
			inputs = x if x is not None else value
			if _layer.training and torch.is_grad_enabled():
				return checkpoint(_forward, inputs, use_reentrant=False)
			return _forward(inputs)

		layer.forward = forward


def model_smoke(model_class, architecture, torch, device, dtype, args):
	probe = dict(architecture)
	probe["seq_len"] = max(8, (architecture["kernel_size"] or 8) * max(architecture["dilation"]))
	model = model_class(**probe).to(device)
	if args.gradient_checkpointing:
		checkpoint_layers(model, torch)
	genotypes = torch.randint(0, 3, (2, probe["seq_len"]), device=device).float()
	genotypes[0, 0] = -1
	with torch.autocast(device_type=device.type, enabled=args.amp != "off", dtype=dtype):
		output = model(genotypes)
		loss = output.float().square().mean()
	if output.shape != (2, architecture["num_phenos"]) or not torch.isfinite(loss):
		raise RuntimeError("The official model failed its shape/finite-value preflight.")
	loss.backward()
	if not all(parameter.grad is None or torch.isfinite(parameter.grad).all() for parameter in model.parameters()):
		raise RuntimeError("The official model produced nonfinite gradients in the device preflight.")
	del model, genotypes, output, loss
	if device.type == "cuda":
		torch.cuda.empty_cache()


class GenotypeRows:
	def __init__(self, filename, rows):
		self.filename = str(filename)
		self.rows = np.asarray(rows, dtype=np.int64)
		self.matrix = None

	def __len__(self):
		return len(self.rows)

	def __getitem__(self, index):
		if self.matrix is None:
			self.matrix = np.load(self.filename, mmap_mode="r", allow_pickle=False)
		row = int(self.rows[index])
		values = np.array(self.matrix[row], dtype=np.float32, copy=True)
		validate_genotypes(values)
		return values, row

	def __getstate__(self):
		state = dict(self.__dict__)
		state["matrix"] = None
		return state


def make_loader(args, rows, torch, shuffle=False, seed=None):
	generator = torch.Generator()
	generator.manual_seed(args.seed if seed is None else seed)
	return torch.utils.data.DataLoader(
		GenotypeRows(args.genotypes, rows), batch_size=args.batch_size, shuffle=shuffle,
		num_workers=args.workers, pin_memory=args.device.startswith("cuda"),
		generator=generator, drop_last=False,
	)


def loss_matrix(output, row_indices, prepared_labels, baseline_logits, traits, torch, device):
	target_array = prepared_labels[row_indices]
	mask = torch.as_tensor(np.isfinite(target_array), device=device)
	target = torch.as_tensor(np.nan_to_num(target_array, nan=0.0), device=device, dtype=torch.float32)
	offsets = torch.as_tensor(baseline_logits[row_indices], device=device, dtype=torch.float32)
	columns = []
	for column, trait in enumerate(traits):
		if SUPPORTED_TRAITS[trait] == "continuous":
			loss = (output[:, column].float() - target[:, column]) ** 2
		else:
			loss = torch.nn.functional.binary_cross_entropy_with_logits(
				output[:, column].float() + offsets[:, column], target[:, column], reduction="none",
			)
		columns.append(loss)
	return torch.stack(columns, dim=1) * mask, mask


def evaluate_loss(model, rows, args, prepared_labels, baseline_logits, traits, torch, device, dtype):
	model.eval()
	totals, counts = np.zeros(len(traits)), np.zeros(len(traits), dtype=np.int64)
	with torch.no_grad():
		for genotypes, row_indices in make_loader(args, rows, torch):
			genotypes = genotypes.to(device, non_blocking=True)
			row_indices = row_indices.numpy()
			with torch.autocast(device_type=device.type, enabled=args.amp != "off", dtype=dtype):
				output = model(genotypes)
				losses, mask = loss_matrix(output, row_indices, prepared_labels, baseline_logits, traits, torch, device)
			if not torch.isfinite(losses).all():
				raise RuntimeError("Nonfinite validation loss; no model is selected using invalid values.")
			totals += losses.sum(dim=0).cpu().numpy()
			counts += mask.sum(dim=0).cpu().numpy()
	return np.divide(totals, counts, out=np.full_like(totals, np.nan), where=counts > 0), counts


def predict_outputs(model, rows, args, torch, device, dtype, tasks):
	model.eval()
	outputs = np.empty((len(rows), tasks), dtype=np.float64)
	position = 0
	with torch.no_grad():
		for genotypes, _ in make_loader(args, rows, torch):
			with torch.autocast(device_type=device.type, enabled=args.amp != "off", dtype=dtype):
				prediction = model(genotypes.to(device, non_blocking=True))
			if not torch.isfinite(prediction).all():
				raise RuntimeError("The model produced nonfinite test predictions.")
			count = len(prediction)
			outputs[position:position + count] = prediction.float().cpu().numpy()
			position += count
	return outputs


def correlation_squared(first, second):
	if len(first) < 2 or np.std(first) < 1e-12 or np.std(second) < 1e-12:
		return np.nan
	return float(np.corrcoef(first, second)[0, 1] ** 2)


def auc(phenotype, prediction):
	positive = phenotype == 1
	n_positive, n_negative = int(positive.sum()), int((~positive).sum())
	if not n_positive or not n_negative:
		return np.nan
	ranks = rankdata(prediction, method="average")
	return float((ranks[positive].sum() - n_positive * (n_positive + 1) / 2) / (n_positive * n_negative))


def average_precision(phenotype, prediction):
	if not np.any(phenotype == 1) or not np.any(phenotype == 0):
		return np.nan
	order = np.argsort(-prediction, kind="mergesort")
	scores, labels = prediction[order], phenotype[order]
	endpoints = np.r_[np.flatnonzero(np.diff(scores) != 0), len(scores) - 1]
	true_positives = np.cumsum(labels)[endpoints]
	precision = true_positives / (endpoints + 1)
	recall = true_positives / labels.sum()
	return float(np.sum(np.diff(np.r_[0, recall]) * precision))


def metrics_for(phenotype, prediction, baseline, genetic_score, trait):
	if not len(phenotype):
		return {}
	if SUPPORTED_TRAITS[trait] == "continuous":
		baseline_error = np.sum((phenotype - baseline) ** 2)
		full_error = np.sum((phenotype - prediction) ** 2)
		total_variance = np.sum((phenotype - phenotype.mean()) ** 2)
		full_r2 = float(1 - full_error / total_variance) if total_variance > 0 else np.nan
		return {
			"residual_correlation_R2": correlation_squared(phenotype - baseline, genetic_score),
			"predictive_R2": full_r2,
			"SSE_partial_R2": float(1 - full_error / baseline_error) if baseline_error > 0 else np.nan,
			"full_R2": full_r2,
			"baseline_R2": float(1 - baseline_error / total_variance) if total_variance > 0 else np.nan,
			"full_RMSE": float(np.sqrt(full_error / len(phenotype))),
			"baseline_RMSE": float(np.sqrt(baseline_error / len(phenotype))),
			"prediction_bias": float(np.mean(prediction - phenotype)),
		}
	full_auc, baseline_auc = auc(phenotype, prediction), auc(phenotype, baseline)
	clipped = np.clip(prediction, 1e-12, 1 - 1e-12)
	return {
		"AUC": full_auc, "baseline_AUC": baseline_auc, "delta_AUC": full_auc - baseline_auc,
		"AUPRC": average_precision(phenotype, prediction), "baseline_AUPRC": average_precision(phenotype, baseline),
		"Brier": float(np.mean((phenotype - prediction) ** 2)),
		"baseline_Brier": float(np.mean((phenotype - baseline) ** 2)),
		"log_loss": float(-np.mean(phenotype * np.log(clipped) + (1 - phenotype) * np.log1p(-clipped))),
		"prevalence": float(phenotype.mean()), "cases": int(phenotype.sum()),
	}


def result_tables(data, labels, rows, raw_outputs, baselines, baseline_logits, traits, preprocessing):
	prediction_tables, metric_records = [], []
	for column, trait in enumerate(traits):
		baseline = baselines[rows, column]
		if SUPPORTED_TRAITS[trait] == "continuous":
			genetic = raw_outputs[:, column] * preprocessing["tasks"][trait]["residual_scale"]
			prediction = baseline + genetic
		else:
			genetic = raw_outputs[:, column]
			prediction = expit(baseline_logits[rows, column] + genetic)
		table = data.iloc[rows][["eid", "target", "split"]].copy()
		table["trait"] = trait
		table["y"] = labels[rows, column]
		table["baseline"], table["prediction"], table["genetic_score"] = baseline, prediction, genetic
		prediction_tables.append(table)
		groups = [("ALL", np.ones(len(table), dtype=bool))]
		groups += [(str(target), table["target"].to_numpy() == target) for target in sorted(table["target"].unique()) if target != "ALL"]
		for target, membership in groups:
			observed = membership & np.isfinite(table["y"].to_numpy())
			values = metrics_for(table["y"].to_numpy()[observed], prediction[observed], baseline[observed], genetic[observed], trait)
			for metric, value in values.items():
				metric_records.append({"split": "test", "target": target, "trait": trait, "n": int(observed.sum()), "metric": metric, "value": value})
	return pd.concat(prediction_tables, ignore_index=True), pd.DataFrame(metric_records)


def atomic_torch_save(torch, payload, filename):
	filename = Path(filename)
	staging = filename.with_name(filename.name + f".tmp.{os.getpid()}")
	try:
		with staging.open("wb") as handle:
			torch.save(payload, handle)
			handle.flush()
			os.fsync(handle.fileno())
		os.replace(staging, filename)
	finally:
		staging.unlink(missing_ok=True)


def atomic_table(table, filename):
	filename = Path(filename)
	staging = filename.with_name(filename.name + f".tmp.{os.getpid()}")
	try:
		table.to_csv(staging, sep="\t", index=False, compression="gzip" if str(filename).endswith(".gz") else None)
		os.replace(staging, filename)
	finally:
		staging.unlink(missing_ok=True)


def stage_evaluation(args, identity, preparation, checkpoint):
	# This completion record is temporary. Publication checks it and exports no JSON.
	genotype_stat = Path(args.genotypes).stat()
	if genotype_stat.st_size != identity["genotypes_size"] or genotype_stat.st_mtime_ns != identity["genotypes_mtime_ns"]:
		raise RuntimeError("The genotype cache changed while the model was running; outputs are not marked complete.")
	if file_sha256(args.data) != identity["data_file_sha256"] or file_sha256(args.variants) != identity["variants_file_sha256"]:
		raise RuntimeError("The input cohort or SNP manifest changed while the model was running; outputs are not marked complete.")
	output_directory = Path(args.out_dir)
	model_stat = (output_directory / "model.pt").stat()
	metadata = {
		"format_version": 1,
		"preparation_signature": preparation.get("signature") if preparation else None,
		"inputs": {
			"data_sha256": identity["data_file_sha256"], "variants_sha256": identity["variants_file_sha256"],
			"canonical_data_sha256": identity["data_sha256"], "canonical_variants_sha256": identity["variants_sha256"],
			"genotypes": {"size": identity["genotypes_size"], "mtime_ns": identity["genotypes_mtime_ns"], "shape": identity["shape"], "dtype": identity["genotypes_dtype"]},
		},
		"outputs_sha256": {name: file_sha256(output_directory / name) for name in ("metrics.tsv", "test_predictions.tsv.gz", "history.tsv")},
		"model": {
			"file": "model.pt", "size": int(model_stat.st_size), "mtime_ns": int(model_stat.st_mtime_ns),
			"training_signature": checkpoint["signature"], "upstream_source_sha256": checkpoint["upstream"]["source_sha256"],
		},
	}
	filename = output_directory / "evaluation.json"
	staging = filename.with_name(filename.name + f".tmp.{os.getpid()}")
	try:
		staging.write_text(json.dumps(metadata, indent=2, allow_nan=False) + "\n")
		os.replace(staging, filename)
	finally:
		staging.unlink(missing_ok=True)


def cpu_state(model):
	return {name: value.detach().cpu() for name, value in model.state_dict().items()}


@contextlib.contextmanager
def output_lock(directory):
	import fcntl
	key = hashlib.sha256(str(Path(directory).resolve()).encode()).hexdigest()[:20]
	with open(f"/tmp/grid-prsformer-{key}.lock", "a+") as handle:
		try:
			fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
		except BlockingIOError as exc:
			raise RuntimeError("Another PRSformer process is writing this output directory.") from exc
		yield


def save_history(history, filename, best_epoch):
	table = pd.DataFrame(history)
	table["selected"] = table["epoch"] == best_epoch
	atomic_table(table, filename)


def run_training(args):
	traits = csv_list(args.traits) or list(SUPPORTED_TRAITS)
	if len(set(traits)) != len(traits) or not set(traits).issubset(SUPPORTED_TRAITS):
		raise ValueError("--traits must be a nonrepeating subset of height,ldl,t2dm.")
	covariates = [] if args.covariates.strip().lower() == "none" else csv_list(args.covariates)
	if len(set(covariates)) != len(covariates):
		raise ValueError("Covariate names must not repeat.")
	if set(covariates) & set(traits):
		raise ValueError("A modeled phenotype cannot also be supplied as a covariate.")
	data, labels, identity, variants = read_inputs(args, traits)
	preparation_path = Path(args.genotypes).resolve().parent / "prepare.json"
	preparation = json.loads(preparation_path.read_text()) if preparation_path.is_file() else None
	architecture = architecture_from_args(args, identity["shape"][1], len(traits))
	preprocessing, prepared_labels, baselines, baseline_logits = fit_baselines(data, labels, traits, covariates)
	torch, device, dtype = runtime(args, args.attention)
	model_class, upstream = load_official_model(args.upstream_dir, args.attention)
	model_smoke(model_class, architecture, torch, device, dtype, args)
	parameter_count = parameter_estimate(architecture)
	summary = {
		"samples": len(data), "variants": identity["shape"][1], "traits": traits,
		"split_counts": {str(key): int(value) for key, value in data["split"].value_counts().items()},
		"observed_training": {trait: preprocessing["tasks"][trait]["n_training"] for trait in traits},
		"parameters": parameter_count, "fp32_parameters_GiB": parameter_count * 4 / 2 ** 30,
		"adam_parameters_gradients_states_minimum_GiB": parameter_count * 16 / 2 ** 30,
		"one_float32_activation_GiB": args.batch_size * identity["shape"][1] * args.embed_dim * 4 / 2 ** 30,
		"device": str(device), "attention": args.attention, "official_forward_backward_preflight": "passed",
		"note": "Activation/workspace memory is additional; parameter estimates do not guarantee a model will fit.",
	}
	print(json.dumps(summary, indent=2), flush=True)
	if args.check:
		return
	if device.type == "cuda":
		free_memory, _ = torch.cuda.mem_get_info(device)
		if parameter_count * 16 > free_memory * 0.8:
			raise RuntimeError("Model weights/gradients/Adam states alone would consume over 80% of free GPU memory; reduce SNPs/model dimensions before running.")
	output_directory = Path(args.out_dir)
	output_directory.mkdir(parents=True, exist_ok=True)
	best_file, training_file = output_directory / "model.pt", output_directory / "training.pt"
	if args.resume:
		if not training_file.is_file() or not best_file.is_file():
			raise ValueError("--resume needs both training.pt and model.pt from an interrupted run.")
	elif best_file.exists() or training_file.exists():
		raise ValueError("Output already contains a model. Use --resume for an interrupted run or choose a new output directory.")
	(output_directory / "evaluation.json").unlink(missing_ok=True)
	train_rows = np.flatnonzero((data["split"].to_numpy() == "train") & np.isfinite(labels).any(axis=1))
	validation_rows = np.flatnonzero(data["split"].to_numpy() == "validation")
	test_rows = np.flatnonzero(data["split"].to_numpy() == "test")
	task_counts = np.isfinite(prepared_labels[train_rows]).sum(axis=0)
	task_adjustment = torch.as_tensor(len(train_rows) / task_counts, device=device, dtype=torch.float32)
	# Uniform subject sampling plus this adjustment estimates equal mean loss per task.
	training_options = {
		name: getattr(args, name) for name in (
			"epochs", "patience", "min_delta", "lr", "weight_decay", "batch_size", "accumulation_steps",
			"clip_grad", "amp", "seed", "gradient_checkpointing",
		)
	}
	adapter_logic = ast.dump(ast.parse(Path(__file__).read_text()), include_attributes=False)
	adapter_digest = hashlib.sha256(adapter_logic.encode()).hexdigest()
	signature = digest_json({
		"identity": identity, "architecture": architecture, "traits": traits, "covariates": covariates,
		"training": training_options, "sources": upstream["source_sha256"], "adapter_logic_sha256": adapter_digest,
		"preprocessing": preprocessing, "preparation": preparation,
	})
	random.seed(args.seed)
	np.random.seed(args.seed)
	torch.manual_seed(args.seed)
	if device.type == "cuda":
		torch.cuda.manual_seed_all(args.seed)
	model = model_class(**architecture).to(device)
	if args.gradient_checkpointing:
		checkpoint_layers(model, torch)
	if sum(parameter.numel() for parameter in model.parameters()) != parameter_count:
		raise RuntimeError("Upstream parameter count changed; review the architecture before proceeding.")
	optimizer = torch.optim.AdamW(model.parameters(), lr=args.lr, betas=(0.9, 0.999), weight_decay=args.weight_decay)
	scheduler = torch.optim.lr_scheduler.CosineAnnealingLR(optimizer, T_max=args.epochs, eta_min=args.lr * 0.02)
	scaler = torch.amp.GradScaler(device.type, enabled=device.type == "cuda" and args.amp == "fp16")
	start_epoch, best_epoch, best_loss, stale_epochs, history = 1, 0, float("inf"), 0, []
	metadata = {
		"format_version": 1, "traits": traits, "architecture": architecture, "attention": args.attention,
		"upstream": upstream, "preprocessing": preprocessing, "training_identity": identity,
		"split_registry": split_registry(data),
		"variants": variants.to_dict(orient="list"), "preparation": preparation,
		"adapter_logic_sha256": adapter_digest,
		"training_options": training_options, "signature": signature, "torch_version": str(torch.__version__),
		"adaptation": "Mixed continuous/binary tasks; training-only covariate baselines; standardized residual MSE and fixed-logistic-offset BCE; equal task mean losses; validation early stopping.",
		"variant_dosage": "Raw count of the ALT allele in the aligned variant manifest; -1 denotes missing.",
		"test_usage": "Held out from neural training, baseline fitting, preprocessing, early stopping, and model selection.",
	}
	if args.resume:
		state = torch.load(training_file, map_location="cpu", weights_only=True)
		if state.get("signature") != signature:
			raise ValueError("Resume input, split, source, preprocessing or model/training settings differ from the interrupted run.")
		model.load_state_dict(state["model_state"])
		optimizer.load_state_dict(state["optimizer_state"])
		scheduler.load_state_dict(state["scheduler_state"])
		scaler.load_state_dict(state["scaler_state"])
		start_epoch, best_epoch, best_loss = state["epoch"] + 1, state["best_epoch"], state["best_loss"]
		stale_epochs, history = state["stale_epochs"], state["history"]
		torch.set_rng_state(state["torch_rng"])
		if device.type == "cuda" and state["cuda_rng"]:
			torch.cuda.set_rng_state_all(state["cuda_rng"])
		del state
	for epoch in range(start_epoch, args.epochs + 1):
		if stale_epochs >= args.patience:
			break
		model.train()
		optimizer.zero_grad(set_to_none=True)
		started = time.monotonic()
		last_progress = started
		loss_totals, observed_counts = np.zeros(len(traits)), np.zeros(len(traits), dtype=np.int64)
		loader = make_loader(args, train_rows, torch, shuffle=True, seed=args.seed + epoch)
		learning_rate = float(optimizer.param_groups[0]["lr"])
		for batch, (genotypes, row_indices) in enumerate(loader):
			row_indices = row_indices.numpy()
			group_start = (batch // args.accumulation_steps) * args.accumulation_steps * args.batch_size
			effective_count = min(args.accumulation_steps * args.batch_size, len(train_rows) - group_start)
			with torch.autocast(device_type=device.type, enabled=args.amp != "off", dtype=dtype):
				output = model(genotypes.to(device, non_blocking=True))
				losses, mask = loss_matrix(output, row_indices, prepared_labels, baseline_logits, traits, torch, device)
				loss = (losses * task_adjustment).sum() / effective_count / len(traits)
			if not torch.isfinite(loss):
				raise RuntimeError("Nonfinite training loss. Reduce the learning rate or use FP32; an invalid model is not saved.")
			scaler.scale(loss).backward()
			loss_totals += losses.detach().sum(dim=0).cpu().numpy()
			observed_counts += mask.sum(dim=0).cpu().numpy()
			if (batch + 1) % args.accumulation_steps == 0 or batch + 1 == len(loader):
				scaler.unscale_(optimizer)
				# FP16 loss scaling can initially overflow; GradScaler records this,
				# skips the update and reduces its scale. FP32/BF16 failures are errors.
				torch.nn.utils.clip_grad_norm_(model.parameters(), args.clip_grad, error_if_nonfinite=not scaler.is_enabled())
				scaler.step(optimizer)
				scaler.update()
				optimizer.zero_grad(set_to_none=True)
			if batch == 0 or time.monotonic() - last_progress >= 60:
				seen = min((batch + 1) * args.batch_size, len(train_rows))
				elapsed = time.monotonic() - started
				print(f"TRAIN epoch {epoch}/{args.epochs}: {seen:,}/{len(train_rows):,} samples; {seen / max(elapsed, 1e-9):.2f} samples/s", flush=True)
				last_progress = time.monotonic()
		train_losses = loss_totals / observed_counts
		print(f"VALIDATION epoch {epoch}: {len(validation_rows):,} samples", flush=True)
		validation_losses, validation_counts = evaluate_loss(model, validation_rows, args, prepared_labels, baseline_logits, traits, torch, device, dtype)
		validation_loss = float(np.mean(validation_losses))
		seconds = float(time.monotonic() - started)
		for split, losses, counts in (("train", train_losses, observed_counts), ("validation", validation_losses, validation_counts)):
			for column, trait in enumerate(traits):
				history.append({"epoch": epoch, "split": split, "trait": trait, "loss": float(losses[column]), "n": int(counts[column]), "learning_rate": learning_rate, "seconds": seconds})
		if validation_loss < best_loss - args.min_delta:
			best_loss, best_epoch, stale_epochs = validation_loss, epoch, 0
			atomic_torch_save(torch, {**metadata, "model_state": cpu_state(model), "best_epoch": best_epoch, "best_validation_loss": best_loss}, best_file)
		else:
			stale_epochs += 1
		scheduler.step()
		save_history(history, output_directory / "history.tsv", best_epoch)
		atomic_torch_save(torch, {
			"signature": signature, "epoch": epoch, "best_epoch": best_epoch, "best_loss": best_loss,
			"stale_epochs": stale_epochs, "history": history, "model_state": cpu_state(model),
			"optimizer_state": optimizer.state_dict(), "scheduler_state": scheduler.state_dict(), "scaler_state": scaler.state_dict(),
			"torch_rng": torch.get_rng_state(), "cuda_rng": torch.cuda.get_rng_state_all() if device.type == "cuda" else [],
		}, training_file)
		print(f"Epoch {epoch}: train={float(np.mean(train_losses)):.6f}; validation={validation_loss:.6f}; best={best_epoch}; seconds={seconds:.1f}", flush=True)
	if not best_file.is_file():
		raise RuntimeError("Training did not produce a valid validation-selected model.")
	best = torch.load(best_file, map_location="cpu", weights_only=True)
	model.load_state_dict(best["model_state"])
	outputs = predict_outputs(model, test_rows, args, torch, device, dtype, len(traits))
	predictions, metrics = result_tables(data, labels, test_rows, outputs, baselines, baseline_logits, traits, preprocessing)
	atomic_table(predictions, output_directory / "test_predictions.tsv.gz")
	atomic_table(metrics, output_directory / "metrics.tsv")
	best["completed_epochs"] = max(record["epoch"] for record in history)
	best["test_evaluated"] = True
	best["history"] = history
	atomic_torch_save(torch, best, best_file)
	stage_evaluation(args, identity, preparation, best)
	training_file.unlink(missing_ok=True)
	print(f"Completed: selected epoch {best_epoch}; test-only predictions for {len(test_rows)} individuals.", flush=True)


def run_prediction(args):
	if not args.checkpoint:
		raise ValueError("predict requires --checkpoint model.pt.")
	checkpoint_path = Path(args.checkpoint).resolve()
	checkpoint_stat = checkpoint_path.stat()
	torch = import_torch()
	checkpoint = torch.load(checkpoint_path, map_location="cpu", weights_only=True)
	traits, architecture = checkpoint["traits"], checkpoint["architecture"]
	if args.traits and csv_list(args.traits) != traits:
		raise ValueError("Prediction traits and their order must match the checkpoint.")
	data, labels, identity, _ = read_inputs(args, traits, require_training=False)
	validate_prediction_holdout(data, identity, checkpoint)
	preparation_path = Path(args.genotypes).resolve().parent / "prepare.json"
	preparation = json.loads(preparation_path.read_text()) if preparation_path.is_file() else None
	if identity["variants_sha256"] != checkpoint["training_identity"]["variants_sha256"]:
		raise ValueError("The ordered SNP/REF/ALT manifest differs from model training; do not score misaligned genotypes.")
	torch, device, dtype = runtime(args, checkpoint["attention"])
	model_class, upstream = load_official_model(args.upstream_dir, checkpoint["attention"])
	if upstream["source_sha256"] != checkpoint["upstream"]["source_sha256"]:
		raise ValueError("Upstream source files changed since model training; restore the recorded revision before prediction.")
	model_smoke(model_class, architecture, torch, device, dtype, args)
	if args.check:
		print("Prediction input identities and official model forward/backward preflight passed.", flush=True)
		return
	model = model_class(**architecture).to(device)
	model.load_state_dict(checkpoint["model_state"])
	baselines, baseline_logits = apply_baselines(data, traits, checkpoint["preprocessing"])
	rows = np.flatnonzero(data["split"].to_numpy() == "test")
	outputs = predict_outputs(model, rows, args, torch, device, dtype, len(traits))
	predictions, metrics = result_tables(data, labels, rows, outputs, baselines, baseline_logits, traits, checkpoint["preprocessing"])
	output_directory = Path(args.out_dir).resolve()
	output_directory.mkdir(parents=True, exist_ok=True)
	(output_directory / "evaluation.json").unlink(missing_ok=True)
	atomic_table(predictions, output_directory / "test_predictions.tsv.gz")
	atomic_table(metrics, output_directory / "metrics.tsv")
	if not checkpoint.get("history"):
		raise ValueError("The checkpoint lacks training history; use a complete checkpoint produced by this adapter.")
	save_history(checkpoint["history"], output_directory / "history.tsv", checkpoint["best_epoch"])
	model_destination = output_directory / "model.pt"
	if checkpoint_path != model_destination:
		staging = model_destination.with_name(model_destination.name + f".tmp.{os.getpid()}")
		try:
			shutil.copy2(checkpoint_path, staging)
			os.replace(staging, model_destination)
		finally:
			staging.unlink(missing_ok=True)
	current_checkpoint_stat = checkpoint_path.stat()
	if current_checkpoint_stat.st_size != checkpoint_stat.st_size or current_checkpoint_stat.st_mtime_ns != checkpoint_stat.st_mtime_ns:
		raise RuntimeError("The source checkpoint changed during prediction; outputs are not marked complete.")
	stage_evaluation(args, identity, preparation, checkpoint)
	print(f"Scored {len(rows)} held-out individuals using the saved model and training-only preprocessing.", flush=True)


def add_model_arguments(parser):
	parser.add_argument("--genotypes", required=True, help="Aligned sample-major .npy dosage memmap")
	parser.add_argument("--data", required=True, help="TSV(.gz): eid,target,split,phenotypes,numeric covariates")
	parser.add_argument("--variants", required=True, help="Ordered TSV(.gz): CHR,BP,SNP,REF,ALT")
	parser.add_argument("--upstream-dir", required=True, help="Official 23andMe/PRSformer checkout")
	parser.add_argument("--out-dir", required=True, help="Private staging directory; publish temporary TSV files using grid.sh prsformer")
	parser.add_argument("--traits", help="Default height,ldl,t2dm; prediction uses the checkpoint task order")
	parser.add_argument("--covariates", default="", help="Comma-separated numeric covariates; transformations fit only on training samples")
	parser.add_argument("--device", default="cuda")
	parser.add_argument("--attention", choices=["neighborhood", "global"], default="neighborhood")
	parser.add_argument("--embed-dim", type=int, default=64)
	parser.add_argument("--heads", type=int, default=4)
	parser.add_argument("--layers", type=int, default=2)
	parser.add_argument("--ff-dim", type=int, default=128)
	parser.add_argument("--kernel-size", type=int, default=385)
	parser.add_argument("--dilation", default="1")
	parser.add_argument("--batch-size", type=int, default=1)
	parser.add_argument("--accumulation-steps", type=int, default=64)
	parser.add_argument("--epochs", type=int, default=30)
	parser.add_argument("--patience", type=int, default=5)
	parser.add_argument("--min-delta", type=float, default=1e-4)
	parser.add_argument("--lr", type=float, default=5e-4)
	parser.add_argument("--weight-decay", type=float, default=0.05)
	parser.add_argument("--clip-grad", type=float, default=1.0)
	parser.add_argument("--amp", choices=["fp16", "bf16", "off"], default="fp16")
	parser.add_argument("--gradient-checkpointing", action=argparse.BooleanOptionalAction, default=True)
	parser.add_argument("--workers", type=int, default=0)
	parser.add_argument("--threads", type=int, default=4)
	parser.add_argument("--seed", type=int, default=12345)
	parser.add_argument("--check", action="store_true", help="Validate aligned inputs, baselines and a small actual official-model forward/backward pass")
	parser.add_argument("--resume", action="store_true", help="Resume an interrupted training run with identical settings")
	parser.add_argument("--checkpoint", help="Saved model.pt for predict")


def run_model_command(args, parser):
	if min(args.batch_size, args.accumulation_steps, args.epochs, args.patience, args.threads) < 1 or args.workers < 0:
		parser.error("Batch/accumulation/epoch/patience/thread counts must be positive; workers must be nonnegative.")
	if not all(math.isfinite(value) for value in (args.lr, args.weight_decay, args.clip_grad, args.min_delta)) or args.lr <= 0 or args.weight_decay < 0 or args.clip_grad <= 0 or args.min_delta < 0:
		parser.error("Learning rate and gradient clip must be positive; weight decay and minimum improvement must be nonnegative.")
	if args.command == "runtime":
		torch, device, dtype = runtime(args, args.attention)
		model_class, upstream = load_official_model(args.upstream_dir, args.attention)
		length = max(8, args.kernel_size * max(int(x) for x in csv_list(args.dilation)))
		architecture = architecture_from_args(args, length, len(csv_list(args.traits or "height,ldl,t2dm")))
		model_smoke(model_class, architecture, torch, device, dtype, args)
		print(json.dumps({"status": "RUNTIME OK: official model forward/backward passed",
			"python": sys.executable, "torch": str(torch.__version__), "cuda": torch.version.cuda,
			"gpu": torch.cuda.get_device_name(device) if device.type == "cuda" else None,
			"attention": args.attention, "amp": args.amp, "architecture": architecture,
			"upstream": upstream}, indent=2), flush=True)
		return
	with output_lock(args.out_dir):
		if args.command == "train":
			run_training(args)
		else:
			run_prediction(args)


def parser():
	p = argparse.ArgumentParser(description = __doc__)
	sub = p.add_subparsers(dest = "command", required = True)
	a = sub.add_parser("prepare", help = "Align individual data, fix splits, apply train-only QC and write dosage memmap")
	a.add_argument("--dir-gen", default = "/mnt/f/gen/ukb/37/hap")
	a.add_argument("--pheno-file", default = "/mnt/d/data/ukb/phe/Rdata/all.rds")
	a.add_argument("--ancestry-file", default = "/mnt/d/data/ukb/pca_proj/ukb.ancestry.auto.tsv.gz")
	a.add_argument("--group-col", default = "genetic_ancestry")
	a.add_argument("--traits", default = "height,ldl,t2dm")
	a.add_argument("--covariates", default = "age,sex,PC1,PC2")
	a.add_argument("--height-col", default = "height")
	a.add_argument("--ldl-col", default = "ldl")
	a.add_argument("--t2dm-col", default = "auto")
	a.add_argument("--snp-list", default = "/mnt/f/refLD/csx/snpinfo_mult_1kg_hm3")
	a.add_argument("--chrs", default = "1-22")
	a.add_argument("--keep", default = "none")
	a.add_argument("--remove", default = "/mnt/d/files/ukb.exclude.id")
	a.add_argument("--split-file", default = "none")
	a.add_argument("--split-group-file", default = "none")
	a.add_argument("--train-fraction", type = float, default = .6)
	a.add_argument("--validation-fraction", type = float, default = .2)
	a.add_argument("--seed", type = int, default = 20260904)
	a.add_argument("--maf", type = float, default = .01)
	a.add_argument("--call-rate", type = float, default = .98)
	a.add_argument("--chunk-variants", type = int, default = 32)
	a.add_argument("--max-samples", type = int, default = 0)
	a.add_argument("--max-variants", type = int, default = 0)
	a.add_argument("--cache-dir", required = True)
	a.add_argument("--check", action = "store_true")
	a.add_argument("--replace", action = "store_true")
	a.set_defaults(func = prepare)
	b = sub.add_parser("publish", help = "Publish held-out RDS predictions, reusable model and aggregate figure/workbook pairs")
	b.add_argument("--run-dir", required = True)
	b.add_argument("--cache-dir", required = True)
	b.add_argument("--output-root", default = "/mnt/d/analysis/grid")
	b.add_argument("--score-dir", default = "/mnt/d/data/ukb/pgs")
	b.add_argument("--traits", default = "height,ldl,t2dm")
	b.add_argument("--replace", action = "store_true")
	b.set_defaults(func = publish)
	for name, help_text in (("train", "Train the official model on a prepared cohort"),
		("predict", "Score held-out individuals with a frozen checkpoint"),
		("runtime", "Check official-model forward/backward without cohort IO")):
		add_model_arguments(sub.add_parser(name, help=help_text))
	return p


def run_data_command(a):
	a.traits = ",".join(x.strip() for x in a.traits.split(","))
	traits = a.traits.split(",")
	if not set(traits).issubset({"height", "ldl", "t2dm"}) or len(traits) != len(set(traits)):
		raise ValueError("Choose unique traits from height,ldl,t2dm")
	if a.command == "prepare":
		a.covariates = ",".join(x.strip() for x in a.covariates.split(","))
		if len(set(a.covariates.split(","))) != len(a.covariates.split(",")):
			raise ValueError("Covariate names must not repeat")
		if not (0 < a.train_fraction < 1 and 0 < a.validation_fraction < 1 - a.train_fraction):
			raise ValueError("Need positive train, validation and test fractions")
		if not (0 <= a.maf < .5 and 0 < a.call_rate <= 1 and a.chunk_variants > 0 and a.max_samples >= 0 and a.max_variants >= 0):
			raise ValueError("Invalid MAF/call-rate/chunk/subset setting")
	a.func(a)


def main(argv=None):
	cli = parser()
	args = cli.parse_args(argv)
	if args.command in {"prepare", "publish"}:
		run_data_command(args)
	else:
		run_model_command(args, cli)


if __name__ == "__main__":
	try:
		main()
	except (ValueError, RuntimeError, FileNotFoundError, ImportError) as error:
		raise SystemExit(f"ERROR: {error}") from error
