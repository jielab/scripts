#!/usr/bin/env python3
"""GRID: fixed-cohort human evolutionary features and individual reference prediction.

Public entry: ../grid.sh. Variant annotation, genotype preparation and person-level
matching have separate helpers because they run with different data and memory
requirements. Only report() reads held-out outcomes after the predictor is frozen.
"""

from __future__ import annotations

import argparse
import base64
import hashlib
import importlib.util
import importlib.metadata
import json
import os
from pathlib import Path
import pickle
import platform
import shutil
import sys
import tempfile
import zlib

sys.dont_write_bytecode = True

if sys.version_info < (3, 11):
	raise SystemExit("ERROR: GRID requires Python >=3.11 for rdata 1.1.0 read/write support. Run install_grid.sh or set --python to a compatible environment.")

import joblib
import numpy as np
import pandas as pd


# 🚩 Runtime and lossless private model/result storage
def log(status, stage, detail = ""):
	print(f"{status} GRID {stage}" + (f": {detail}" if detail else ""), flush = True)


def load_module(name):
	key = "grid_grid_" + name
	if key not in sys.modules:
		spec = importlib.util.spec_from_file_location(key, Path(__file__).with_name(name + ".py"))
		module = importlib.util.module_from_spec(spec)
		sys.modules[key] = module
		spec.loader.exec_module(module)
	return sys.modules[key]


def sha256(path):
	digest = hashlib.sha256()
	with Path(path).open("rb") as stream:
		for block in iter(lambda: stream.read(4 * 1024 * 1024), b""):
			digest.update(block)
	return digest.hexdigest()


def frame_digest(frame):
	ordered = frame.sort_values("eid", kind = "stable") if "eid" in frame else frame
	digest = hashlib.sha256("\t".join(map(str, ordered.columns)).encode())
	digest.update(pd.util.hash_pandas_object(ordered, index = False).to_numpy().tobytes())
	return digest.hexdigest()


def json_safe(value):
	if isinstance(value, pd.DataFrame):
		return {"columns": list(map(str, value.columns)), "records": json_safe(value.to_dict(orient = "records"))}
	if isinstance(value, pd.Series):
		return json_safe(value.to_list())
	if isinstance(value, dict):
		return {str(k): json_safe(v) for k, v in value.items()}
	if isinstance(value, (tuple, list)):
		return [json_safe(v) for v in value]
	if isinstance(value, np.ndarray):
		return json_safe(value.tolist())
	if isinstance(value, np.generic):
		return json_safe(value.item())
	if isinstance(value, float) and not np.isfinite(value):
		return None
	if isinstance(value, Path):
		return str(value)
	return value


def r_safe(value):
	"""rdata writes native R integer/double arrays; preserve matrices and string IDs."""
	if isinstance(value, pd.DataFrame):
		out = value.copy()
		for name in out:
			if isinstance(out[name].dtype, pd.CategoricalDtype):
				out[name] = out[name].astype(object)
			elif pd.api.types.is_float_dtype(out[name].dtype):
				out[name] = out[name].astype(np.float64)
			elif pd.api.types.is_integer_dtype(out[name].dtype):
				if out[name].dropna().abs().max() > np.iinfo(np.int32).max:
					out[name] = out[name].astype(str)
		return out
	if isinstance(value, dict):
		return {str(k): r_safe(v) for k, v in value.items()}
	if isinstance(value, (list, tuple)):
		return [r_safe(v) for v in value]
	if isinstance(value, np.ndarray):
		if value.dtype.kind in "iu":
			if value.size and (value.max() > np.iinfo(np.int32).max or value.min() < np.iinfo(np.int32).min + 1):
				return value.astype(str)
			return value.astype(np.int32, copy = False)
		if value.dtype.kind == "f":
			return value.astype(np.float64, copy = False)
		if value.dtype.kind == "c":
			return value.astype(np.complex128, copy = False)
		return value
	if isinstance(value, np.generic):
		return r_safe(value.item())
	if isinstance(value, int) and abs(value) > np.iinfo(np.int32).max:
		return str(value)
	if isinstance(value, Path):
		return str(value)
	return value


def write_rds(value, destination):
	import rdata

	destination = Path(destination)
	destination.parent.mkdir(parents = True, exist_ok = True)
	fd, name = tempfile.mkstemp(prefix = ".grid-", suffix = ".rds", dir = destination.parent)
	os.close(fd)
	try:
		rdata.write_rds(name, r_safe(value))
		os.replace(name, destination)
	finally:
		Path(name).unlink(missing_ok = True)


def read_rds(path):
	import rdata

	return rdata.read_rds(path)


def scalar(value):
	if isinstance(value, np.ndarray) and value.size == 1:
		return value.ravel()[0].item() if isinstance(value.ravel()[0], np.generic) else value.ravel()[0]
	return value


def runtime_signature():
	import scipy
	import sklearn
	try:
		torch_version = importlib.metadata.version("torch")
	except importlib.metadata.PackageNotFoundError:
		torch_version = None
	return {"torch": torch_version, "python": ".".join(map(str, sys.version_info[:2])), "numpy": np.__version__,
		"scipy": scipy.__version__, "sklearn": sklearn.__version__, "joblib": joblib.__version__}


def write_model(bundle, metadata, path):
	# R can inspect metadata directly. Python estimators are retained losslessly in
	# one compressed payload; packed donor explanations are separate native R lists.
	payload = base64.b64encode(zlib.compress(pickle.dumps(bundle, protocol = 5), level = 3)).decode("ascii")
	write_rds({"format": "GRID-model-1", "metadata_json": json.dumps(json_safe(metadata), ensure_ascii = False),
		"python_payload_codec": "base64-zlib-pickle5", "python_payload": payload}, path)


def read_model(path):
	load_module("abm")
	model = read_rds(path)
	if not isinstance(model, dict) or scalar(model.get("format")) != "GRID-model-1":
		raise ValueError("Expected a model created by this GRID implementation")
	if scalar(model.get("python_payload_codec")) != "base64-zlib-pickle5":
		raise ValueError("Unrecognized model encoding")
	metadata = json.loads(scalar(model["metadata_json"]))
	if any(metadata.get("sources", {}).get(name) != source_signature().get(name) for name in ("abm.py", "attention.py")):
		raise ValueError("Saved GRID model uses a different matching implementation; use its original code or refit with this release")
	if metadata.get("runtime") != runtime_signature():
		raise ValueError("Saved GRID model runtime differs or was not recorded; use the recorded Python/library versions or refit")
	if metadata.get("training_diagnostics", {}).get("abm_backend") == "selective_attention":
		load_module("attention")
	bundle = pickle.loads(zlib.decompress(base64.b64decode(scalar(model["python_payload"]))))
	return bundle, metadata


def save_cache(value, path):
	path = Path(path)
	path.parent.mkdir(parents = True, exist_ok = True)
	fd, temporary = tempfile.mkstemp(prefix = ".grid-", dir = path.parent)
	os.close(fd)
	try:
		joblib.dump(value, temporary, compress = 1, protocol = 5)
		os.replace(temporary, path)
	finally:
		Path(temporary).unlink(missing_ok = True)


def named_table(path):
	path = Path(path)
	if path.suffix.lower() == ".rds":
		obj = read_rds(path)
		if isinstance(obj, pd.DataFrame):
			return obj
		if isinstance(obj, dict):
			frames = [v for v in obj.values() if isinstance(v, pd.DataFrame)]
			if len(frames) == 1:
				return frames[0]
		raise ValueError(f"Expected one individual table in {path}")
	return pd.read_csv(path, sep = "\t", dtype = {"eid": str, "IID": str, "#IID": str, "family_id": str, "group": str})


def identified(frame, context):
	d = frame.copy()
	if "eid" not in d:
		column = next((c for c in ("IID", "#IID", "ID") if c in d), None)
		if column is None:
			raise ValueError(f"{context}: missing eid/IID column")
		d = d.rename(columns = {column: "eid"})
	if d.eid.isna().any():
		raise ValueError(f"{context}: missing IDs")
	d["eid"] = d.eid.astype(str).str.replace(r"\.0$", "", regex = True)
	if d.eid.duplicated().any() or d.eid.eq("").any():
		raise ValueError(f"{context}: duplicate/empty IDs")
	return d


# 🚩 Fixed data preparation and frozen model fitting
def source_signature():
	return {name: sha256(Path(__file__).with_name(name))
		for name in ("main.py", "abm.py", "attention.py", "data.py", "evolution.py")}


def settings(args):
	exclude = {"stage", "check_device", "replace", "model_file", "feature_manifest", "out_root", "prsformer_root", "bootstrap", "threads", "plot_dpi", "quiet", "publish_only"}
	return {k: json_safe(v) for k, v in vars(args).items() if k not in exclude}


def validate_input_sources(prepared):
	metadata = prepared.get("metadata", {})
	sources = list(metadata.get("sources", []))
	for trait in metadata.get("evolution", {}).values():
		sources.extend(trait.get("weights_sources", []))
	seen = set()
	for source in sources:
		path = Path(source["path"])
		if str(path) in seen:
			continue
		seen.add(str(path))
		if not path.is_file():
			raise ValueError(f"Prepared input is no longer available: {path}; prepare again")
		stat = path.stat()
		if stat.st_size != source["size"] or stat.st_mtime_ns != source["mtime_ns"]:
			raise ValueError(f"Prepared input changed: {path}; prepare again instead of reusing stale features")


def freeze_scoring_files(prepared):
	"""Keep reusable scoring data separate from private individual/model results."""
	identities = {}
	for trait, artifacts in prepared.get("scoring_artifacts", {}).items():
		files = [(name, artifacts[name], name + ".tsv.gz") for name in
			("canonical_weights", "modules", "training_frequencies")]
		for chrom, values in artifacts["chromosomes"].items():
			files.extend([(f"chr{chrom}.wide_weights", values["wide_weights"], f"chr{chrom}.weights.tsv.gz"),
				(f"chr{chrom}.train_afreq", values["train_afreq"], f"chr{chrom}.train.afreq"),
				(f"chr{chrom}.extract", values["extract"], f"chr{chrom}.snps.txt")])
		identities[trait] = []
		for role, source, relative_name in files:
			path = Path(source)
			if not path.is_file():
				raise ValueError(f"Missing reusable scoring data: {path}")
			identities[trait].append({"role": role, "source": str(path),
				"relative_path": "scoring/" + relative_name, "size": path.stat().st_size, "sha256": sha256(path)})
	return identities


def publish_scoring_files(prepared, trait, stage):
	files = prepared.get("scoring_files", {}).get(trait, [])
	for item in files:
		source = Path(item["source"])
		if not source.is_file() or source.stat().st_size != item["size"]:
			raise ValueError(f"Frozen scoring material is missing or changed: {source}; rebuild evolutionary features")
		output = stage / item["relative_path"]
		output.parent.mkdir(parents = True, exist_ok = True)
		shutil.copyfile(source, output)
		if sha256(output) != item["sha256"]:
			raise ValueError(f"Frozen scoring material changed: {source}; do not publish mismatched weights")
	return [{k: v for k, v in item.items() if k != "source"} for item in files]


def prepare(args):
	data = load_module("data")
	cache_path = Path(args.cache_dir) / "prepared.joblib"
	if cache_path.is_file() and not args.replace:
		previous = joblib.load(cache_path)
		if previous.get("settings") != settings(args):
			raise ValueError("This cache belongs to different preparation/model options; choose a new cache or --replace")
	log("START", "prepare", ",".join(args.traits))
	prepared = data.prepare_data(args)
	prepared["settings"] = settings(args)
	prepared["source_signature"] = source_signature()
	prepared["state"] = "prepared"
	save_cache(prepared, Path(args.cache_dir) / "prepared.joblib")
	log("DONE", "prepare", f"cohort={len(prepared['cohort']):,}; common 50/50 roster fixed")
	return prepared


def load_prepared(args, require_features = False):
	path = Path(args.cache_dir) / "prepared.joblib"
	if not path.is_file():
		raise ValueError("No prepared cohort; run grid.sh abm --stage prepare with the same options")
	prepared = joblib.load(path)
	if prepared.get("settings") != settings(args):
		old = prepared.get("settings", {})
		changed = [k for k, v in settings(args).items() if old.get(k) != v]
		raise ValueError("Preparation options changed: " + ", ".join(changed) + "; reuse identical options or prepare a separate cache")
	if prepared.get("source_signature") != source_signature():
		raise ValueError("GRID source changed since preparation; prepare again so stale features are not reused")
	if require_features and prepared.get("state") != "features_ready":
		raise ValueError("Evolutionary score features are missing; run grid.sh abm --stage evolution")
	validate_input_sources(prepared)
	return prepared


def evolution(args, prepared = None):
	if prepared is None:
		prepared = load_prepared(args)
	log("START", "evolution")
	prepared = load_module("data").evolution_data(args, prepared)
	prepared["scoring_files"] = freeze_scoring_files(prepared)
	prepared["state"] = "features_ready"
	save_cache(prepared, Path(args.cache_dir) / "prepared.joblib")
	log("DONE", "evolution", f"mode={prepared['metadata'].get('evolution_mode', 'canonical_annotations')}; feature preparation complete")
	return prepared


def model_configuration(args, groups, trait, require_features = True):
	primary = ("GRID_evolution" if groups.get("evolution") else
		"GRID_frequency_only" if groups.get("frequency") else "GRID_no_evolution")
	return {
		**{name: getattr(args, name) for name in ("abm_backend", "device", "attention_epochs", "attention_patience",
			"attention_batch", "attention_width", "attention_heads", "attention_layers", "attention_dropout",
			"attention_lr", "attention_reconstruction", "donor_block")},
		"trait_type": "binary" if trait == "t2dm" else "continuous", "seed": args.seed,
		"coverage": args.coverage, "coverages": tuple(args.coverages), "primary_arm": primary,
		"train_role_fractions": tuple(args.train_role_fractions), "folds": args.folds,
		"k_grid": tuple(args.k_grid), "alpha_grid": tuple(args.alpha_grid), "radius_grid": tuple(args.radius_grid),
		"max_dims": args.max_dims, "retrieval": args.retrieval, "query_batch": args.query_batch,
		"ann_recall_min": args.ann_recall_min, "ann_audit_n": args.ann_audit_n,
		"hgb_iterations": args.hgb_iterations, "gate_iterations": args.gate_iterations,
		"audit_bootstrap": args.audit_bootstrap, "retain_match_arms": "primary",
		"min_role_n": args.min_role_n, "verbose": not args.quiet,
		# prepare_data validates the shared outer roster before trait-specific
		# missing outcomes are removed. Their loss must not redefine that split.
		"enforce_half_split": False,
	}


def feature_contract(prepared, trait):
	"""Identify the construction rules that a precomputed feature provider must use.

	A matching declaration prevents accidental mixing of runs. It cannot prove
	that an external producer actually followed the declared computation.
	"""
	files = [{k: item[k] for k in ("role", "relative_path", "sha256")}
		for item in prepared.get("scoring_files", {}).get(trait, [])]
	construction = {"format": "GRID-feature-contract-1", "trait": trait,
		"feature_groups": prepared["feature_groups"][trait], "build": prepared["settings"].get("build"),
		"matching_pcs": "same_reference_projection_as_training; do not recompute target PCA",
		"centering": "saved_training_frequencies; no new-cohort recentering",
		"scoring_files": files, "input_mode": prepared["metadata"].get("input_mode"),
		"training_source_identity": prepared["metadata"].get("sources", []),
		"sources": prepared.get("source_signature", {})}
	payload = json.dumps(json_safe(construction), sort_keys=True, separators=(",", ":"))
	return {**construction, "contract_sha256": hashlib.sha256(payload.encode()).hexdigest()}


def validate_feature_manifest(path, input_file, metadata):
	if not path:
		raise ValueError("predict requires --feature-manifest: declare the saved feature contract and the exact precomputed input file")
	manifest = json.loads(Path(path).read_text())
	expected = metadata.get("feature_contract", {}).get("contract_sha256")
	if not expected or manifest.get("format") != "GRID-features-1" or manifest.get("trait") != metadata.get("trait"):
		raise ValueError("Invalid feature manifest format/trait or model without a feature contract")
	if manifest.get("feature_contract_sha256") != expected:
		raise ValueError("Feature manifest belongs to different weights, centering, PCA or feature definitions")
	if manifest.get("data_sha256") != sha256(input_file):
		raise ValueError("Feature manifest does not identify this exact data file")
	return {"status": "provider_declared_contract_and_file_hash_matched",
		"feature_contract_sha256": expected, "data_sha256": manifest["data_sha256"],
		"interpretation": "External feature calculation is declared, not independently reconstructed by predict"}


def fit(args, prepared = None):
	if prepared is None:
		prepared = load_prepared(args, require_features = True)
	validate_input_sources(prepared)
	abm = load_module("abm")
	if args.abm_backend == "selective_attention":
		load_module("attention")
	for trait in args.traits:
		d = prepared["tables"][trait].copy()
		# Eligibility depends on whether the endpoint is observed, never its value.
		d = d[pd.to_numeric(d.y, errors = "coerce").notna()].reset_index(drop = True)
		groups = prepared["feature_groups"][trait]
		config = model_configuration(args, groups, trait)
		identity = {"training_data": frame_digest(d[d.split == "train"]),
			"test_predictors": frame_digest(d.loc[d.split == "test"].drop(columns = "y")),
			"feature_groups": groups, "config": config, "sources": source_signature()}
		path = Path(args.cache_dir) / trait / "fit.joblib"
		if path.is_file() and not args.replace:
			previous = joblib.load(path)
			if previous.get("identity") != identity:
				raise ValueError(f"{trait}: fitted model/input mismatch; choose another cache or --replace")
			log("DONE", "fit", f"{trait}: reused identical frozen model")
			continue
		log("START", "fit", f"{trait}; primary={config['primary_arm']}; test Y excluded from fitting")
		result = abm.fit_predict(d, groups, config)
		result["identity"] = identity
		result["metadata"] = {
			"method": "GRID", "version": "1.3-optional-annotations", "trait": trait,
			"endpoint": ("baseline_t2dm_from_Yr2e_Yt2e" if args.t2dm_col == "auto" else args.t2dm_col)
				if trait == "t2dm" else getattr(args, trait + "_col"),
			"input_metadata": prepared.get("metadata", {}), "covariates": args.covariates,
			"outer_cohort_n": len(prepared["cohort"]), "analyzable_trait_n": len(d),
			"cohort_roster_sha256": frame_digest(prepared["cohort"][["eid", "split", "family_id"]]),
			"training_data_sha256": identity["training_data"], "test_predictor_sha256": identity["test_predictors"],
			"feature_groups": groups, "sources": source_signature(), "python": platform.python_version(),
			"runtime": runtime_signature(), "feature_contract": feature_contract(prepared, trait),
			"numpy": np.__version__, "pandas": pd.__version__, "training_diagnostics": result["training_diagnostics"],
			"primary_model": config["primary_arm"], "primary_policy": "GRID_policy",
			"primary_endpoint": "Brier" if trait == "t2dm" else "MSE",
			"method_label_budget": result.get("method_label_budget", pd.DataFrame()),
			"gate_reference": "CSx fitted on build; not CSx_full_training or PRSformer",
			"scope": "associative prediction and reference evidence; no reconstructed demographic history or causal mechanism",
		}
		# Persist the frozen predictor and label-free test predictions before report
		# is allowed to join any test outcomes.
		save_cache(result, path)
		log("DONE", "fit", f"{trait}; donor bank={len(result['bundle']['donor_eids']):,}")


# 🚩 Evaluation on the same frozen test individuals and subset masks
def prediction_metrics(y, prediction, covariate, binary):
	valid = np.isfinite(y) & np.isfinite(prediction) & np.isfinite(covariate)
	y, prediction, covariate = y[valid], prediction[valid], covariate[valid]
	if not len(y):
		return {"n": 0}
	error = y - prediction
	result = {"n": len(y), "MSE": float(np.mean(error ** 2)), "RMSE": float(np.sqrt(np.mean(error ** 2))),
		"MAE": float(np.mean(np.abs(error))), "outcome_variance": float(np.var(y)),
		"mean_outcome": float(np.mean(y)), "mean_prediction": float(np.mean(prediction))}
	if binary:
		from sklearn.metrics import average_precision_score, roc_auc_score
		p = np.clip(prediction, 1e-7, 1 - 1e-7)
		result.update(Brier = result["MSE"], cases = int(y.sum()), controls = int(len(y) - y.sum()),
			log_loss = float(-np.mean(y * np.log(p) + (1 - y) * np.log1p(-p))))
		result["AUC"] = float(roc_auc_score(y, prediction)) if len(np.unique(y)) == 2 else np.nan
		result["average_precision"] = float(average_precision_score(y, prediction)) if len(np.unique(y)) == 2 else np.nan
	else:
		baseline_error = y - covariate
		genetic_increment = prediction - covariate
		sse = float(np.sum(error ** 2))
		sst = float(np.sum((y - y.mean()) ** 2))
		bsse = float(np.sum(baseline_error ** 2))
		result["total_R2"] = 1 - sse / sst if sst else np.nan
		result["SSE_partial_R2"] = 1 - sse / bsse if bsse else np.nan
		result["predictive_R2"] = result["total_R2"]
		result["residual_correlation_R2"] = float(np.corrcoef(baseline_error, genetic_increment)[0, 1] ** 2) \
			if len(y) > 2 and np.std(baseline_error) > 0 and np.std(genetic_increment) > 0 else np.nan
	result.update(calibration_diagnostics(y, prediction, binary))
	return result


def calibration_diagnostics(y, prediction, binary):
	"""Held-out calibration diagnostics; these never recalibrate predictions."""
	y, prediction = np.asarray(y, float), np.asarray(prediction, float)
	answer = {"calibration_intercept": np.nan, "calibration_slope": np.nan,
		"calibration_status": "not_identifiable", "calibration_refits_prediction": False}
	if len(y) < 3 or np.std(prediction) <= 1e-12:
		return answer
	if not binary:
		variance = float(np.mean((prediction-prediction.mean())**2))
		slope = float(np.mean((prediction-prediction.mean())*(y-y.mean()))/variance)
		return {**answer, "calibration_intercept": float(y.mean()-slope*prediction.mean()),
			"calibration_slope": slope, "calibration_status": "estimated"}
	if not np.isin(y, [0,1]).all() or len(np.unique(y)) != 2:
		return answer
	from scipy.optimize import minimize
	from scipy.special import expit
	p = np.clip(prediction, 1e-7, 1-1e-7)
	x = np.log(p) - np.log1p(-p)
	# In one predictor dimension, disjoint/touching class ranges imply complete
	# or quasi separation. A small optimizer gradient can otherwise masquerade
	# as a finite calibration MLE while the true slope is unbounded.
	zero, one = x[y == 0], x[y == 1]
	if zero.max() <= one.min() or one.max() <= zero.min():
		return {**answer, "calibration_status": "separated"}
	def objective(beta):
		eta = beta[0] + beta[1]*x
		error = expit(eta)-y
		return float(np.mean(np.logaddexp(0,eta)-y*eta)), np.array([error.mean(), np.mean(error*x)])
	fit = minimize(objective, np.array([0.,1.]), jac=True, method="BFGS", options={"maxiter":200,"gtol":1e-7})
	if fit.success and np.isfinite(fit.x).all() and np.max(np.abs(fit.x)) < 1e3:
		answer.update(calibration_intercept=float(fit.x[0]), calibration_slope=float(fit.x[1]), calibration_status="estimated")
	else:
		answer["calibration_status"] = "not_converged_or_separated"
	return answer


def add_prsformer(args, trait, result, cohort, table, features=None):
	root = Path(args.prsformer_root)
	roster_path = root / "prsformer/3.prsformer.split.rds"
	score_path = root / "scores" / trait / "3.prsformer.scores.rds"
	if not roster_path.exists() and not score_path.exists():
		return table, {"method": "PRSformer", "status": "not_run_on_common_split", "n": 0}
	if not roster_path.is_file() or not score_path.is_file():
		raise ValueError("Incomplete PRSformer benchmark: both common-split roster and trait scores are required")
	roster = identified(named_table(roster_path), "PRSformer roster")
	if "split" not in roster or not roster.split.isin(["train", "validation", "test"]).all():
		raise ValueError("PRSformer roster has invalid split labels")
	if set(roster.eid) != set(cohort.eid):
		raise ValueError("PRSformer and GRID must use exactly the same outer cohort")
	outer = roster[["eid", "split"]].merge(cohort[["eid", "split", "family_id"]], on = "eid", suffixes = ("_pf", "_grid"), validate = "one_to_one")
	if not np.array_equal(outer.split_pf.eq("test"), outer.split_grid.eq("test")):
		raise ValueError("PRSformer test set differs from the fixed GRID test half; rerun with the common roster")
	if outer.groupby("family_id").split_pf.nunique().max() > 1:
		raise ValueError("A family crosses PRSformer train/validation/test partitions")
	scores = identified(named_table(score_path), "PRSformer scores")
	if "split" not in scores or not scores.split.eq("test").all() or not {"prediction", "outcome"}.issubset(scores):
		raise ValueError("PRSformer comparator must contain actual held-out outcome-scale predictions")
	if set(scores.eid) != set(cohort.loc[cohort.split == "test", "eid"]):
		raise ValueError("PRSformer prediction IDs do not cover the entire common test half")
	x = table.merge(scores[["eid", "prediction", "outcome"]], on = "eid", how = "left", validate = "one_to_one")
	if not np.isfinite(pd.to_numeric(x.prediction, errors = "coerce")).all():
		raise ValueError("PRSformer test predictions are incomplete/nonfinite")
	if trait == "t2dm" and not x.prediction.between(0, 1).all():
		raise ValueError("T2DM PRSformer predictions must be probabilities in [0,1], not logits")
	if not np.allclose(x.y, x.outcome, equal_nan = True, atol = 1e-7, rtol = 1e-7):
		raise ValueError("PRSformer and GRID endpoints or units differ")
	if "covariate_names" in scores:
		actual = set(scores.covariate_names.dropna().astype(str))
		if actual != {args.covariates}:
			raise ValueError("PRSformer and GRID covariate definitions differ")
	else:
		raise ValueError("PRSformer scores lack covariate provenance; regenerate with the included adapter")
	covariates = load_module("data").names(args.covariates)
	if covariates and (features is None or any("covariate."+name not in scores for name in covariates)):
		raise ValueError("PRSformer scores lack original covariate values; republish with this adapter for a verified comparison")
	if covariates:
		observed = identified(features, "GRID comparison covariates")[["eid", *covariates]]
		compare = table[["eid"]].merge(observed,on="eid",how="left",validate="one_to_one").merge(
			scores[["eid", *["covariate."+name for name in covariates]]],on="eid",how="left",validate="one_to_one")
		for name in covariates:
			if not np.allclose(pd.to_numeric(compare[name], errors="coerce"),
				pd.to_numeric(compare["covariate."+name], errors="coerce"), equal_nan=False, atol=1e-8, rtol=1e-7):
				raise ValueError(f"PRSformer and GRID covariate values differ for {name}; same column names are insufficient")
	if "endpoint_definition" in scores and set(scores.endpoint_definition.dropna().astype(str)) != {result["metadata"]["endpoint"]}:
		raise ValueError("PRSformer and GRID endpoint definitions differ")
	x = x.rename(columns = {"prediction": "PRSformer"}).drop(columns = "outcome")
	return x, {"method": "PRSformer", "status": "common_cohort_test_verified", "n": len(x),
		"score_file_sha256": sha256(score_path), "roster_file_sha256": sha256(roster_path),
		"covariate_values": "matched_by_eid",
		"comparison": "same outer development budget; separate supervised architecture and external GWAS information budget"}


def paired_loss_comparison(table, methods, bootstrap, seed, reference = "CSx_full_training"):
	"""Same-mask paired, family-cluster nonparametric bootstrap of loss differences."""
	if len(table) < 2:
		return []
	y = table.y.to_numpy(float)
	if reference not in table:
		reference = "CSx"
	base = (y - table[reference].to_numpy(float)) ** 2
	levels, inverse = np.unique(table.family_id.astype(str), return_inverse = True)
	count = np.bincount(inverse).astype(float)
	comparisons = []
	for name in methods:
		if name in {reference, "covariate_baseline"}:
			continue
		prediction = table[name].to_numpy(float)
		if not np.isfinite(prediction).all():
			raise ValueError(f"Nonfinite {name} predictions in a common comparison subset")
		delta = (y - prediction) ** 2 - base
		group_sum = np.bincount(inverse, weights = delta)
		point = float(np.mean(delta))
		lo, hi = np.nan, np.nan
		if bootstrap >= 20 and len(levels) >= 2:
			rng = np.random.default_rng(seed)
			draws = np.empty(bootstrap)
			for b in range(bootstrap):
				ix = rng.integers(0, len(levels), size = len(levels))
				draws[b] = group_sum[ix].sum() / count[ix].sum()
			lo, hi = map(float, np.quantile(draws, [.025, .975]))
		comparisons.append({"method": name, "reference": reference, "n": len(y), "families": len(levels),
			"delta_loss": point, "lower": lo, "upper": hi, "relative_loss_reduction": -point / np.mean(base) if np.mean(base) else np.nan,
			"bootstrap": bootstrap, "better_direction": "negative", "interval": "paired_family_percentile_95pct"})
	return comparisons


def evaluate(table, methods, args, trait):
	binary = trait == "t2dm"
	selected = table.selected.astype(bool).to_numpy()
	populations = ["ALL", *sorted(table.ancestry.astype(str).unique())]
	metrics, paired, coverage, contrasts = [], [], [], []
	primary = "GRID_evolution" if "GRID_evolution" in methods else "GRID_frequency_only" if "GRID_frequency_only" in methods else "GRID_no_evolution"
	for population in populations:
		population_mask = np.ones(len(table), dtype = bool) if population == "ALL" else table.ancestry.eq(population).to_numpy()
		low_error = table.selected_absolute_error.astype(bool).to_numpy()
		for subset, mask in {"all": population_mask, "selected": population_mask & selected,
			"rejected": population_mask & ~selected, "selected_low_error": population_mask & low_error,
			"rejected_low_error": population_mask & ~low_error}.items():
			x = table.loc[mask]
			if x.empty:
				continue
			for method in methods:
				row = prediction_metrics(x.y.to_numpy(float), x[method].to_numpy(float), x.covariate_baseline.to_numpy(float), binary)
				metrics.append({"trait": trait, "ancestry": population, "subset": subset, "method": method, **row})
			paired.extend({"trait": trait, "ancestry": population, "subset": subset,
				"loss": "Brier" if binary else "MSE", **row} for row in paired_loss_comparison(x, methods, args.bootstrap, args.seed + 811))
			for reference in ["PRSformer", "GRID_permuted_evolution", "GRID_frequency_only", "GRID_no_evolution", "Ridge_evolution_full_training", "HGB_evolution_full_training"]:
				if reference not in methods or reference == primary:
					continue
				candidates = [primary, "GRID_policy"] if reference == "PRSformer" else [primary]
				contrasts.extend({"trait": trait, "ancestry": population, "subset": subset,
					"loss": "Brier" if binary else "MSE", **row} for row in paired_loss_comparison(x, candidates, args.bootstrap, args.seed + 811, reference))
			for kind in ("Ridge", "HGB"):
				for suffix in ("", "_full_training"):
					candidate, reference = kind+"_evolution"+suffix, kind+"_no_evolution"+suffix
					if {candidate, reference}.issubset(methods):
						contrasts.extend({"trait": trait, "ancestry": population, "subset": subset,
							"comparison_type": "global_annotation_increment_same_model_budget",
							"loss": "Brier" if binary else "MSE", **row}
							for row in paired_loss_comparison(x, [candidate], args.bootstrap, args.seed + 811, reference))
		x = table.loc[population_mask]
		coverage.append({"trait": trait, "ancestry": population, "n": len(x), "selected": int(x.selected.sum()),
			"coverage": float(x.selected.mean()), "released": int(x.released.sum()), "release_coverage": float(x.released.mean()),
			"matched_support": float(x.matching_supported.mean()), "median_matched_count": float(x.matched_count.median()),
			"median_family_ess": float(x.family_ess.median()),
			"case_coverage": float(x.loc[x.y == 1, "selected"].mean()) if binary and (x.y == 1).any() else np.nan})
	curves = []
	for column in [c for c in table if c.startswith("selected.") and ".q" in c and not c.startswith("selected.q")]:
		selector, target = column.removeprefix("selected.").rsplit(".q", 1)
		mask = table[column].astype(bool)
		x = table.loc[mask]
		if x.empty:
			continue
		for method in methods:
			values = prediction_metrics(x.y.to_numpy(float), x[method].to_numpy(float), x.covariate_baseline.to_numpy(float), binary)
			curves.append({"trait": trait, "selector": selector, "target_coverage": float(target), "actual_coverage": float(mask.mean()),
				"case_coverage": float(mask[table.y == 1].mean()) if binary and (table.y == 1).any() else np.nan,
				"method": method, **values})
	# The gain selector uses selected.qX rather than selected.gain.qX.
	for column in [c for c in table if c.startswith("selected.q")]:
		mask = table[column].astype(bool)
		x = table.loc[mask]
		if x.empty:
			continue
		for method in methods:
			curves.append({"trait": trait, "selector": "expected_gain", "target_coverage": float(column.split(".q")[1]),
				"actual_coverage": float(mask.mean()), "case_coverage": float(mask[table.y == 1].mean()) if binary and (table.y == 1).any() else np.nan,
				"method": method, **prediction_metrics(x.y.to_numpy(float), x[method].to_numpy(float), x.covariate_baseline.to_numpy(float), binary)})
	return {"metrics": pd.DataFrame(metrics), "paired": pd.DataFrame(paired), "coverage": pd.DataFrame(coverage),
		"curves": pd.DataFrame(curves), "contrasts": pd.DataFrame(contrasts)}



def confidence_evaluation(table, result, args, trait):
	primary = result["bundle"]["primary_arm"]
	binary = trait == "t2dm"
	if "confidence_percentile" not in table:
		return {"summary": pd.DataFrame(), "curve": pd.DataFrame(), "quintiles": pd.DataFrame()}

	y = table.y.to_numpy(float)
	pred = table[primary].to_numpy(float)
	base = table.covariate_baseline.to_numpy(float)
	conf = pd.to_numeric(table.confidence_percentile, errors="coerce").to_numpy(float)
	loss = (y - pred) ** 2
	family = table.family_id.astype(str).to_numpy()
	valid = np.isfinite(conf) & np.isfinite(loss)
	if not valid.any():
		return {"summary": pd.DataFrame(), "curve": pd.DataFrame(), "quintiles": pd.DataFrame()}

	all_loss = float(np.mean(loss[valid]))
	rows = []
	for coverage, cutoff in [(0.2,.8),(0.4,.6),(0.6,.4),(0.8,.2),(1.0,0.)]:
		m = valid & (conf >= cutoff)
		if not m.any():
			continue
		pm = prediction_metrics(y[m], pred[m], base[m], binary)
		rows.append({
			"trait": trait, "model": primary, "target_coverage": coverage,
			"confidence_cutoff": cutoff, "n": int(m.sum()),
			"actual_coverage": float(m.sum()/valid.sum()),
			"mean_confidence_percentile": float(np.mean(conf[m])),
			"mean_expected_rmse": float(np.nanmean(table.expected_rmse.to_numpy(float)[m])),
			"observed_loss": float(np.mean(loss[m])),
			"delta_loss_vs_all": float(np.mean(loss[m]) - all_loss),
			"cases": int(np.sum(y[m] == 1)) if binary else np.nan,
			**pm
		})
	curve = pd.DataFrame(rows)

	qrows = []
	edges = [(0,.2,"Q1 lowest"),(.2,.4,"Q2"),(.4,.6,"Q3"),(.6,.8,"Q4"),(.8,1.000001,"Q5 highest")]
	for lo, hi, label in edges:
		m = valid & (conf >= lo) & (conf < hi)
		if not m.any():
			continue
		pm = prediction_metrics(y[m], pred[m], base[m], binary)
		qrows.append({
			"trait": trait, "confidence_bin": label, "n": int(m.sum()),
			"mean_confidence_percentile": float(np.mean(conf[m])),
			"observed_loss": float(np.mean(loss[m])),
			"cases": int(np.sum(y[m] == 1)) if binary else np.nan,
			**pm
		})
	quintiles = pd.DataFrame(qrows)

	top = valid & (conf >= .8) & table.technical_qc_pass.to_numpy(bool) & table.matching_supported.to_numpy(bool)
	delta = float(np.mean(loss[top]) - all_loss) if top.any() else np.nan
	lower = upper = np.nan
	if args.bootstrap and top.any():
		# Resample entire families using sufficient statistics. This is exactly
		# the row-expanded cluster bootstrap without a full table scan per family.
		levels, inverse = np.unique(family[valid], return_inverse=True)
		counts = np.bincount(inverse)
		sums = np.bincount(inverse, weights=loss[valid])
		top_counts = np.bincount(inverse, weights=top[valid].astype(float))
		top_sums = np.bincount(inverse, weights=loss[valid] * top[valid])
		rng = np.random.default_rng(args.seed + 1701)
		draws = []
		if len(levels) >= 2:
			for _ in range(args.bootstrap):
				ix = rng.integers(0, len(levels), size=len(levels))
				denom = top_counts[ix].sum()
				if denom:
					draws.append(float(top_sums[ix].sum()/denom - sums[ix].sum()/counts[ix].sum()))
		if len(draws) >= 20:
			lower, upper = map(float, np.quantile(draws, [.025,.975]))

	internal = result["bundle"].get("low_error_audit", {}).get("status", "not_evaluated")
	confidence_audit = result["bundle"].get("confidence_audit", {}).get("status", "not_evaluated")
	mono = False
	if len(curve) >= 3:
		v = curve.sort_values("target_coverage").observed_loss.to_numpy(float)
		mono = bool(np.all(np.diff(v) >= -1e-12))

	if internal != "supported_lower_group_error" or confidence_audit != "supported_lower_group_error":
		conclusion = "candidate_only_not_validated"
	elif np.isfinite(upper) and upper < 0:
		conclusion = "validated_high_confidence_subgroup"
	else:
		conclusion = "development_validated_but_not_confirmed_in_test"

	summary = pd.DataFrame([{
		"trait": trait, "model": primary,
		"loss": "Brier" if binary else "MSE",
		"internal_low_error_audit": internal,
		"internal_confidence_audit": confidence_audit,
		"high_confidence_definition": "development percentile >= 0.8; technical QC and matching support required",
		"test_n": int(valid.sum()),
		"test_high_confidence_n": int(top.sum()),
		"test_high_confidence_coverage": float(top.sum()/valid.sum()),
		"test_all_loss": all_loss,
		"test_high_confidence_loss": float(np.mean(loss[top])) if top.any() else np.nan,
		"test_delta_loss_high_vs_all": delta,
		"test_delta_lower95": lower,
		"test_delta_upper95": upper,
		"monotonic_loss_with_expanding_coverage": mono,
		"conclusion": conclusion,
		"test_outcomes_used_for_training_or_thresholds": False
	}])
	return {"summary": summary, "curve": curve, "quintiles": quintiles}


def make_confidence_figure(confidence, path, trait):
	if confidence["curve"].empty:
		return
	import matplotlib
	matplotlib.use("Agg")
	import matplotlib.pyplot as plt
	curve = confidence["curve"].sort_values("actual_coverage")
	q = confidence["quintiles"]
	status = confidence["summary"].iloc[0]["conclusion"]
	fig, axes = plt.subplots(1,2,figsize=(12,4.8))
	axes[0].plot(curve.actual_coverage, curve.observed_loss, "o-")
	axes[0].set(xlabel="Fraction of test population retained", ylabel="Observed loss", title="a  Frozen confidence coverage")
	if not q.empty:
		axes[1].bar(q.confidence_bin, q.observed_loss)
		axes[1].tick_params(axis="x", rotation=30)
	axes[1].set(ylabel="Observed loss", title="b  Confidence quintiles")
	fig.suptitle(f"GRID confidence — {trait}: {status}", fontsize=13)
	fig.tight_layout(rect=(0,0,1,.94))
	fig.savefig(path / "grid.confidence.png")
	plt.close(fig)


# 🚩 Aggregate figures with matching Excel tables; no participant IDs in workbooks
def workbook(tables, path):
	from openpyxl import Workbook
	from openpyxl.styles import Font, PatternFill

	book = Workbook()
	book.remove(book.active)
	for name, table in tables.items():
		if not isinstance(table, pd.DataFrame) or table.empty:
			continue
		if {"eid", "IID", "donor_eid", "family_id"} & set(table.columns):
			raise ValueError("Participant-level tables must be stored in private RDS files")
		if len(table) >= 1048576:
			raise ValueError("Aggregate workbook table exceeds an Excel worksheet")
		sheet = book.create_sheet(name[:31])
		sheet.append([str(x) for x in table.columns])
		for cell in sheet[1]:
			cell.font = Font(bold = True, color = "FFFFFF")
			cell.fill = PatternFill("solid", fgColor = "31586B")
		for row in table.itertuples(index = False, name = None):
			values = []
			for value in row:
				if isinstance(value, np.generic):
					value = value.item()
				if isinstance(value, (list, tuple, dict)):
					value = json.dumps(json_safe(value))
				if value is None or (not isinstance(value, str) and pd.isna(value)) or (isinstance(value, float) and not np.isfinite(value)):
					value = None
				values.append(value)
			sheet.append(values)
		sheet.freeze_panes = "A2"
		sheet.auto_filter.ref = sheet.dimensions
		for cells in sheet.columns:
			letter = cells[0].column_letter
			sheet.column_dimensions[letter].width = min(52, max(14, max(len(str(c.value or "")) for c in cells[:101]) + 2))
	if not book.sheetnames:
		raise ValueError("No substantive aggregate data for workbook")
	book.save(path)


def make_figures(evaluation, table, result, path, args, trait):
	import matplotlib
	matplotlib.use("Agg")
	import matplotlib.pyplot as plt

	plt.rcParams.update({"font.size": 9, "axes.spines.top": False, "axes.spines.right": False, "savefig.dpi": args.plot_dpi})
	metrics, paired, cov = evaluation["metrics"], evaluation["paired"], evaluation["coverage"]
	primary = result["bundle"]["primary_arm"]
	loss = "Brier" if trait == "t2dm" else "MSE"
	m = metrics[(metrics.ancestry == "ALL") & (metrics.subset == "all") & (metrics.method != "covariate_baseline")]
	fig, axes = plt.subplots(2, 2, figsize = (16, max(10, 5 + .45*len(m))))
	colors = ["#C16A42" if x in (primary, "GRID_policy") else "#548399" for x in m.method]
	axes[0, 0].barh(m.method, m[loss], color = colors)
	axes[0, 0].invert_yaxis()
	axes[0, 0].set(xlabel = f"{loss} (lower is better)", title = "a  Entire held-out test half")
	p = paired[(paired.ancestry == "ALL") & (paired.subset == "all")].reset_index(drop = True)
	for i, row in p.iterrows():
		if np.isfinite(row.lower) and np.isfinite(row.upper):
			axes[0, 1].plot([row.lower, row.upper], [i, i], color = "#31586B")
		axes[0, 1].plot(row.delta_loss, i, "o", color = "#C16A42" if row.method in (primary, "GRID_policy") else "#31586B")
	axes[0, 1].set_yticks(np.arange(len(p)), p.method)
	axes[0, 1].invert_yaxis()
	axes[0, 1].axvline(0, color = "grey", linewidth = 1)
	baseline_name = p.reference.iloc[0] if len(p) else "CSx_full_training"
	axes[0, 1].set(xlabel = f"Paired Δ{loss} vs {baseline_name} (negative is better)", title = "b  Same individuals; family bootstrap")
	q = cov[cov.ancestry != "ALL"]
	positions = np.arange(len(q))
	axes[1, 0].bar(positions - .18, q.coverage, width = .36, label = "Gain selection vs build CSx", color = "#548399")
	axes[1, 0].bar(positions + .18, q.release_coverage, width = .36, label = "Global internal-audit policy", color = "#C16A42")
	axes[1, 0].set_xticks(positions, q.ancestry)
	axes[1, 0].set(ylim = (0, 1), ylabel = "Fraction of test individuals", title = "c  Frozen thresholds; actual test coverage")
	axes[1, 0].legend(fontsize = 8)
	selected_methods = [x for x in ("CSx_full_training", "CSx", "DiscoDivas_calibrated_reference_full_training", "PRSformer", primary, "GRID_policy") if x in metrics.method.values]
	for method in selected_methods:
		x = metrics[(metrics.ancestry == "ALL") & (metrics.method == method)].set_index("subset")
		x = x.reindex(["all", "selected", "rejected"])
		axes[1, 1].plot([0, 1, 2], x[loss], "o-", label = method)
	axes[1, 1].set_xticks([0, 1, 2], ["All", "Selected", "Rejected"])
	axes[1, 1].set(ylabel = loss, title = "d  Compare methods inside identical subsets")
	axes[1, 1].legend(fontsize = 8)
	fig.suptitle(f"GRID — {trait}; outer test labels used only for this evaluation", fontsize = 14)
	fig.tight_layout(rect = (0, 0, 1, .96))
	fig.savefig(path / "grid.performance.png")
	plt.close(fig)
	workbook({"metrics": metrics, "paired_loss": paired, "coverage": cov}, path / "grid.performance.xlsx")
	contrasts = evaluation["contrasts"]
	if not contrasts.empty:
		n_contrasts = len(contrasts[(contrasts.ancestry == "ALL") & (contrasts.subset == "all")])
		fig, axes = plt.subplots(1, 2, figsize = (16, max(5.5, .60*n_contrasts)))
		for ax, subset in zip(axes, ["all", "selected"]):
			x = contrasts[(contrasts.ancestry == "ALL") & (contrasts.subset == subset)].reset_index(drop = True)
			for i, row in x.iterrows():
				if np.isfinite(row.lower) and np.isfinite(row.upper):
					ax.plot([row.lower, row.upper], [i, i], color = "#31586B")
				ax.plot(row.delta_loss, i, "o", color = "#C16A42")
			ax.set_yticks(np.arange(len(x)), [f"{r.method}\nvs {r.reference}" for r in x.itertuples()])
			ax.invert_yaxis()
			ax.axvline(0, color = "grey", linewidth = 1)
			ax.set(title = "All test individuals" if subset == "all" else "Same gain-selected subset", xlabel = f"Direct paired Δ{loss}; negative is better")
		fig.tight_layout()
		fig.savefig(path / "grid.contrasts.png")
		plt.close(fig)
		workbook({"direct_paired_contrasts": contrasts}, path / "grid.contrasts.xlsx")
	curves = evaluation["curves"]
	if not curves.empty:
		fig, axes = plt.subplots(1, 2, figsize = (12, 4.5))
		for method in selected_methods:
			x = curves[(curves.selector == "expected_gain") & (curves.method == method)].sort_values("actual_coverage")
			axes[0].plot(x.actual_coverage, x[loss], "o-", label = method)
		for selector, x in curves[curves.method == primary].groupby("selector"):
			x = x.sort_values("actual_coverage")
			axes[1].plot(x.actual_coverage, x[loss], "o-", label = selector)
		axes[0].set(title = "a  Same gain-selected people", xlabel = "Actual coverage", ylabel = loss)
		axes[1].set(title = "b  Selector controls for the primary candidate", xlabel = "Actual coverage", ylabel = loss)
		for ax in axes:
			ax.legend(fontsize = 8)
		fig.tight_layout()
		fig.savefig(path / "grid.coverage.png")
		plt.close(fig)
		workbook({"coverage_curves": curves}, path / "grid.coverage.xlsx")
	counts = table.groupby("matched_count", dropna = False).size().reset_index(name = "n")
	ess_bins = np.array([0, 1, 2, 4, 8, 16, 32, 64, np.inf])
	ess_group = pd.cut(table.family_ess, ess_bins, right = False, include_lowest = True)
	ess = pd.DataFrame({"family_ess_bin": ess_group.astype(str)}).groupby("family_ess_bin", sort = False).size().reset_index(name = "n")
	fig, axes = plt.subplots(1, 2, figsize = (12, 4.5))
	axes[0].bar(counts.matched_count, counts.n, color = "#548399")
	axes[0].set(xlabel = "Actual matched references within caliper", ylabel = "Test individuals", title = "a  No forced fixed number of twins")
	axes[1].bar(ess.family_ess_bin, ess.n, color = "#C16A42")
	axes[1].tick_params(axis = "x", rotation = 35)
	axes[1].set(xlabel = "Effective independent family count", ylabel = "Test individuals", title = "b  Concentrated weights reduce effective support")
	fig.tight_layout()
	fig.savefig(path / "grid.support.png")
	plt.close(fig)
	workbook({"matched_count": counts, "family_ess": ess}, path / "grid.support.xlsx")
	confidence = confidence_evaluation(table, result, args, trait)
	if not confidence["summary"].empty:
		make_confidence_figure(confidence, path, trait)
		workbook({"summary": confidence["summary"], "coverage_curve": confidence["curve"],
			"quintiles": confidence["quintiles"]}, path / "grid.confidence.xlsx")


def report(args, prepared = None):
	# Fresh stage processes need the dynamically named pickle classes registered.
	load_module("abm")
	if args.abm_backend == "selective_attention":
		load_module("attention")
	if prepared is None:
		prepared = load_prepared(args, require_features = True)
	validate_input_sources(prepared)
	root = Path(args.out_root)
	root.mkdir(parents = True, exist_ok = True)
	for trait in args.traits:
		log("START", "report", trait)
		cache = Path(args.cache_dir) / trait / "fit.joblib"
		if not cache.is_file():
			raise ValueError(f"{trait}: frozen model is missing; run --stage fit")
		result = joblib.load(cache)
		d = prepared["tables"][trait]
		eligible = d[pd.to_numeric(d.y, errors = "coerce").notna()].reset_index(drop = True)
		expected_identity = {
			"training_data": frame_digest(eligible[eligible.split == "train"]),
			"test_predictors": frame_digest(eligible.loc[eligible.split == "test"].drop(columns = "y")),
			"feature_groups": prepared["feature_groups"][trait],
			"config": model_configuration(args, prepared["feature_groups"][trait], trait),
			"sources": source_signature(),
		}
		if result.get("identity") != expected_identity:
			raise ValueError(f"{trait}: frozen model no longer matches the prepared predictors or training data")
		y = d.loc[(d.split == "test") & pd.to_numeric(d.y, errors = "coerce").notna(), ["eid", "y"]]
		table = result["predictions"].merge(y, on = "eid", how = "inner", validate = "one_to_one")
		if len(table) != len(result["predictions"]):
			raise ValueError("Frozen prediction cohort and observed endpoint cohort differ")
		methods = [*result["bundle"]["models"], *result["bundle"]["banks"], "GRID_policy"]
		pf_status = {"status": "comparison_in_final"}
		if not args.publish_only:
			table, pf_status = add_prsformer(args, trait, result, prepared["cohort"], table, d)
			if "PRSformer" in table:
				methods.append("PRSformer")
			evaluation = evaluate(table, methods, args, trait)
		destination = root / trait
		if destination.exists() and any(destination.iterdir()) and not args.replace:
			raise ValueError(f"Formal {trait} results already exist; choose another --out-root or explicitly --replace")
		destination.mkdir(parents = True, exist_ok = True)
		with tempfile.TemporaryDirectory(prefix = "grid-report-", dir = "/tmp") as temporary:
			stage = Path(temporary)
			if args.publish_only:
				confidence = confidence_evaluation(table, result, args, trait)
				make_confidence_figure(confidence, stage, trait)
				if not confidence["curve"].empty:
					workbook(confidence, stage / "grid.confidence.xlsx")
			else:
				make_figures(evaluation, table, result, stage, args, trait)
			result["metadata"]["scoring_files"] = publish_scoring_files(prepared, trait, stage)
			write_model(result["bundle"], result["metadata"], stage / "grid.model.rds")
			write_rds({"format": "GRID-matches-1", "primary_arm": result["bundle"]["primary_arm"],
				"matches": result["matches"], "index_base": 0, "unmatched_index": -1,
				"interpretation": "weighted OOF residual references; not biological twins"}, stage / "grid.matches.rds")
			write_rds(table, stage / "grid.test_individuals.rds")
			write_rds(result["roles"], stage / "grid.training_roles.rds")
			feature_names = list(dict.fromkeys(name for values in prepared["feature_groups"][trait].values() for name in values))
			donor = result["donor_oof"].merge(d[["eid", *feature_names]], on = "eid", how = "left", validate = "one_to_one")
			write_rds(donor, stage / "grid.donor_residuals.rds")
			explanations = table[[c for c in table if c not in methods or c in ["CSx", "GRID_policy", result["bundle"]["primary_arm"]]]]
			explanations = explanations.merge(d[["eid", *feature_names]], on = "eid", how = "left", validate = "one_to_one")
			write_rds(explanations, stage / "grid.individual_explanations.rds")
			contract = result["metadata"]["feature_contract"]
			(stage / "grid.feature_contract.json").write_text(json.dumps(json_safe(contract), ensure_ascii=False, indent=2) + "\n")
			(stage / "grid.individual_explanations.manifest.json").write_text(json.dumps({
				"format": "GRID-features-1", "trait": trait,
				"feature_contract_sha256": contract["contract_sha256"],
				"data_sha256": sha256(stage / "grid.individual_explanations.rds"),
				"producer": "GRID_report_frozen_test_features"}, indent=2) + "\n")
			score_columns = ["eid", "split", "CSx", result["bundle"]["primary_arm"], "GRID_policy", "selected", "released"]
			for col in ("confidence_percentile","confidence_validated","high_confidence","confidence_label",
				"confidence_candidate","expected_rmse","eligible_match_count","matched_count","family_ess","support"):
				if col in table:
					score_columns.append(col)
			write_rds(table[score_columns], stage / "grid.scores.rds")
			workbook({"role_counts": result["role_counts"], "model_tuning": result["tuning"],
				"donor_crossfit": result["crossfit_audit"], "retrieval": result["retrieval_audit"],
				"baseline_groups": result.get("baseline_groups", pd.DataFrame()),
				"method_label_budget": result.get("method_label_budget", pd.DataFrame()),
				"metric_features": result["metric_features"], "independent_audit": result["calibration_audit"],
				"low_error_audit": result.get("low_error_audit", pd.DataFrame()),
				"confidence_audit": result.get("confidence_audit", pd.DataFrame()),
				"annotation_qc": prepared.get("metadata", {}).get("evolution", {}).get(trait, {}).get("annotation_qc", pd.DataFrame()),
				"comparators": pd.DataFrame([pf_status])}, stage / "grid.training.xlsx")
			for path in stage.rglob("*"):
				# Stage files are complete before replacing a same-named formal result.
				if not path.is_file():
					continue
				target = destination / path.relative_to(stage)
				target.parent.mkdir(parents = True, exist_ok = True)
				shutil.copyfile(path, target)
		write_rds(prepared["cohort"][["eid", "family_id", "split", "ancestry"]], root / "grid.split.rds")
		log("DONE", "report", f"{trait}; n={len(table):,}; PRSformer={pf_status['status']}")


def project(args):
	if not args.model_file or not args.data_file or len(args.traits) != 1:
		raise ValueError("predict requires --model-file, --data-file and one --trait")
	bundle, metadata = read_model(args.model_file)
	if metadata.get("trait") != args.traits[0]:
		raise ValueError(f"Requested trait {args.traits[0]} differs from saved model trait {metadata.get('trait')}")
	provenance = validate_feature_manifest(args.feature_manifest, args.data_file, metadata)
	d = identified(named_table(args.data_file), "Projection")
	if "family_id" not in d:
		d["family_id"] = d.eid
	if "ancestry" not in d:
		d["ancestry"] = "UNASSIGNED"
	result = load_module("abm").predict_new(bundle, d)
	result["predictions"]["feature_provenance_status"] = provenance["status"]
	output = Path(args.out_root) / args.traits[0] / "projection"
	if output.exists() and any(output.iterdir()) and not args.replace:
		raise ValueError("Projection output already exists; choose another --out-root or --replace")
	output.mkdir(parents = True, exist_ok = True)
	write_rds(result["predictions"], output / "grid.predictions.rds")
	write_rds({"matches": result["matches"], "index_base": 0, "unmatched_index": -1,
		"model_training_sha256": metadata["training_data_sha256"]}, output / "grid.matches.rds")
	log("DONE", "predict", f"{len(d):,} individuals; no outcome column required")


# 🚩 One CLI shared by the shell entry and reproducible stage calls
def csv_values(text, cast = float):
	return [cast(x.strip()) for x in text.split(",") if x.strip()]


def parser():
	p = argparse.ArgumentParser(description = __doc__)
	p.add_argument("--stage", choices = ["prepare", "evolution", "fit", "report", "all", "check", "predict"], default = "all")
	p.add_argument("--check-device", action="store_true", help="Exercise GRID Transformer and donor-attention forward/backward without reading cohort data")
	p.add_argument("--abm-backend", choices=["selective_attention", "reference"], default="selective_attention")
	p.add_argument("--device", default="cuda", help="GRID ABM device (cuda[:index] default); no silent CPU fallback")
	p.add_argument("--attention-epochs", type=int, default=20)
	p.add_argument("--attention-patience", type=int, default=4)
	p.add_argument("--attention-batch", type=int, default=256)
	p.add_argument("--attention-width", type=int, default=64)
	p.add_argument("--attention-heads", type=int, default=4)
	p.add_argument("--attention-layers", type=int, default=2)
	p.add_argument("--attention-dropout", type=float, default=.1)
	p.add_argument("--attention-lr", type=float, default=.0005)
	p.add_argument("--attention-reconstruction", type=float, default=.1)
	p.add_argument("--donor-block", type=int, default=8192)
	p.add_argument("--traits", "--trait", default = "height,ldl,t2dm")
	p.add_argument("--cache-dir", default = "/tmp/grid-cache/grid")
	p.add_argument("--out-root", default = "/mnt/d/analysis/grid/grid")
	p.add_argument("--prsformer-root", default = None)
	p.add_argument("--publish-only", action="store_true", help="Publish ABM models, predictions and confidence diagnostics; cross-method performance is evaluated by grid.sh final")
	p.add_argument("--pheno-file", default = "/mnt/d/data/ukb/phe/Rdata/all.rds")
	p.add_argument("--score-dir", default = "/mnt/d/data/ukb/pgs")
	p.add_argument("--pca-file", default = "/mnt/d/data/ukb/pca_proj/ukb.discodivas.pca.tsv.gz")
	p.add_argument("--ancestry-file", default = "/mnt/d/data/ukb/pca_proj/ukb.ancestry.auto.tsv.gz")
	p.add_argument("--group-col", default = "genetic_ancestry")
	p.add_argument("--dir-gen", default = "/mnt/f/gen/ukb/37/hap")
	p.add_argument("--gwas-dir", default = "/mnt/f/gwas/4grid/common")
	p.add_argument("--snpinfo", default = "/mnt/f/refLD/csx/snpinfo_mult_1kg_hm3")
	p.add_argument("--weights-file")
	p.add_argument("--data-file")
	p.add_argument("--split-file")
	p.add_argument("--split-group-file")
	p.add_argument("--keep")
	p.add_argument("--remove", default = "/mnt/d/files/ukb.exclude.id")
	p.add_argument("--covariates", default = "age,sex,PC1,PC2")
	p.add_argument("--distance-pcs", type = int, default = 10)
	p.add_argument("--height-col", default = "height")
	p.add_argument("--ldl-col", default = "ldl")
	p.add_argument("--t2dm-col", default = "auto", help = "Explicit baseline 0/1 column, or derive from t2dm.Yr2e/Yt2e")
	p.add_argument("--chrs", default = "1-22")
	p.add_argument("--annotation-file", help="Optional canonical annotations; omitted by default for CSx/PC reference matching")
	p.add_argument("--build", choices = ["GRCh37"], default = "GRCh37")
	p.add_argument("--allow-proxy-only", action = "store_true", help="For an explicit annotation-file: permit inadequate age coverage")
	p.add_argument("--include-proxy", action = "store_true")
	p.add_argument("--min-age-variants", type=int, default=1,
		help="Minimum nonzero-weight age assignments; a data-availability threshold, not sufficient evidence of predictive value")
	p.add_argument("--age-permutation-mode", choices=["auto", "chr-maf", "chr-only"], default="auto",
		help="auto labels missing-MAF controls exploratory; chr-maf requires complete reference MAF")
	p.add_argument("--plink2", default = "plink2")
	p.add_argument("--threads", type = int, default = 4)
	p.add_argument("--score-memory", type = int, default = 2048)
	p.add_argument("--seed", type = int, default = 20260904)
	p.add_argument("--coverage", type = float, default = .5)
	p.add_argument("--coverages", type = csv_values, default = [.2, .4, .5, .6, .8, 1.])
	p.add_argument("--train-role-fractions", type = csv_values, default = [.60, .15, .15, .10])
	p.add_argument("--folds", type = int, default = 5)
	p.add_argument("--k-grid", type = lambda x: csv_values(x, int), default = [16, 32, 64])
	p.add_argument("--alpha-grid", type = csv_values, default = [0., .25, .5, .75, 1.])
	p.add_argument("--radius-grid", type = csv_values, default = [1., 2., 4.])
	p.add_argument("--max-dims", type = int, choices = [12], default = 12, help = "Fixed 2+4+2+4 block slots for controlled comparisons")
	p.add_argument("--retrieval", choices = ["cuda", "kd_tree", "hnsw"], default = "cuda")
	p.add_argument("--query-batch", type = int, default = 1024)
	p.add_argument("--ann-recall-min", type = float, default = .95)
	p.add_argument("--ann-audit-n", type = int, default = 64)
	p.add_argument("--hgb-iterations", type = int, default = 150)
	p.add_argument("--gate-iterations", type = int, default = 100)
	p.add_argument("--min-role-n", type = int, default = 20)
	p.add_argument("--audit-bootstrap", type = int, default = 500)
	p.add_argument("--bootstrap", type = int, default = 200, help = "Paired test family bootstrap draws; 0 skips test intervals")
	p.add_argument("--plot-dpi", type = int, default = 160)
	p.add_argument("--model-file")
	p.add_argument("--feature-manifest", help="Predict: JSON with trait, saved feature_contract_sha256 and data_sha256")
	p.add_argument("--replace", action = "store_true")
	p.add_argument("--quiet", action = "store_true")
	return p


def main():
	args = parser().parse_args()
	args.traits = [x.strip().lower() for x in args.traits.split(",") if x.strip()]
	if not args.traits or len(set(args.traits)) != len(args.traits) or set(args.traits) - {"height", "ldl", "t2dm"}:
		raise ValueError("Use unique traits from height,ldl,t2dm")
	args.covariates = ",".join(x.strip() for x in args.covariates.split(",") if x.strip())
	if args.prsformer_root is None:
		args.prsformer_root = str(Path(args.out_root) / "benchmark")
	if args.bootstrap != 0 and args.bootstrap < 20:
		raise ValueError("Use --bootstrap 0 or at least 20 draws")
	if args.threads < 1 or args.distance_pcs < 1:
		raise ValueError("threads and distance-pcs must be positive")
	if args.min_age_variants < 1:
		raise ValueError("min-age-variants must be at least one")
	os.umask(0o077)
	# Check cheap native prerequisites before importing GPU libraries or reading
	# the large phenotype RDS. --check-device is deliberately independent of data.
	if not args.check_device and args.stage in {"all", "check"} and not args.data_file:
		missing = []
		if not args.weights_file:
			for trait in args.traits:
				path = Path(args.score_dir) / trait / "1csx.scores.provenance.json"
				if not path.is_file():
					missing.append(str(path))
		if missing:
			raise ValueError("Native input prerequisites missing before cohort loading: " + ", ".join(missing) +
				". Create score provenance with 1.csx.sh --stage score using existing weights.")
	from threadpoolctl import threadpool_limits
	with threadpool_limits(limits = args.threads):
		if args.check_device or args.stage in {"all", "fit", "check"}:
			config = model_configuration(args, {}, args.traits[0], require_features=False)
			log("DONE", "device_check", json.dumps(load_module("abm").preflight(config)))
			if args.check_device:
				return
		if args.stage == "predict":
			project(args)
		elif args.stage == "check":
			# The temporary roster from this check is deliberately isolated from the
			# real run; no model or formal result is fitted or overwritten.
			with tempfile.TemporaryDirectory(prefix = "grid-check-", dir = "/tmp") as temporary:
				args.cache_dir = temporary
				prepared = load_module("data").prepare_data(args)
				for trait in args.traits:
					groups = prepared["feature_groups"][trait]
					load_module("abm")._configuration(model_configuration(args, groups, trait, require_features = bool(args.data_file)))
			log("DONE", "check", "cohort, split and dependencies checked; optional annotation/SNP alignment is checked during evolution")
		elif args.stage == "all":
			prepared = prepare(args)
			prepared = evolution(args, prepared)
			fit(args, prepared)
			report(args, prepared)
		elif args.stage == "prepare":
			prepare(args)
		elif args.stage == "evolution":
			evolution(args)
		elif args.stage == "fit":
			fit(args)
		elif args.stage == "report":
			report(args)


if __name__ == "__main__":
	try:
		main()
	except (ValueError, FileNotFoundError, ImportError) as error:
		print(f"ERROR: {error}", file = sys.stderr)
		raise SystemExit(2)
