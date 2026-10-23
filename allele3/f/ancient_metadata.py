"""Strict, offline sample-metadata preparation for Akbari et al. (2026).

The published STable1 and the supplied PSAM define sample identity. All PSAM
rows are retained in PSAM order; eligibility is an explicit ``include`` flag.
qpAdm validity never changes eligibility for geographic trajectories.
"""
from __future__ import annotations

from collections import Counter
import csv
import gzip
import hashlib
import json
import math
import os
from pathlib import Path


SHEET = "STable1(22274.individuals)"
COL = {
    "sample_id": "Genetic version identifier",
    "individual_id": "Unique individual identifier",
    "age_bp": "Date mean in BP in years before 1950 CE",
    "region": "Broad geographic region (N=North, C=Central, SW=Southwest, SE=Southeast, E=East)",
    "unrelated": "In unrelated set?",
    "data_source": "Data source",
    "data_types": "Contributing data types",
    "target_coverage": "Mean coverage on 1.15M autosomal targets",
}
MODERN_SOURCE = "1. Previously published data (modern)"
ANCIENT_SOURCES = {
    "1. Previously published data (ancient)",
    "2. Previously published sample with new shotgun data",
    "3. Previously published sample with new capture data",
    "4. New sample (capture)",
    "4. New sample (shotgun)",
}
IQS_HEADERS = ("IQS", "Imputation quality score")
REGIONS = {"N", "C", "SW", "SE", "E"}
MISSING = {"", ".", "..", "na", "n/a", "nan", "none", "null"}
COMPONENTS = {"3way": ("WHG", "EEF", "STEPPE"),
              "4way": ("ANF", "EHG", "WHG", "ICR")}
CORE_FIELDS = [
    "sample_id", "cohort", "age_bp", "include", "stratum", "region",
    "individual_id", "unrelated", "data_source", "data_types", "target_coverage",
    "iqs", "sex_psam", "exclusion_reasons", "qpadm_model", "qpadm_valid",
    "qpadm_reason", "qpadm_left_or_right", "qpadm_pval",
]
QC_FIELDS = ["sample_id", "cohort", "include", "exclusion_reasons", "qpadm_valid", "qpadm_reason"]
SOURCES = {
    "paper": "https://doi.org/10.1038/s41586-026-10358-1",
    "dataset": "https://doi.org/10.7910/DVN/7RVV9N",
    "readme": "https://dataverse.harvard.edu/api/access/datafile/13596811",
    "public_supplement": "https://media.springernature.com/original/springer-static/esm/art%3A10.1038%2Fs41586-026-10358-1/MediaObjects/41586_2026_10358_MOESM3_ESM.xlsx",
}


def _text(value):
    return "" if value is None else str(value).strip()


def _missing(value):
    return _text(value).lower() in MISSING


def _finite(value):
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _number(value, where, optional=False, minimum=0):
    if optional and _missing(value):
        return "."
    number = _finite(value)
    if number is None or number < minimum:
        raise ValueError(f"Expected a finite number >= {minimum} for {where}: {value!r}")
    return format(number, ".15g")


def _flag(value, where):
    value = _text(value).lower()
    if value in {"1", "1.0", "yes", "true"}:
        return "1"
    if value in {"0", "0.0", "no", "false"}:
        return "0"
    raise ValueError(f"Expected official 0/1 (or Yes/No) for {where}: {value!r}")


def _identifier(value, where):
    value = _text(value)
    if _missing(value) or any(c.isspace() for c in value):
        raise ValueError(f"Missing or whitespace-containing sample identifier at {where}")
    return value


def _open_text(path):
    return (gzip.open(path, "rt", encoding="utf-8-sig", newline="")
            if str(path).endswith((".gz", ".bgz"))
            else open(path, "rt", encoding="utf-8-sig", newline=""))


def _read_psam(path):
    samples, seen, header = [], set(), None
    with _open_text(path) as handle:
        for line_number, line in enumerate(handle, 1):
            if not line.strip() or line.startswith("##"):
                continue
            fields = line.rstrip("\r\n").split("\t")
            if header is None:
                header = [x.lstrip("#") for x in fields]
                if len(set(header)) != len(header) or "IID" not in header:
                    raise ValueError(f"PSAM needs one exact IID/#IID column: {path}")
                continue
            if len(fields) != len(header):
                raise ValueError(f"PSAM column count mismatch at {path}:{line_number}")
            row = dict(zip(header, fields))
            sid = _identifier(row["IID"], f"{path}:{line_number}")
            if sid in seen:
                raise ValueError(f"Duplicate PSAM sample ID: {sid}")
            seen.add(sid)
            samples.append((sid, row.get("SEX", ".")))
    if not samples:
        raise ValueError(f"PSAM contains no samples: {path}")
    return samples


def _read_supplement(path):
    try:
        import openpyxl
    except ImportError as exc:
        raise ValueError("Reading the public supplementary XLSX requires openpyxl") from exc
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET not in book.sheetnames:
            raise ValueError(f"Expected exact public supplementary sheet {SHEET!r}")
        records = iter(book[SHEET].iter_rows(values_only=True))
        next(records, None)  # Published title occupies row 1.
        names = [_text(v) for v in next(records, ())]
        meaningful = [name for name in names if name]
        if len(set(meaningful)) != len(meaningful):
            raise ValueError("Duplicate metadata header columns in supplementary row 2")
        missing = set(COL.values()) - set(names)
        if missing:
            raise ValueError("Missing exact supplementary row-2 columns: " + "; ".join(sorted(missing)))
        iqs_columns = [name for name in IQS_HEADERS if name in names]
        if len(iqs_columns) > 1:
            raise ValueError("Ambiguous IQS columns; provide one explicitly named IQS column")
        iqs_column = iqs_columns[0] if iqs_columns else None
        result = {}
        for row_number, cells in enumerate(records, 3):
            if all(_missing(value) for value in cells):
                continue
            row = {name: value for name, value in zip(names, cells) if name}
            sid = _identifier(row[COL["sample_id"]], f"{SHEET}, row {row_number}")
            if sid in result:
                raise ValueError(f"Duplicate supplementary sample ID: {sid}")
            result[sid] = row
        return result, iqs_column, meaningful
    finally:
        book.close()


def _read_qpadm(path, model):
    if path is None:
        if model is not None:
            raise ValueError("ancestry_model requires qpadm_path")
        return {}, None
    if model is not None and model not in COMPONENTS:
        raise ValueError("ancestry_model must be '3way' or '4way'")
    with _open_text(path) as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        names = reader.fieldnames or []
        if len(set(names)) != len(names):
            raise ValueError("Duplicate qpAdm header columns")
        matches = [m for m, components in COMPONENTS.items()
                   if set(components).issubset(names)]
        if len(matches) != 1:
            raise ValueError("qpAdm table must contain exactly one of the official 3way/4way bases")
        inferred = matches[0]
        if model is not None and model != inferred:
            raise ValueError(f"qpAdm model mismatch: declared {model}, columns describe {inferred}")
        model = inferred
        required = {"Genetic version identifier", "Left or Right", "pval",
                    *COMPONENTS[model], *(x + ".SE" for x in COMPONENTS[model])}
        if not required.issubset(names):
            raise ValueError("Missing exact qpAdm columns: " + ", ".join(sorted(required - set(names))))
        result = {}
        for row_number, row in enumerate(reader, 2):
            if None in row or any(value is None for value in row.values()):
                raise ValueError(f"Malformed qpAdm row {row_number}")
            sid = _identifier(row["Genetic version identifier"], f"qpAdm row {row_number}")
            if sid in result:
                raise ValueError(f"Duplicate qpAdm sample ID: {sid}")
            result[sid] = row
    return result, model


def _qpadm_fields(sid, rows, model):
    result = dict(qpadm_model=model or ".", qpadm_valid="0", qpadm_reason="not_requested",
                  qpadm_left_or_right=".", qpadm_pval=".")
    if model is None:
        return result
    for component in COMPONENTS[model]:
        result["qpadm_" + component.lower()] = "."
        result["qpadm_" + component.lower() + "_se"] = "."
    if sid not in rows:
        result["qpadm_reason"] = "missing_qpadm_row"
        return result
    row, reasons = rows[sid], []
    role = row["Left or Right"]
    result.update(qpadm_left_or_right=role, qpadm_pval=row["pval"])
    if role == "Yes":
        reasons.append("left_or_right_reference_role")
    elif role != "No":
        reasons.append("unknown_left_or_right_role")
    pval = _finite(row["pval"])
    if pval is None or not 0 <= pval <= 1:
        reasons.append("missing_or_invalid_pval")
    elif pval <= 0.05:
        reasons.append("pval_not_above_0.05")
    for component in COMPONENTS[model]:
        result["qpadm_" + component.lower()] = row[component]
        result["qpadm_" + component.lower() + "_se"] = row[component + ".SE"]
        number = _finite(row[component])
        if number is None:
            reasons.append("nonfinite_proportion:" + component)
        elif not 0 <= number <= 1:
            reasons.append("out_of_range_proportion:" + component)
    result.update(qpadm_valid=str(int(not reasons)), qpadm_reason=";".join(reasons) or "valid")
    return result


def _fingerprint(path):
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return {"path": str(path), "size": path.stat().st_size, "sha256": digest.hexdigest()}


def prepare_metadata(psam_path, supplement_xlsx, out_dir, qpadm_path=None, ancestry_model=None):
    """Return manifest and paths; write every PSAM sample, in exact PSAM order.

    ``include=1`` requires the official unrelated flag. If an explicitly named
    IQS column exists, additionally require finite IQS > 0.9. The current public
    STable1 contains no IQS, so no per-sample IQS threshold can be recomputed.
    An ancient sample needs a valid broad region for geographic stratification.
    qpAdm p > 0.05, finite [0,1] proportions and a non-reference target role are
    required only for ``qpadm_valid=1``, never for geographic ``include``.
    """
    psam = Path(psam_path).expanduser().resolve()
    supplement = Path(supplement_xlsx).expanduser().resolve()
    qpadm = Path(qpadm_path).expanduser().resolve() if qpadm_path is not None else None
    for path in (psam, supplement, qpadm):
        if path is not None and not path.is_file():
            raise FileNotFoundError(path)
    samples = _read_psam(psam)
    source_rows, iqs_column, source_columns = _read_supplement(supplement)
    absent = [sid for sid, _ in samples if sid not in source_rows]
    if absent:
        raise ValueError(f"{len(absent)} PSAM IDs missing from supplementary metadata: {', '.join(absent[:10])}")
    ancestry, model = _read_qpadm(qpadm, ancestry_model)
    rows, counts, reasons_count = [], Counter({key: 0 for key in (
        "samples_total", "ancient_total", "modern_total", "included", "excluded",
        "ancient_included", "modern_included", "ancient_excluded", "modern_excluded",
        "modern_without_region", "qpadm_valid", "qpadm_valid_included")}), Counter()
    for sid, sex in samples:
        source = source_rows[sid]
        data_source = _text(source[COL["data_source"]])
        if data_source == MODERN_SOURCE:
            cohort = "modern"
        elif data_source in ANCIENT_SOURCES:
            cohort = "ancient"
        else:
            raise ValueError(f"Unrecognized official Data source for {sid}: {data_source!r}")
        age = _number(source[COL["age_bp"]], f"age_bp of {sid}")
        unrelated = _flag(source[COL["unrelated"]], f"unrelated flag of {sid}")
        region = _text(source[COL["region"]])
        region = "." if _missing(region) else region
        if region != "." and region not in REGIONS:
            raise ValueError(f"Unknown broad region for {sid}: {region!r}")
        reasons = [] if unrelated == "1" else ["related_in_official_table"]
        if region == "." and cohort == "ancient":
            reasons.append("ancient_region_missing")
        if region == "." and cohort == "modern":
            counts["modern_without_region"] += 1
        iqs = "."
        if iqs_column:
            iqs = _number(source.get(iqs_column), f"IQS of {sid}", optional=True)
            if iqs == ".":
                reasons.append("iqs_missing")
            elif not 0 <= float(iqs) <= 1:
                raise ValueError(f"IQS outside [0,1] for {sid}")
            elif float(iqs) <= 0.9:
                reasons.append("iqs_not_above_0.9")
        row = dict(sample_id=sid, cohort=cohort, age_bp=age, include=str(int(not reasons)),
                   stratum=region if region != "." else ("modern" if cohort == "modern" else "."),
                   region=region, individual_id=_text(source[COL["individual_id"]]) or ".",
                   unrelated=unrelated, data_source=data_source,
                   data_types=_text(source[COL["data_types"]]) or ".",
                   target_coverage=_number(source[COL["target_coverage"]], f"target coverage of {sid}", optional=True),
                   iqs=iqs, sex_psam=sex, exclusion_reasons=";".join(reasons) or ".")
        row.update(_qpadm_fields(sid, ancestry, model))
        rows.append(row)
        counts["samples_total"] += 1
        counts[cohort + "_total"] += 1
        counts["included" if not reasons else "excluded"] += 1
        counts[cohort + ("_included" if not reasons else "_excluded")] += 1
        counts["qpadm_valid"] += int(row["qpadm_valid"])
        if row["qpadm_valid"] == "1" and not reasons:
            counts["qpadm_valid_included"] += 1
        reasons_count.update(reasons)
    destination = Path(out_dir).expanduser().resolve()
    paths = {"metadata_path": destination / "sample_metadata.tsv",
             "qc_path": destination / "metadata_qc.tsv",
             "manifest_path": destination / "metadata_qc.json"}
    if set(paths.values()) & {psam, supplement, qpadm}:
        raise ValueError("Metadata output paths must not overwrite inputs")
    destination.mkdir(parents=True, exist_ok=True)
    partials = {key: Path(str(path) + ".partial") for key, path in paths.items()}
    if any(path.exists() for path in partials.values()):
        raise ValueError("A metadata .partial output exists; resolve the interrupted run before retrying")
    fields = CORE_FIELDS + (["qpadm_" + x.lower() for x in COMPONENTS[model]] +
                           ["qpadm_" + x.lower() + "_se" for x in COMPONENTS[model]] if model else [])
    manifest = {
        "schema_version": 1, **{key: str(path) for key, path in paths.items()},
        "inputs": {"psam": _fingerprint(psam), "supplement_xlsx": _fingerprint(supplement),
                   "qpadm": _fingerprint(qpadm) if qpadm else None},
        "sheet": SHEET, "header_row_1based": 2, "source_columns": source_columns,
        "counts": dict(counts), "exclusions": dict(reasons_count), "fields": fields,
        "supplement_rows_not_in_psam": len(source_rows) - len(samples),
        "qpadm_rows_not_in_psam": len(set(ancestry) - {sid for sid, _ in samples}),
        "ancestry_model": model, "iqs_column": iqs_column,
        "qc_rule": "Official unrelated flag equals 1; ancient region required; explicit IQS > 0.9 only if a named IQS column exists.",
        "qpadm_rule": "pval > 0.05, all proportions finite and in [0,1], Left or Right == No. No clipping. Invalid qpAdm never excludes geographic trajectories.",
        "sample_order": "exact_psam_order", "date_units": "years BP before 1950 CE; point estimate",
        "sources": SOURCES,
        "limitations": [
            "Public STable1 supplies point dates, not individual date distributions; geographic strata do not fully control ancestry or migration.",
            "Target coverage is mean depth on 1.15M capture targets, not genome-wide depth or IQS.",
            "The public released dataset contains ancient samples and EUR1KG controls; the original study also used 5935 UK Biobank samples absent from this release.",
            "Current public STable1 has no IQS column. Upstream study sample QC included IQS > 0.9, but individual values cannot be independently rechecked here." if not iqs_column else "IQS filtering used the explicitly named supplied column; its provenance must be documented.",
            "qpAdm reference-role rows are preserved and flagged; composition estimates are model-specific and never hard ancestry labels.",
        ],
    }
    try:
        for key, names in (("metadata_path", fields), ("qc_path", QC_FIELDS)):
            with partials[key].open("w", encoding="utf-8", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=names, delimiter="\t", lineterminator="\n", extrasaction="ignore")
                writer.writeheader()
                writer.writerows(rows)
        with partials["manifest_path"].open("w", encoding="utf-8") as handle:
            json.dump(manifest, handle, ensure_ascii=False, indent=2, allow_nan=False)
            handle.write("\n")
        for key in ("metadata_path", "qc_path", "manifest_path"):
            os.replace(partials[key], paths[key])
    finally:
        for path in partials.values():
            if path.exists():
                path.unlink()
    return manifest
