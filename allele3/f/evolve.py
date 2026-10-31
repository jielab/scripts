#!/usr/bin/env python3
"""Ancient metadata, trajectories, PCA and optional Atlas analysis."""
from __future__ import annotations

# The shared utility filename is sorted first; use a stable import name in Python.
import importlib.util as _common_import
import sys as _common_sys
from pathlib import Path as _CommonPath
if "allele_common" not in _common_sys.modules:
    _common_spec = _common_import.spec_from_file_location(
        "allele_common", _CommonPath(__file__).with_name("0.common.py"))
    _common_module = _common_import.module_from_spec(_common_spec)
    _common_sys.modules[_common_spec.name] = _common_module
    try:
        _common_spec.loader.exec_module(_common_module)
    except BaseException:
        _common_sys.modules.pop(_common_spec.name, None)
        raise


# ========================================================================
# evolve_ancient
# ========================================================================

# Ancient metadata, selection joins, candidate selection and count audits.





# ========================================================================
# Ancient metadata
# ========================================================================

# Strict, offline sample-metadata preparation for Akbari et al. (2026).
# 
# The published STable1 and the supplied PSAM define sample identity. All PSAM
# rows are retained in PSAM order; eligibility is an explicit ``include`` flag.
# qpAdm validity never changes eligibility for geographic trajectories.


from collections import Counter
import csv
import gzip
import hashlib
import json
import math
import os
from pathlib import Path


SHEET = "STable1(22274.individuals)"
COL = {
    "sample_id": "Genetic version identifier",
    "individual_id": "Unique individual identifier",
    "age_bp": "Date mean in BP in years before 1950 CE",
    "region": "Broad geographic region (N=North, C=Central, SW=Southwest, SE=Southeast, E=East)",
    "unrelated": "In unrelated set?",
    "data_source": "Data source",
    "data_types": "Contributing data types",
    "target_coverage": "Mean coverage on 1.15M autosomal targets",
}
MODERN_SOURCE = "1. Previously published data (modern)"
ANCIENT_SOURCES = {
    "1. Previously published data (ancient)",
    "2. Previously published sample with new shotgun data",
    "3. Previously published sample with new capture data",
    "4. New sample (capture)",
    "4. New sample (shotgun)",
}
IQS_HEADERS = ("IQS", "Imputation quality score")
REGIONS = {"N", "C", "SW", "SE", "E"}
MISSING = {"", ".", "..", "na", "n/a", "nan", "none", "null"}
COMPONENTS = {"3way": ("WHG", "EEF", "STEPPE"),
              "4way": ("ANF", "EHG", "WHG", "ICR")}
CORE_FIELDS = [
    "sample_id", "cohort", "age_bp", "include", "stratum", "region",
    "individual_id", "unrelated", "data_source", "data_types", "target_coverage",
    "iqs", "sex_psam", "exclusion_reasons", "qpadm_model", "qpadm_valid",
    "qpadm_reason", "qpadm_left_or_right", "qpadm_pval",
]
QC_FIELDS = ["sample_id", "cohort", "include", "exclusion_reasons", "qpadm_valid", "qpadm_reason"]
SOURCES = {
    "paper": "https://doi.org/10.1038/s41586-026-10358-1",
    "dataset": "https://doi.org/10.7910/DVN/7RVV9N",
    "readme": "https://dataverse.harvard.edu/api/access/datafile/13596811",
    "public_supplement": "https://media.springernature.com/original/springer-static/esm/art%3A10.1038%2Fs41586-026-10358-1/MediaObjects/41586_2026_10358_MOESM3_ESM.xlsx",
}


def _text(value):
    return "" if value is None else str(value).strip()


def _missing(value):
    return _text(value).lower() in MISSING


def _finite(value):
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _number(value, where, optional=False, minimum=0):
    if optional and _missing(value):
        return "."
    number = _finite(value)
    if number is None or number < minimum:
        raise ValueError(f"Expected a finite number >= {minimum} for {where}: {value!r}")
    return format(number, ".15g")


def _flag(value, where):
    value = _text(value).lower()
    if value in {"1", "1.0", "yes", "true"}:
        return "1"
    if value in {"0", "0.0", "no", "false"}:
        return "0"
    raise ValueError(f"Expected official 0/1 (or Yes/No) for {where}: {value!r}")


def _identifier(value, where):
    value = _text(value)
    if _missing(value) or any(c.isspace() for c in value):
        raise ValueError(f"Missing or whitespace-containing sample identifier at {where}")
    return value


def _open_text(path):
    return (gzip.open(path, "rt", encoding="utf-8-sig", newline="")
            if str(path).endswith((".gz", ".bgz"))
            else open(path, "rt", encoding="utf-8-sig", newline=""))


def _read_psam(path):
    samples, seen, header = [], set(), None
    with _open_text(path) as handle:
        for line_number, line in enumerate(handle, 1):
            if not line.strip() or line.startswith("##"):
                continue
            fields = line.rstrip("\r\n").split("\t")
            if header is None:
                header = [x.lstrip("#") for x in fields]
                if len(set(header)) != len(header) or "IID" not in header:
                    raise ValueError(f"PSAM needs one exact IID/#IID column: {path}")
                continue
            if len(fields) != len(header):
                raise ValueError(f"PSAM column count mismatch at {path}:{line_number}")
            row = dict(zip(header, fields))
            sid = _identifier(row["IID"], f"{path}:{line_number}")
            if sid in seen:
                raise ValueError(f"Duplicate PSAM sample ID: {sid}")
            seen.add(sid)
            samples.append((sid, row.get("SEX", ".")))
    if not samples:
        raise ValueError(f"PSAM contains no samples: {path}")
    return samples


def _read_supplement(path):
    try:
        import openpyxl
    except ImportError as exc:
        raise ValueError("Reading the public supplementary XLSX requires openpyxl") from exc
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET not in book.sheetnames:
            raise ValueError(f"Expected exact public supplementary sheet {SHEET!r}")
        records = iter(book[SHEET].iter_rows(values_only=True))
        next(records, None)  # Published title occupies row 1.
        names = [_text(v) for v in next(records, ())]
        meaningful = [name for name in names if name]
        if len(set(meaningful)) != len(meaningful):
            raise ValueError("Duplicate metadata header columns in supplementary row 2")
        missing = set(COL.values()) - set(names)
        if missing:
            raise ValueError("Missing exact supplementary row-2 columns: " + "; ".join(sorted(missing)))
        iqs_columns = [name for name in IQS_HEADERS if name in names]
        if len(iqs_columns) > 1:
            raise ValueError("Ambiguous IQS columns; provide one explicitly named IQS column")
        iqs_column = iqs_columns[0] if iqs_columns else None
        result = {}
        for row_number, cells in enumerate(records, 3):
            if all(_missing(value) for value in cells):
                continue
            row = {name: value for name, value in zip(names, cells) if name}
            sid = _identifier(row[COL["sample_id"]], f"{SHEET}, row {row_number}")
            if sid in result:
                raise ValueError(f"Duplicate supplementary sample ID: {sid}")
            result[sid] = row
        return result, iqs_column, meaningful
    finally:
        book.close()


def _read_qpadm(path, model):
    if path is None:
        if model is not None:
            raise ValueError("ancestry_model requires qpadm_path")
        return {}, None
    if model is not None and model not in COMPONENTS:
        raise ValueError("ancestry_model must be '3way' or '4way'")
    with _open_text(path) as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        names = reader.fieldnames or []
        if len(set(names)) != len(names):
            raise ValueError("Duplicate qpAdm header columns")
        matches = [m for m, components in COMPONENTS.items()
                   if set(components).issubset(names)]
        if len(matches) != 1:
            raise ValueError("qpAdm table must contain exactly one of the official 3way/4way bases")
        inferred = matches[0]
        if model is not None and model != inferred:
            raise ValueError(f"qpAdm model mismatch: declared {model}, columns describe {inferred}")
        model = inferred
        required = {"Genetic version identifier", "Left or Right", "pval",
                    *COMPONENTS[model], *(x + ".SE" for x in COMPONENTS[model])}
        if not required.issubset(names):
            raise ValueError("Missing exact qpAdm columns: " + ", ".join(sorted(required - set(names))))
        result = {}
        for row_number, row in enumerate(reader, 2):
            if None in row or any(value is None for value in row.values()):
                raise ValueError(f"Malformed qpAdm row {row_number}")
            sid = _identifier(row["Genetic version identifier"], f"qpAdm row {row_number}")
            if sid in result:
                raise ValueError(f"Duplicate qpAdm sample ID: {sid}")
            result[sid] = row
    return result, model


def _qpadm_fields(sid, rows, model):
    result = dict(qpadm_model=model or ".", qpadm_valid="0", qpadm_reason="not_requested",
                  qpadm_left_or_right=".", qpadm_pval=".")
    if model is None:
        return result
    for component in COMPONENTS[model]:
        result["qpadm_" + component.lower()] = "."
        result["qpadm_" + component.lower() + "_se"] = "."
    if sid not in rows:
        result["qpadm_reason"] = "missing_qpadm_row"
        return result
    row, reasons = rows[sid], []
    role = row["Left or Right"]
    result.update(qpadm_left_or_right=role, qpadm_pval=row["pval"])
    if role == "Yes":
        reasons.append("left_or_right_reference_role")
    elif role != "No":
        reasons.append("unknown_left_or_right_role")
    pval = _finite(row["pval"])
    if pval is None or not 0 <= pval <= 1:
        reasons.append("missing_or_invalid_pval")
    elif pval <= 0.05:
        reasons.append("pval_not_above_0.05")
    for component in COMPONENTS[model]:
        result["qpadm_" + component.lower()] = row[component]
        result["qpadm_" + component.lower() + "_se"] = row[component + ".SE"]
        number = _finite(row[component])
        if number is None:
            reasons.append("nonfinite_proportion:" + component)
        elif not 0 <= number <= 1:
            reasons.append("out_of_range_proportion:" + component)
    result.update(qpadm_valid=str(int(not reasons)), qpadm_reason=";".join(reasons) or "valid")
    return result


def _metadata_fingerprint(path):
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return {"path": str(path), "size": path.stat().st_size, "sha256": digest.hexdigest()}


def prepare_metadata(psam_path, supplement_xlsx, out_dir, qpadm_path=None, ancestry_model=None):
    """Return manifest and paths; write every PSAM sample, in exact PSAM order.

    ``include=1`` requires the official unrelated flag. If an explicitly named
    IQS column exists, additionally require finite IQS > 0.9. The current public
    STable1 contains no IQS, so no per-sample IQS threshold can be recomputed.
    An ancient sample needs a valid broad region for geographic stratification.
    qpAdm p > 0.05, finite [0,1] proportions and a non-reference target role are
    required only for ``qpadm_valid=1``, never for geographic ``include``.
    """
    psam = Path(psam_path).expanduser().resolve()
    supplement = Path(supplement_xlsx).expanduser().resolve()
    qpadm = Path(qpadm_path).expanduser().resolve() if qpadm_path is not None else None
    for path in (psam, supplement, qpadm):
        if path is not None and not path.is_file():
            raise FileNotFoundError(path)
    samples = _read_psam(psam)
    source_rows, iqs_column, source_columns = _read_supplement(supplement)
    absent = [sid for sid, _ in samples if sid not in source_rows]
    if absent:
        raise ValueError(f"{len(absent)} PSAM IDs missing from supplementary metadata: {', '.join(absent[:10])}")
    ancestry, model = _read_qpadm(qpadm, ancestry_model)
    rows, counts, reasons_count = [], Counter({key: 0 for key in (
        "samples_total", "ancient_total", "modern_total", "included", "excluded",
        "ancient_included", "modern_included", "ancient_excluded", "modern_excluded",
        "modern_without_region", "qpadm_valid", "qpadm_valid_included")}), Counter()
    for sid, sex in samples:
        source = source_rows[sid]
        data_source = _text(source[COL["data_source"]])
        if data_source == MODERN_SOURCE:
            cohort = "modern"
        elif data_source in ANCIENT_SOURCES:
            cohort = "ancient"
        else:
            raise ValueError(f"Unrecognized official Data source for {sid}: {data_source!r}")
        age = _number(source[COL["age_bp"]], f"age_bp of {sid}")
        unrelated = _flag(source[COL["unrelated"]], f"unrelated flag of {sid}")
        region = _text(source[COL["region"]])
        region = "." if _missing(region) else region
        if region != "." and region not in REGIONS:
            raise ValueError(f"Unknown broad region for {sid}: {region!r}")
        reasons = [] if unrelated == "1" else ["related_in_official_table"]
        if region == "." and cohort == "ancient":
            reasons.append("ancient_region_missing")
        if region == "." and cohort == "modern":
            counts["modern_without_region"] += 1
        iqs = "."
        if iqs_column:
            iqs = _number(source.get(iqs_column), f"IQS of {sid}", optional=True)
            if iqs == ".":
                reasons.append("iqs_missing")
            elif not 0 <= float(iqs) <= 1:
                raise ValueError(f"IQS outside [0,1] for {sid}")
            elif float(iqs) <= 0.9:
                reasons.append("iqs_not_above_0.9")
        row = dict(sample_id=sid, cohort=cohort, age_bp=age, include=str(int(not reasons)),
                   stratum=region if region != "." else ("modern" if cohort == "modern" else "."),
                   region=region, individual_id=_text(source[COL["individual_id"]]) or ".",
                   unrelated=unrelated, data_source=data_source,
                   data_types=_text(source[COL["data_types"]]) or ".",
                   target_coverage=_number(source[COL["target_coverage"]], f"target coverage of {sid}", optional=True),
                   iqs=iqs, sex_psam=sex, exclusion_reasons=";".join(reasons) or ".")
        row.update(_qpadm_fields(sid, ancestry, model))
        rows.append(row)
        counts["samples_total"] += 1
        counts[cohort + "_total"] += 1
        counts["included" if not reasons else "excluded"] += 1
        counts[cohort + ("_included" if not reasons else "_excluded")] += 1
        counts["qpadm_valid"] += int(row["qpadm_valid"])
        if row["qpadm_valid"] == "1" and not reasons:
            counts["qpadm_valid_included"] += 1
        reasons_count.update(reasons)
    destination = Path(out_dir).expanduser().resolve()
    paths = {"metadata_path": destination / "sample_metadata.tsv",
             "qc_path": destination / "metadata_qc.tsv",
             "manifest_path": destination / "metadata_qc.json"}
    if set(paths.values()) & {psam, supplement, qpadm}:
        raise ValueError("Metadata output paths must not overwrite inputs")
    destination.mkdir(parents=True, exist_ok=True)
    partials = {key: Path(str(path) + ".partial") for key, path in paths.items()}
    if any(path.exists() for path in partials.values()):
        raise ValueError("A metadata .partial output exists; resolve the interrupted run before retrying")
    fields = CORE_FIELDS + (["qpadm_" + x.lower() for x in COMPONENTS[model]] +
                           ["qpadm_" + x.lower() + "_se" for x in COMPONENTS[model]] if model else [])
    manifest = {
        "schema_version": 1, **{key: str(path) for key, path in paths.items()},
        "inputs": {"psam": _metadata_fingerprint(psam), "supplement_xlsx": _metadata_fingerprint(supplement),
                   "qpadm": _metadata_fingerprint(qpadm) if qpadm else None},
        "sheet": SHEET, "header_row_1based": 2, "source_columns": source_columns,
        "counts": dict(counts), "exclusions": dict(reasons_count), "fields": fields,
        "supplement_rows_not_in_psam": len(source_rows) - len(samples),
        "qpadm_rows_not_in_psam": len(set(ancestry) - {sid for sid, _ in samples}),
        "ancestry_model": model, "iqs_column": iqs_column,
        "qc_rule": "Official unrelated flag equals 1; ancient region required; explicit IQS > 0.9 only if a named IQS column exists.",
        "qpadm_rule": "pval > 0.05, all proportions finite and in [0,1], Left or Right == No. No clipping. Invalid qpAdm never excludes geographic trajectories.",
        "sample_order": "exact_psam_order", "date_units": "years BP before 1950 CE; point estimate",
        "sources": SOURCES,
        "limitations": [
            "Public STable1 supplies point dates, not individual date distributions; geographic strata do not fully control ancestry or migration.",
            "Target coverage is mean depth on 1.15M capture targets, not genome-wide depth or IQS.",
            "The public released dataset contains ancient samples and EUR1KG controls; the original study also used 5935 UK Biobank samples absent from this release.",
            "Current public STable1 has no IQS column. Upstream study sample QC included IQS > 0.9, but individual values cannot be independently rechecked here." if not iqs_column else "IQS filtering used the explicitly named supplied column; its provenance must be documented.",
            "qpAdm reference-role rows are preserved and flagged; composition estimates are model-specific and never hard ancestry labels.",
        ],
    }
    try:
        for key, names in (("metadata_path", fields), ("qc_path", QC_FIELDS)):
            with partials[key].open("w", encoding="utf-8", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=names, delimiter="\t", lineterminator="\n", extrasaction="ignore")
                writer.writeheader()
                writer.writerows(rows)
        with partials["manifest_path"].open("w", encoding="utf-8") as handle:
            json.dump(manifest, handle, ensure_ascii=False, indent=2, allow_nan=False)
            handle.write("\n")
        for key in ("metadata_path", "qc_path", "manifest_path"):
            os.replace(partials[key], paths[key])
    finally:
        for path in partials.values():
            if path.exists():
                path.unlink()
    return manifest


# ========================================================================
# Ancient join
# ========================================================================

# Streaming selection joins and bounded candidate design for allele3 evolution.
# 
# The sampling universe is modern multi-ALT sites. It is NOT an all-site
# case/control enrichment analysis. SQLite bounds RAM; no genotype matrix or
# AlphaGenome request occurs in this module.


from array import array
from bisect import bisect_left
from collections import Counter, OrderedDict, defaultdict
import csv
import gzip
import heapq
import itertools
import json
import math
import os
from pathlib import Path
import sqlite3
import sys
import time

from allele_common import dump_json, fingerprint, fmt, open_text, stable_hash, tsv_writer
from allele_common import canonical_chrom
from allele_common import MAPPING_AUTOSOMES as AUTOSOMES, BASES, GenomeMapper, MappingError, orient

SCHEMA = "allele3-evolution-1"
SELECTION_FIELDS = ("CHROM", "POS", "REF", "ALT", "ANC", "ID", "RSID", "AF", "S", "SE",
                    "X", "P_X", "POSTERIOR", "FDR", "CHI2_BE", "FILTER")
JOIN_FIELDS = ["site_id", "chrom", "pos", "ref", "chrom37", "pos37", "ref37", "alt37",
               "strand37to38", "reference_changed", "source_ref38", "anchor_ref", "anchor_alt",
               "extra_alt", "extra_ac", "extra_an", "extra_af", "extra_nfe_ac", "extra_nfe_an",
               "extra_nfe_af", "anchor_alt_modern_status", "anchor_alt_modern_af", "selection_id",
               "rsid", "ancestral_base37", "selection_pooled_af", "s", "se", "x", "p_x", "posterior",
               "fdr", "chi2_be", "selection_filter", "context", "canonical_context", "region_flags",
               "has_non_snv", "an_relative_spread", "n_modern_alt", "consequence_terms", "gene_symbols",
               "estimated_increasing_base", "estimated_decreasing_base", "extra_selection_status", "population_class"]
CANDIDATE_FIELDS = JOIN_FIELDS + ["group", "match_set", "match_stratum", "locus_block",
                                "selection_evidence", "comparison_universe"]


def table_rows(path, required=()):
    """Read public TSVs with ## metadata and a #CHROM or ordinary header."""
    with open_text(path) as handle:
        lines = (line for line in handle if line.strip() and not line.startswith("##"))
        reader = csv.DictReader(lines, delimiter="\t")
        if not reader.fieldnames:
            raise ValueError(f"Empty table: {path}")
        reader.fieldnames = [name.lstrip("#") for name in reader.fieldnames]
        missing = set(required) - set(reader.fieldnames)
        if missing:
            raise ValueError(f"Missing columns in {path}: {', '.join(sorted(missing))}")
        for row in reader:
            if None in row or any(v is None for v in row.values()):
                raise ValueError(f"Malformed/truncated TSV row in {path}: line {reader.line_num}")
            yield row


def number(value, default=None):
    try:
        x = float(value)
    except (TypeError, ValueError):
        return default
    return x if math.isfinite(x) else default


def _modern_sites(shards):
    for shard in shards:
        site_file = Path(shard) / "multiallelic_sites.tsv.gz"
        allele_file = Path(shard) / "multiallelic_alleles.tsv.gz"
        allele_groups = itertools.groupby(table_rows(allele_file, ("chrom", "pos", "ref", "alt", "status")),
                                          key=lambda r: (r["chrom"], r["pos"], r["ref"]))
        for site in table_rows(site_file, ("chrom", "pos", "ref", "n_alt")):
            try:
                key, group = next(allele_groups)
            except StopIteration as exc:
                raise ValueError("Modern site/allele files have different site counts") from exc
            if key != (site["chrom"], site["pos"], site["ref"]):
                raise ValueError("Modern site/allele file order or keys disagree")
            alleles = [r for r in group if r["status"] == "observed_pass"]
            if len(alleles) != int(site["n_alt"]):
                raise ValueError("Modern occupancy differs between site and allele tables")
            yield site, alleles
        if next(allele_groups, None) is not None:
            raise ValueError("Unmatched trailing modern allele records")


def _stage_identity(kind, inputs, parameters):
    return dict(schema=SCHEMA, join_schema=2 if kind == "join" else None, kind=kind, inputs={k: fingerprint(v) for k, v in inputs.items()},
                parameters=parameters)


def completed_stage(out, identity):
    path = Path(out) / "manifest.json"
    if not path.exists():
        return None
    result = json.loads(path.read_text())
    if result.get("identity") != identity:
        raise ValueError(f"Inputs/settings changed at {out}; choose a new output directory")
    for name, fp in result.get("outputs", {}).items():
        p = Path(out) / name
        if not p.exists() or fingerprint(p) != fp:
            raise ValueError(f"Completed output changed: {p}; restore it or choose a new output directory")
    return result


def finish_stage(out, identity, result, filenames):
    out = Path(out)
    result.update(identity=identity, outputs={name: fingerprint(out / name) for name in filenames},
                  manifest_path=str((out / "manifest.json").resolve()))
    dump_json(out / "manifest.json", result)
    return result


def _contrast_rows(row, mapped, site, alleles, qc):
    """Shared exact allele/numeric rules for both join strategies. REF is checked by the caller."""
    ref, alt = row["REF"], row["ALT"]
    chrom = canonical_chrom(row["CHROM"])
    strand = mapped["strand37to38"]
    anchor_ref, anchor_alt = orient(ref, strand), orient(alt, strand)
    by_base = {a["alt"]: a for a in alleles}
    anchor_modern = by_base.get(anchor_alt)
    s, posterior, pooled_af = number(row["S"]), number(row["POSTERIOR"]), number(row["AF"])
    x, se = number(row["X"]), number(row["SE"])
    if (s is None or posterior is None or pooled_af is None or x is None or se is None or se <= 0
            or not 0 <= posterior <= 1 or not 0 <= pooled_af <= 1):
        qc["invalid_selection_numeric"] += 1
        return
    for allele in alleles:
        extra = allele["alt"]
        if extra in (anchor_ref, anchor_alt):
            continue
        qc["extra_allele_contrasts"] += 1
        result = dict(site_id=mapped["site_id"], chrom=mapped["chrom38"], pos=mapped["pos38"],
                      ref=mapped["ref38"], chrom37=chrom, pos37=row["POS"], ref37=ref, alt37=alt,
                      strand37to38=strand, reference_changed=mapped["reference_changed"],
                      source_ref38=mapped["source_ref38"], anchor_ref=anchor_ref, anchor_alt=anchor_alt,
                      extra_alt=extra, extra_ac=allele["ac"], extra_an=allele["an"], extra_af=allele["af"],
                      extra_nfe_ac=allele.get("ac_nfe", "."), extra_nfe_an=allele.get("an_nfe", "."),
                      extra_nfe_af=allele.get("af_nfe", "."),
                      anchor_alt_modern_status=("observed_pass" if anchor_modern else
                                               "source_reference_frequency_unknown" if anchor_alt == mapped["source_ref38"]
                                               else "not_observed_pass"),
                      anchor_alt_modern_af=anchor_modern["af"] if anchor_modern else ".",
                      selection_id=row["ID"], rsid=row["RSID"], ancestral_base37=row["ANC"],
                      selection_pooled_af=row["AF"], s=row["S"], se=row["SE"], x=row["X"],
                      p_x=row["P_X"], posterior=row["POSTERIOR"], fdr=row["FDR"], chi2_be=row["CHI2_BE"],
                      selection_filter=row["FILTER"], context=mapped["context38"],
                      canonical_context=mapped["canonical_context"], region_flags=site["region_flags"],
                      has_non_snv=site["has_non_snv"], an_relative_spread=site["an_relative_spread"],
                      n_modern_alt=site["n_alt"], consequence_terms=allele["consequence_terms"],
                      gene_symbols=allele["gene_symbols"],
                      estimated_increasing_base=anchor_alt if s > 0 else anchor_ref if s < 0 else ".",
                      estimated_decreasing_base=anchor_ref if s > 0 else anchor_alt if s < 0 else ".",
                      extra_selection_status="not_estimated_by_anchor_contrast",
                      population_class=site.get("population_class", "NA"))
        yield result


def _join_modern_first(modern_out, selection_path, out_dir, fasta38, fasta37,
                       chain38to37, chain37to38, *, mapper=None):
    """Join all QC published selection contrasts to modern multi-ALT loci.

    Different ancient ALT contrasts at one coordinate remain different rows.
    A contrast's s is NEVER assigned to extra_alt. Missing summary entries are
    not declared absent from PVAR. Only point SNVs and autosomes are accepted.
    """
    from scan import load_run
    modern = load_run(modern_out)
    source_build = modern["options"]["build"]
    inputs = dict(modern_manifest=Path(modern_out) / "run.json", selection=selection_path,
                  fasta38=fasta38, fasta37=fasta37, chain38to37=chain38to37, chain37to38=chain37to38)
    inputs.update({"fai38": str(fasta38) + ".fai", "fai37": str(fasta37) + ".fai"})
    identity = _stage_identity("join", inputs, dict(modern_signature=modern["signature"], source_build=source_build,
                                                                 strategy="modern-first"))
    out = Path(out_dir).resolve()
    out.mkdir(parents=True, exist_ok=True)
    if old := completed_stage(out, identity):
        return old
    start = time.monotonic()
    db_path = out / "modern_index.sqlite"
    tmp = out / "modern_index.sqlite.partial"
    if tmp.exists():
        tmp.unlink()
    db = sqlite3.connect(tmp)
    db.execute("PRAGMA journal_mode=OFF")
    db.execute("PRAGMA synchronous=OFF")
    db.execute("PRAGMA cache_size=-65536")
    db.execute("CREATE TABLE modern (chrom37 TEXT, pos37 INTEGER, payload TEXT, matched INTEGER DEFAULT 0, "
               "PRIMARY KEY(chrom37,pos37)) WITHOUT ROWID")
    qc = Counter()
    own_mapper = mapper is None
    mapper = mapper or GenomeMapper(fasta38, fasta37, chain38to37, chain37to38)
    rejects_handle, rejects = tsv_writer(out / "mapping_rejections.tsv.gz", ["chrom", "pos", "ref", "reason"])
    try:
        with rejects_handle:
            for site, alleles in _modern_sites(modern["shards"]):
                qc["modern_multi_alt_sites"] += 1
                try:
                    mapped = mapper.map(site["chrom"], int(site["pos"]), site["ref"], source_build)
                except MappingError as exc:
                    qc["mapping_" + exc.reason] += 1
                    rejects.writerow({**{k: site[k] for k in ("chrom", "pos", "ref")}, "reason": exc.reason})
                    continue
                oriented = []
                for allele in alleles:
                    a = dict(allele)
                    a["alt"] = orient(a["alt"], mapped["source_to38"])
                    oriented.append(a)
                payload = dict(mapped=mapped, site=site, alleles=oriented)
                try:
                    db.execute("INSERT INTO modern(chrom37,pos37,payload) VALUES(?,?,?)",
                               (mapped["chrom37"], mapped["pos37"], json.dumps(payload, separators=(",", ":"))))
                except sqlite3.IntegrityError as exc:
                    raise ValueError("Two modern sites map to the same ancient coordinate; ambiguous source merge") from exc
                qc["mapped_sites"] += 1
                qc["reference_changed_sites"] += mapped["reference_changed"]
                if qc["modern_multi_alt_sites"] % 100000 == 0:
                    db.commit()
                    print(f"[map] {qc['modern_multi_alt_sites']:,} multi-ALT sites", flush=True)
        db.commit()
        # Reuse the exact, bounded coordinate cache to avoid millions of
        # payload SELECT misses on mounted filesystems. Mapping still covers
        # every modern site; no selection statistic or allele is prefiltered.
        db.execute("CREATE TEMP VIEW coordinates AS SELECT chrom37 AS source_chrom, "
                   "pos37 AS source_pos, NULL AS reason FROM modern")
        cache = _CoordinateCache(db)
        # Keep misses inside a transaction too. Before the first matching UPDATE,
        # implicit sqlite transactions otherwise acquire a file lock for every
        # SELECT; millions of misses are especially slow on WSL mounted drives.
        db.execute("BEGIN")
        handle, writer = tsv_writer(out / "selection_intersection.tsv.gz", JOIN_FIELDS)
        with handle:
            for row in table_rows(selection_path, SELECTION_FIELDS):
                qc["selection_rows"] += 1
                if qc["selection_rows"] % 1000000 == 0:
                    db.commit()
                    db.execute("BEGIN")
                    print(f"[selection] {qc['selection_rows']:,} published rows streamed", flush=True)
                ref, alt = row["REF"], row["ALT"]
                if ref not in BASES or alt not in BASES or ref == alt:
                    qc["selection_non_snv"] += 1
                    continue
                chrom = canonical_chrom(row["CHROM"])
                if not cache.contains(chrom, int(row["POS"])):
                    continue
                record = db.execute("SELECT payload FROM modern WHERE chrom37=? AND pos37=?",
                                    (chrom, int(row["POS"]))).fetchone()
                if record is None:
                    continue
                payload = json.loads(record[0])
                mapped, site, alleles = payload["mapped"], payload["site"], payload["alleles"]
                if ref != mapped["ref37"]:
                    qc["selection_ref37_mismatch"] += 1
                    continue
                db.execute("UPDATE modern SET matched=1 WHERE chrom37=? AND pos37=?", (chrom, int(row["POS"])))
                qc["matched_contrasts"] += 1
                writer.writerows(_contrast_rows(row, mapped, site, alleles, qc))
        db.commit()
        qc["mapped_sites_with_summary"] = db.execute("SELECT COUNT(*) FROM modern WHERE matched=1").fetchone()[0]
        qc["mapped_sites_without_summary"] = qc["mapped_sites"] - qc["mapped_sites_with_summary"]
        db.close()
        os.replace(tmp, db_path)
    except BaseException:
        db.close()
        raise
    finally:
        if own_mapper:
            mapper.close()
    return finish_stage(out, identity, dict(qc=dict(qc), elapsed_seconds=time.monotonic() - start,
                         strategy="modern-first", qc_scope="All modern multi-ALT sites receive full mapping; reference implementation",
                         performance=dict(fullmap_calls=qc["modern_multi_alt_sites"], index_bytes=db_path.stat().st_size,
                                          coordinate_cache=cache.metrics()),
                         intersection_path=str(out / "selection_intersection.tsv.gz"),
                         index_path=str(db_path), universe="modern multi-ALT sites; no all-site enrichment denominator; "
                         "extra must be a directly observed passing source ALT, never inferred from the source REF"),
                        ["selection_intersection.tsv.gz", "mapping_rejections.tsv.gz", "modern_index.sqlite"])


def _process_memory():
    """Observed process memory, not a promise about chain-library allocations."""
    result = {"rss_bytes": None, "lifetime_peak_rss_bytes": None}
    try:
        with open("/proc/self/statm", encoding="ascii") as handle:
            result["rss_bytes"] = int(handle.read().split()[1]) * os.sysconf("SC_PAGE_SIZE")
    except (OSError, ValueError, IndexError):
        pass
    try:
        import resource
        peak = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
        result["lifetime_peak_rss_bytes"] = int(peak * (1 if sys.platform == "darwin" else 1024))
    except (ImportError, OSError, ValueError):
        pass
    return result


class _CoordinateCache:
    """Exact per-chromosome uint arrays, with a strict array-byte LRU budget.

    No Python set/list of millions of positions is constructed. An individual
    chromosome larger than the budget uses indexed SQLite membership instead,
    preserving every target. Array capacity is allocated exactly before filling.
    The tiny chromosome/count dictionary and LRU bookkeeping are reported apart
    from array bytes. Source shards need not be globally chromosome sorted.
    """
    def __init__(self, db, max_bytes=64 * 1024 * 1024, max_chromosomes=2):
        if max_bytes < 1 or max_chromosomes < 1:
            raise ValueError("Coordinate cache limits must be positive")
        self.db, self.max_bytes, self.max_chromosomes = db, max_bytes, max_chromosomes
        self.arrays, self.array_bytes, self.stats = OrderedDict(), 0, Counter()
        self.info = {}
        for chrom, count, largest in db.execute(
                "SELECT source_chrom,COUNT(*),MAX(source_pos) FROM coordinates "
                "WHERE reason IS NULL GROUP BY source_chrom"):
            code = "I" if largest <= (1 << (8 * array("I").itemsize)) - 1 else "Q"
            self.info[chrom] = (count, code)
        self.stats["all_coordinates_uncompressed_uint_bytes"] = sum(
            count * array(code).itemsize for count, code in self.info.values())

    def _update_peak(self):
        self.stats["peak_array_bytes"] = max(self.stats["peak_array_bytes"], self.array_bytes)
        overhead = sys.getsizeof(self.arrays) + sys.getsizeof(self.info)
        overhead += sum(sys.getsizeof(k) + sys.getsizeof(v) for k, v in self.info.items())
        overhead += sum(sys.getsizeof(k) for k in self.arrays)
        self.stats["peak_bookkeeping_bytes"] = max(self.stats["peak_bookkeeping_bytes"], overhead)

    def contains(self, chrom, pos):
        self.stats["membership_lookups"] += 1
        info = self.info.get(chrom)
        if info is None:
            return False
        values = self.arrays.get(chrom)
        if values is not None:
            self.arrays.move_to_end(chrom)
            self.stats["cached_array_lookups"] += 1
        else:
            count, code = info
            size = sys.getsizeof(array(code)) + count * array(code).itemsize
            if size > self.max_bytes:
                self.stats["disk_membership_queries"] += 1
                return self.db.execute("SELECT 1 FROM coordinates WHERE source_chrom=? AND source_pos=? "
                                       "AND reason IS NULL", (chrom, pos)).fetchone() is not None
            while self.arrays and (len(self.arrays) >= self.max_chromosomes or
                                   self.array_bytes + size > self.max_bytes):
                _, removed = self.arrays.popitem(last=False)
                self.array_bytes -= sys.getsizeof(removed)
                del removed
                self.stats["array_evictions"] += 1
            values = array(code, [0]) * count
            for index, (position,) in enumerate(self.db.execute(
                    "SELECT source_pos FROM coordinates WHERE source_chrom=? AND reason IS NULL ORDER BY source_pos",
                    (chrom,))):
                values[index] = position
            self.arrays[chrom] = values
            self.array_bytes += sys.getsizeof(values)
            if self.array_bytes > self.max_bytes:
                raise RuntimeError("Coordinate array allocator exceeded its declared budget")
            self.stats["array_loads"] += 1
            self._update_peak()
        index = bisect_left(values, pos)
        return index < len(values) and values[index] == pos

    def metrics(self):
        self._update_peak()
        return dict(self.stats, array_budget_bytes=self.max_bytes,
                    max_cached_chromosomes=self.max_chromosomes,
                    retained_array_bytes=self.array_bytes,
                    indexed_coordinates=sum(v[0] for v in self.info.values()),
                    indexed_chromosomes=len(self.info))


def _join_selection_first(modern_out, selection_path, out_dir, fasta38, fasta37,
                          chain38to37, chain37to38, *, mapper=None):
    from scan import load_run
    modern = load_run(modern_out)
    source_build = modern["options"]["build"]
    if source_build not in ("GRCh37", "GRCh38"):
        raise ValueError("source_build must be GRCh37 or GRCh38")
    inputs = dict(modern_manifest=Path(modern_out) / "run.json", selection=selection_path,
                  fasta38=fasta38, fasta37=fasta37, chain38to37=chain38to37, chain37to38=chain37to38)
    inputs.update({"fai38": str(fasta38) + ".fai", "fai37": str(fasta37) + ".fai"})
    identity = _stage_identity("join", inputs, dict(modern_signature=modern["signature"], source_build=source_build,
                                                   strategy="selection-first"))
    out = Path(out_dir).resolve()
    out.mkdir(parents=True, exist_ok=True)
    if old := completed_stage(out, identity):
        return old
    start = time.monotonic()
    before_memory = _process_memory()
    db_path, tmp = out / "selection_index.sqlite", out / "selection_index.sqlite.partial"
    if tmp.exists():
        tmp.unlink()
    db = sqlite3.connect(tmp)
    db.execute("PRAGMA journal_mode=OFF")
    db.execute("PRAGMA synchronous=OFF")
    db.execute("PRAGMA cache_size=-65536")
    db.execute("PRAGMA temp_store=FILE")
    db.execute("PRAGMA mmap_size=0")
    db.execute("CREATE TABLE coordinates (chrom37 TEXT, pos37 INTEGER, source_chrom TEXT, source_pos INTEGER, "
               "strand TEXT, reason TEXT, modern_seen INTEGER DEFAULT 0, "
               "PRIMARY KEY(chrom37,pos37)) WITHOUT ROWID")
    db.execute("CREATE TABLE contrasts (chrom37 TEXT, pos37 INTEGER, row_no INTEGER, payload TEXT, "
               "PRIMARY KEY(chrom37,pos37,row_no)) WITHOUT ROWID")
    qc = Counter()
    own_mapper = mapper is None
    payload_bytes = 0
    try:
        mapper = mapper or GenomeMapper(fasta38, fasta37, chain38to37, chain37to38)
        coord_handle, coord_rejects = tsv_writer(out / "selection_coordinate_rejections.tsv.gz",
                                                ["chrom37", "pos37", "reason"])
        previous_key, previous_state = None, None
        with coord_handle:
            for row_no, row in enumerate(table_rows(selection_path, SELECTION_FIELDS), 1):
                qc["selection_rows"] += 1
                if row_no % 1000000 == 0:
                    db.commit()
                    print(f"[selection index] {row_no:,} published rows streamed", flush=True)
                ref, alt = row["REF"], row["ALT"]
                if ref not in BASES or alt not in BASES or ref == alt:
                    qc["selection_non_snv"] += 1
                    continue
                qc["selection_snv_rows"] += 1
                chrom37, pos37 = canonical_chrom(row["CHROM"]), int(row["POS"])
                key = (chrom37, pos37)
                state = (previous_state if key == previous_key else db.execute(
                    "SELECT source_chrom,source_pos,strand,reason FROM coordinates WHERE chrom37=? AND pos37=?",
                    key).fetchone())
                if state is None:
                    qc["selection_unique_coordinates"] += 1
                    try:
                        if chrom37 not in AUTOSOMES:
                            raise MappingError("non_autosomal_source")
                        if pos37 < 1:
                            raise MappingError("invalid_source_position")
                        if source_build == "GRCh38":
                            qc["selection_coordinate_map_calls"] += 1
                            coords = mapper.map_coordinates(chrom37, pos37, "GRCh37")
                            state = (coords["chrom38"], coords["pos38"], coords["strand37to38"], None)
                        else:
                            # Identity lookup in the input build; do both chains only on a modern hit.
                            qc["selection_coordinate_direct_grch37"] += 1
                            state = (chrom37, pos37, None, None)
                        qc["selection_coordinates_indexed"] += 1
                    except MappingError as exc:
                        state = (None, None, None, exc.reason)
                        qc["selection_coordinate_rejected_" + exc.reason] += 1
                        coord_rejects.writerow(dict(chrom37=chrom37, pos37=pos37, reason=exc.reason))
                    db.execute("INSERT INTO coordinates(chrom37,pos37,source_chrom,source_pos,strand,reason) "
                               "VALUES(?,?,?,?,?,?)", (*key, *state))
                previous_key, previous_state = key, state
                if state[3] is not None:
                    qc["selection_snv_rows_without_usable_coordinate"] += 1
                    continue
                # Fixed-order values avoid copying repeated column names millions of times.
                # Numeric/filter/posterior validity is intentionally NOT a prefilter.
                payload = json.dumps([row[k] for k in SELECTION_FIELDS], separators=(",", ":"))
                payload_bytes += len(payload)  # ensure_ascii=True, hence length is exact UTF-8 byte count.
                db.execute("INSERT INTO contrasts VALUES(?,?,?,?)", (*key, row_no, payload))
                qc["selection_snv_rows_indexed"] += 1
        db.commit()
        db.execute("CREATE UNIQUE INDEX coordinates_source ON coordinates(source_chrom,source_pos) WHERE reason IS NULL")
        db.commit()
        cache = _CoordinateCache(db)
        after_index_memory = _process_memory()
        index_elapsed = time.monotonic() - start
        rejects_handle, rejects = tsv_writer(out / "mapping_rejections.tsv.gz", ["chrom", "pos", "ref", "reason"])
        handle, writer = tsv_writer(out / "selection_intersection.tsv.gz", JOIN_FIELDS)
        with rejects_handle, handle:
            for site, alleles in _modern_sites(modern["shards"]):
                qc["modern_multi_alt_sites"] += 1
                if qc["modern_multi_alt_sites"] % 1000000 == 0:
                    db.commit()
                    print(f"[modern intersect] {qc['modern_multi_alt_sites']:,} sites; "
                          f"{qc['modern_coordinate_hits']:,} coordinate hits", flush=True)
                chrom, pos = canonical_chrom(site["chrom"]), int(site["pos"])
                if chrom not in AUTOSOMES:
                    qc["modern_non_autosomal_sites"] += 1
                    continue
                if not cache.contains(chrom, pos):
                    qc["modern_sites_without_coordinate_hit"] += 1
                    continue
                qc["modern_coordinate_hits"] += 1
                chrom37, pos37, strand, seen = db.execute(
                    "SELECT chrom37,pos37,strand,modern_seen FROM coordinates WHERE source_chrom=? "
                    "AND source_pos=? AND reason IS NULL", (chrom, pos)).fetchone()
                qc["fullmap_calls"] += 1
                try:
                    if source_build == "GRCh38":
                        coords = dict(chrom38=chrom, pos38=pos, chrom37=chrom37, pos37=pos37,
                                      strand37to38=strand)
                        mapped = mapper.map_from_coordinates(coords, site["ref"], source_build)
                    else:
                        mapped = mapper.map(chrom, pos, site["ref"], source_build)
                except MappingError as exc:
                    qc["hit_mapping_" + exc.reason] += 1
                    rejects.writerow({**{k: site[k] for k in ("chrom", "pos", "ref")}, "reason": exc.reason})
                    continue
                if seen:
                    raise ValueError("Two modern sites map to the same ancient coordinate; ambiguous source merge")
                db.execute("UPDATE coordinates SET modern_seen=1 WHERE chrom37=? AND pos37=?", (chrom37, pos37))
                qc["modern_fullmap_passed"] += 1
                qc["hit_reference_changed_sites"] += mapped["reference_changed"]
                oriented = [dict(a, alt=orient(a["alt"], mapped["source_to38"])) for a in alleles]
                matched_ref = False
                for (payload,) in db.execute("SELECT payload FROM contrasts WHERE chrom37=? AND pos37=? ORDER BY row_no",
                                             (chrom37, pos37)):
                    row = dict(zip(SELECTION_FIELDS, json.loads(payload)))
                    if row["REF"] != mapped["ref37"]:
                        qc["selection_ref37_mismatch"] += 1
                        continue
                    matched_ref = True
                    qc["matched_contrasts"] += 1
                    writer.writerows(_contrast_rows(row, mapped, site, oriented, qc))
                qc["modern_hit_sites_with_ref_matched_summary" if matched_ref else
                   "modern_validated_hits_without_ref_matched_summary"] += 1
        db.commit()
        memory = dict(before=before_memory, after_coordinate_index=after_index_memory,
                      after_modern_stream=_process_memory(), coordinate_cache=cache.metrics(),
                      sqlite_page_cache_target_bytes=64 * 1024 * 1024,
                      note="Array budget excludes small Python bookkeeping, FASTA caches and both pyliftover chain indexes. "
                           "RSS is measured for the whole process; lifetime peak can include preceding stages.")
        db.close()
        os.replace(tmp, db_path)
    finally:
        db.close()
        if own_mapper and mapper is not None:
            mapper.close()
    performance = dict(index_bytes=db_path.stat().st_size, selection_payload_bytes=payload_bytes,
                       selection_index_seconds=index_elapsed,
                       coordinate_hits=qc["modern_coordinate_hits"], fullmap_calls=qc["fullmap_calls"],
                       selection_coordinate_map_calls=qc["selection_coordinate_map_calls"], memory=memory)
    return finish_stage(out, identity, dict(qc=dict(qc), elapsed_seconds=time.monotonic() - start,
                         strategy="selection-first", performance=performance,
                         qc_scope="All canonical published SNVs are coordinate-indexed without AF, posterior, S or FILTER preselection. "
                         "All modern multi-ALT rows are streamed. Full mapping/REF QC and mapping_rejections apply ONLY to "
                         "modern coordinate hits. Nonhits are untested, not mapping failures; no global mapped-without-summary count. "
                         "For GRCh37 modern input the coordinate index is direct; reciprocal chains are checked on hits. "
                         "Modern duplicate-coordinate rejection is likewise restricted to validated coordinate hits.",
                         intersection_path=str(out / "selection_intersection.tsv.gz"), index_path=str(db_path),
                         universe="modern multi-ALT sites; no all-site enrichment denominator; "
                         "extra must be a directly observed passing source ALT, never inferred from the source REF"),
                        ["selection_intersection.tsv.gz", "mapping_rejections.tsv.gz",
                         "selection_coordinate_rejections.tsv.gz", "selection_index.sqlite"])


def join_selection(modern_out, selection_path, out_dir, fasta38, fasta37,
                   chain38to37, chain37to38, *, mapper=None, strategy="selection-first"):
    """Intersect complete canonical selection SNVs with complete modern multi-ALT tables.

    The default uses a reciprocal coordinate prefilter and validates FASTA/REF on
    modern hits only. ``modern-first`` preserves the slower reference algorithm.
    Their accepted allele-contrast multisets are identical for valid modern runs;
    row order and mapping-QC scope differ. Both retain the same explicit observed
    source-ALT restriction, numeric validation and selection direction semantics.
    Changing strategy requires a new join directory, not a new modern scan.
    """
    methods = {"selection-first": _join_selection_first, "modern-first": _join_modern_first}
    if strategy not in methods:
        raise ValueError("join strategy must be selection-first or modern-first")
    return methods[strategy](modern_out, selection_path, out_dir, fasta38, fasta37,
                             chain38to37, chain37to38, mapper=mapper)


def _match_stratum(row):
    # Explicit coarse bins, identical between case and control. Matching the
    # extra AF is for FUNCTION comparisons only, not AF-as-outcome analyses.
    af = float(row["selection_pooled_af"])
    maf = min(af, 1 - af)
    maf_bin = min(9, int(maf / .05))
    extra_af = max(float(row["extra_af"]), 1e-12)
    an = max(float(row["extra_an"]), 1.)
    log_af_bin = math.floor(math.log10(extra_af) * 2)
    log_an_bin = math.floor(math.log10(an) * 4)
    coding = ("unknown" if row["consequence_terms"] in ("", ".", "NA") else
              str(int(any(t in row["consequence_terms"] for t in
                          ("missense", "synonymous", "stop_", "splice_", "frameshift")))))
    return f"{row['canonical_context']}|maf{maf_bin}|af{log_af_bin}|an{log_an_bin}|coding{coding}"


def _best_extra_rows(path, seed):
    """The join emits contiguous extra rows for each precise selection contrast.

    Use the same highest-AC extra rule for cases and controls before sampling.
    The expected ordering is part of the join's published file contract.
    """
    rows = table_rows(path, JOIN_FIELDS)
    groups = itertools.groupby(rows, key=lambda r: (r["site_id"], r["selection_id"]))
    for _, group in groups:
        yield max(group, key=lambda r: (float(r["extra_ac"]),
                  stable_hash([seed, r["site_id"], r["selection_id"], r["extra_alt"]])))


def _physical_block(row, size):
    return f"{row['chrom37']}:{(int(row['pos37']) - 1) // size}"


def select_candidates(intersection_path, out_dir, *, max_sites=2000, strong_posterior=.99,
                      weak_posterior=.5, min_extra_ac=5, block_bp=1000000,
                      exclude_mhc=True, allow_filter_warnings=False, seed=20261021):
    """Deterministic, bounded, pre-functional-score case/control selection.

    One strong site per proximity block, one highest-count extra per contrast,
    one matching weak-signal site from a different unused block. Proximity blocks
    are NOT LD-clumped loci; the weak-signal group is NOT proven neutral.
    Unmatched cases are retained with an explicit status and are not paired tests.
    """
    if max_sites < 2 or block_bp < 1 or min_extra_ac < 1 or not 0 <= weak_posterior < strong_posterior <= 1:
        raise ValueError("Invalid selection budget, posterior thresholds, block size or AC floor")
    parameters = dict(max_sites=max_sites, strong_posterior=strong_posterior, weak_posterior=weak_posterior,
                      min_extra_ac=min_extra_ac, block_bp=block_bp, exclude_mhc=exclude_mhc,
                      allow_filter_warnings=allow_filter_warnings, seed=seed)
    identity = _stage_identity("candidate_design", dict(intersection=intersection_path), parameters)
    out = Path(out_dir).resolve()
    out.mkdir(parents=True, exist_ok=True)
    if old := completed_stage(out, identity):
        return old
    qc, strong_blocks = Counter(), {}
    summary = defaultdict(lambda: [0, 0., 0.])

    def eligible(row):
        qc["intersection_rows"] += 1
        if number(row["extra_ac"], 0) < min_extra_ac:
            qc["below_extra_ac"] += 1
            return False
        if row["selection_filter"] != "PASS" and not allow_filter_warnings:
            qc["selection_filter_excluded"] += 1
            return False
        if row["has_non_snv"] == "1" or row["region_flags"] not in (".", "", "none"):
            qc["complex_or_flagged_site_excluded"] += 1
            return False
        if row["chrom37"] == "6" and 25000000 <= int(row["pos37"]) <= 34000000 and exclude_mhc:
            qc["mhc_excluded"] += 1
            return False
        if float(row["s"]) == 0:
            qc["zero_direction_excluded"] += 1
            return False
        return True

    def rank(row):
        return (float(row["posterior"]), abs(float(row["x"])), float(row["extra_ac"]),
                stable_hash([seed, row["site_id"], row["selection_id"], row["extra_alt"]]))

    for row in _best_extra_rows(intersection_path, seed):
        if not eligible(row):
            continue
        p = float(row["posterior"])
        group = "strong" if p >= strong_posterior else "weak_signal" if p <= weak_posterior else "intermediate"
        key = (group, row["canonical_context"])
        summary[key][0] += 1
        summary[key][1] += float(row["extra_af"])
        summary[key][2] += float(row["extra_ac"])
        qc["eligible_" + group] += 1
        if group == "strong":
            block = _physical_block(row, block_bp)
            if block not in strong_blocks or rank(row) > rank(strong_blocks[block]):
                strong_blocks[block] = row
    cases = sorted(strong_blocks.values(), key=rank, reverse=True)[:max_sites // 2]
    qc["strong_proximity_blocks"] = len(strong_blocks)
    qc["case_sites_selected"] = len(cases)
    # Keep a bounded deterministic pool of controls only for strata actually
    # needed. Reservoirs are capped at max_sites rows per stratum; total <=
    # max_sites * number of case strata (not all genome rows).
    wanted = Counter(_match_stratum(row) for row in cases)
    pools = defaultdict(list)
    case_sites = {row["site_id"] for row in cases}
    for row in _best_extra_rows(intersection_path, seed):
        # Do not double-count QC from the first pass.
        if float(row["posterior"]) > weak_posterior or row["site_id"] in case_sites:
            continue
        if number(row["extra_ac"], 0) < min_extra_ac or float(row["s"]) == 0:
            continue
        if (row["selection_filter"] != "PASS" and not allow_filter_warnings) or row["has_non_snv"] == "1" or row["region_flags"] not in (".", "", "none"):
            continue
        if exclude_mhc and row["chrom37"] == "6" and 25000000 <= int(row["pos37"]) <= 34000000:
            continue
        key = _match_stratum(row)
        if key not in wanted:
            continue
        # Smallest stable hash is an outcome-independent deterministic sample.
        h = int(stable_hash([seed, row["site_id"], row["selection_id"], row["extra_alt"]]), 16)
        item = (-h, row["site_id"] + row["selection_id"] + row["extra_alt"], row)
        pool = pools[key]
        capacity = min(max_sites, max(40, wanted[key] * 20))
        if len(pool) < capacity:
            heapq.heappush(pool, item)
        elif item[:2] > pool[0][:2]:
            heapq.heapreplace(pool, item)
    sorted_pools = {key: [item[2] for item in sorted(values, reverse=True)] for key, values in pools.items()}
    used_blocks = {_physical_block(row, block_bp) for row in cases}
    used_sites = set(case_sites)
    selected, matched = [], 0
    for index, case in enumerate(cases, 1):
        key = _match_stratum(case)
        control = next((row for row in sorted_pools.get(key, [])
                        if row["site_id"] not in used_sites and _physical_block(row, block_bp) not in used_blocks), None)
        match_set = f"pair{index:05d}" if control else "."
        for row, group in ((case, "strong"), (control, "weak_signal")):
            if row is None:
                continue
            block = _physical_block(row, block_bp)
            selected.append({**row, "group": group, "match_set": match_set, "match_stratum": key,
                             "locus_block": block, "selection_evidence": "published_strong" if group == "strong" else "weak_signal_not_neutral",
                             "comparison_universe": "matched_modern_multi_alt_functional_comparison"})
            used_blocks.add(block)
            used_sites.add(row["site_id"])
        matched += int(control is not None)
    qc["matched_pairs"] = matched
    qc["unmatched_cases"] = len(cases) - matched
    qc["selected_sites"] = len(selected)
    handle, writer = tsv_writer(out / "candidates.tsv", CANDIDATE_FIELDS)
    with handle:
        writer.writerows(selected)
    handle, writer = tsv_writer(out / "selection_context_summary.tsv", ["selection_group", "canonical_context", "extra_contrast_rows", "mean_extra_af", "mean_extra_ac"])
    with handle:
        for (group, context), (n, af_sum, ac_sum) in sorted(summary.items()):
            writer.writerow(dict(selection_group=group, canonical_context=context, extra_contrast_rows=n,
                                 mean_extra_af=af_sum / n, mean_extra_ac=ac_sum / n))
    return finish_stage(out, identity, dict(qc=dict(qc), candidate_path=str(out / "candidates.tsv"),
                         caveat="Conditional on modern multi-ALT sites; physical blocks do not establish LD independence. "
                                "Extra AF is matched for functional comparisons, so do not test AF differences in paired candidates."),
                        ["candidates.tsv", "selection_context_summary.tsv"])


def index_target_variants(candidate_path, pvar_path, high_quality_path, out_dir):
    """One sequential PVAR pass; preserve original indices for targeted PGEN reads.

    Targets include each ancient anchor and the extra base versus ancient REF.
    A missing extra is 'not_represented_in_panel', NEVER dosage zero.
    """
    identity = _stage_identity("panel_index", dict(candidates=candidate_path, pvar=pvar_path,
                                                   high_quality=high_quality_path), {})
    out = Path(out_dir).resolve()
    out.mkdir(parents=True, exist_ok=True)
    if old := completed_stage(out, identity):
        return old
    targets, links = {}, []
    for row in table_rows(candidate_path, ("site_id", "chrom37", "pos37", "ref37", "alt37", "extra_alt", "strand37to38")):
        extra37 = orient(row["extra_alt"], row["strand37to38"])
        for role, alt in (("anchor", row["alt37"]), ("extra", extra37)):
            key = (canonical_chrom(row["chrom37"]), int(row["pos37"]), row["ref37"], alt)
            if alt == key[2]:
                raise ValueError("Extra base equals ancient REF; invalid harmonization")
            targets.setdefault(key, dict(variant_id="_".join(map(str, key)), chrom=key[0], pos=key[1],
                                         ref=key[2], alt=key[3], pvar_index=".", status="not_represented_in_panel"))
            links.append(dict(site_id=row["site_id"], role=role, target_id="_".join(map(str, key)),
                              base38=row["anchor_alt"] if role == "anchor" else row["extra_alt"]))
    hq, seen = set(), set()
    for row in (table_rows(high_quality_path, ("CHROM", "POS", "ID", "REF", "ALT")) if targets else ()):
        key = (canonical_chrom(row["CHROM"]), int(row["POS"]), row["REF"], row["ALT"])
        if key in targets:
            hq.add(key)
    rows = 0
    for index, row in enumerate(table_rows(pvar_path, ("CHROM", "POS", "ID", "REF", "ALT")) if targets else ()):
        rows += 1
        key = (canonical_chrom(row["CHROM"]), int(row["POS"]), row["REF"], row["ALT"])
        if "," in row["ALT"]:
            for alt in row["ALT"].split(","):
                multi_key = (*key[:3], alt)
                if multi_key in targets:
                    # pgenlib allele_idx=1 cannot represent arbitrary multiallelic
                    # rows in this biallelic module. Do not call this absent.
                    if multi_key in seen:
                        raise ValueError(f"Duplicate exact/multiallelic PVAR contrast: {multi_key}")
                    seen.add(multi_key)
                    targets[multi_key].update(variant_id=row["ID"] + ":" + alt, pvar_index=".",
                                              status="represented_multiallelic_unsupported")
        if key in targets:
            if key in seen:
                raise ValueError(f"Duplicate exact PVAR contrast: {key}")
            seen.add(key)
            targets[key].update(variant_id=row["ID"], pvar_index=index,
                                status="assayed_pass" if key in hq else "not_in_high_quality")
        if rows % 5000000 == 0:
            print(f"[PVAR] {rows:,} records streamed; {len(seen):,}/{len(targets):,} requested contrasts located", flush=True)
    fields = ["variant_id", "pvar_index", "chrom", "pos", "ref", "alt", "status"]
    handle, writer = tsv_writer(out / "selected_variants.tsv", fields)
    with handle:
        writer.writerows(sorted(targets.values(), key=lambda r: (int(r["chrom"]), r["pos"], r["alt"])))
    link_by_id = {"_".join(map(str, k)): v for k, v in targets.items()}
    handle, writer = tsv_writer(out / "site_variant_links.tsv", ["site_id", "role", "target_id", "base38", "variant_id", "status"])
    with handle:
        for link in links:
            target = link_by_id[link["target_id"]]
            writer.writerow({**link, "variant_id": target["variant_id"], "status": target["status"]})
    counts = Counter(row["status"] for row in targets.values())
    return finish_stage(out, identity, dict(pvar_records=rows if targets else None,
                         scan_skipped_no_targets=not bool(targets), target_counts=dict(counts),
                         selected_variants_path=str(out / "selected_variants.tsv"),
                         links_path=str(out / "site_variant_links.tsv")), ["selected_variants.tsv", "site_variant_links.tsv"])


# ========================================================================
# Counts audit
# ========================================================================

# Read-only arithmetic audit and descriptive report for legacy counts.tsv.
# 
# These are conditional counts of legacy PASS/positive-AC ALT position groups.
# REF occurrence, callable sequence, release/build and evolutionary processes
# cannot be recovered from this summary. No hypothesis tests are performed.


from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation, localcontext
import base64
import csv
import gzip
import hashlib
import html
import json
import math
import os
from pathlib import Path
import re
import tempfile


TOOL_ID = "allele3.counts_audit"
SCHEMA_VERSION = 1
CATEGORIES = ("singleton", "non_singleton_biallelic", "tri_allelic", "four_plus_allelic")
TYPES = ("SNP", "INDEL")
REQUIRED = ("CHR", "TYPE", "CATEGORY", "COUNT", "TOTAL_SITES", "PERCENT")
COUNTS_AUTOSOMES = tuple(f"chr{i}" for i in range(1, 23))
STANDARD_CHROMS = COUNTS_AUTOSOMES + ("chrX", "chrY")
OUTPUT_NAMES = ("counts_audit.json", "counts_summary.tsv", "counts_by_chromosome.tsv",
                "counts_review.html", "counts_review.png", "manifest.json")
DEFINITIONS = {
    "expected_producer": "Reviewed legacy check implementation; a matching TSV header alone does not prove its producer or settings.",
    "retained_alt": "ALT retained by the legacy producer after PASS and AC > 0 filters.",
    "singleton": "Exactly one retained ALT and its AC equals 1.",
    "non_singleton_biallelic": "Exactly one retained ALT and its AC exceeds 1; the name does not establish observed REF copies.",
    "tri_allelic": "Exactly two retained ALT entries/states at the legacy position group; REF is an annotation label, not evidence of an observed third state.",
    "four_plus_allelic": "At least three retained ALT entries/states. Canonical A/C/G/T SNPs have at most three ALT bases relative to REF; observed REF is not established.",
    "INDEL": "Legacy groups use the start coordinate and were not fully FASTA-left-aligned; category size is not a reliable count of normalized polymorphic states.",
    "multi_alt": "tri_allelic + four_plus_allelic, i.e. legacy groups with at least two retained ALTs.",
    "alt_3plus": "four_plus_allelic, i.e. legacy groups with at least three retained ALTs.",
    "denominator": "TOTAL_SITES within each CHR/TYPE; aggregate fractions use summed counts divided by summed totals.",
}
LIMITATIONS = [
    "The counts file contains no release, reference-build, sample-size, population, or complete producer-configuration metadata.",
    "Rates are conditional fractions among legacy position groups, not densities per callable genomic base.",
    "REF labels do not establish observed REF copies. Two retained ALTs do not by themselves prove three observed alleles.",
    "INDEL start-coordinate groups are not fully reference-normalized; do not interpret their categories as normalized polymorphic-state counts.",
    "Missing chromosome/type groups are unavailable, not zero. X and Y are reported separately from autosomes.",
    "SNP and INDEL totals can overlap genomic positions; their sum is not asserted to count unique coordinates.",
    "Leave-one-chromosome ratios are descriptive and unadjusted for callability, ancestry, allele frequency, sequence context, or QC.",
    "No P values, confidence intervals, selection estimates, pleiotropy effects, or ancient-DNA conclusions are inferred.",
]
SUMMARY_FIELDS = ["scope", "type", "status", "observed_chromosomes", "missing_chromosomes", "total_sites"]
for _category in CATEGORIES:
    SUMMARY_FIELDS += [_category + "_count", _category + "_percent"]
SUMMARY_FIELDS += ["multi_alt_count", "multi_alt_percent", "alt_3plus_count", "alt_3plus_percent", "alt_3plus_within_multi_percent"]
CHROM_FIELDS = ["chrom", "type", "status", "total_sites"] + SUMMARY_FIELDS[6:] + [
    "reference_scope", "reference_total_sites", "multi_alt_reference_percent",
    "multi_alt_ratio_to_reference", "alt_3plus_reference_percent", "alt_3plus_ratio_to_reference",
]


def _open(path):
    if str(path).endswith((".gz", ".bgz")):
        return gzip.open(path, "rt", encoding="utf-8-sig", newline="")
    return open(path, "rt", encoding="utf-8-sig", newline="")


def _sha(path):
    hasher = hashlib.sha256()
    with open(path, "rb") as handle:
        for block in iter(lambda: handle.read(1048576), b""):
            hasher.update(block)
    return hasher.hexdigest()


def _counts_fingerprint(path):
    path = Path(path).resolve()
    return {"path": str(path), "size": path.stat().st_size, "sha256": _sha(path)}


def _chrom(value):
    value = value.strip()
    short = value[3:] if value.lower().startswith("chr") else value
    if short.isdigit() and 1 <= int(short) <= 22:
        return "chr" + str(int(short))
    if short.upper() in {"X", "Y"}:
        return "chr" + short.upper()
    if not value or any(c.isspace() for c in value):
        raise ValueError(f"Invalid chromosome label: {value!r}")
    return value


def _order(chrom):
    if chrom in STANDARD_CHROMS:
        return (STANDARD_CHROMS.index(chrom), "")
    return (100, chrom)


def _integer(value, where):
    if not re.fullmatch(r"[0-9]+", value.strip()):
        raise ValueError(f"{where}: expected a nonnegative integer, got {value!r}")
    return int(value)


def _read_counts(path):
    grouped = defaultdict(list)
    with _open(path) as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        fields = reader.fieldnames or []
        if len(fields) != len(set(fields)) or not set(REQUIRED).issubset(fields):
            raise ValueError("Counts table requires unique CHR,TYPE,CATEGORY,COUNT,TOTAL_SITES,PERCENT columns")
        for line, raw in enumerate(reader, 2):
            if None in raw or any(v is None for v in raw.values()):
                raise ValueError(f"Counts line {line}: wrong number of fields")
            row = {key: value.strip() for key, value in raw.items()}
            chrom, typ, category = _chrom(row["CHR"]), row["TYPE"], row["CATEGORY"]
            if typ not in TYPES or category not in CATEGORIES:
                raise ValueError(f"Counts line {line}: unknown TYPE/CATEGORY {typ!r}/{category!r}")
            count = _integer(row["COUNT"], f"line {line} COUNT")
            total = _integer(row["TOTAL_SITES"], f"line {line} TOTAL_SITES")
            percent_text = row["PERCENT"]
            if total == 0 and percent_text in {".", "NA", "N/A"}:
                percent = None
            else:
                if not re.fullmatch(r"[0-9]+(?:\.[0-9]+)?", percent_text):
                    raise ValueError(f"line {line} PERCENT: expected a finite plain decimal in [0,100]")
                try:
                    percent = Decimal(percent_text)
                except InvalidOperation as exc:
                    raise ValueError(f"line {line}: invalid PERCENT") from exc
                if not percent.is_finite() or not 0 <= percent <= 100:
                    raise ValueError(f"line {line} PERCENT: expected a finite decimal in [0,100]")
            grouped[(chrom, typ)].append(dict(line=line, category=category, count=count,
                                              total=total, percent=percent, percent_text=percent_text))
    if not grouped:
        raise ValueError("Counts table contains no data")
    groups, audits, errors = [], [], []
    for (chrom, typ), rows in sorted(grouped.items(), key=lambda item: (_order(item[0][0]), item[0][1])):
        label = f"{chrom}/{typ}"
        if Counter(r["category"] for r in rows) != Counter(CATEGORIES):
            errors.append(f"{label}: each of the four categories must occur exactly once")
            continue
        totals = {r["total"] for r in rows}
        if len(totals) != 1:
            errors.append(f"{label}: repeated TOTAL_SITES values disagree")
            continue
        total = next(iter(totals))
        if sum(r["count"] for r in rows) != total:
            errors.append(f"{label}: sum(COUNT) does not equal TOTAL_SITES")
            continue
        checks = []
        for row in rows:
            if total == 0:
                if row["percent"] not in (None, Decimal(0)):
                    errors.append(f"{label}/{row['category']}: zero denominator cannot have a nonzero PERCENT")
                exact, delta, tolerance = None, None, None
            else:
                places = max(0, -row["percent"].as_tuple().exponent)
                with localcontext() as ctx:
                    ctx.prec = max(60, len(str(total)) + 20)
                    exact = Decimal(row["count"]) * 100 / Decimal(total)
                    delta = row["percent"] - exact
                    tolerance = Decimal(1).scaleb(-places) / 2
                    if abs(delta) > tolerance:
                        errors.append(f"{label}/{row['category']}: PERCENT disagrees with COUNT/TOTAL beyond rounding tolerance")
            checks.append(dict(category=row["category"], line=row["line"], reported_percent=row["percent_text"],
                               exact_percent=str(exact) if exact is not None else None,
                               error_percentage_points=str(delta) if delta is not None else None,
                               rounding_tolerance_percentage_points=str(tolerance) if tolerance is not None else None))
        counts = {r["category"]: r["count"] for r in rows}
        group = dict(chrom=chrom, type=typ, total_sites=total, counts=counts)
        groups.append(group)
        audits.append(dict(chrom=chrom, type=typ, category_count=4, category_sum=sum(counts.values()),
                           total_sites=total, reported_percent_sum=str(sum((r["percent"] or Decimal(0)) for r in rows)),
                           zero_denominator=total == 0, percent_checks=checks))
    if errors:
        raise ValueError("Counts audit failed: " + "; ".join(errors[:12]) + (f"; and {len(errors)-12} more" if len(errors)>12 else ""))
    return groups, audits, fields


def _percent(count, total):
    return 100.0 * count / total if count is not None and total else None


def _values(total, counts):
    result = {"total_sites": total}
    for category in CATEGORIES:
        count = counts.get(category) if counts is not None else None
        result[category + "_count"] = count
        result[category + "_percent"] = _percent(count, total)
    multi = counts["tri_allelic"] + counts["four_plus_allelic"] if counts is not None else None
    third = counts["four_plus_allelic"] if counts is not None else None
    result.update(multi_alt_count=multi, multi_alt_percent=_percent(multi, total),
                  alt_3plus_count=third, alt_3plus_percent=_percent(third, total),
                  alt_3plus_within_multi_percent=_percent(third, multi))
    return result


def _aggregate(groups, chromosomes, typ, scope):
    selected = [r for r in groups if r["chrom"] in chromosomes and r["type"] == typ]
    observed = sorted({r["chrom"] for r in selected}, key=_order)
    missing = [c for c in chromosomes if c not in observed]
    if not selected:
        status, total, counts = "missing", None, None
    else:
        total = sum(r["total_sites"] for r in selected)
        counts = {c: sum(r["counts"][c] for r in selected) for c in CATEGORIES}
        status = "zero_denominator" if total == 0 else "partial_coverage" if missing else "observed"
    return dict(scope=scope, type=typ, status=status, observed_chromosomes=observed,
                missing_chromosomes=missing, **_values(total, counts))


def _summaries(groups):
    observed = sorted({r["chrom"] for r in groups}, key=_order)
    scopes = [("autosomes", list(COUNTS_AUTOSOMES)), ("X", ["chrX"]), ("Y", ["chrY"]), ("all_observed", observed)]
    other = [c for c in observed if c not in STANDARD_CHROMS]
    if other:
        scopes.append(("other_contigs", other))
    summaries = [_aggregate(groups, chroms, typ, scope) for scope, chroms in scopes for typ in TYPES]
    by_chromosome = []
    for chrom in sorted(set(STANDARD_CHROMS) | set(observed), key=_order):
        for typ in TYPES:
            row = _aggregate(groups, [chrom], typ, chrom)
            row.pop("scope")
            row["chrom"] = chrom
            if chrom in COUNTS_AUTOSOMES:
                ref_chroms = [c for c in COUNTS_AUTOSOMES if c != chrom]
                reference_scope = "other_observed_autosomes"
            elif chrom in {"chrX", "chrY"}:
                ref_chroms = list(COUNTS_AUTOSOMES)
                reference_scope = "observed_autosomes"
            else:
                ref_chroms, reference_scope = [], "not_defined_for_other_contigs"
            reference = _aggregate(groups, ref_chroms, typ, reference_scope)
            row.update(reference_scope=reference_scope, reference_total_sites=reference["total_sites"],
                       reference_missing_chromosomes=reference["missing_chromosomes"])
            comparisons = {}
            for metric in (*CATEGORIES, "multi_alt", "alt_3plus"):
                numerator = row[metric + "_percent"]
                denominator = reference[metric + "_percent"]
                ratio = numerator / denominator if numerator is not None and denominator else None
                comparisons[metric] = dict(percent=numerator, reference_percent=denominator, ratio=ratio,
                                           difference_percentage_points=numerator-denominator if numerator is not None and denominator is not None else None)
                if metric in {"multi_alt", "alt_3plus"}:
                    row[metric + "_reference_percent"] = denominator
                    row[metric + "_ratio_to_reference"] = ratio
            row["comparisons"] = comparisons
            by_chromosome.append(row)
    extremes = {}
    for typ in TYPES:
        eligible = [r for r in by_chromosome if r["type"] == typ and r["chrom"] in COUNTS_AUTOSOMES and r["total_sites"]]
        extremes[typ] = {}
        for metric in ("multi_alt_percent", "alt_3plus_percent", "multi_alt_count", "alt_3plus_count", "total_sites"):
            extremes[typ][metric] = ({side: {"chrom": row["chrom"], "value": row[metric]}
                                     for side, row in (("min", min(eligible, key=lambda r:r[metric])),
                                                       ("max", max(eligible, key=lambda r:r[metric])))} if eligible else None)
    return summaries, by_chromosome, extremes


def _plot_backend(make_plots):
    if not make_plots:
        return None, "not_requested", None
    try:
        import matplotlib
        from matplotlib.figure import Figure
        from matplotlib.backends.backend_agg import FigureCanvasAgg
    except ImportError:
        return None, "unavailable_matplotlib", None
    return (Figure, FigureCanvasAgg), "available", matplotlib.__version__


def _make_plot(path, by_chromosome, summaries, backend):
    Figure, FigureCanvasAgg = backend
    fig = Figure(figsize=(14, 9.2), facecolor="white")
    FigureCanvasAgg(fig)
    axes = fig.subplots(2, 2)
    lookup = {(r["chrom"], r["type"]):r for r in by_chromosome}
    baseline = {r["type"]:r for r in summaries if r["scope"] == "autosomes"}
    positions = {c:i for i,c in enumerate(COUNTS_AUTOSOMES)} | {"chrX":23, "chrY":24.5}
    tick_labels = [str(i) for i in range(1,23)] + ["X", "Y"]
    colors = {"SNP":"#2166ac", "INDEL":"#16856c"}
    for col, typ in enumerate(TYPES):
        for row, metric in enumerate(("multi_alt_percent", "alt_3plus_percent")):
            ax = axes[row,col]
            records = [lookup[(chrom,typ)] for chrom in STANDARD_CHROMS]
            vals = [r[metric] for r in records if r[metric] is not None]
            for chrom in COUNTS_AUTOSOMES:
                value = lookup[(chrom,typ)][metric]
                if value is not None:
                    ax.scatter(positions[chrom], value, s=39, color=colors[typ], zorder=3)
            for chrom in ("chrX","chrY"):
                value = lookup[(chrom,typ)][metric]
                if value is not None:
                    ax.scatter(positions[chrom], value, marker="D", s=46, color="#d87920", zorder=4)
                else:
                    ax.text(positions[chrom], .025, "NA", transform=ax.get_xaxis_transform(),
                            ha="center", va="bottom", color="#9ca3af", fontsize=8)
            value = baseline[typ][metric]
            if value is not None:
                ax.axhline(value, color="#64748b", lw=1.15, ls=(0,(5,3)), zorder=1)
                ax.text(.02,.97,f"Autosomes weighted: {value:.4f}%", transform=ax.transAxes,
                        va="top", fontsize=10, color="#4b5563")
            valid_auto = [r for r in records if r["chrom"] in COUNTS_AUTOSOMES and r[metric] is not None]
            if valid_auto:
                highest, lowest = max(valid_auto,key=lambda r:r[metric]), min(valid_auto,key=lambda r:r[metric])
                for target, dy in ((highest,11),(lowest,-18)):
                    if highest is lowest and dy < 0:
                        continue
                    label = target["chrom"].removeprefix("chr")
                    ax.annotate(f"{label}: {target[metric]:.4f}%", (positions[target["chrom"]],target[metric]),
                                xytext=(0,dy), textcoords="offset points", ha="center", fontsize=9,
                                color=colors[typ], bbox=dict(facecolor="white",edgecolor="none",alpha=.9,pad=1))
            ax.axvline(22, color="#d3dce5", lw=1)
            ax.set_xticks([positions[c] for c in STANDARD_CHROMS],tick_labels, fontsize=8.5)
            for tick, chrom in zip(ax.get_xticklabels(), STANDARD_CHROMS):
                if lookup[(chrom,typ)]["status"] == "missing":
                    tick.set_color("#a4adb7")
            ax.set_xlim(-.8,25.3)
            ymax = max(vals) * 1.28 if vals and max(vals)>0 else 1.0
            ax.set_ylim(0,ymax)
            ax.set_ylabel("Percent of legacy position groups", fontsize=10)
            ax.set_xlabel("Chromosome",fontsize=10)
            label = "≥2 retained ALT" if row==0 else "≥3 retained ALT"
            ax.set_title(f"{'ABCD'[row*2+col]}   {typ}  |  {label}",loc="left",fontsize=13,fontweight="bold",pad=10)
            ax.grid(axis="y",color="#e5ebf0",lw=.7)
            ax.set_axisbelow(True)
            ax.spines[["top","right"]].set_visible(False)
            ax.spines[["left","bottom"]].set_color("#c4cdd6")
    fig.suptitle("PASS, AC>0 ALT categories by chromosome",fontsize=19,fontweight="bold",x=.5,y=.982)
    fig.text(.5,.942,"Autosomes: panel colour     X/Y: orange diamond     Dashed line: observed-autosome weighted fraction",ha="center",fontsize=10,color="#475569")
    fig.text(.055,.057,"Legacy conditional fractions, not callable-base densities. REF occurrence is not established.",fontsize=10,color="#475569")
    fig.text(.055,.034,"INDELs share start-coordinate groups without full reference left-alignment; their categories are not normalized allele-state counts.",fontsize=9.5,color="#475569")
    fig.text(.055,.012,"NA = chromosome/type not provided or zero denominator. No P values, confidence intervals, or evolutionary conclusions are shown.",fontsize=9.5,color="#475569")
    fig.subplots_adjust(top=.88,bottom=.15,left=.06,right=.98,hspace=.40,wspace=.20)
    fig.savefig(path,dpi=175,facecolor="white")


def _fmt(value, percent=False):
    if value is None:
        return "NA"
    if isinstance(value,float):
        return f"{value:.4f}%" if percent else f"{value:.6g}"
    return f"{value:,}" if isinstance(value,int) else str(value)


def _table(headers, rows):
    return '<div class="table-scroll"><table><thead><tr>'+''.join('<th>'+html.escape(h)+'</th>' for h in headers)+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+html.escape(str(v))+'</td>' for v in row)+'</tr>' for row in rows)+'</tbody></table></div>'


def _report(audit, image_path, plot_status):
    summary = audit["summaries"]
    table_rows = []
    scope_names = {"autosomes":"常染色体（已提供）", "X":"X", "Y":"Y", "all_observed":"全部已提供染色体", "other_contigs":"其他 contig"}
    for r in summary:
        table_rows.append([scope_names[r["scope"]],r["type"],r["status"],_fmt(r["total_sites"]),
                           _fmt(r["singleton_count"]),_fmt(r["non_singleton_biallelic_count"]),
                           _fmt(r["tri_allelic_count"]),_fmt(r["four_plus_allelic_count"]),
                           _fmt(r["multi_alt_percent"],True),_fmt(r["alt_3plus_percent"],True)])
    by_rows = []
    for r in audit["by_chromosome"]:
        by_rows.append([r["chrom"],r["type"],r["status"],_fmt(r["total_sites"]),
                        _fmt(r["multi_alt_percent"],True),_fmt(r["alt_3plus_percent"],True),
                        _fmt(r["multi_alt_ratio_to_reference"]),_fmt(r["alt_3plus_ratio_to_reference"])])
    image_block = ('<img class="figure" alt="SNP and INDEL retained ALT category percentages by chromosome" src="data:image/png;base64,'+
                   base64.b64encode(Path(image_path).read_bytes()).decode()+'">' if image_path else
                   '<p class="note">图未生成：'+html.escape(plot_status)+'。下方表格和 JSON/TSV 汇总仍完整可用。</p>')
    qc = audit["qc"]
    missing = '; '.join(typ+': '+(', '.join(qc['missing_standard_chromosomes_by_type'][typ]) or '无') for typ in TYPES)
    status_counts = json.dumps(qc["reported_percent_sum_distribution"],ensure_ascii=False)
    document = '''<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>counts.tsv 条件计数审计</title><style>
body{margin:0;background:#edf2f6;color:#183247;font:15px/1.65 system-ui,-apple-system,"Segoe UI",sans-serif}main{max-width:1440px;margin:auto;padding:24px}header{padding:24px 30px;background:#17374f;color:white;border-radius:12px}h1{margin:0 0 10px;font-size:29px}h2{font-size:21px;margin:0 0 14px}.card{background:white;border:1px solid #d7e1e9;border-radius:10px;padding:22px;margin:18px 0}.note{background:#fff5df;border-left:4px solid #d69b38;padding:14px 17px}.figure{display:block;width:100%;height:auto}.table-scroll{overflow:auto}table{width:100%;border-collapse:collapse;font-size:13px;font-variant-numeric:tabular-nums}th,td{padding:9px 10px;border-bottom:1px solid #e0e8ee;text-align:right;white-space:nowrap}th{background:#edf4f8;color:#24445d}th:first-child,td:first-child{text-align:left}code{font-family:ui-monospace,monospace;word-break:break-all}small{color:#577084}li{margin:7px 0}.pill{display:inline-block;border:1px solid #8fc9ba;background:#e6f7f1;color:#175d4d;border-radius:5px;padding:3px 9px}a{color:#2166ac}footer{font-size:13px;color:#526c7e;margin:20px 4px}
</style></head><body><main><header><h1>counts.tsv 条件计数审计</h1><p>按已核验 legacy check 定义：统计 PASS 且 AC&gt;0 的 ALT 数量。所有比例采用各类型自己的 TOTAL_SITES 分母。</p></header>'''
    document += '<section class="card"><h2>数值核对</h2><p><span class="pill">算术校验通过</span> '+f"{qc['rows']} 个数据行，{qc['groups']} 个 CHR/TYPE 组；四类唯一、同组 TOTAL 一致、COUNT 合计和百分比舍入全部通过。"+'</p>'
    document += '<p>缺失标准染色体：'+html.escape(missing)+'。缺失项保留 NA，不补零。</p><p>四类已报告百分比之和（值→组数）：<code>'+html.escape(status_counts)+'</code>；逐项按其显示小数精度检查舍入。</p>'
    document += '<p><small>输入：'+html.escape(audit['input']['path'])+'<br>SHA-256：<code>'+audit['input']['sha256']+'</code></small></p></section>'
    document += '<section class="card"><h2>四类究竟表示什么</h2><ul><li><b>singleton</b>：仅一个保留 ALT，且其 AC=1。</li><li><b>non_singleton_biallelic</b>：仅一个保留 ALT，且 AC&gt;1。</li><li><b>tri_allelic</b>：两个保留 ALT。</li><li><b>four_plus_allelic</b>：至少三个保留 ALT；标准 A/C/G/T SNP 相对 REF 最多只有三个 ALT 碱基。</li></ul><p class="note"><b>REF 只是标签，不证明样本中存在 REF 拷贝。</b> 因而两个 ALT 不能自动称为三个已观测等位状态。INDEL 是同起始坐标分组，未完整结合 FASTA 左对齐，类别数也不能直接当作规范化多态状态数。</p><p>以上是此模块所要求的已核验 legacy producer 语义；其他来源即使列名相同也需核对源码。输入表本身没有 release/build、样本量、群体或完整过滤参数元数据。</p></section>'
    document += '<section class="card"><h2>染色体分布</h2>'+image_block+'<p><small>≥2 ALT = tri_allelic + four_plus_allelic；≥3 ALT = four_plus_allelic。虚线为已提供常染色体按 TOTAL 加权的比例，X/Y 单独展示。图的分母不是可调用碱基数。</small></p></section>'
    document += '<section class="card"><h2>汇总计数与加权比例</h2>'+_table(['范围','TYPE','状态','TOTAL','singleton','单ALT, AC>1','两ALT','至少三ALT','≥2 ALT 比例','≥3 ALT 比例'],table_rows)+'<p>完整四类加权比例另见 counts_summary.tsv；各类型 totals 可能覆盖相同坐标，其相加不代表唯一基因组位置数。</p></section>'
    document += '<section class="card"><h2>逐染色体与描述性比例比</h2><p>常染色体的参考为其余已提供常染色体；X/Y 参考为全部已提供常染色体。比例比是未校正的描述量，不是选择、功能或富集检验。</p>'+_table(['CHR','TYPE','状态','TOTAL','≥2 ALT','≥3 ALT','≥2 ALT / 参考','≥3 ALT / 参考'],by_rows)+'</section>'
    document += '<section class="card"><h2>结果可以支持的范围</h2><p>这些数据支持分类内计数、条件比例、染色体描述性比较和候选计算规模估算。它们不能单独给出可调用密度、变异产生率、祖先状态、首次出现时间、平衡选择或 pleiotropy 关系；未计算 P 值、置信区间或古代 DNA 结论。</p><p>较高的条件比例可用于安排后续序列背景、原始位点、规范化及 QC 审计；不能仅凭本表判定其生物学原因。</p></section>'
    return document+'<footer>独立离线 HTML：图像以 base64 内嵌，无 CDN、脚本或外部字体请求。完整机器可读审计见 counts_audit.json 和 manifest.json。</footer></main></body></html>'


def _write_tsv(path, fields, rows):
    with open(path,"w",encoding="utf-8",newline="") as handle:
        writer = csv.DictWriter(handle,fields,delimiter="\t",lineterminator="\n")
        writer.writeheader()
        for row in rows:
            clean = {}
            for name in fields:
                value = row.get(name)
                clean[name] = "." if value is None else ",".join(value) if isinstance(value,list) else format(value,".12g") if isinstance(value,float) else value
            writer.writerow(clean)


def _existing_manifest(destination):
    manifest_path = destination / "manifest.json"
    if manifest_path.exists():
        try:
            old = json.loads(manifest_path.read_text())
        except (ValueError,OSError) as exc:
            raise ValueError("Existing manifest.json is not a valid owned counts-audit manifest") from exc
        if old.get("tool") != TOOL_ID or old.get("schema_version") != SCHEMA_VERSION or not isinstance(old.get("outputs"),dict):
            raise ValueError("Refusing to overwrite outputs without an owned counts-audit manifest")
        owned = set(old["outputs"]) | {"manifest.json"}
        if not owned.issubset(OUTPUT_NAMES):
            raise ValueError("Existing counts-audit manifest lists unexpected output names")
    else:
        old, owned = None, set()
    for name in OUTPUT_NAMES:
        if (destination/name).exists() and name not in owned:
            raise ValueError(f"Refusing to overwrite unowned output: {destination/name}")
    return old


def audit_counts(counts_path, out_dir, make_plots=True):
    """Validate, summarize and publish owned offline outputs; never edit input.

    Invalid input raises ValueError before any output is replaced. Identical
    inputs/settings reuse checksum-verified outputs. Changed runs replace only
    outputs declared by this module's manifest, using staged atomic file moves.
    Matplotlib is optional; TSV/JSON/HTML are still produced if it is absent.
    """
    if not isinstance(make_plots,bool):
        raise ValueError("make_plots must be True or False")
    source = Path(counts_path).expanduser().resolve()
    destination = Path(out_dir).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    if source in {(destination/name).resolve() for name in OUTPUT_NAMES}:
        raise ValueError("Counts output must not overwrite its input")
    source_info = _counts_fingerprint(source)
    groups, group_audit, source_columns = _read_counts(source)
    if _sha(source) != source_info["sha256"]:
        raise ValueError("Counts input changed while it was being read")
    summaries, by_chromosome, extremes = _summaries(groups)
    backend, backend_status, version = _plot_backend(make_plots)
    settings = dict(make_plots=make_plots,plot_backend_status=backend_status,matplotlib_version=version)
    code_hash = _sha(Path(__file__))
    signature = hashlib.sha256(json.dumps(dict(input=source_info,settings=settings,code_sha256=code_hash),sort_keys=True).encode()).hexdigest()
    destination.mkdir(parents=True,exist_ok=True)
    old = _existing_manifest(destination)
    if old and old.get("signature") == signature and all((destination/name).is_file() and _sha(destination/name)==spec["sha256"] for name,spec in old["outputs"].items()):
        return old
    missing = {typ:[c for c in STANDARD_CHROMS if not any(r["chrom"]==c and r["type"]==typ for r in groups)] for typ in TYPES}
    other = sorted({r["chrom"] for r in groups if r["chrom"] not in STANDARD_CHROMS})
    errors = [abs(float(c["error_percentage_points"])) for r in group_audit for c in r["percent_checks"] if c["error_percentage_points"] is not None]
    qc = dict(status="passed",rows=len(groups)*4,groups=len(groups),four_categories_unique=True,
              totals_consistent=True,category_sums_match=True,percentages_within_display_rounding=True,
              max_absolute_percent_error_pp=max(errors) if errors else None,
              reported_percent_sum_distribution=dict(Counter(r["reported_percent_sum"] for r in group_audit)),
              zero_denominator_groups=sum(r["total_sites"]==0 for r in groups),
              missing_standard_chromosomes_by_type=missing,other_contigs=other,
              missing_types=[typ for typ in TYPES if not any(r["type"]==typ for r in groups)])
    audit = dict(tool=TOOL_ID,schema_version=SCHEMA_VERSION,input=source_info,source_columns=source_columns,
                 release=None,reference_build=None,producer_metadata_in_input=False,
                 definitions=DEFINITIONS,limitations=LIMITATIONS,qc=qc,group_audit=group_audit,
                 summaries=summaries,by_chromosome=by_chromosome,autosomal_extremes=extremes)
    output_files = [name for name in OUTPUT_NAMES if name not in {"manifest.json","counts_review.png"}]
    plot_status = backend_status
    with tempfile.TemporaryDirectory(prefix=".counts-audit-",dir=destination) as temporary:
        stage = Path(temporary)
        plot_path = None
        if backend is not None:
            plot_path = stage/"counts_review.png"
            _make_plot(plot_path,by_chromosome,summaries,backend)
            output_files.append("counts_review.png")
            plot_status = "generated"
        (stage/"counts_audit.json").write_text(json.dumps(audit,ensure_ascii=False,indent=2,allow_nan=False)+"\n",encoding="utf-8")
        _write_tsv(stage/"counts_summary.tsv",SUMMARY_FIELDS,summaries)
        _write_tsv(stage/"counts_by_chromosome.tsv",CHROM_FIELDS,by_chromosome)
        (stage/"counts_review.html").write_text(_report(audit,plot_path,plot_status),encoding="utf-8")
        if _sha(source) != source_info["sha256"]:
            raise ValueError("Counts input changed during report generation")
        outputs = {name:{"path":str(destination/name),"size":(stage/name).stat().st_size,"sha256":_sha(stage/name)} for name in output_files}
        manifest = dict(tool=TOOL_ID,schema_version=SCHEMA_VERSION,signature=signature,input=source_info,
                        code_sha256=code_hash,settings=settings,plot_status=plot_status,outputs=outputs,
                        manifest_path=str(destination/"manifest.json"),report_path=str(destination/"counts_review.html"),
                        audit_path=str(destination/"counts_audit.json"),summary_path=str(destination/"counts_summary.tsv"),
                        by_chromosome_path=str(destination/"counts_by_chromosome.tsv"),
                        plot_path=str(destination/"counts_review.png") if plot_path else None,
                        qc=qc,definitions=DEFINITIONS,limitations=LIMITATIONS)
        (stage/"manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2,allow_nan=False)+"\n",encoding="utf-8")
        for name in output_files:
            os.replace(stage/name,destination/name)
        if old:
            for name in set(old["outputs"])-set(output_files):
                (destination/name).unlink(missing_ok=True)
        os.replace(stage/"manifest.json",destination/"manifest.json")
    return manifest


# ========================================================================
# evolve_atlas
# ========================================================================

# Bounded AlphaGenome Atlas queries and functional summaries.



# ========================================================================
# Sparse atlas
# ========================================================================

# Bounded, optional AlphaGenome Atlas lookup and same-scorer geometry.
# 
# The default gnomAD scan never imports this module or calls an API. Planning
# requires only the standard library. Actual lookup uses the verified Atlas
# GetDenseVariantScores RPC (precomputed data, not model inference). Geometry
# uses NumPy and cached raw effects; it does not make network requests.


from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
from contextlib import contextmanager
import csv
import dataclasses
from datetime import datetime, timezone
import enum
import fcntl
import gzip
import hashlib
import importlib.metadata
import json
import math
import os
from pathlib import Path
import re
import shutil
import tempfile
import threading
import time
from typing import Any


ATLAS_SCHEMA = "allele3-sparse-atlas-v1"
ENDPOINT = "dns:///gdmscience.googleapis.com:443"
VERIFIED_SDK_COMMIT = "55943112b83439e0c2512110f1ae607d4b5b492a"
DEFAULT_SCORERS = ("ATAC", "DNASE", "CHIP_TF", "CHIP_HISTONE", "CAGE", "PROCAP")
ATLAS_BASES = "ACGT"
REQUIRED_COLUMNS = ("site_id", "chrom", "pos", "ref", "anchor_ref", "anchor_alt",
                    "extra_alt", "s", "posterior")
SOURCE_CAVEAT = (
    "Atlas precomputed lookup; the service exposes no dataset-release pin. "
    "This cache records endpoint, SDK version, query settings, timestamps, "
    "fetch sessions and metadata hashes, not an immutable server release. "
    "Resume can span sessions; geometry requires matching complete axes and "
    "intervals. Matching metadata cannot prove unchanged model weights."
)
GEOMETRY_CAVEAT = (
    "Raw effects are compared separately for each scorer with a shared "
    "GRCh38 reference. A REF effect is the zero vector by definition. "
    "Distances are descriptive molecular predictions, not fitness effects. "
    "Response breadth is a fraction of predicted response dimensions above "
    "an absolute raw-effect threshold, not a count of biological functions "
    "or a validated measure of pleiotropy. The direction inferred from s "
    "remains an estimate, including when posterior confidence is low."
)


def _utc() -> str:
    return datetime.now(timezone.utc).isoformat()


def _hash(value: Any) -> str:
    return hashlib.sha256(json.dumps(value, sort_keys=True, separators=(",", ":"),
                                     ensure_ascii=False, allow_nan=False).encode()).hexdigest()


def _file_hash(path: Path) -> str:
    result = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            result.update(block)
    return result.hexdigest()


def _json_write(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix=path.name + ".", suffix=".tmp", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as handle:
            json.dump(data, handle, indent=2, ensure_ascii=False, allow_nan=False)
            handle.write("\n")
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def _read_json(path: Path) -> dict:
    with path.open(encoding="utf-8") as handle:
        return json.load(handle)


def _positive_int(value: int, label: str) -> int:
    if isinstance(value, bool) or not isinstance(value, int) or value < 1:
        raise ValueError(f"{label} must be a positive integer")
    return value


def _names(values, defaults, label) -> list[str]:
    values = defaults if values is None else values
    if isinstance(values, str):
        values = values.split(",")
    result = sorted({str(value).strip() for value in values if str(value).strip()})
    if not result and label == "scorers":
        raise ValueError("At least one scorer must be requested")
    pattern = r"[A-Za-z][A-Za-z0-9_]*" if label == "scorers" else r"(?:UBERON|CL|CLO|EFO):[0-9]+"
    if any(not re.fullmatch(pattern, name) for name in result):
        raise ValueError(f"Invalid {label}: {result}")
    return result


def _atlas_chrom(value: str) -> str:
    value = str(value).strip()
    value = value[3:] if value.lower().startswith("chr") else value
    if value.isdigit():
        value = str(int(value))
    else:
        value = value.upper()
    if value == "MT":
        value = "M"
    if value not in {*(str(i) for i in range(1, 23)), "X", "Y", "M"}:
        raise ValueError(f"Sparse Atlas requires a primary human chromosome, got {value!r}")
    return "chr" + value


def _variant(site: dict, alt: str) -> dict:
    return dict(chrom=site["chrom"], pos=site["pos"], ref=site["ref"], alt=alt)


def _variant_id(variant: dict) -> str:
    return f"{variant['chrom']}:{variant['pos']}:{variant['ref']}>{variant['alt']}"


def _variants(sites: list[dict]) -> list[dict]:
    unique = {}
    for site in sites:
        for alt in ATLAS_BASES:
            if alt != site["ref"]:
                value = _variant(site, alt)
                unique[_variant_id(value)] = value
    return [unique[key] for key in sorted(unique)]


def _site(row: dict, line: int) -> dict:
    if None in row or any(row.get(name) is None for name in REQUIRED_COLUMNS):
        raise ValueError(f"Malformed candidate row {line}")
    result = dict(row)
    result["site_id"] = str(row["site_id"]).strip()
    if not result["site_id"]:
        raise ValueError(f"Empty site_id at row {line}")
    result["chrom"] = _atlas_chrom(row["chrom"])
    try:
        result["pos"] = int(row["pos"])
        result["s"] = float(row["s"])
        result["posterior"] = float(row["posterior"])
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Invalid position, s or posterior at row {line}") from exc
    if result["pos"] < 1 or not math.isfinite(result["s"]) or not math.isfinite(result["posterior"]):
        raise ValueError(f"Position must be 1-based and s/posterior finite at row {line}")
    if not 0 <= result["posterior"] <= 1:
        raise ValueError(f"posterior must lie in [0,1] at row {line}")
    for name in ("ref", "anchor_ref", "anchor_alt", "extra_alt"):
        result[name] = str(row[name]).strip().upper()
        if result[name] not in ATLAS_BASES or len(result[name]) != 1:
            raise ValueError(f"{name} must be one forward-strand A/C/G/T at row {line}")
    if len({result["anchor_ref"], result["anchor_alt"], result["extra_alt"]}) != 3:
        raise ValueError(f"anchor_ref, anchor_alt and extra_alt must be three distinct bases at row {line}")
    for name in ("build", "assembly", "genome_build"):
        if name in row and str(row[name]).strip().lower() not in {"grch38", "hg38"}:
            raise ValueError(f"Sparse Atlas only supports GRCh38; row {line} declares {row[name]!r}")
    if "strand" in row and str(row["strand"]).strip() not in {"+", "1", "forward"}:
        raise ValueError(f"Candidate row {line} must already be mapped to the GRCh38 forward strand")
    return result


def build_plan(candidate_tsv, out_dir, max_sites=2000, max_variants=6000,
               scorers=None, ontology_terms=None) -> dict:
    """Write an entirely offline bounded plan; never truncate an over-budget list.

    Each row denotes one site and supplies the ancient pair plus the observed
    extra base after mapping to GRCh38 forward-strand coordinates. All three
    non-REF substitutions are planned, even when one is not observed in gnomAD.
    Additional candidate columns are retained. No API key or SDK is needed.
    """
    _positive_int(max_sites, "max_sites")
    _positive_int(max_variants, "max_variants")
    candidate = Path(candidate_tsv).expanduser().resolve()
    scorers = _names(scorers, DEFAULT_SCORERS, "scorers")
    terms = _names(ontology_terms, (), "ontology_terms")
    sites, ids, loci = [], set(), set()
    opener = gzip.open if candidate.suffix in {".gz", ".bgz"} else open
    with opener(candidate, "rt", encoding="utf-8-sig", newline="") as handle:
        def lines():
            for line in handle:
                if line.startswith("##"):
                    if re.match(r"##\s*(build|assembly|genome_build)\s*[=:]", line, re.I):
                        declared = re.split(r"[=:]", line, maxsplit=1)[1].strip().lower()
                        if declared not in {"grch38", "hg38"}:
                            raise ValueError(f"Sparse Atlas only supports GRCh38, not {declared}")
                    continue
                if line.strip():
                    yield line[1:] if line.startswith("#site_id\t") else line
        reader = csv.DictReader(lines(), delimiter="\t")
        fields = list(reader.fieldnames or [])
        missing = set(REQUIRED_COLUMNS) - set(fields)
        if missing or len(set(fields)) != len(fields):
            raise ValueError(f"Candidate TSV needs unique column names and required columns; missing {sorted(missing)}")
        if any(name.startswith("atlas_") for name in fields):
            raise ValueError("Candidate columns starting atlas_ are reserved for Atlas outputs")
        for line, row in enumerate(reader, 2):
            site = _site(row, line)
            if site["site_id"] in ids:
                raise ValueError(f"Duplicate site_id: {site['site_id']}")
            locus = (site["chrom"], site["pos"])
            if locus in loci:
                raise ValueError(f"More than one candidate row at {locus}; supply one row per site")
            ids.add(site["site_id"])
            loci.add(locus)
            sites.append(site)
            if len(sites) > max_sites:
                raise ValueError(f"Candidate site budget exceeded: more than max_sites={max_sites}; no requests made")
            if 3 * len(sites) > max_variants:
                raise ValueError(f"Variant budget exceeded: three non-REF variants/site exceed max_variants={max_variants}; no requests made")
    variants = _variants(sites)
    body = dict(schema=ATLAS_SCHEMA, build="GRCh38", coordinate_system="1-based forward strand",
                candidate_columns=fields, candidates_sha256=_file_hash(candidate),
                sites=sites, variants=variants, scorers=scorers, ontology_terms=terms)
    output = Path(out_dir).expanduser().resolve()
    path = output / "plan.json"
    manifest = dict(body, plan_id=_hash(body), created_utc=_utc(),
                    candidate_tsv=str(candidate), plan_path=str(path),
                    n_sites=len(sites), n_variants=len(variants),
                    max_sites=max_sites, max_variants=max_variants,
                    network_requests_made=0, variant_rpc_count=len(variants),
                    implicit_reference="REF raw effect = zero; REF is not queried",
                    source_caveat=SOURCE_CAVEAT)
    if path.exists():
        existing = _load_plan(path)
        if existing["plan_id"] != manifest["plan_id"]:
            raise ValueError(f"Different plan already exists at {path}; choose a new output directory")
        return existing
    _json_write(path, manifest)
    return manifest


def _load_plan(path) -> dict:
    data = _read_json(Path(path))
    keys = ("schema", "build", "coordinate_system", "candidate_columns", "candidates_sha256",
            "sites", "variants", "scorers", "ontology_terms")
    if data.get("schema") != ATLAS_SCHEMA or data.get("build") != "GRCh38":
        raise ValueError("Unsupported sparse Atlas plan schema/build")
    if data.get("plan_id") != _hash({key: data[key] for key in keys}):
        raise ValueError("Sparse Atlas plan content or fingerprint changed")
    sites = [_site(row, i + 2) for i, row in enumerate(data["sites"])]
    if len({(row["chrom"], row["pos"]) for row in sites}) != len(sites):
        raise ValueError("Sparse Atlas plan contains duplicate loci")
    if len({row["site_id"] for row in sites}) != len(sites):
        raise ValueError("Sparse Atlas plan contains duplicate site IDs")
    if data["variants"] != _variants(sites) or data["n_sites"] != len(sites) or data["n_variants"] != len(data["variants"]):
        raise ValueError("Sparse Atlas variant list/count is inconsistent with candidate sites")
    if data["scorers"] != _names(data["scorers"], (), "scorers") or data["ontology_terms"] != _names(data["ontology_terms"], (), "ontology_terms"):
        raise ValueError("Sparse Atlas plan has invalid query settings")
    return data


def _settings(plan: dict) -> dict:
    return dict(schema=ATLAS_SCHEMA, build=plan["build"], scorers=plan["scorers"],
                ontology_terms=plan["ontology_terms"])


def _empty_catalog(plan: dict) -> dict:
    """An empty plan is a successful offline result, with no claimed SDK use."""
    settings = _settings(plan)
    source = dict(endpoint=ENDPOINT, sdk_version="not_loaded_empty_plan",
                  protocol="AtlasService.GetDenseVariantScores", lookup="precomputed",
                  verified_source_commit=VERIFIED_SDK_COMMIT, server_release="not_exposed")
    return dict(schema=ATLAS_SCHEMA, settings=settings, source=source,
                namespace=_hash(dict(settings=settings, source=source)), created_utc=_utc(),
                caveat=SOURCE_CAVEAT)


def _entry_path(cache: Path, catalog: dict, variant: dict) -> Path:
    return cache / _hash(dict(namespace=catalog["namespace"], variant=variant))


def _cache_entry(cache: Path, catalog: dict, variant: dict) -> tuple[Path, dict] | None:
    path = _entry_path(cache, catalog, variant)
    marker = path / "COMPLETE.json"
    if not marker.exists():
        if path.exists():
            raise ValueError(f"Unrecognized incomplete Atlas cache: {path}")
        return None
    complete = _read_json(marker)
    if complete.get("namespace") != catalog["namespace"] or complete.get("variant") != variant:
        raise ValueError(f"Atlas cache identity mismatch at {path}")
    for name in ("response.pb", "arrays.npz", "metadata.json"):
        file = path / name
        if not file.is_file() or _file_hash(file) != complete.get("files", {}).get(name):
            raise ValueError(f"Missing or changed completed Atlas cache file: {file}")
    return path, _read_json(path / "metadata.json")


def _json_value(value):
    """Preserve typed SDK metadata in JSON, without pickled Python objects."""
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else {"nonfinite_float": str(value)}
    if isinstance(value, enum.Enum):
        return dict(enum=type(value).__name__, name=value.name, value=_json_value(value.value))
    if dataclasses.is_dataclass(value):
        return {key: _json_value(item) for key, item in dataclasses.asdict(value).items()}
    if isinstance(value, dict):
        return {str(key): _json_value(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_value(item) for item in value]
    if hasattr(value, "item"):
        return _json_value(value.item())
    # pandas.NA/NaT can occur in optional metadata columns.
    if type(value).__name__ in {"NAType", "NaTType"}:
        return None
    raise TypeError(f"Cannot losslessly export SDK metadata type {type(value).__name__}")


def _atlas_table(frame) -> dict:
    return dict(index=[_json_value(x) for x in frame.index.tolist()],
                index_names=[_json_value(x) for x in frame.index.names],
                columns=[str(x) for x in frame.columns],
                dtypes=[str(x) for x in frame.dtypes],
                data=[[_json_value(x) for x in row]
                      for row in frame.itertuples(index=False, name=None)])


def _response_packet(response, requested: dict) -> dict:
    """Use the official converter and retain the original protobuf verbatim."""
    import numpy as np
    from alphagenome.atlas import atlas
    from google.protobuf.json_format import MessageToDict

    actual = dict(chrom=_atlas_chrom(response.variant.chromosome), pos=response.variant.position,
                  ref=response.variant.reference_bases, alt=response.variant.alternate_bases)
    if actual != requested:
        raise ValueError(f"Atlas response variant mismatch: requested {requested}, received {actual}")
    raw_proto = response.SerializeToString()
    structure = MessageToDict(response, preserving_proto_field_name=True)
    for score in structure.get("scores", []):
        score.pop("scores", None)
        score.pop("calibrated_scores", None)
    by_name = {}
    for score in response.scores:
        name = score.variant_scorer.name
        if name in by_name:
            raise ValueError(f"Duplicate scorer in Atlas response: {name}")
        by_name[name] = score
    converted = atlas.convert_variant_scores_to_anndata([response], None)
    arrays, scorers = {}, []
    for i, (name, adata) in enumerate(sorted(converted.items())):
        raw_key = f"raw_{i}"
        arrays[raw_key] = np.asarray(adata.X).copy()
        item = dict(name=name, is_signed=bool(by_name[name].variant_scorer.is_signed),
                    raw_key=raw_key, shape=list(adata.X.shape), obs=_atlas_table(adata.obs),
                    var=_atlas_table(adata.var), uns=_json_value(dict(adata.uns)))
        if "quantiles" in adata.layers:
            key = f"quantiles_{i}"
            arrays[key] = np.asarray(adata.layers["quantiles"]).copy()
            item["quantiles_key"] = key
        scorers.append(item)
    return dict(protobuf=raw_proto, arrays=arrays, scorers=scorers,
                interval=structure.get("interval"), response_structure=structure,
                variant=actual)


class _FetchError(RuntimeError):
    def __init__(self, code: str, detail: str, *, retryable=False, fatal=False):
        self.code, self.retryable, self.fatal = code, retryable, fatal
        super().__init__(f"{code}: {detail}")


class _AtlasTransport:
    """Exactly one verified RPC per fetch; application code controls retries."""
    endpoint = ENDPOINT

    def __init__(self, api_key: str, timeout: float):
        try:
            import grpc
            from alphagenome.atlas import atlas_utils
            from alphagenome.data import genome
            from alphagenome.protos import atlas_service_pb2, atlas_service_pb2_grpc
            from alphagenome.protos.v1 import dna_model_pb2
        except ImportError as exc:
            raise RuntimeError("Optional Atlas query requires alphagenome==0.10.0; planning/scanning do not") from exc
        self.sdk_version = importlib.metadata.version("alphagenome")
        self.grpc, self.atlas_utils, self.genome = grpc, atlas_utils, genome
        self.messages, self.dna_messages = atlas_service_pb2, dna_model_pb2
        self.metadata = (("x-goog-api-key", api_key),)
        self.channel = grpc.secure_channel(
            self.endpoint, grpc.ssl_channel_credentials(),
            options=(("grpc.enable_retries", 0),),
        )
        try:
            grpc.channel_ready_future(self.channel).result(timeout=timeout)
        except BaseException:
            self.channel.close()
            raise
        self.stub = atlas_service_pb2_grpc.AtlasServiceStub(self.channel)

    def fetch(self, variant: dict, scorers: list[str], ontology_terms: list[str], timeout: float) -> dict:
        request = self.messages.GetDenseVariantScoresRequest(
            variant=self.genome.Variant(chromosome=variant["chrom"], position=variant["pos"],
                                        reference_bases=variant["ref"], alternate_bases=variant["alt"]).to_proto(),
            organism=self.dna_messages.ORGANISM_HOMO_SAPIENS,
            filter=self.atlas_utils.build_filter(requested_scorers=scorers,
                                                 ontology_terms=ontology_terms or None),
        )
        try:
            response = self.stub.GetDenseVariantScores(request, metadata=self.metadata, timeout=timeout)
        except self.grpc.RpcError as exc:
            code = exc.code().name
            # Never expose authentication material or server text in logs.
            raise _FetchError(code, "Atlas read failed",
                              retryable=code in {"UNAVAILABLE", "DEADLINE_EXCEEDED", "RESOURCE_EXHAUSTED", "ABORTED"},
                              fatal=code in {"UNAUTHENTICATED", "PERMISSION_DENIED", "UNIMPLEMENTED", "INVALID_ARGUMENT"}) from None
        return _response_packet(response, variant)

    def close(self):
        self.channel.close()


def _create_transport(api_key: str, timeout: float):
    return _AtlasTransport(api_key, timeout)


def _api_key(api_key_file, api_key_env) -> str:
    if api_key_file:
        value = Path(api_key_file).expanduser().read_text(encoding="utf-8").strip()
    else:
        value = os.environ.get(api_key_env, "").strip()
    if not value or any(character.isspace() for character in value):
        raise ValueError(f"A single API key is required in --api-key-file or environment {api_key_env}")
    return value


class _Limiter:
    def __init__(self, rate: float):
        self.interval, self.next_time = 1 / rate, 0.0
        self.lock = threading.Lock()

    def acquire(self, stop: threading.Event):
        with self.lock:
            now = time.monotonic()
            begin = max(now, self.next_time)
            self.next_time = begin + self.interval
        if stop.wait(max(0.0, begin - now)):
            raise _FetchError("CANCELLED", "Stopped before a new request", fatal=True)


def _retry_delay(attempt: int) -> float:
    return min(30.0, 2.0 ** min(attempt, 5))


def _save_packet(cache: Path, catalog: dict, variant: dict, packet: dict, session: str) -> dict:
    import numpy as np

    if packet.get("variant") != variant or not isinstance(packet.get("protobuf"), bytes):
        raise ValueError("Atlas response packet has missing raw protobuf or mismatched variant")
    if not packet["protobuf"]:
        raise ValueError("Atlas response protobuf is empty")
    arrays = packet["arrays"]
    names = set()
    for scorer in packet["scorers"]:
        if scorer["name"] in names or not isinstance(scorer.get("is_signed"), bool):
            raise ValueError("Duplicate scorer or missing signedness in Atlas packet")
        names.add(scorer["name"])
        raw = np.asarray(arrays[scorer["raw_key"]])
        if raw.ndim != 2 or raw.dtype.kind != "f" or list(raw.shape) != scorer["shape"]:
            raise ValueError("Atlas raw score shape/dtype is inconsistent")
        if raw.shape != (len(scorer["obs"]["data"]), len(scorer["var"]["data"])):
            raise ValueError("Atlas raw scores and labels have different shapes")
    path = _entry_path(cache, catalog, variant)
    partial = Path(tempfile.mkdtemp(prefix=".partial-", dir=cache))
    metadata = dict(schema=ATLAS_SCHEMA, variant=variant, scorers=packet["scorers"],
                    interval=packet.get("interval"), response_structure=packet.get("response_structure"),
                    namespace=catalog["namespace"], source=catalog["source"],
                    fetch_session=session, fetched_utc=_utc())
    metadata["axis_metadata_sha256"] = _hash(metadata["scorers"])
    try:
        (partial / "response.pb").write_bytes(packet["protobuf"])
        np.savez_compressed(partial / "arrays.npz", **arrays)
        _json_write(partial / "metadata.json", metadata)
        _json_write(partial / "COMPLETE.json", dict(namespace=catalog["namespace"], variant=variant,
                    files={name: _file_hash(partial / name) for name in ("response.pb", "arrays.npz", "metadata.json")}))
        os.replace(partial, path)
    finally:
        if partial.exists():
            shutil.rmtree(partial)
    return metadata


@contextmanager
def _run_lock(output: Path):
    output.mkdir(parents=True, exist_ok=True)
    with (output / ".atlas.lock").open("a") as lock:
        fcntl.flock(lock, fcntl.LOCK_EX)
        yield


def run_sparse(plan_path, out_dir, api_key_file=None, api_key_env="ALPHAGENOME_API_KEY",
               max_variants=6000, workers=2, retries=3, timeout=60, *,
               requests_per_second=2.0, max_consecutive_failures=3) -> dict:
    """Explicit opt-in execution; cache completed variants and stop bounded failures.

    The hard budget is unique variants in the entire plan, including cached
    variants. At most retries+1 RPC attempts are made per missing variant in
    this invocation. Authentication/configuration errors stop immediately.
    Partial completion raises RuntimeError after writing query.json; resume
    with the same plan/out_dir reuses every verified complete cache entry.
    """
    _positive_int(max_variants, "max_variants")
    _positive_int(workers, "workers")
    _positive_int(max_consecutive_failures, "max_consecutive_failures")
    if isinstance(retries, bool) or not isinstance(retries, int) or retries < 0:
        raise ValueError("retries must be a nonnegative integer")
    if not math.isfinite(timeout) or timeout <= 0 or not math.isfinite(requests_per_second) or requests_per_second <= 0:
        raise ValueError("timeout and requests_per_second must be finite and positive")
    plan = _load_plan(plan_path)
    if plan["n_variants"] > max_variants:
        raise ValueError(f"Variant budget exceeded before connection: {plan['n_variants']} > {max_variants}")
    output = Path(out_dir).expanduser().resolve()
    with _run_lock(output):
        return _run_locked(plan, output, api_key_file, api_key_env, workers, retries, timeout,
                           requests_per_second, max_consecutive_failures)


def _run_locked(plan, output, api_key_file, api_key_env, workers, retries, timeout, rate, max_failures):
    cache = output / "cache"
    cache.mkdir(exist_ok=True)
    query_path, catalog_path = output / "query.json", cache / "catalog.json"
    if query_path.exists() and _read_json(query_path).get("plan_id") != plan["plan_id"]:
        raise ValueError("Sparse Atlas output belongs to another plan; use a separate output directory")
    settings = _settings(plan)
    catalog = _read_json(catalog_path) if catalog_path.exists() else None
    if catalog and catalog.get("settings") != settings:
        raise ValueError("Atlas cache request settings differ from this plan")
    if catalog and catalog.get("namespace") != _hash(dict(settings=settings, source=catalog["source"])):
        raise ValueError("Atlas cache namespace fingerprint changed")
    if catalog is None and any(cache.iterdir()):
        raise ValueError("Unrecognized existing Atlas cache; refusing to adopt it")
    complete, missing = {}, []
    for variant in plan["variants"]:
        entry = _cache_entry(cache, catalog, variant) if catalog else None
        if entry:
            complete[_variant_id(variant)] = entry[1]["fetch_session"]
        else:
            missing.append(variant)
    session = _hash(dict(plan=plan["plan_id"], time=_utc()))[:20]
    result = dict(schema=ATLAS_SCHEMA, plan_id=plan["plan_id"], plan_path=str(Path(plan["plan_path"]).resolve()),
                  cache_dir=str(cache), query_json=str(query_path), started_utc=_utc(),
                  fetch_session=session, n_sites=plan["n_sites"], n_variants=plan["n_variants"],
                  reused_variants=len(complete), requested_variants=len(missing),
                  max_variant_rpc_attempts=len(missing) * (retries + 1), rpc_attempts=0,
                  workers=workers, retries=retries, timeout=timeout, requests_per_second=rate,
                  max_consecutive_failures=max_failures, failures={}, source_caveat=SOURCE_CAVEAT)
    if not missing:
        if catalog is None:
            catalog = _empty_catalog(plan)
            _json_write(catalog_path, catalog)
        result.update(status="complete", completed_variants=len(complete), pending_variants=0,
                      finished_utc=_utc(), source=catalog["source"],
                      cache_fetch_sessions=sorted(set(complete.values())), empty_plan=not plan["sites"])
        _json_write(query_path, result)
        return result
    transport = _create_transport(_api_key(api_key_file, api_key_env), timeout)
    source = dict(endpoint=transport.endpoint, sdk_version=transport.sdk_version,
                  protocol="AtlasService.GetDenseVariantScores", lookup="precomputed",
                  verified_source_commit=VERIFIED_SDK_COMMIT, server_release="not_exposed")
    try:
        if catalog is not None and catalog["source"] != source:
            raise ValueError("SDK/endpoint/protocol changed since cached queries; use a separate output directory")
        if catalog is None:
            catalog = dict(schema=ATLAS_SCHEMA, settings=settings, source=source,
                           namespace=_hash(dict(settings=settings, source=source)), created_utc=_utc(),
                           caveat=SOURCE_CAVEAT)
            _json_write(catalog_path, catalog)
        result.update(status="running", source=source)
        _json_write(query_path, result)
        stop, limiter, counter_lock = threading.Event(), _Limiter(rate), threading.Lock()

        def fetch_one(variant):
            for attempt in range(retries + 1):
                limiter.acquire(stop)
                with counter_lock:
                    result["rpc_attempts"] += 1
                try:
                    packet = transport.fetch(variant, plan["scorers"], plan["ontology_terms"], timeout)
                    return _save_packet(cache, catalog, variant, packet, session)
                except _FetchError as exc:
                    if exc.fatal:
                        stop.set()
                    if not exc.retryable or attempt == retries or stop.wait(_retry_delay(attempt)):
                        raise
                except BaseException:
                    stop.set()
                    raise
            raise AssertionError("unreachable retry loop")

        todo = iter(missing)
        consecutive = 0
        interrupted = None
        with ThreadPoolExecutor(max_workers=workers) as pool:
            pending = {}
            def schedule_one():
                if stop.is_set():
                    return
                variant = next(todo, None)
                if variant is not None:
                    pending[pool.submit(fetch_one, variant)] = variant
            for _ in range(min(workers, len(missing))):
                schedule_one()
            try:
                while pending:
                    done, _ = wait(pending, return_when=FIRST_COMPLETED)
                    for future in done:
                        variant = pending.pop(future)
                        key = _variant_id(variant)
                        try:
                            metadata = future.result()
                            complete[key] = metadata["fetch_session"]
                            consecutive = 0
                        except _FetchError as exc:
                            if exc.code != "CANCELLED":
                                result["failures"][key] = dict(code=exc.code, retryable=exc.retryable)
                                consecutive += 1
                            if exc.fatal or consecutive >= max_failures:
                                stop.set()
                        except Exception as exc:
                            # Exception class only: a transport error string could contain a key.
                            result["failures"][key] = dict(code=type(exc).__name__, retryable=False)
                            stop.set()
                        schedule_one()
                    result.update(completed_variants=len(complete), pending_variants=plan["n_variants"] - len(complete))
                    _json_write(query_path, result)
            except BaseException as exc:
                stop.set()
                interrupted = exc
        # A Ctrl-C may race with a successful in-flight response; recover it from disk.
        complete = {}
        for variant in plan["variants"]:
            entry = _cache_entry(cache, catalog, variant)
            if entry:
                complete[_variant_id(variant)] = entry[1]["fetch_session"]
        result.update(completed_variants=len(complete), pending_variants=plan["n_variants"] - len(complete),
                      cache_fetch_sessions=sorted(set(complete.values())), finished_utc=_utc(),
                      status="complete" if len(complete) == plan["n_variants"] else "partial")
        _json_write(query_path, result)
        if interrupted is not None:
            raise interrupted
        if result["status"] != "complete":
            raise RuntimeError(f"Sparse Atlas query incomplete: {len(complete)}/{plan['n_variants']} variants cached; see {query_path}; resume the same plan")
        return result
    finally:
        transport.close()


def _axis_rows(table: dict, ignored=()) -> list[str]:
    columns = table["columns"]
    if len(set(columns)) != len(columns):
        raise ValueError("duplicate axis metadata columns")
    rows = []
    for row in table["data"]:
        if len(row) != len(columns):
            raise ValueError("axis metadata row length mismatch")
        labels = {key: value for key, value in zip(columns, row) if key not in ignored}
        rows.append(json.dumps(labels, sort_keys=True, separators=(",", ":"), ensure_ascii=False))
    if len(set(rows)) != len(rows):
        raise ValueError("ambiguous duplicate biological axis labels")
    return rows


def _vector(entry: tuple[Path, dict], scorer_name: str):
    import numpy as np

    path, metadata = entry
    matches = [item for item in metadata["scorers"] if item["name"] == scorer_name]
    if len(matches) != 1:
        return None, "missing_scorer", "Requested scorer absent from this cached response"
    item = matches[0]
    if not item["is_signed"]:
        return None, "unsigned_scorer", "Unsigned scores do not define a directional response vector"
    try:
        # obs variant/interval identify the input, not a gene/junction axis.
        rows = _axis_rows(item["obs"], ignored=("variant", "interval"))
        columns = _axis_rows(item["var"])
    except ValueError as exc:
        return None, "ambiguous_axes", str(exc)
    with np.load(path / "arrays.npz", allow_pickle=False) as arrays:
        raw = np.asarray(arrays[item["raw_key"]], dtype=np.float64)
    if raw.shape != (len(rows), len(columns)) or not raw.size:
        return None, "empty_or_invalid_shape", "Raw array has empty or incompatible axes"
    if not np.isfinite(raw).all():
        return None, "nonfinite_raw", "At least one raw effect is not finite"
    features = [(row, column) for row in rows for column in columns]
    order = sorted(range(len(features)), key=features.__getitem__)
    return dict(values=raw.reshape(-1)[order], axes=[features[i] for i in order],
                interval=metadata.get("interval"), session=metadata["fetch_session"]), "ok", ""


def _direction_metrics(extra, start, end) -> dict:
    import numpy as np

    delta, displacement = end - start, extra - start
    squared = float(np.dot(delta, delta))
    if squared == 0:
        return dict(cosine=None, projection_fraction=None, projection_raw=None, orthogonal_l2=None)
    norm = math.sqrt(squared)
    projection_raw = float(np.dot(displacement, delta)) / norm
    fraction = projection_raw / norm
    length = float(np.linalg.norm(displacement))
    cosine = max(-1.0, min(1.0, projection_raw / length)) if length > 0 else None
    return dict(cosine=cosine, projection_fraction=fraction, projection_raw=projection_raw,
                orthogonal_l2=float(np.linalg.norm(displacement - fraction * delta)))


GEOMETRY_COLUMNS = [
    "atlas_scorer", "atlas_status", "atlas_reason", "atlas_n_features", "atlas_axis_sha256",
    "atlas_distance_to_anchor_ref", "atlas_distance_to_anchor_alt", "atlas_anchor_delta_l2",
    "atlas_signed_cosine_from_anchor_ref", "atlas_projection_fraction_from_anchor_ref",
    "atlas_projection_raw_from_anchor_ref", "atlas_orthogonal_residual_l2",
    "atlas_estimated_increasing_base", "atlas_estimated_decreasing_base", "atlas_direction_status",
    "atlas_distance_to_estimated_increasing", "atlas_distance_to_estimated_decreasing",
    "atlas_signed_cosine_toward_estimated_increasing", "atlas_projection_fraction_toward_estimated_increasing",
    "atlas_projection_raw_toward_estimated_increasing", "atlas_fetch_sessions",
]
BREADTH_COLUMNS = ["atlas_scorer", "atlas_role", "atlas_base", "atlas_threshold", "atlas_n_features",
                   "atlas_n_above_threshold", "atlas_predicted_response_breadth", "atlas_axis_sha256"]


def summarize_geometry(plan_path, cache_dir, out_dir, effect_thresholds=(0.1, 0.5, 1.0)) -> dict:
    """Export descriptive geometry from complete matched axes; never query Atlas.

    s>0 means the mapped ancient anchor_alt is the estimated increasing base;
    s<0 means anchor_ref. Posterior is retained, not used to reverse this sign.
    At s=0, selection-oriented labels/metrics are absent. Zero anchor deltas
    retain distances and breadth but leave direction metrics undefined.
    """
    import numpy as np

    plan = _load_plan(plan_path)
    thresholds = sorted({float(value) for value in effect_thresholds})
    if not thresholds or any(not math.isfinite(value) or value <= 0 for value in thresholds):
        raise ValueError("effect_thresholds must be finite and positive")
    cache = Path(cache_dir).expanduser().resolve()
    if not (cache / "catalog.json").exists() and (cache / "cache" / "catalog.json").exists():
        cache /= "cache"
    if not (cache / "catalog.json").exists() and plan["n_variants"]:
        raise ValueError(f"No Atlas cache catalog at {cache}; execute an explicit query first")
    catalog = _read_json(cache / "catalog.json") if (cache / "catalog.json").exists() else _empty_catalog(plan)
    if catalog.get("settings") != _settings(plan) or catalog.get("namespace") != _hash(dict(settings=catalog.get("settings"), source=catalog.get("source"))):
        raise ValueError("Atlas cache settings or namespace differ from the plan")
    output = Path(out_dir).expanduser().resolve()
    output.mkdir(parents=True, exist_ok=True)
    manifest_path = output / "geometry.json"
    if manifest_path.exists() and _read_json(manifest_path).get("plan_id") != plan["plan_id"]:
        raise ValueError("Geometry output belongs to a different plan")
    geometry_rows, breadth_rows, counts, all_sessions = [], [], {}, set()
    for site in plan["sites"]:
        entries = {base: _cache_entry(cache, catalog, _variant(site, base))
                   for base in (site["anchor_ref"], site["anchor_alt"], site["extra_alt"])
                   if base != site["ref"]}
        for scorer in plan["scorers"]:
            row = dict(site, atlas_scorer=scorer, atlas_status="ok", atlas_reason="")
            vectors = {}
            for base, entry in entries.items():
                if entry is None:
                    row.update(atlas_status="missing_variant", atlas_reason=f"No complete cache for REF>{base}")
                    break
                vector, status, reason = _vector(entry, scorer)
                if status != "ok":
                    row.update(atlas_status=status, atlas_reason=f"REF>{base}: {reason}")
                    break
                vectors[base] = vector
            if row["atlas_status"] == "ok":
                exemplar = next(iter(vectors.values()))
                if any(value["axes"] != exemplar["axes"] for value in vectors.values()):
                    row.update(atlas_status="axis_mismatch", atlas_reason="Complete obs/var biological labels differ between alleles")
                elif any(value["interval"] != exemplar["interval"] for value in vectors.values()):
                    row.update(atlas_status="interval_mismatch", atlas_reason="Atlas scoring intervals differ between alleles")
            if row["atlas_status"] == "ok":
                raw = {base: value["values"] for base, value in vectors.items()}
                raw[site["ref"]] = np.zeros_like(exemplar["values"])
                anc_ref, anc_alt, extra = (raw[site[key]] for key in ("anchor_ref", "anchor_alt", "extra_alt"))
                n_features = len(extra)
                axis_hash = _hash(exemplar["axes"])
                sessions = sorted({value["session"] for value in vectors.values()})
                all_sessions.update(sessions)
                row.update(atlas_n_features=n_features, atlas_axis_sha256=axis_hash,
                           atlas_distance_to_anchor_ref=float(np.linalg.norm(extra - anc_ref)),
                           atlas_distance_to_anchor_alt=float(np.linalg.norm(extra - anc_alt)),
                           atlas_anchor_delta_l2=float(np.linalg.norm(anc_alt - anc_ref)),
                           atlas_fetch_sessions=",".join(sessions))
                direction = _direction_metrics(extra, anc_ref, anc_alt)
                row.update(atlas_signed_cosine_from_anchor_ref=direction["cosine"],
                           atlas_projection_fraction_from_anchor_ref=direction["projection_fraction"],
                           atlas_projection_raw_from_anchor_ref=direction["projection_raw"],
                           atlas_orthogonal_residual_l2=direction["orthogonal_l2"])
                if site["s"] == 0:
                    row["atlas_direction_status"] = "undefined_s_zero"
                else:
                    increasing, decreasing = ((site["anchor_alt"], site["anchor_ref"]) if site["s"] > 0
                                              else (site["anchor_ref"], site["anchor_alt"]))
                    oriented = _direction_metrics(extra, raw[decreasing], raw[increasing])
                    row.update(atlas_estimated_increasing_base=increasing, atlas_estimated_decreasing_base=decreasing,
                               atlas_direction_status="estimated_from_s_sign_posterior_retained",
                               atlas_distance_to_estimated_increasing=float(np.linalg.norm(extra - raw[increasing])),
                               atlas_distance_to_estimated_decreasing=float(np.linalg.norm(extra - raw[decreasing])),
                               atlas_signed_cosine_toward_estimated_increasing=oriented["cosine"],
                               atlas_projection_fraction_toward_estimated_increasing=oriented["projection_fraction"],
                               atlas_projection_raw_toward_estimated_increasing=oriented["projection_raw"])
                if row["atlas_anchor_delta_l2"] == 0:
                    row.update(atlas_status="degenerate_anchor_delta", atlas_reason="Anchor effect vectors coincide; directional geometry is undefined")
                for role, field in (("anchor_ref", "anchor_ref"), ("anchor_alt", "anchor_alt"), ("extra", "extra_alt")):
                    base = site[field]
                    for threshold in thresholds:
                        above = int(np.count_nonzero(np.abs(raw[base]) > threshold))
                        breadth_rows.append(dict(site, atlas_scorer=scorer, atlas_role=role, atlas_base=base,
                                                 atlas_threshold=threshold, atlas_n_features=n_features,
                                                 atlas_n_above_threshold=above,
                                                 atlas_predicted_response_breadth=above / n_features,
                                                 atlas_axis_sha256=axis_hash))
            counts[row["atlas_status"]] = counts.get(row["atlas_status"], 0) + 1
            geometry_rows.append(row)
    paths = {"geometry": output / "geometry.tsv", "breadth": output / "breadth.tsv"}
    for key, rows, columns in (("geometry", geometry_rows, GEOMETRY_COLUMNS), ("breadth", breadth_rows, BREADTH_COLUMNS)):
        fd, temporary = tempfile.mkstemp(prefix=key + ".", suffix=".tmp", dir=output)
        try:
            with os.fdopen(fd, "w", newline="", encoding="utf-8") as handle:
                writer = csv.DictWriter(handle, fieldnames=plan["candidate_columns"] + columns,
                                        delimiter="\t", lineterminator="\n", extrasaction="raise")
                writer.writeheader()
                writer.writerows(rows)
            os.replace(temporary, paths[key])
        finally:
            if os.path.exists(temporary):
                os.unlink(temporary)
    result = dict(schema=ATLAS_SCHEMA, plan_id=plan["plan_id"], created_utc=_utc(), cache_dir=str(cache),
                  source=catalog["source"], cache_namespace=catalog["namespace"], source_caveat=SOURCE_CAVEAT,
                  caveat=GEOMETRY_CAVEAT, effect_thresholds=thresholds, threshold_rule="abs(raw_effect) > threshold",
                  n_sites=plan["n_sites"], n_geometry_rows=len(geometry_rows), n_breadth_rows=len(breadth_rows),
                  status_counts=counts, cache_fetch_sessions=sorted(all_sessions),
                  outputs={key: str(path) for key, path in paths.items()}, manifest_path=str(manifest_path),
                  files={key: _file_hash(path) for key, path in paths.items()})
    _json_write(manifest_path, result)
    return result


# ========================================================================
# evolve
# ========================================================================

# Ancient trajectories, shared PCA and evolution command handling.





# ========================================================================
# Trajectory
# ========================================================================

# Allele-specific, descriptive trajectories from targeted ancient PGEN rows.
# 
# This module does not estimate selection coefficients, reconstruct unrepresented
# alleles, infer REF ancestry, or claim to reproduce the published selection GLMM.
# The aggregation layer needs only the standard library.  PGEN access lazily
# imports numpy and the official pgenlib Python API.


from bisect import bisect_right
from collections import Counter
from contextlib import ExitStack
import csv
from datetime import datetime, timezone
import gzip
import hashlib
import importlib.metadata
import json
import math
import os
from pathlib import Path
import shutil
import tempfile


EVOLUTION_SCHEMA_VERSION = "allele3-trajectory-1"
DEFAULT_BINS = (0, 2000, 4000, 6000, 8000, 10000)
STATUSES = {"assayed_pass", "not_in_high_quality", "not_represented_in_panel",
            "represented_multiallelic_unsupported"}
KINDS = {"hardcall", "dosage"}
META_REQUIRED = {"sample_id", "cohort", "age_bp", "include", "stratum"}
SELECTED_REQUIRED = {"variant_id", "pvar_index", "chrom", "pos", "ref", "alt", "status"}
BASE_FIELDS = ["variant_id", "build", "chrom", "pos", "ref", "alt", "pvar_index"]
TIME_FIELDS = ["cohort", "time_bin", "start_bp", "end_bp", "right_inclusive"]
BIN_FIELDS = BASE_FIELDS + TIME_FIELDS + [
    "group_type", "stratum", "count_kind", "dosage_evidence", "n_eligible",
    "n_nonmissing", "call_rate", "ac", "expected_ac", "an", "af", "bin_status",
]
STD_FIELDS = BASE_FIELDS + TIME_FIELDS + [
    "count_kind", "dosage_evidence", "n_reference_strata", "n_supported_strata",
    "supported_weight", "missing_strata", "af_standardized", "standardization_status",
]
AVAIL_FIELDS = BASE_FIELDS + [
    "status", "count_kind", "dosage_evidence", "n_eligible", "n_nonmissing",
    "trajectory_available", "availability_note",
]
QPADM_COMPONENTS = {"3way": ("whg", "eef", "steppe"),
                    "4way": ("anf", "ehg", "whg", "icr")}
EVOLUTION_OUTPUT_NAMES = ("trajectory_bins.tsv", "trajectory_standardized.tsv",
                "sample_bin_composition.tsv", "sample_inclusion.tsv",
                "allele_availability.tsv", "trajectory_manifest.json")


def _evolution_open_text(path):
    path = Path(path)
    opener = gzip.open if path.suffix.lower() in {".gz", ".bgz", ".bgzip"} else open
    return opener(path, "rt", encoding="utf-8-sig", newline="")


def _trajectory_rows(path, required):
    with _evolution_open_text(path) as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        fields = reader.fieldnames or []
        if len(fields) != len(set(fields)) or not required.issubset(fields):
            raise ValueError(f"{path}: requires unique TSV columns {sorted(required)}")
        for line, row in enumerate(reader, 2):
            if None in row or any(value is None for value in row.values()):
                raise ValueError(f"{path}:{line}: wrong number of TSV fields")
            yield {key: value.strip() for key, value in row.items()}


def _evolution_missing(value):
    return value is None or (isinstance(value, str) and value.strip().lower() in
                             {"", ".", "na", "nan", "none"})


def _trajectory_number(value, label):
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label}: expected a finite number, got {value!r}") from exc
    if not math.isfinite(result):
        raise ValueError(f"{label}: expected a finite number, got {value!r}")
    return result


def _evolution_integer(value, label, minimum=0):
    result = _trajectory_number(value, label)
    if result != int(result) or result < minimum:
        raise ValueError(f"{label}: expected an integer >= {minimum}, got {value!r}")
    return int(result)


def _sample_ids(values):
    ids = list(values)
    if not ids or any(not isinstance(x, str) or _evolution_missing(x) or x != x.strip() for x in ids):
        raise ValueError("sample_ids must be a nonempty sequence of exact, nonmissing strings")
    if len(ids) != len(set(ids)):
        raise ValueError("Duplicate sample_id; IID alone must uniquely identify every PSAM row")
    return ids


def read_psam_ids(psam_path):
    """Read unique IID values in original PSAM order; never guess IID from FID."""
    ids, fields = [], None
    with _evolution_open_text(psam_path) as handle:
        for line_no, line in enumerate(handle, 1):
            line = line.strip()
            if not line or line.startswith("##"):
                continue
            parts = line.split()
            if fields is None:
                fields = [x.lstrip("#") for x in parts]
                if "IID" not in fields or len(fields) != len(set(fields)):
                    raise ValueError(f"{psam_path}: PSAM must have a unique #IID or #FID ... IID header")
                iid_col = fields.index("IID")
                continue
            if line.startswith("#") or len(parts) != len(fields):
                raise ValueError(f"{psam_path}:{line_no}: malformed PSAM row")
            ids.append(parts[iid_col])
    return _sample_ids(ids)


def _evolution_variant(row):
    result = dict(row)
    missing = SELECTED_REQUIRED - result.keys()
    if missing:
        raise ValueError(f"Variant is missing fields: {sorted(missing)}")
    vid = result["variant_id"]
    if not isinstance(vid, str) or _evolution_missing(vid):
        raise ValueError("Each target needs a nonmissing variant_id")
    result["build"] = result.get("build", "GRCh37")
    if result["build"] != "GRCh37":
        raise ValueError(f"{vid}: this PGEN trajectory adapter requires explicitly harmonized GRCh37 targets")
    chrom = str(result["chrom"]).removeprefix("chr")
    if chrom not in {str(x) for x in range(1, 23)}:
        raise ValueError(f"{vid}: only autosomes 1-22 are supported, got {chrom!r}")
    result["chrom"] = chrom
    result["pos"] = _evolution_integer(result["pos"], f"{vid} pos", 1)
    for key in ("ref", "alt"):
        result[key] = str(result[key]).upper()
        if result[key] not in {"A", "C", "G", "T"}:
            raise ValueError(f"{vid}: only single A/C/G/T alleles are supported")
    if result["ref"] == result["alt"]:
        raise ValueError(f"{vid}: REF and ALT must differ")
    status = result["status"]
    if status not in STATUSES:
        raise ValueError(f"{vid}: unsupported availability status {status!r}")
    index = result["pvar_index"]
    result["pvar_index"] = None if _evolution_missing(index) else _evolution_integer(index, f"{vid} pvar_index")
    if status in {"assayed_pass", "not_in_high_quality"} and result["pvar_index"] is None:
        raise ValueError(f"{vid}: a represented panel row needs its original pvar_index")
    if status in {"not_represented_in_panel", "represented_multiallelic_unsupported"} and result["pvar_index"] is not None:
        raise ValueError(f"{vid}: an unrepresented/unsupported allele must not be assigned a readable PVAR row")
    return result


def _load_pgen_backend():
    try:
        import numpy as np
        import pgenlib
    except ImportError as exc:
        raise RuntimeError("Targeted PGEN reading needs numpy and pgenlib. Install into the active "
                           "environment with: python -m pip install numpy pgenlib") from exc
    return np, pgenlib


def _file_info(path, digest=False):
    path = Path(path).expanduser().resolve()
    info = {"path": str(path), "size": path.stat().st_size,
            "mtime_ns": path.stat().st_mtime_ns}
    if digest:
        hasher = hashlib.sha256()
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1048576), b""):
                hasher.update(chunk)
        info["sha256"] = hasher.hexdigest()
    return info


def iter_pgen_variants(pgen_path, psam_path, selected_variants_path, count_kind="hardcall"):
    """Return (sample_ids, iterator) for selected original zero-based PVAR rows.

    Selected TSV: variant_id,pvar_index,chrom,pos,ref,alt,status.  The caller
    must have verified those REF/ALT labels against the original, biallelic
    PVAR; PGEN does not contain the nucleotide labels.  No filtered/reindexed
    PVAR may be substituted.  Read allele_idx=1 explicitly.  Every yielded
    dosages list is independently owned and is in the complete PSAM order.

    ``dosage`` explicitly selects read_dosages.  That API can use hardcalls
    when explicit dosages are absent, so dosage_evidence is always recorded;
    integer-only outputs do not establish that posterior dosages were stored.
    """
    if count_kind not in KINDS:
        raise ValueError("count_kind must be hardcall or dosage")
    sample_ids = read_psam_ids(psam_path)
    selected, seen_ids, seen_indexes = [], set(), set()
    for raw in _trajectory_rows(selected_variants_path, SELECTED_REQUIRED):
        row = _evolution_variant(raw)
        if row["variant_id"] in seen_ids:
            raise ValueError(f"Duplicate selected variant_id: {row['variant_id']}")
        seen_ids.add(row["variant_id"])
        index = row["pvar_index"]
        if index is not None:
            if index in seen_indexes:
                raise ValueError(f"Original PVAR row {index} was assigned to more than one target; "
                                 "this adapter accepts one biallelic ALT per original row")
            seen_indexes.add(index)
        selected.append(row)

    def generate():
        reader = None
        source = {"reader": "pgenlib", "psam": _file_info(psam_path, True),
                  "selected_variants": _file_info(selected_variants_path, True),
                  "requested_count_kind": count_kind, "allele_idx": 1,
                  "index_basis": "zero_based_original_unfiltered_PVAR_data_rows"}
        try:
            for selected_row in selected:
                row = dict(selected_row, count_kind=count_kind, dosages=None,
                           dosage_evidence="not_read", source=source)
                if row["status"] != "assayed_pass":
                    yield row
                    continue
                if reader is None:
                    np, pgenlib = _load_pgen_backend()
                    source["pgen"] = _file_info(pgen_path)
                    try:
                        source["pgenlib_version"] = importlib.metadata.version("pgenlib")
                    except importlib.metadata.PackageNotFoundError:
                        source["pgenlib_version"] = getattr(pgenlib, "__version__", "unknown")
                    reader = pgenlib.PgenReader(os.fsencode(Path(pgen_path).expanduser()),
                                               raw_sample_ct=len(sample_ids))
                    if reader.get_raw_sample_ct() != len(sample_ids):
                        raise ValueError("PGEN sample count disagrees with the full PSAM")
                    if count_kind == "dosage" and not callable(getattr(reader, "read_dosages", None)):
                        raise RuntimeError("This pgenlib has no read_dosages API; upgrade pgenlib or "
                                           "explicitly choose hardcall. No automatic fallback was used.")
                    buffer = np.empty(len(sample_ids), dtype=np.int8 if count_kind == "hardcall" else np.float64)
                    variant_ct = reader.get_variant_ct()
                index = row["pvar_index"]
                if index >= variant_ct:
                    raise ValueError(f"{row['variant_id']}: pvar_index {index} is outside PGEN's "
                                     f"{variant_ct} original rows")
                try:
                    if count_kind == "hardcall":
                        reader.read(index, buffer, allele_idx=1)
                    else:
                        reader.read_dosages(index, buffer, allele_idx=1)
                except Exception as exc:
                    raise RuntimeError(f"PGEN read failed for {row['variant_id']} at original row {index}: {exc}") from exc
                row["dosages"] = [None if x == -9 else (int(x) if count_kind == "hardcall" else float(x))
                                   for x in buffer]
                if count_kind == "dosage":
                    present = [x for x in row["dosages"] if x is not None]
                    row["dosage_evidence"] = ("fractional_values_present" if any(x != int(x) for x in present)
                                              else "integer_only_storage_unverified" if present else "all_missing")
                else:
                    row["dosage_evidence"] = "hardcalls_only"
                yield row
        finally:
            if reader is not None:
                reader.close()

    return sample_ids, generate()


def _trajectory_metadata(path, sample_ids, bins):
    by_id = {}
    for row in _trajectory_rows(path, META_REQUIRED):
        sid = row["sample_id"]
        if _evolution_missing(sid) or sid in by_id:
            raise ValueError(f"Duplicate or missing metadata sample_id: {sid!r}")
        if row["include"] not in {"0", "1"}:
            raise ValueError(f"{sid}: include must be exactly 0 or 1")
        if row["cohort"] not in {"ancient", "modern"}:
            raise ValueError(f"{sid}: cohort must be ancient or modern")
        row["included"] = row["include"] == "1"
        row["age"] = None
        if not _evolution_missing(row["age_bp"]):
            row["age"] = _trajectory_number(row["age_bp"], f"{sid} age_bp")
        if row["included"] and _evolution_missing(row["stratum"]):
            raise ValueError(f"{sid}: included samples need a predeclared stratum")
        if row["included"] and row["cohort"] == "ancient" and (row["age"] is None or row["age"] < 0):
            raise ValueError(f"{sid}: included ancient sample needs a finite nonnegative age_bp")
        row["bin_index"] = None
        if not row["included"]:
            row["inclusion_status"] = "metadata_excluded"
        elif row["cohort"] == "modern":
            row["bin_index"] = -1
            row["inclusion_status"] = "included"
        elif not bins[0] <= row["age"] <= bins[-1]:
            row["inclusion_status"] = "outside_time_range"
        else:
            row["bin_index"] = min(bisect_right(bins, row["age"]) - 1, len(bins) - 2)
            row["inclusion_status"] = "included"
        by_id[sid] = row
    expected, found = set(sample_ids), set(by_id)
    if expected != found:
        raise ValueError("Metadata IDs must exactly match PSAM/sample_ids: "
                         f"missing={sorted(expected-found)[:5]}, extra={sorted(found-expected)[:5]}")
    return [by_id[sid] for sid in sample_ids]


def _weights(path):
    if path is None:
        return None
    weights = {}
    for row in _trajectory_rows(path, {"stratum", "weight"}):
        key = row["stratum"]
        if _evolution_missing(key) or key in weights:
            raise ValueError(f"Duplicate or missing weight stratum: {key!r}")
        weight = _trajectory_number(row["weight"], f"{key} weight")
        if weight < 0:
            raise ValueError("Reference weights must be nonnegative")
        weights[key] = weight
    total = math.fsum(weights.values())
    if not math.isclose(total, 1.0, rel_tol=0, abs_tol=1e-8):
        raise ValueError(f"Reference weights must sum to 1, got {total}")
    return {key: weight/total for key, weight in weights.items() if weight > 0}


def _time_fields(cohort, index, bins):
    if cohort == "modern":
        return dict(cohort=cohort, time_bin="modern", start_bp=".", end_bp=".", right_inclusive=".")
    start, end = bins[index:index+2]
    last = index == len(bins) - 2
    return dict(cohort=cohort, time_bin=f"[{start:g},{end:g}{']' if last else ')'}",
                start_bp=start, end_bp=end, right_inclusive=int(last))


def _writer(stack, stage, name, fields):
    handle = stack.enter_context((stage / name).open("w", newline="", encoding="utf-8"))
    writer = csv.DictWriter(handle, fieldnames=fields, delimiter="\t", lineterminator="\n")
    writer.writeheader()
    return writer


def _write(writer, row):
    writer.writerow({key: "." if value is None else value for key, value in row.items()})


def _composition(metadata, members):
    rows = [metadata[i] for i in members]
    ages = [r["age"] for r in rows if r["age"] is not None and r["cohort"] == "ancient"]
    result = {"n_eligible": len(rows), "mean_age_bp": math.fsum(ages)/len(ages) if ages else None}
    for source in ("region", "batch", "data_source", "data_types"):
        result[source + "_counts_json"] = json.dumps(dict(sorted(Counter(r.get(source, ".") or "."
                                                               for r in rows).items())), ensure_ascii=False)
    for model, components in QPADM_COMPONENTS.items():
        valid = [r for r in rows if r.get("qpadm_valid") == "1" and r.get("qpadm_model") == model]
        values = []
        for row in valid:
            vector = [_trajectory_number(row.get("qpadm_" + c), f"{row['sample_id']} qpadm_{c}") for c in components]
            if any(not 0 <= x <= 1 for x in vector) or not math.isclose(sum(vector), 1, abs_tol=0.02):
                raise ValueError(f"{row['sample_id']}: qpadm_valid=1 contradicts invalid {model} proportions")
            values.append(vector)
        result[f"n_qpadm_{model}_valid"] = len(values)
        for j, component in enumerate(components):
            result[f"mean_qpadm_{model}_{component}"] = math.fsum(v[j] for v in values)/len(values) if values else None
    return result


def _call(value, kind, label):
    if value is None:
        return None
    # The generic interface uses None for missingness; pgenlib's -9 is
    # translated by the reader.  Reject NaN and all other sentinels here.
    value = _trajectory_number(value, label)
    if not 0 <= value <= 2 or (kind == "hardcall" and value != int(value)):
        raise ValueError(f"{label}: {kind} needs {'0, 1, or 2' if kind == 'hardcall' else 'a value in [0,2]'} or None")
    return value


def aggregate_trajectories(variants, sample_ids, metadata_path, out_dir,
                           bins=DEFAULT_BINS, min_samples=10, weights_path=None):
    """Write streaming, descriptive ALT trajectories and return a manifest.

    ``variants`` yields dictionaries from iter_pgen_variants (or the same
    protocol).  Low-sample raw bins retain counts and descriptive AF, with
    bin_status=below_min_samples.  Standardized AF requires every positively
    weighted stratum to have >=min_samples nonmissing calls.  Missing or
    unrepresented alleles never become a synthetic all-zero trajectory.
    All output files are committed only after successful validation.
    """
    sample_ids = _sample_ids(sample_ids)
    edges = tuple(_trajectory_number(x, "bin edge") for x in bins)
    if len(edges) < 2 or edges[0] < 0 or any(a >= b for a, b in zip(edges, edges[1:])):
        raise ValueError("bins must contain at least two strictly increasing nonnegative boundaries")
    min_samples = _evolution_integer(min_samples, "min_samples", 1)
    metadata = _trajectory_metadata(metadata_path, sample_ids, edges)
    weights = _weights(weights_path)
    time_keys = [("ancient", index) for index in range(len(edges)-1)] + [("modern", -1)]
    strata = {cohort: sorted({r["stratum"] for r in metadata
                             if r["cohort"] == cohort and r["inclusion_status"] == "included"})
              for cohort in ("ancient", "modern")}
    groups = [(cohort, index, None) for cohort, index in time_keys]
    groups += [(cohort, index, stratum) for cohort, index in time_keys for stratum in strata[cohort]]
    members = {key: [] for key in groups}
    sample_groups = []
    for index, row in enumerate(metadata):
        keys = []
        if row["inclusion_status"] == "included":
            keys = [(row["cohort"], row["bin_index"], None),
                    (row["cohort"], row["bin_index"], row["stratum"])]
            for key in keys:
                members[key].append(index)
        sample_groups.append(keys)
    active = sum(bool(keys) for keys in sample_groups)
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    existing = [name for name in EVOLUTION_OUTPUT_NAMES if (out / name).exists()]
    if existing:
        raise ValueError(f"Trajectory output already exists ({existing[0]}); use a fresh output directory")
    stage = Path(tempfile.mkdtemp(prefix=".trajectory-", dir=out))
    counts = Counter({key: 0 for key in ("n_targets", "n_bin_rows", "n_standardized_rows", *sorted(STATUSES))})
    kinds, sources = Counter(), {}
    seen_ids, seen_keys = set(), set()
    iterator = iter(variants)
    try:
        with ExitStack() as stack:
            bins_writer = _writer(stack, stage, "trajectory_bins.tsv", BIN_FIELDS)
            std_writer = _writer(stack, stage, "trajectory_standardized.tsv", STD_FIELDS)
            avail_writer = _writer(stack, stage, "allele_availability.tsv", AVAIL_FIELDS)
            composition_fields = TIME_FIELDS + ["group_type", "stratum", "n_eligible", "mean_age_bp"]
            composition_fields += [x + "_counts_json" for x in ("region", "batch", "data_source", "data_types")]
            for model, components in QPADM_COMPONENTS.items():
                composition_fields += [f"n_qpadm_{model}_valid"] + [f"mean_qpadm_{model}_{c}" for c in components]
            comp_writer = _writer(stack, stage, "sample_bin_composition.tsv", composition_fields)
            for cohort, index, stratum in groups:
                _write(comp_writer, dict(_time_fields(cohort, index, edges),
                                         group_type="raw_all" if stratum is None else "by_stratum",
                                         stratum=stratum, **_composition(metadata, members[(cohort, index, stratum)])))
            inclusion = _writer(stack, stage, "sample_inclusion.tsv",
                                ["sample_id", "cohort", "age_bp", "include", "stratum", "inclusion_status"])
            for row in metadata:
                _write(inclusion, {key: row[key] for key in inclusion.fieldnames})
            for raw in iterator:
                row = _evolution_variant(raw)
                vid = row["variant_id"]
                key = tuple(row[x] for x in ("build", "chrom", "pos", "ref", "alt"))
                if vid in seen_ids or key in seen_keys:
                    raise ValueError(f"Duplicate trajectory target ID or exact allele: {vid}")
                seen_ids.add(vid)
                seen_keys.add(key)
                kind = row.get("count_kind")
                if kind not in KINDS:
                    raise ValueError(f"{vid}: count_kind must explicitly be hardcall or dosage")
                evidence = row.get("dosage_evidence", "hardcalls_only" if kind == "hardcall" else "external_dosage_provenance_unverified")
                counts["n_targets"] += 1
                counts[row["status"]] += 1
                kinds[kind] += 1
                if row.get("source"):
                    encoded = json.dumps(row["source"], sort_keys=True)
                    sources[encoded] = dict(row["source"])
                base = {x: row[x] for x in BASE_FIELDS}
                availability = dict(base, status=row["status"], count_kind=kind,
                                    dosage_evidence=evidence, n_eligible=active)
                if row["status"] != "assayed_pass":
                    if row.get("dosages") is not None:
                        raise ValueError(f"{vid}: unavailable/failed-QC allele must have dosages=None")
                    note = {"not_represented_in_panel": "unrepresented_allele_not_measured_not_zero",
                            "not_in_high_quality": "excluded_by_variant_quality_control",
                            "represented_multiallelic_unsupported": "represented_but_multiallelic_reader_not_supported"}[row["status"]]
                    _write(avail_writer, dict(availability, n_nonmissing=None, trajectory_available=0,
                                              availability_note=note))
                    continue
                calls = row.get("dosages")
                if calls is None or len(calls) != len(sample_ids):
                    raise ValueError(f"{vid}: dosages length must equal the complete sample_ids length")
                totals = {group: [0, 0.0] for group in groups}
                nonmissing = 0
                for sample_index, value in enumerate(calls):
                    value = _call(value, kind, f"{vid}/{sample_ids[sample_index]}")
                    if value is None or not sample_groups[sample_index]:
                        continue
                    nonmissing += 1
                    for group in sample_groups[sample_index]:
                        totals[group][0] += 1
                        totals[group][1] += value
                _write(avail_writer, dict(availability, n_nonmissing=nonmissing,
                                          trajectory_available=int(nonmissing > 0),
                                          availability_note="descriptive_biallelic_panel_estimate" if nonmissing else "no_nonmissing_included_calls"))
                for cohort, index, stratum in groups:
                    group = (cohort, index, stratum)
                    n, allele_sum = totals[group]
                    eligible = len(members[group])
                    bin_status = ("no_eligible_samples" if not eligible else "all_calls_missing" if not n
                                  else "below_min_samples" if n < min_samples else "ok")
                    record = dict(base, **_time_fields(cohort, index, edges),
                                  group_type="raw_all" if stratum is None else "by_stratum", stratum=stratum,
                                  count_kind=kind, dosage_evidence=evidence, n_eligible=eligible, n_nonmissing=n,
                                  call_rate=n/eligible if eligible else None,
                                  ac=int(allele_sum) if kind == "hardcall" else None,
                                  expected_ac=allele_sum if kind == "dosage" else None,
                                  an=2*n, af=allele_sum/(2*n) if n else None, bin_status=bin_status)
                    _write(bins_writer, record)
                    counts["n_bin_rows"] += 1
                if weights is not None:
                    for cohort, index in time_keys:
                        supported, missing, contributions = [], [], []
                        for stratum, weight in weights.items():
                            n, allele_sum = totals.get((cohort, index, stratum), (0, 0))
                            if n < min_samples:
                                missing.append(stratum)
                            else:
                                supported.append(weight)
                                contributions.append(weight*allele_sum/(2*n))
                        _write(std_writer, dict(base, **_time_fields(cohort, index, edges), count_kind=kind,
                                                dosage_evidence=evidence, n_reference_strata=len(weights),
                                                n_supported_strata=len(supported), supported_weight=math.fsum(supported),
                                                missing_strata=json.dumps(sorted(missing), ensure_ascii=False),
                                                af_standardized=math.fsum(contributions) if not missing else None,
                                                standardization_status="complete_support" if not missing else "insufficient_stratum_support"))
                        counts["n_standardized_rows"] += 1
        manifest = {
            "schema_version": EVOLUTION_SCHEMA_VERSION, "status": "complete",
            "created_utc": datetime.now(timezone.utc).isoformat(), "build": "GRCh37",
            "scope": "autosomal_diploid_descriptive_allele_specific_trajectories",
            "metadata": _file_info(metadata_path, True),
            "sample_order_sha256": hashlib.sha256("\n".join(sample_ids).encode()).hexdigest(),
            "n_samples": len(sample_ids), "n_included_in_time_scope": active,
            "sample_inclusion_counts": dict(Counter(r["inclusion_status"] for r in metadata)),
            "bins_bp": list(edges), "bin_convention": "left_closed_right_open_except_final_right_closed",
            "modern_endpoint": "separate_cohort_never_pooled_into_youngest_ancient_bin",
            "min_samples": min_samples, "raw_low_sample_af": "retained_with_below_min_samples_status",
            "weights_input": _file_info(weights_path, True) if weights_path else None,
            "reference_weights": weights,
            "standardization": "fixed_weights_complete_support_only" if weights else "not_requested",
            "counts": dict(counts), "count_kind_counts": dict(kinds), "sources": list(sources.values()),
            "outputs": {name: str((out / name).resolve()) for name in EVOLUTION_OUTPUT_NAMES},
            "limitations": [
                "No selection coefficients or inferential P values are estimated.",
                "Unrepresented alleles are unmeasured, never zero-frequency trajectories.",
                "AN is twice nonmissing diploid calls within the panel's allele model, not independent verification of all bases.",
                "Dosage reading can include hardcall fallback; integer values alone do not verify stored posterior dosages.",
                "PGEN does not preserve genotype probability triplets or per-call read/quality information.",
                "Fixed stratum weights control only the specified observed composition; they do not reproduce the published GLMM.",
                "Point sample dates are used; date and imputation uncertainty are not propagated.",
                "A published biallelic selection coefficient must not be transferred to another ALT at the same coordinate.",
                "Directional selection does not establish balancing selection or pleiotropic maintenance.",
            ],
        }
        (stage / "trajectory_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False)+"\n", encoding="utf-8")
        for name in EVOLUTION_OUTPUT_NAMES:
            os.replace(stage / name, out / name)
        return manifest
    finally:
        close = getattr(iterator, "close", None)
        if callable(close):
            close()
        shutil.rmtree(stage, ignore_errors=True)


# ========================================================================
# Pca time
# ========================================================================

# Fixed-basis PCA planning and offline PC1-PC2-time visualization.
# 
# Planning never executes PLINK or downloads inputs.  The generated workflow fits
# one background-marker basis and projects every plotted sample through the same
# score transform.  Time and target-allele dosage are overlays, not PCA features.


from bisect import bisect_right
from collections import Counter, defaultdict
import csv
import gzip
import hashlib
import html
import json
import math
import os
from pathlib import Path
import shlex
import sys


META_COLUMNS = {"sample_id", "cohort", "age_bp", "include", "stratum"}
LONG_LD_SOURCE = "https://genome.sph.umich.edu/wiki/Regions_of_high_linkage_disequilibrium_(LD)"
# Explicit GRCh37 table, positions interpreted as 1-based inclusive intervals.
# These coarse historical masks are a sensitivity preset, not a completeness claim.
LONG_LD_GRCH37 = [
    (1, 48000000, 52000000), (2, 86000000, 100500000),
    (2, 134500000, 138000000), (2, 183000000, 190000000),
    (3, 47500000, 50000000), (3, 83500000, 87000000), (3, 89000000, 97500000),
    (5, 44500000, 50500000), (5, 98000000, 100500000),
    (5, 129000000, 132000000), (5, 135500000, 138500000),
    (6, 25000000, 35000000), (6, 57000000, 64000000), (6, 140000000, 142500000),
    (7, 55000000, 66000000), (8, 7000000, 13000000),
    (8, 43000000, 50000000), (8, 112000000, 115000000),
    (10, 37000000, 43000000), (11, 46000000, 57000000),
    (11, 87500000, 90500000), (12, 33000000, 40000000),
    (12, 109500000, 112000000), (20, 32000000, 34500000),
]


def _evolution_open(path):
    return (gzip.open(path, "rt", encoding="utf-8-sig", newline="")
            if str(path).endswith((".gz", ".bgz"))
            else open(path, "rt", encoding="utf-8-sig", newline=""))


def _pca_rows(path, required):
    with _evolution_open(path) as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        names = reader.fieldnames or []
        if len(names) != len(set(names)) or not required.issubset(names):
            raise ValueError(f"{path}: requires unique exact columns {sorted(required)}")
        for line, row in enumerate(reader, 2):
            if None in row or any(value is None for value in row.values()):
                raise ValueError(f"{path}:{line}: wrong number of fields")
            yield row


def _pca_number(value, label):
    try:
        value = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label}: finite number required") from exc
    if not math.isfinite(value):
        raise ValueError(f"{label}: finite number required")
    return value


def _pca_metadata(path):
    result = {}
    for row in _pca_rows(path, META_COLUMNS):
        sid = row["sample_id"]
        if not sid or sid == "." or any(x.isspace() for x in sid) or sid in result:
            raise ValueError(f"Duplicate or invalid metadata sample_id: {sid!r}")
        if row["include"] not in {"0", "1"} or row["cohort"] not in {"ancient", "modern"}:
            raise ValueError(f"{sid}: invalid include/cohort")
        age = _pca_number(row["age_bp"], f"{sid} age_bp")
        if age < 0:
            raise ValueError(f"{sid}: age_bp must be nonnegative")
        if row["include"] == "1" and row["stratum"] in {"", "."}:
            raise ValueError(f"{sid}: included sample needs stratum")
        row["age"] = age
        result[sid] = row
    if not result or not any(r["include"] == "1" for r in result.values()):
        raise ValueError("No included metadata samples")
    return result


def _trajectory_module():
    import evolve as trajectory
    return trajectory


def _fingerprint(path):
    path = Path(path).expanduser().resolve()
    stat = path.stat()
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        digest.update(handle.read(65536))
        if stat.st_size > 65536:
            handle.seek(max(65536, stat.st_size - 65536))
            digest.update(handle.read(65536))
    result = {"path": str(path), "size": stat.st_size, "mtime_ns": stat.st_mtime_ns,
              "edge_sha256": digest.hexdigest()}
    if path.suffix in {'.tsv', '.json', '.allele', '.acount', '.keep', '.ids'} and stat.st_size <= 16 * 1024 * 1024:
        result['sha256'] = _evolution_sha(path)
    return result


def _evolution_sha(path):
    h = hashlib.sha256()
    with open(path, "rb") as handle:
        for block in iter(lambda: handle.read(1048576), b""):
            h.update(block)
    return h.hexdigest()


def _atomic_text(path, contents):
    path = Path(path)
    partial = Path(str(path) + ".partial")
    if partial.exists():
        raise ValueError(f"Interrupted output exists: {partial}")
    try:
        with partial.open("w", encoding="utf-8", newline="") as handle:
            handle.write(contents)
        os.replace(partial, path)
    finally:
        if partial.exists():
            partial.unlink()


def _evolution_chrom(value):
    value = str(value)
    if value.lower().startswith("chr"):
        value = value[3:]
    if value.isdigit() and 1 <= int(value) <= 22:
        return int(value)
    return None


def _masks(exclude_bed, target_variants, flank, exclude_long_ld):
    intervals = defaultdict(list)
    if exclude_long_ld:
        for chrom, start, end in LONG_LD_GRCH37:
            intervals[chrom].append((start - 1, end))
    else:
        # Keep the MHC exclusion even in the reduced sensitivity preset.
        intervals[6].append((24999999, 35000000))
    if exclude_bed:
        with _evolution_open(exclude_bed) as handle:
            for line_no, line in enumerate(handle, 1):
                if not line.strip() or line.startswith(("#", "track", "browser")):
                    continue
                values = line.split()
                if len(values) < 3:
                    raise ValueError(f"BED line {line_no}: need chrom,start0,end")
                chrom = _evolution_chrom(values[0])
                try:
                    start, end = int(values[1]), int(values[2])
                except ValueError as exc:
                    raise ValueError(f"BED line {line_no}: integer bounds required") from exc
                if start < 0 or end <= start:
                    raise ValueError(f"BED line {line_no}: invalid interval")
                if chrom is not None:
                    intervals[chrom].append((start, end))
    if target_variants:
        for row in _pca_rows(target_variants, {"chrom", "pos"}):
            if row.get("build", "GRCh37") not in {"GRCh37", "hg19"}:
                raise ValueError("PCA target exclusions must already be harmonized to GRCh37")
            chrom = _evolution_chrom(row["chrom"])
            pos = int(row["pos"])
            if pos <= 0:
                raise ValueError("Target POS must be positive")
            if chrom is not None:
                intervals[chrom].append((max(0, pos - 1 - flank), pos + flank))
    merged = {}
    for chrom, values in intervals.items():
        compact = []
        for start, end in sorted(values):
            if compact and start <= compact[-1][1]:
                compact[-1] = (compact[-1][0], max(end, compact[-1][1]))
            else:
                compact.append((start, end))
        merged[chrom] = compact
    return merged


def _background_ids(hq_path, path, masks):
    counts = Counter()
    starts = {chrom: [pair[0] for pair in values] for chrom, values in masks.items()}
    partial = Path(str(path) + ".partial")
    if partial.exists():
        raise ValueError(f"Interrupted output exists: {partial}")
    try:
        with _evolution_open(hq_path) as source, partial.open("w", encoding="utf-8") as dest:
            reader = csv.DictReader(source, delimiter="\t")
            required = {"#CHROM", "POS", "ID", "REF", "ALT"}
            if not required.issubset(reader.fieldnames or []):
                raise ValueError("HQ table requires exact #CHROM,POS,ID,REF,ALT header")
            for line, row in enumerate(reader, 2):
                if None in row or any(value is None for value in row.values()):
                    raise ValueError(f"Malformed HQ row {line}")
                counts["hq_rows"] += 1
                chrom = _evolution_chrom(row["#CHROM"])
                if chrom is None or len(row["REF"]) != 1 or len(row["ALT"]) != 1 or row["REF"] not in "ACGT" or row["ALT"] not in "ACGT" or row["REF"] == row["ALT"]:
                    counts["non_autosomal_biallelic_snv"] += 1
                    continue
                pos = int(row["POS"])
                if pos <= 0:
                    raise ValueError(f"Invalid HQ POS at row {line}")
                j = bisect_right(starts.get(chrom, []), pos - 1) - 1
                if j >= 0 and pos - 1 < masks[chrom][j][1]:
                    counts["excluded_region"] += 1
                    continue
                vid = row["ID"]
                if not vid or vid == "." or any(c.isspace() for c in vid):
                    raise ValueError(f"Invalid HQ variant ID at row {line}")
                dest.write(vid + "\n")
                counts["background_ids"] += 1
        if counts["background_ids"] == 0:
            raise ValueError("No autosomal biallelic background SNVs remain in HQ table")
        os.replace(partial, path)
    finally:
        if partial.exists():
            partial.unlink()
    return dict(counts)


def pca_mode(n_basis):
    if not isinstance(n_basis,int) or n_basis<1: raise ValueError('Positive PCA basis size required')
    return 'standard' if n_basis<=5000 else 'approx'


def select_pca_basis(metadata, psam_ids, *, basis, max_per_group, time_bin_years,
                     seed, min_age_bp=None, max_age_bp=None):
    """Select a fixed training set; every included person remains projectable.

    The ancient reference requires an explicit unrelated flag. Capping strata
    limits the influence of oversampled epochs, but does not create equal weights.
    Age determines sampling strata only; it never enters the genotype PCA matrix.
    """
    if basis not in {"joint_balanced", "modern_projection", "ancient_balanced"}:
        raise ValueError("Unknown PCA basis")
    for label, value in (("min_age_bp", min_age_bp), ("max_age_bp", max_age_bp)):
        if value is not None and (not math.isfinite(value) or value < 0):
            raise ValueError(f"{label} must be finite and nonnegative")
    if min_age_bp is not None and max_age_bp is not None and min_age_bp > max_age_bp:
        raise ValueError("PCA basis minimum age exceeds maximum age")
    if basis == "modern_projection" and (min_age_bp is not None or max_age_bp is not None):
        raise ValueError("Ancient training-age limits do not apply to modern_projection")
    included = [sid for sid in psam_ids if metadata[sid]["include"] == "1"]
    groups, training, exclusions = defaultdict(list), defaultdict(list), Counter()
    for sid in included:
        row = metadata[sid]
        ancient = row["cohort"] == "ancient"
        epoch = str(int(row["age"] // time_bin_years)) if ancient else "modern"
        group = row["cohort"], row["stratum"], epoch
        groups[group].append(sid)
        if basis == "modern_projection" and ancient:
            exclusions["ancient_projection_only"] += 1
            continue
        if basis == "ancient_balanced" and not ancient:
            exclusions["modern_projection_only"] += 1
            continue
        if basis == "ancient_balanced" and row.get("unrelated") != "1":
            exclusions["ancient_relationship_not_explicitly_unrelated"] += 1
            continue
        if ancient and ((min_age_bp is not None and row["age"] < min_age_bp) or
                        (max_age_bp is not None and row["age"] > max_age_bp)):
            exclusions["ancient_outside_training_age_range"] += 1
            continue
        training[group].append(sid)
    chosen = set()
    for values in training.values():
        if basis == "modern_projection":
            chosen.update(values)
        else:
            chosen.update(sorted(values, key=lambda sid: hashlib.sha256(f"{seed}\t{sid}".encode()).digest())[:max_per_group])
    basis_ids = [sid for sid in included if sid in chosen]
    return included, basis_ids, groups, dict(exclusions)


def prepare_pca_plan(pgen_prefix, metadata_path, hq_variants_path, out_dir, *,
                     basis="joint_balanced", maf=0.05, geno=0.1, n_components=10,
                     max_per_group=200, time_bin_years=2000, seed=2026,
                     exclude_bed=None, target_variants=None, target_flank_bp=1000000,
                     plink2="plink2", exclude_long_ld=True, threads=4, memory_mb=8192,
                     basis_min_age_bp=None, basis_max_age_bp=None):
    """Write keep files, streamed HQ IDs, a reviewable script and a manifest.

    ``joint_balanced`` caps each cohort/region/time group; it does not give every
    group equal sample size. ``modern_projection`` fits on all included modern
    samples. Both project all included people through one identical --score
    call. The script has safe complete-run reuse; interrupted stages are rerun.
    """
    if basis not in {"joint_balanced", "modern_projection", "ancient_balanced"}:
        raise ValueError("basis must be joint_balanced, modern_projection or ancient_balanced")
    if not 0 < maf < .5 or not 0 <= geno < 1:
        raise ValueError("Require 0 < maf < 0.5 and 0 <= geno < 1")
    for label, value, minimum in [("n_components", n_components, 2), ("max_per_group", max_per_group, 1),
                                  ("time_bin_years", time_bin_years, 1), ("threads", threads, 1),
                                  ("memory_mb", memory_mb, 640), ("target_flank_bp", target_flank_bp, 0)]:
        if not isinstance(value, int) or isinstance(value, bool) or value < minimum:
            raise ValueError(f"{label} must be an integer >= {minimum}")
    if not isinstance(seed, int) or not 0 <= seed <= 4294967295:
        raise ValueError("seed must be an unsigned 32-bit integer")
    prefix = Path(pgen_prefix).expanduser().resolve()
    meta_path = Path(metadata_path).expanduser().resolve()
    hq_path = Path(hq_variants_path).expanduser().resolve()
    files = {"metadata": meta_path, "high_quality": hq_path,
             **{ext: Path(str(prefix) + "." + ext) for ext in ("pgen", "pvar", "psam")}}
    if exclude_bed:
        files["exclude_bed"] = Path(exclude_bed).expanduser().resolve()
    if target_variants:
        files["target_variants"] = Path(target_variants).expanduser().resolve()
    inputs = {key: _fingerprint(path) for key, path in files.items()}
    metadata = _pca_metadata(meta_path)
    psam_ids = _trajectory_module().read_psam_ids(files["psam"])
    if set(psam_ids) != set(metadata):
        raise ValueError("PCA metadata must cover exactly the complete original PSAM sample IDs")
    included, basis_ids, groups, basis_exclusions = select_pca_basis(metadata, psam_ids,
        basis=basis, max_per_group=max_per_group, time_bin_years=time_bin_years, seed=seed,
        min_age_bp=basis_min_age_bp, max_age_bp=basis_max_age_bp)
    if len(basis_ids) < max(50, n_components + 1):
        raise ValueError("Fixed PCA basis needs at least max(50, n_components+1) included samples; do not fit tiny time-specific PCAs")
    mode = pca_mode(len(basis_ids))
    settings = dict(pca_mode=mode, basis=basis, maf=maf, geno=geno, n_components=n_components,
                    max_per_group=max_per_group, time_bin_years=time_bin_years, seed=seed,
                    target_flank_bp=target_flank_bp, exclude_long_ld=exclude_long_ld,
                    plink2=str(plink2), threads=threads, memory_mb=memory_mb,
                    ld_window_kb=200, ld_step=1, ld_r2=0.1)
    # Keep existing joint/modern plan signatures unchanged when no new limits
    # are requested. A new ancient basis or explicit age contract is new work.
    if basis_min_age_bp is not None or basis_max_age_bp is not None:
        settings.update(basis_min_age_bp=basis_min_age_bp, basis_max_age_bp=basis_max_age_bp)
    if basis == "ancient_balanced":
        settings["ancient_relationship_rule"] = "unrelated=1 explicitly required for training"
    signature = hashlib.sha256(json.dumps({"schema": 2, "inputs": inputs, "settings": settings}, sort_keys=True).encode()).hexdigest()
    destination = Path(out_dir).expanduser().resolve()
    destination.mkdir(parents=True, exist_ok=True)
    manifest_path = destination / "pca_plan.json"
    if manifest_path.exists():
        prior = json.loads(manifest_path.read_text())
        if prior.get("signature") != signature:
            raise ValueError("PCA plan inputs/settings changed; use a new output directory to preserve provenance")
        for key, digest in prior["plan_file_sha256"].items():
            path = Path(key)
            if not path.is_file() or _evolution_sha(path) != digest:
                raise ValueError(f"Existing PCA plan file is missing or modified: {path}")
        return prior
    # Existing work with no manifest cannot safely be assumed to belong here.
    if any(destination.iterdir()):
        raise ValueError("New PCA plan requires an empty output directory")
    masks = _masks(exclude_bed, target_variants, target_flank_bp, exclude_long_ld)
    all_keep, basis_keep = destination / "all.keep", destination / "basis.keep"
    marker_path, mask_path = destination / "background.ids", destination / "excluded_regions.GRCh37.bed"
    _atomic_text(all_keep, "#IID\n" + "\n".join(included) + "\n")
    _atomic_text(basis_keep, "#IID\n" + "\n".join(basis_ids) + "\n")
    _atomic_text(mask_path, "# GRCh37; 0-based half-open; historical LD/MHC plus target/custom masks\n" +
                 "".join(f"{chrom}\t{start}\t{end}\n" for chrom in sorted(masks) for start, end in masks[chrom]))
    counts = _background_ids(hq_path, marker_path, masks)
    if counts["background_ids"] < n_components + 1:
        raise ValueError("Fewer HQ background SNVs than required PCA dimensions")
    common, ld, pruned, ref, projected = [destination / x for x in ("basis_common", "ld", "all_pruned", "ref", "all")]
    shared = [str(plink2), "--threads", str(threads), "--memory", str(memory_mb)]
    commands = [
        shared + ["--pfile", str(prefix), "--keep", str(basis_keep), "--extract", str(marker_path),
                  "--chr", "1-22", "--snps-only", "just-acgt", "--min-alleles", "2", "--max-alleles", "2",
                  "--maf", str(maf), "--geno", str(geno), "--make-pgen", "--out", str(common)],
        shared + ["--pfile", str(common), "--indep-pairwise", "200kb", "1", "0.1", "--out", str(ld)],
        shared + ["--pfile", str(prefix), "--keep", str(all_keep), "--extract", str(ld) + ".prune.in",
                  "--make-pgen", "--out", str(pruned)],
        shared + ["--pfile", str(pruned), "--keep", str(basis_keep), "--freq", "counts", "--pca",
                  str(n_components), *(["approx"] if mode == "approx" else []), "allele-wts", "vcols=chrom,ref,alt", "--seed", str(seed), "--out", str(ref)],
        shared + ["--pfile", str(pruned), "--read-freq", str(ref) + ".acount", "--score",
                  str(ref) + ".eigenvec.allele", "2", "5", "header-read", "variance-standardize",
                  "cols=+scoresums", "--score-col-nums", f"6-{5+n_components}", "--out", str(projected)],
    ]
    script_path = destination / "commands.sh"
    projected_path = Path(str(projected) + ".sscore")
    checksums = destination / "completed.sha256"
    marker = destination / "PCA_COMPLETE"
    # Check original inputs before any reuse/execution, without rehashing huge PGEN files.
    verifier = """import hashlib,json,pathlib
p=json.loads(pathlib.Path(__import__('sys').argv[1]).read_text())
for spec in p['inputs'].values():
 f=pathlib.Path(spec['path']); s=f.stat(); h=hashlib.sha256()
 with f.open('rb') as r:
  h.update(r.read(65536))
  if s.st_size>65536:
   r.seek(max(65536,s.st_size-65536)); h.update(r.read(65536))
 if spec.get('sha256'):
  full=hashlib.sha256()
  with f.open('rb') as source:
   for block in iter(lambda:source.read(1048576),b''): full.update(block)
  if full.hexdigest()!=spec['sha256']: raise SystemExit('PCA input full SHA256 changed: '+str(f))
 if s.st_size!=spec['size'] or s.st_mtime_ns!=spec['mtime_ns'] or h.hexdigest()!=spec['edge_sha256']:
  raise SystemExit('PCA input changed; generate a new plan: '+str(f))
for name,want in p['plan_file_sha256'].items():
 h=hashlib.sha256()
 with open(name,'rb') as r:
  for b in iter(lambda:r.read(1048576),b''): h.update(b)
 if h.hexdigest()!=want: raise SystemExit('PCA plan modified: '+name)
"""
    quote = shlex.quote
    script = ["#!/usr/bin/env bash", "set -euo pipefail", "# Reviewable offline PCA workflow; no downloads.",
              "# Incomplete runs recompute stages; only a checksum-verified full result is reused.",
              f"python3 - {quote(str(manifest_path))} <<'ALLELE3_VERIFY_PCA_INPUTS'", verifier.rstrip(),
              "ALLELE3_VERIFY_PCA_INPUTS", f"if [[ -f {quote(str(marker))} ]]; then",
              f"  [[ $(cat {quote(str(marker))}) == {quote(signature)} ]] || exit 1",
              f"  sha256sum --check --status {quote(str(checksums))}", "  exit 0", "fi",
              shlex.join([str(plink2), "--version"])]
    for index, command in enumerate(commands, 1):
        script += [f"# Stage {index}: see pca_plan.json for scientific assumptions.", shlex.join(command)]
    completed = [projected_path, Path(str(ref) + ".acount"), Path(str(ref) + ".eigenvec.allele"), Path(str(ref) + ".eigenval")]
    script += ["sha256sum " + " ".join(quote(str(p)) for p in completed) + " > " + quote(str(checksums)),
               f"printf '%s\\n' {quote(signature)} > {quote(str(marker) + '.partial')}",
               f"mv {quote(str(marker) + '.partial')} {quote(str(marker))}"]
    _atomic_text(script_path, "\n".join(script) + "\n")
    script_path.chmod(0o755)
    basis_set = set(basis_ids)
    group_counts = [{"cohort": key[0], "stratum": key[1], "time_group": key[2],
                     "eligible": len(values), "basis": sum(sid in basis_set for sid in values)}
                    for key, values in sorted(groups.items())]
    manifest = dict(schema_version=1, signature=signature, build="GRCh37", inputs=inputs,
                    settings=settings, script_path=str(script_path), projected_path=str(projected_path),
                    manifest_path=str(manifest_path), commands=commands,
                    plan_file_sha256={str(p): _evolution_sha(p) for p in (all_keep, basis_keep, marker_path, mask_path, script_path)},
                    counts={**counts, "samples_eligible": len(included), "samples_basis": len(basis_ids)},
                    basis_exclusions=basis_exclusions,
                    group_counts=group_counts, long_ld_source=LONG_LD_SOURCE,
                    coordinate_kind="all samples use PC*_AVG from the same --score transform; never mix with raw eigenvec units",
                    limitations=[
                        "Age is only an overlay; no separate per-epoch PCA is fitted.",
                        "joint_balanced caps cohort/region/time groups, rather than guaranteeing equal group weights.",
                        "ancient_balanced fits only explicitly unrelated ancient people within configured ages; modern and excluded training groups remain projections.",
                        "The public Reich modern component is 503 EUR individuals, not the full 26-population 1000 Genomes reference. Reference geography must be reported.",
                        "Background common SNVs are not proven neutral; use target/MHC/long-LD exclusions as sensitivity analyses.",
                        "Projection can shrink out-of-basis samples; modern-only PCs may miss ancient variation.",
                        "MAF and genotype-missingness filters are evaluated on the fixed basis; LD pruning uses hardcall correlations.",
                        "PGEN imputation/batch differences may affect PCs. Inspect data-source, coverage and region overlays.",
                        "basis_common and all_pruned can occupy several GiB; incomplete stages are rerun, completed outputs are checksum-verified.",
                        "Huge inputs use size/mtime/edge hashes, not full-file cryptographic validation. Do not mutate original files during execution.",
                        "No target exclusion file supplied; provide GRCh37 selected variants before interpreting target overlays." if target_variants is None else "Targets and configured flanks were excluded from the background marker list.",
                    ])
    _atomic_text(manifest_path, json.dumps(manifest, ensure_ascii=False, indent=2, allow_nan=False) + "\n")
    return manifest


def _coordinates(path):
    values, names, columns = {}, None, None
    with _evolution_open(path) as handle:
        for line_no, line in enumerate(handle, 1):
            if not line.strip() or line.startswith("##"):
                continue
            cells = line.split()
            if names is None:
                names = [cell.lstrip("#") for cell in cells]
                if len(names) != len(set(names)) or "IID" not in names:
                    raise ValueError("PCA coordinates need a header with unique IID")
                present = [(a, b, kind) for a, b, kind in [
                    ("PC1_AVG", "PC2_AVG", "projected_average"),
                    ("PC1_SUM", "PC2_SUM", "projected_sum"),
                    ("PC1", "PC2", "eigenvector")]
                    if a in names and b in names]
                if not present:
                    raise ValueError("Coordinates require matching PC1/PC2, PC1_AVG/PC2_AVG, or PC1_SUM/PC2_SUM columns")
                columns = present[0]
                if any(kind == "eigenvector" for _, _, kind in present) and len(present) > 1:
                    raise ValueError("Do not mix eigenvector and projected score units in one coordinate table")
                continue
            if len(cells) != len(names):
                raise ValueError(f"PCA coordinate column count mismatch at line {line_no}")
            row = dict(zip(names, cells))
            sid = row["IID"]
            if sid in values or not sid or sid == ".":
                raise ValueError(f"Duplicate or invalid coordinate IID: {sid}")
            values[sid] = [_pca_number(row[columns[0]], sid + " PC1"), _pca_number(row[columns[1]], sid + " PC2")]
    if not values:
        raise ValueError("No PCA coordinates")
    return values, columns[2]


def render_pca_time(eigenvec_path, metadata_path, out_html, *, basis_label="fixed shared PCA",
                    pgen_path=None, psam_path=None, selected_variants_path=None,
                    max_variants=12, simulated=False, title=None, count_kind="hardcall", candidate_path=None,
                    layer_years=1000, time_view="layered", color_by="region"):
    """Render one self-contained HTML.  All included sample IDs must have PCs.

    Optional PGEN overlays read only 1..12 preselected, verified original PVAR
    rows.  ALT dosage 0/1/2 never denotes three alleles. Unrepresented panel
    alleles and missing genotypes stay unavailable; no artificial third state.
    """
    if not isinstance(max_variants, int) or not 1 <= max_variants <= 12:
        raise ValueError("max_variants must be an integer from 1 to 12")
    metadata = _pca_metadata(metadata_path)
    coords, coordinate_kind = _coordinates(eigenvec_path)
    extra = set(coords) - set(metadata)
    included = [sid for sid, row in metadata.items() if row["include"] == "1"]
    missing = set(included) - set(coords)
    if extra or missing:
        raise ValueError(f"Coordinate/metadata sample mismatch: {len(extra)} unknown, {len(missing)} included samples missing")
    requested = [pgen_path is not None, psam_path is not None, selected_variants_path is not None]
    if any(requested) and not all(requested):
        raise ValueError("PGEN overlays require pgen_path, psam_path and selected_variants_path together")
    overlays = []
    if all(requested):
        target_count = 0
        for _ in _pca_rows(selected_variants_path, {"variant_id", "pvar_index", "chrom", "pos", "ref", "alt", "status"}):
            target_count += 1
            if target_count > max_variants:
                raise ValueError(f"Overlay file exceeds {max_variants} selected variants; supply a bounded target list")
        if target_count == 0:
            raise ValueError("Overlay target file is empty")
        sample_ids, iterator = _trajectory_module().iter_pgen_variants(pgen_path, psam_path, selected_variants_path, count_kind=count_kind)
        if set(sample_ids) != set(metadata):
            iterator.close()
            raise ValueError("Overlay PSAM must match the complete metadata sample set")
        indexes = {sid: i for i, sid in enumerate(sample_ids)}
        try:
            for row in iterator:
                if len(row["ref"]) != 1 or len(row["alt"]) != 1 or row["ref"] not in "ACGT" or row["alt"] not in "ACGT":
                    raise ValueError("PCA allele overlays currently accept biallelic SNV anchors only")
                dosage = None if row["dosages"] is None else [row["dosages"][indexes[sid]] for sid in included]
                if dosage is not None and any(x is not None and (not math.isfinite(x) or not 0 <= x <= 2) for x in dosage):
                    raise ValueError("ALT dosage must be finite in [0,2], or missing")
                overlays.append({k: row[k] for k in ("variant_id", "chrom", "pos", "ref", "alt", "status", "count_kind", "dosage_evidence")})
                overlays[-1]["values"] = dosage
        finally:
            iterator.close()
    points = []
    for sid in included:
        row = metadata[sid]
        points.append(dict(sample_id=sid, pc1=coords[sid][0], pc2=coords[sid][1], age_bp=row["age"],
                           cohort=row["cohort"], region=row.get("region", row["stratum"]),
                           source=row.get("data_source", row["cohort"]), data_types=row.get("data_types", "."),
                           coverage=row.get("target_coverage", "."), iqs=row.get("iqs", ".")))
    try:
        from plotly.offline import get_plotlyjs
        import plotly
    except ImportError as exc:
        raise RuntimeError("Self-contained PCA HTML needs plotly in the active environment") from exc
    from report import pca_display, static_pca
    display = pca_display(points, layer_years, time_view, color_by)
    payload = dict(points=points, variants=overlays, simulated=bool(simulated), display=display,
                   coordinate_kind=coordinate_kind, basis_label=str(basis_label))
    data_json = json.dumps(payload, ensure_ascii=False, separators=(",", ":"), allow_nan=False).replace("<", "\\u003c").replace("\u2028", "\\u2028").replace("\u2029", "\\u2029")
    title = title or "Ancient DNA · PCA 时间分层"
    document = _HTML.replace("__TITLE__", html.escape(title)).replace("__PLOTLY__", get_plotlyjs()).replace("__PAYLOAD__", data_json)
    destination = Path(out_html).expanduser().resolve()
    input_paths = {Path(p).expanduser().resolve() for p in
                   (eigenvec_path, metadata_path, pgen_path, psam_path, selected_variants_path)
                   if p is not None}
    if destination in input_paths or Path(str(destination) + ".json") in input_paths:
        raise ValueError("HTML output must not overwrite inputs")
    destination.parent.mkdir(parents=True, exist_ok=True)
    _atomic_text(destination, document)
    manifest = dict(html_path=str(destination), samples=len(points), variants=len(overlays),
                    coordinate_kind=coordinate_kind, basis_label=str(basis_label), simulated=bool(simulated),
                    display=display,
                    plotly_version=plotly.__version__, self_contained=True,
                    excluded_metadata_samples=len(metadata) - len(points),
                    third_allele_status="not_inferred_from_0_1_2; inspect_each_target_availability",
                    inputs={"coordinates": _fingerprint(eigenvec_path), "metadata": _fingerprint(metadata_path)},
                    warnings=["ALT copy number 0/1/2 is not three different alleles.",
                              "Coordinates use one fixed basis; positive BP points upward; REF is not assumed ancestral.",
                              "gnomAD aggregate counts are not individual coordinates."])
    manifest["static_figures"] = static_pca(payload, destination.parent, candidate_path)
    manifest_path = Path(str(destination) + ".json")
    _atomic_text(manifest_path, json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    manifest["manifest_path"] = str(manifest_path)
    return manifest


_HTML = r'''<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>__TITLE__</title><style>
:root{font-family:system-ui,-apple-system,"Segoe UI",sans-serif;color:#152e45;background:#edf3f7;font-size:15px}body{margin:0}header{padding:24px 32px 16px;background:#102d45;color:#f8fbff}h1{font-size:26px;font-weight:650;margin:0 0 10px}header p{margin:7px 0;color:#cadce8;line-height:1.6}.badge{background:#f5b95f;color:#3b2506;padding:8px 14px;border-radius:6px;font-weight:800;margin:6px 0}.wrap{display:grid;grid-template-columns:minmax(450px,1fr) 310px;gap:16px;padding:18px 24px}.panel{background:white;border:1px solid #d5e0e9;border-radius:10px;box-shadow:0 3px 12px #102d4507}.chart{min-height:650px;height:76vh}aside{padding:20px}h2{font-size:17px;margin:0 0 15px}.field{margin:16px 0}.field label{display:block;font-weight:650;margin-bottom:7px}select,input[type=number]{box-sizing:border-box;padding:8px;border:1px solid #bdccd8;border-radius:6px;background:#f8fbfd;width:100%;color:#17344d}input[type=range]{width:100%;accent-color:#2478a2}.checks label{display:inline;font-weight:400;margin-right:14px}.note{font-size:13px;line-height:1.65;color:#496174;border-top:1px solid #e1e9ee;margin-top:18px;padding-top:14px}.status{padding:12px;background:#edf5f8;border-radius:6px;line-height:1.6}.warn{background:#fff3d9;color:#705015;padding:10px;border-radius:6px;line-height:1.5;font-size:13px}.numbers{display:grid;grid-template-columns:1fr 1fr;gap:12px}.sub{font-size:12px;color:#59758b}button{background:#146c94;color:white;border:0;padding:8px 12px;border-radius:6px;cursor:pointer}footer{padding:0 28px 25px;color:#456178;font-size:13px;line-height:1.6}@media(max-width:900px){.wrap{grid-template-columns:1fr;padding:12px}.chart{height:65vh;min-height:470px}header{padding:20px}aside{padding:18px}}
</style><script>__PLOTLY__</script></head><body>
<header><div id="simulated" class="badge" hidden>模拟示意数据 · 不是真实古代 DNA 结果</div><h1>__TITLE__</h1><p>同一组 PC1、PC2，逐层查看古代人群。越高越古老；颜色可切换区域、年代等信息，菱形标记现代样本。</p><p id="basis"></p></header>
<main class="wrap"><section class="panel"><div id="plotError" class="warn" hidden></div><div id="plot" class="chart"></div></section><aside class="panel">
<h2>查看同一个坐标系</h2><div class="field"><label for="view">视图</label><select id="view"><option value="layered">3D 时间分层</option><option value="continuous">3D 连续年代</option><option value="2d">2D 时间筛选截面</option></select></div>
<div class="field" id="layerControls"><label for="layerYears">每层年数</label><select id="layerYears"><option value="500">500 年</option><option value="1000">1000 年</option><option value="2000">2000 年</option><option value="5000">5000 年</option></select><label style="margin-top:10px;font-weight:400"><input id="planes" type="checkbox" checked> 显示半透明层面与人数</label><span class="sub">古代样本置于所属区间的中点；现代单独放在 0 BP。悬停与时间筛选使用真实年代。</span></div>
<div class="field checks"><label><input id="ancient" type="checkbox" checked> 古代</label><label><input id="modern" type="checkbox" checked> 现代</label></div>
<div class="field"><label for="color">着色</label><select id="color"><option value="region">大区域</option><option value="layer">时间层</option><option value="age">真实年代 BP</option><option value="cohort">古代 / 现代</option><option value="source">数据来源</option></select></div>
<div id="regionLegend" class="field" style="max-height:160px;overflow:auto"></div>
<div class="field"><label for="region">区域</label><select id="region"><option value="all">全部区域</option></select></div>
<div class="field"><label for="source">数据来源</label><select id="source"><option value="all">全部来源</option></select></div>
<div class="field"><label>时间范围：<span id="rangeLabel"></span> BP</label><input id="young" type="range" min="0" step="100" value="0"><input id="old" type="range" min="0" step="100"><span class="sub">上滑条：年轻边界；下滑条：古老边界。仅改变可见点，坐标轴保持固定。</span></div>
<div class="field"><button id="reset">重置筛选</button></div><div id="counts" class="status"></div><div id="frequency" class="status" hidden></div><p id="alleleNote" class="warn"></p>
<div class="note" id="alleleCopyNote">蓝 → 橙的 0/1/2 表示同一个 ALT 的拷贝数：REF/REF、REF/ALT、ALT/ALT。它们不是三个碱基。灰色表示缺失或当前目标未测定。第三碱基需要其独立目标记录和测定状态，不能从所选 ALT 的 0/1/2 推断。</div>
</aside></main><footer>REF 不等于祖先等位基因；最古老样本也可以携带 ALT。图中分离或频率差异可能来自迁移、抽样、imputation 或批次；不能仅凭图判断选择或 pleiotropy。gnomAD 汇总频率不被转换成虚构个体点。此 HTML 含全部绘图依赖，可离线打开。</footer>
<script id="allele3-data" type="application/json">__PAYLOAD__</script><script>
'use strict';
const DATA=JSON.parse(document.getElementById('allele3-data').textContent), pts=DATA.points;
const display=DATA.display;
const el=id=>document.getElementById(id), esc=v=>String(v).replace(/[&<>"']/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]));
el('simulated').hidden=!DATA.simulated;
el('alleleNote').hidden=el('alleleCopyNote').hidden=DATA.variants.length===0;
el('basis').textContent=DATA.basis_label+' · '+({projected_average:'全体统一投影分数（AVG）',projected_sum:'全体统一投影分数（SUM）',eigenvector:'共享 eigenvector 坐标'}[DATA.coordinate_kind]);
const maxBP=pts.reduce((v,p)=>Math.max(v,p.age_bp),100), upper=Math.ceil(maxBP/100)*100;
for(const id of ['young','old'])el(id).max=upper;el('old').value=upper;
const regions=[...new Set(pts.map(p=>p.region))].sort(),sources=[...new Set(pts.map(p=>p.source))].sort();
for(const [id,values] of [['region',regions],['source',sources]])for(const v of values){let o=document.createElement('option');o.value=v;o.textContent=v;el(id).appendChild(o)}
DATA.variants.forEach((v,i)=>{let o=document.createElement('option');o.value='v'+i;o.textContent=v.variant_id+' · '+v.ref+'→'+v.alt;el('color').appendChild(o)});
if(![...el('layerYears').options].some(o=>Number(o.value)===display.layer_years)){let o=document.createElement('option');o.value=display.layer_years;o.textContent=display.layer_years+' 年';el('layerYears').appendChild(o)}
el('layerYears').value=display.layer_years;el('view').value=display.time_view;el('color').value=display.color_by;
const regionColors=display.region_colors,sourceColors=display.source_colors;
const layerIndex=(p,width)=>p.cohort==='modern'?-1:Math.floor(p.age_bp/width);
const layerZ=(p,width)=>p.cohort==='modern'?0:(layerIndex(p,width)+.5)*width;
const layerLabel=(k,width)=>k<0?'现代 · 0 BP':(k*width/1000).toLocaleString()+'–'+((k+1)*width/1000).toLocaleString()+' ka BP';
function timeColor(t){const stops=[[68,1,84],[59,82,139],[33,145,140],[94,201,98],[253,231,37]],v=Math.max(0,Math.min(1,t))*4,i=Math.min(3,Math.floor(v)),f=v-i;return 'rgb('+stops[i].map((c,j)=>Math.round(c*(1-f)+stops[i+1][j]*f)).join(',')+')'}
function extent(values){let a=Infinity,b=-Infinity;for(const v of values){if(Number.isFinite(v)){a=Math.min(a,v);b=Math.max(b,v)}}if(!Number.isFinite(a))return[-1,1];let m=Math.max((b-a)*.06,Math.abs(a)*.01,.00001);return[a-m,b+m]}
const xRange=extent(pts.map(p=>p.pc1)),yRange=extent(pts.map(p=>p.pc2));
const coordLabel=DATA.coordinate_kind==='eigenvector'?'':' 投影分数';
const layout={margin:{l:15,r:20,b:20,t:35},paper_bgcolor:'#ffffff',font:{family:'system-ui',color:'#29475d'},legend:{orientation:'h',x:.04,y:1.02},uirevision:'fixed-shared-pca',scene:{xaxis:{title:{text:'PC1'+coordLabel},range:xRange,autorange:false,showbackground:false,gridcolor:'#e8eef1'},yaxis:{title:{text:'PC2'+coordLabel},range:yRange,autorange:false,showbackground:false,gridcolor:'#e8eef1'},zaxis:{title:{text:'年代 BP ↑ 越高越古老'},range:[-upper*.025,upper*1.03],autorange:false,showbackground:false,gridcolor:'#e8eef1'},aspectratio:{x:1.2,y:1,z:1.7},camera:{projection:{type:'orthographic'},eye:{x:1.55,y:-1.85,z:.85}}}};
const plotConfig={responsive:true,displaylogo:false,modeBarButtonsToRemove:['toImage']};
let renderedView=null,pendingDraw=null,drawing=false;
async function drawPlot(traces,displayLayout,view){
 pendingDraw={traces,displayLayout,view};if(drawing)return;drawing=true;
 try{
  while(pendingDraw){
   const job=pendingDraw;pendingDraw=null;
   // Plotly.react can retain incompatible subplot internals across 3D/2D.
   // Recreate the graph when its subplot family changes; retain fixed ranges.
   if(renderedView!==job.view){if(renderedView!==null)Plotly.purge('plot');await Plotly.newPlot('plot',job.traces,job.displayLayout,plotConfig)}
   else await Plotly.react('plot',job.traces,job.displayLayout,plotConfig);
   renderedView=job.view;el('plotError').hidden=true;
  }
 }catch(error){renderedView=null;el('plotError').hidden=false;el('plotError').textContent='绘图失败：'+String(error.message||error);console.error('PCA plot failed',error)}
 finally{drawing=false}
}
// Fixed-seed stratified display sampling only. Full points remain in DATA.
function displaySubset(rows,limit=20000){
 if(rows.length<=limit)return rows;
 const hash=s=>{let h=20261028;for(let i=0;i<s.length;i++)h=Math.imul(h^s.charCodeAt(i),16777619);return h>>>0};
 const groups=new Map();for(const row of rows){const p=row[0],k=p.cohort+'|'+p.region+'|'+Math.floor(p.age_bp/1000);if(!groups.has(k))groups.set(k,[]);groups.get(k).push(row)}
 let result=[];for(const group of groups.values()){group.sort((a,b)=>hash(a[0].sample_id)-hash(b[0].sample_id));const n=Math.max(1,Math.floor(limit*group.length/rows.length));result=result.concat(group.slice(0,n))}
 if(result.length>limit)result.sort((a,b)=>hash(a[0].sample_id)-hash(b[0].sample_id));return result.slice(0,limit);
}
function update(){
 let lo=Number(el('young').value),hi=Number(el('old').value);if(lo>hi){[lo,hi]=[hi,lo]}
 el('rangeLabel').textContent=lo.toLocaleString()+' – '+hi.toLocaleString();
 const mode=el('color').value,variant=mode.startsWith('v')?DATA.variants[Number(mode.slice(1))]:null,is2d=el('view').value==='2d';
 const layered=el('view').value==='layered',width=Number(el('layerYears').value);
 const layerTop=pts.reduce((v,p)=>p.cohort==='ancient'?Math.max(v,(layerIndex(p,width)+1)*width):v,width);
 const zTop=layered?Math.max(upper,layerTop):upper;
 el('layerControls').hidden=!layered;
 const statisticalSelected=pts.map((p,i)=>[p,i]).filter(([p])=>p.age_bp>=lo&&p.age_bp<=hi&&el(p.cohort).checked&&(el('region').value==='all'||p.region===el('region').value)&&(el('source').value==='all'||p.source===el('source').value));
 const selected=displaySubset(statisticalSelected);
 const traces=[],annotations=[];let present=0,hasScale=false;
 const layerCounts=new Map();for(const [p] of statisticalSelected){const k=layerIndex(p,width);layerCounts.set(k,(layerCounts.get(k)||0)+1)}
 if(layered&&el('planes').checked){
  for(const [k,n] of [...layerCounts].sort((a,b)=>a[0]-b[0])){
   const z=k<0?0:(k+.5)*width,c=mode==='layer'?timeColor(z/zTop):'#6A9FB5';
   traces.push({type:'mesh3d',x:[xRange[0],xRange[1],xRange[1],xRange[0]],y:[yRange[0],yRange[0],yRange[1],yRange[1]],z:[z,z,z,z],i:[0,0],j:[1,2],k:[2,3],color:c,opacity:.055,hoverinfo:'skip',showlegend:false,flatshading:true});
   traces.push({type:'scatter3d',mode:'lines',x:[xRange[0],xRange[1],xRange[1],xRange[0],xRange[0]],y:[yRange[0],yRange[0],yRange[1],yRange[1],yRange[0]],z:[z,z,z,z,z],line:{color:c,width:1},opacity:.32,hoverinfo:'skip',showlegend:false});
   if(layerCounts.size<=25)annotations.push({x:xRange[0],y:yRange[0],z,text:layerLabel(k,width)+' · n='+n.toLocaleString(),showarrow:false,xanchor:'right',font:{size:10,color:'#4c6576'},bgcolor:'rgba(255,255,255,.78)',borderpad:2});
  }
 }
 for(const cohort of ['ancient','modern'])for(const unavailable of [false,true]){
  const subset=selected.filter(([p,i])=>p.cohort===cohort&&Boolean(variant&&(variant.values===null||variant.values[i]===null))===unavailable);if(!subset.length)continue;
  let numeric=Boolean(variant&&!unavailable)||mode==='age'||mode==='layer';const colors=subset.map(([p,i])=>unavailable?'#a5b3be':variant?variant.values[i]:mode==='age'?p.age_bp:mode==='layer'?layerZ(p,width):mode==='cohort'?(cohort==='ancient'?'#0072B2':'#D55E00'):mode==='source'?sourceColors[p.source]:regionColors[p.region]);
  if(variant&&!unavailable)present+=subset.length;
  const hover=subset.map(([p,i])=>{let s=esc(p.sample_id)+'<br>'+esc(p.cohort)+' · '+esc(p.region)+'<br>真实年代 '+p.age_bp.toLocaleString()+' BP'+(layered?'<br>显示层 '+layerLabel(layerIndex(p,width),width):'')+'<br>来源 '+esc(p.source)+'<br>类型 '+esc(p.data_types)+'<br>capture targets 覆盖 '+esc(p.coverage)+' · IQS '+esc(p.iqs);if(variant){const d=variant.values===null?null:variant.values[i];s+='<br>'+esc(variant.variant_id)+' '+esc(variant.ref)+'→'+esc(variant.alt)+'<br>ALT '+(variant.count_kind==='hardcall'?'拷贝数':'dosage')+': '+(d===null?'缺失 / 未测定':d);if(d!==null&&variant.count_kind==='hardcall')s+=' ('+[variant.ref+'/'+variant.ref,variant.ref+'/'+variant.alt,variant.alt+'/'+variant.alt][d]+')';s+='<br>0/1/2 不编码第三碱基；请查看独立目标测定状态'}return s});
  let marker={size:cohort==='modern'?(is2d?6:4):(is2d?4:2.8),opacity:.82,symbol:cohort==='modern'?'diamond':'circle',color:colors};if(numeric){marker.colorscale=variant?[[0,'#2452a2'],[.5,'#e7e7df'],[1,'#df7b25']]:'Viridis';marker.cmin=0;marker.cmax=variant?2:mode==='layer'?Math.max(upper,layerTop):upper;marker.showscale=!hasScale;hasScale=true;marker.colorbar={title:{text:variant?'ALT 拷贝数 / dosage':mode==='layer'?'时间层中点 BP':'真实年代 BP'},x:1.01,len:.6,tickvals:variant?[0,1,2]:undefined}}
  traces.push({type:is2d?'scattergl':'scatter3d',mode:'markers',name:(cohort==='ancient'?'古代':'现代')+(unavailable?' · 未知':''),x:subset.map(([p])=>p.pc1),y:subset.map(([p])=>p.pc2),...(is2d?{}:{z:subset.map(([p])=>layered?layerZ(p,width):p.age_bp)}),text:hover,hoverinfo:'text',marker});
 }
 const displayLayout=is2d?{...layout,scene:undefined,margin:{l:65,r:50,b:60,t:35},xaxis:{title:{text:'PC1'+coordLabel},range:xRange,autorange:false,gridcolor:'#e3ebf0'},yaxis:{title:{text:'PC2'+coordLabel},range:yRange,autorange:false,gridcolor:'#e3ebf0'},uirevision:'fixed-shared-pca-2d'}:{...layout,scene:{...layout.scene,annotations,zaxis:{...layout.scene.zaxis,title:{text:layered?'时间层 BP ↑ 越高越古老':'真实年代 BP ↑ 越高越古老'},range:[-.03*zTop,1.04*zTop]}}};
 // Plotly may mutate layout objects. Keep the immutable axis template separate.
 drawPlot(traces,JSON.parse(JSON.stringify(displayLayout)),is2d?'2d':'3d');
 const anc=statisticalSelected.filter(([p])=>p.cohort==='ancient').length,mod=statisticalSelected.length-anc;
 el('counts').textContent='筛选统计 N='+statisticalSelected.length.toLocaleString()+'：古代 '+anc.toLocaleString()+'，现代 '+mod.toLocaleString()+'；图上显示 '+selected.length.toLocaleString()+' / 全部 '+pts.length.toLocaleString()+(el('modern').checked&&mod===0?'；当前筛选无现代参照':'')+(layered?'；'+layerCounts.size+' 个时间层':'')+(variant?'；显示样本可用 ALT 值 '+present.toLocaleString():'');
 const legendColors=mode==='region'?regionColors:mode==='source'?sourceColors:null;el('regionLegend').hidden=!legendColors;el('regionLegend').innerHTML=legendColors?Object.entries(legendColors).map(([r,c])=>'<span style="display:inline-block;margin:4px 12px 4px 0"><span style="display:inline-block;width:10px;height:10px;border-radius:50%;background:'+c+';margin-right:5px"></span>'+esc(r)+'</span>').join(''):'';
 el('frequency').hidden=!variant;if(variant){const vals=variant.values===null?[]:selected.map(([,i])=>variant.values[i]).filter(v=>v!==null);const n=vals.length,ac=vals.reduce((a,b)=>a+b,0),an=2*n;el('frequency').textContent='当前筛选个体中的 '+variant.alt+'：'+(n?'AF = '+(ac/an).toFixed(4)+'；'+(variant.count_kind==='dosage'?'expected AC':'AC')+' = '+Number(ac.toFixed(3))+'，AN = '+an:'AF / AC / AN 不可用')+'。这是所显示样本的描述性频率，不是总体频率估计或选择检验。'}
 el('alleleNote').textContent=variant?variant.variant_id+'：'+variant.ref+' / '+variant.alt+'；'+(variant.status==='assayed_pass'?'已通过面板位点筛选，灰色点为个体缺失。':'该目标未通过当前面板可测定条件：'+variant.status+'，不填 0。'):'选择目标位点后可叠加 ALT 拷贝数；横纵 PCA 坐标保持不变。第三碱基不能从 0/1/2 编码推断。';
}
for(const id of ['view','color','region','source','ancient','modern','young','old','layerYears','planes'])el(id).addEventListener('input',update);
el('reset').onclick=()=>{el('ancient').checked=el('modern').checked=true;el('region').value=el('source').value='all';el('young').value=0;el('old').value=upper;update()};
update();
</script></body></html>'''


# ========================================================================
# Evolve cli
# ========================================================================

# Ancient-DNA, fixed PCA/time and optional sparse Atlas workflow.


import argparse
import fcntl
import json
import os
from pathlib import Path
import subprocess
import sys
import time

from allele_common import VERSION, dump_json, fingerprint

PREFIX = "AncientDNA_plus_EUR1KG_AllVariants_NoQC_GRCh37_autosomes"
DEFAULT_ANCIENT = "/mnt/f/gen/ancient"


def parser():
    p = argparse.ArgumentParser(description="allele3 evolution: local ancient trajectories, PCA × time, bounded optional Atlas",
                                formatter_class=argparse.ArgumentDefaultsHelpFormatter)
    p.add_argument("command", choices=["evolve", "prepare-ancient", "join-selection", "select-candidates",
                   "panel-index", "trajectories", "atlas-plan", "atlas-query", "atlas-summarize",
                   "compare-functions", "pca-plan", "pca-time", "report-evolution", "audit-ancient", "audit-counts"])
    p.add_argument("--out", "--dir-out", default="/mnt/d/analysis/allele3/evolution")
    p.add_argument("--ancient-dir", default=DEFAULT_ANCIENT)
    p.add_argument("--modern-out", default="/mnt/d/analysis/allele3/gnomad_local")
    p.add_argument("--pgen-prefix", help="Default: ANCIENT_DIR/AncientDNA_plus_EUR1KG_AllVariants_NoQC_GRCh37_autosomes")
    p.add_argument("--selection")
    p.add_argument("--join-strategy", choices=["selection-first", "modern-first"], default="selection-first",
                   help="Selection-first indexes the complete published SNV coordinate universe, then validates actual modern intersections")
    p.add_argument("--counts-tsv", help="Optional legacy check counts.tsv for descriptive audit; no replacement for per-site VCF data")
    p.add_argument("--no-plots", action="store_true", help="Skip the optional counts-audit plot")
    p.add_argument("--hq-variants")
    p.add_argument("--supplement-xlsx")
    p.add_argument("--qpadm", help="Optional one verified 3way or 4way TSV; no basis mixing")
    p.add_argument("--ancestry-model", choices=["3way", "4way"])
    p.add_argument("--metadata", help="Canonical metadata from prepare-ancient; otherwise prepare automatically for evolve")
    p.add_argument("--fasta38", default=os.environ.get("ALLELE_FASTA", "/mnt/f/gen/1hgp/38/GRCH38.fasta.gz"))
    p.add_argument("--fasta37", help="GRCh37 / hs37d5 FASTA with existing .fai; required for join")
    p.add_argument("--chain38to37", help="Local hg38ToHg19.over.chain.gz")
    p.add_argument("--chain37to38", help="Local hg19ToHg38.over.chain.gz")
    p.add_argument("--intersection", help="Existing join/selection_intersection.tsv.gz")
    p.add_argument("--candidates", help="Candidate TSV; default OUT/candidates/candidates.tsv")
    p.add_argument("--selected-variants", help="Indexed targets TSV; default OUT/panel/selected_variants.tsv")
    p.add_argument("--max-sites", type=int, default=2000, help="Total strong+weak candidate sites; Atlas requests <= 3 times this")
    p.add_argument("--max-variants", type=int, default=9000, help="Hard budget for UNIQUE Atlas variants, not retry attempts")
    p.add_argument("--strong-posterior", type=float, default=.99)
    p.add_argument("--weak-posterior", type=float, default=.5)
    p.add_argument("--min-extra-ac", type=int, default=5)
    p.add_argument("--block-bp", type=int, default=1000000)
    p.add_argument("--include-mhc", action="store_true", help="Sensitivity analysis: include chr6:25–34 Mb GRCh37")
    p.add_argument("--allow-selection-warnings", action="store_true")
    p.add_argument("--seed", type=int, default=20261021)
    p.add_argument("--bins", default="0,2000,4000,6000,8000,10000")
    p.add_argument("--min-samples", type=int, default=10)
    p.add_argument("--weights", help="Fixed stratum/weight TSV; unsupported bins stay NA")
    p.add_argument("--count-kind", choices=["hardcall", "dosage"], default="hardcall")
    p.add_argument("--skip-trajectories", action="store_true", help="Plan/join without opening PGEN; leaves explicit missing trajectory stage")
    p.add_argument("--with-atlas", action="store_true", help="Explicitly execute the bounded Atlas plan after local analysis")
    p.add_argument("--atlas-plan", help="Existing plan.json; default OUT/atlas_plan/plan.json")
    p.add_argument("--atlas-cache", help="Existing run directory or cache directory for geometry")
    p.add_argument("--scorers", help="Comma-separated scorer enum names; default six signed accessibility/TF/transcription scorers")
    p.add_argument("--ontology-terms", help="Optional comma-separated ontology IDs for a smaller hypothesis-driven panel")
    p.add_argument("--api-key-file", help="Secret file; contents never written into manifests")
    p.add_argument("--api-key-env", default="ALPHAGENOME_API_KEY")
    p.add_argument("--workers", type=int, default=2)
    p.add_argument("--requests-per-second", type=float, default=2.)
    p.add_argument("--retries", type=int, default=3)
    p.add_argument("--timeout", type=float, default=60.)
    p.add_argument("--geometry", help="Existing geometry TSV for compare-functions")
    p.add_argument("--with-pca", action="store_true", help="Execute generated local PLINK2 PCA/projection script and render 3D view")
    p.add_argument("--eigenvec", help="Existing eigenvec or common-basis .sscore for all plotted samples")
    p.add_argument("--pca-basis", choices=["joint_balanced", "modern_projection"], default="joint_balanced")
    p.add_argument("--pca-basis-label", help="Provenance label for an externally provided shared PCA coordinate table")
    p.add_argument("--pca-maf", type=float, default=.05)
    p.add_argument("--pca-geno", type=float, default=.1)
    p.add_argument("--pca-max-per-group", type=int, default=200)
    p.add_argument("--pca-components", type=int, default=10)
    p.add_argument("--pca-target-flank-bp", type=int, default=1000000)
    p.add_argument("--dir-pgen-cache", help="Optional verified execution copies of PGEN/PVAR/PSAM on a fast local filesystem; outputs remain under --out")
    p.add_argument("--pca-threads", type=int, default=4)
    p.add_argument("--pca-memory-mb", type=int, default=8192)
    p.add_argument("--pca-exclude-bed", help="Optional GRCh37 BED for long-range LD or other background exclusions")
    p.add_argument("--pca-max-variants", type=int, default=12, help="Maximum candidate genotypes embedded in 3D HTML")
    p.add_argument("--pca-layer-years", type=int, default=1000, help="Display-only time-layer width; does not change the PCA basis")
    p.add_argument("--pca-time-view", choices=["layered", "continuous"], default="layered")
    p.add_argument("--pca-color-by", choices=["region", "layer", "age", "cohort", "source"], default="region")
    p.add_argument("--plink2", default="plink2")
    p.add_argument("--html", help="Output HTML for pca-time; default OUT/pca/pca_time.html")
    p.add_argument("--simulated", action="store_true", help="Prominent synthetic-data label for examples only")
    p.add_argument("--with-universes", action="store_true")
    p.add_argument("--u1-sites", type=int, default=20000)
    p.add_argument("--u3-sites", type=int, default=3000)
    p.add_argument("--controls-per-case", type=int, default=4)
    p.add_argument("--ld-block-dir", default="/mnt/f/LD/block")
    p.add_argument("--with-ukb", action="store_true")
    p.add_argument("--ukb-mfi-dir", default="/mnt/f/gen/ukb/37/imp_mfi")
    p.add_argument("--metag-census", help="Completed metag genome census JSON, reused read-only")
    p.add_argument("--version", action="version", version=VERSION)
    return p


def paths(a):
    ancient, out = Path(a.ancient_dir).expanduser().resolve(), Path(a.out).expanduser().resolve()
    prefix = Path(a.pgen_prefix).expanduser().resolve() if a.pgen_prefix else ancient / PREFIX
    return dict(out=out, prefix=prefix, pgen=Path(str(prefix) + ".pgen"), psam=Path(str(prefix) + ".psam"),
                pvar=Path(str(prefix) + ".pvar"),
                selection=Path(a.selection) if a.selection else ancient / "Selection_Summary_Statistics_01OCT2025.tsv.gz",
                hq=Path(a.hq_variants) if a.hq_variants else ancient / "HighQualityVariants.tsv",
                supplement=Path(a.supplement_xlsx) if a.supplement_xlsx else ancient / "41586_2026_10358_MOESM3_ESM.xlsx",
                metadata=Path(a.metadata) if a.metadata else out / "metadata" / "sample_metadata.tsv",
                candidates=Path(a.candidates) if a.candidates else out / ("universes" if a.with_universes else "candidates") / "candidates.tsv",
                selected=Path(a.selected_variants) if a.selected_variants else out / "panel" / "selected_variants.tsv",
                intersection=Path(a.intersection) if a.intersection else out / "join" / "selection_intersection.tsv.gz",
                atlas_plan=Path(a.atlas_plan) if a.atlas_plan else out / "atlas_plan" / "plan.json")


def _require(value, name):
    if not value:
        raise ValueError(f"{name} is required for reciprocal build harmonization")
    return value


def _prepare(a, d):
    pass
    out = d["out"] / "metadata"
    inputs = dict(psam=d["psam"], supplement=d["supplement"])
    if a.qpadm:
        inputs["qpadm"] = a.qpadm
    identity = _stage_identity("metadata", inputs, dict(ancestry_model=a.ancestry_model))
    if old := completed_stage(out, identity):
        return old
    result = prepare_metadata(d["psam"], d["supplement"], out, qpadm_path=a.qpadm,
                              ancestry_model=a.ancestry_model)
    return finish_stage(out, identity, result, ["sample_metadata.tsv", "metadata_qc.tsv", "metadata_qc.json"])


def _join(a, d):
    return join_selection(a.modern_out, d["selection"], d["out"] / "join", a.fasta38,
                          _require(a.fasta37, "--fasta37"), _require(a.chain38to37, "--chain38to37"),
                          _require(a.chain37to38, "--chain37to38"), strategy=a.join_strategy)


def _counts(a, d):
    pass
    if not a.counts_tsv:
        raise ValueError("audit-counts requires --counts-tsv")
    return audit_counts(a.counts_tsv, d["out"] / "counts_audit", make_plots=not a.no_plots)


def _select(a, d):
    if a.with_universes:
        from scan import build_universes
        return build_universes(d['intersection'],d['out']/'universes',ld_block_dir=a.ld_block_dir,
            u1_sites=a.u1_sites,strong_posterior=a.strong_posterior,weak_posterior=a.weak_posterior,
            min_extra_ac=a.min_extra_ac,controls_per_case=a.controls_per_case,seed=a.seed,
            exclude_mhc=not a.include_mhc,allow_filter_warnings=a.allow_selection_warnings)
    return select_candidates(d["intersection"], d["out"] / "candidates", max_sites=a.max_sites,
                             strong_posterior=a.strong_posterior, weak_posterior=a.weak_posterior,
                             min_extra_ac=a.min_extra_ac, block_bp=a.block_bp,
                             exclude_mhc=not a.include_mhc, allow_filter_warnings=a.allow_selection_warnings, seed=a.seed)


def _panel(a, d):
    return index_target_variants(d["candidates"], d["pvar"], d["hq"], d["out"] / "panel")


def _trajectories(a, d):
    pass
    out = d["out"] / "trajectories"
    inputs = {key: d[key] for key in ("pgen", "psam", "selected", "metadata")}
    if a.weights:
        inputs["weights"] = a.weights
    identity = _stage_identity("trajectories", inputs,
                               dict(bins=a.bins, min_samples=a.min_samples, count_kind=a.count_kind))
    if old := completed_stage(out, identity):
        return old
    samples, variants = iter_pgen_variants(d["pgen"], d["psam"], d["selected"], count_kind=a.count_kind)
    result = aggregate_trajectories(variants, samples, d["metadata"], out,
                                    bins=tuple(float(x) for x in a.bins.split(",")), min_samples=a.min_samples,
                                    weights_path=a.weights)
    return finish_stage(out, identity, result, ["trajectory_bins.tsv", "trajectory_standardized.tsv",
                        "sample_bin_composition.tsv", "sample_inclusion.tsv", "allele_availability.tsv", "trajectory_manifest.json"])


def _atlas_plan(a, d):
    pass
    source=d['out']/'universes'/'U3.tsv' if a.with_universes else d['candidates']
    return build_plan(source, d["out"] / "atlas_plan", max_sites=a.u3_sites if a.with_universes else a.max_sites,
                      max_variants=a.max_variants, scorers=a.scorers,
                      ontology_terms=a.ontology_terms)


def _atlas_query(a, d):
    pass
    return run_sparse(d["atlas_plan"], d["out"] / "atlas", api_key_file=a.api_key_file,
                      api_key_env=a.api_key_env, max_variants=a.max_variants, workers=a.workers,
                      retries=a.retries, timeout=a.timeout, requests_per_second=a.requests_per_second)


def _atlas_summarize(a, d):
    pass
    return summarize_geometry(d["atlas_plan"], a.atlas_cache or d["out"] / "atlas", d["out"] / "geometry")


def _pca_plan(a, d):
    pass
    prefix = d["prefix"]
    if a.dir_pgen_cache:
        from allele_common import cache_pgen_prefix
        prefix, cache_record = cache_pgen_prefix(prefix, a.dir_pgen_cache)
        dump_json(d["out"] / "pca_input_cache.json", cache_record)
    pca_targets=d['selected']
    if a.with_universes:
        # U1 covers the genome: expanding 20k sites by 1Mb would erase all PCA markers.
        # Exclude U2 strong loci and flanks; U1 sites are only sampling/background.
        from allele_common import tsv_writer
        pca_targets=d['out']/'pca_exclusion_targets.tsv'
        target_rows={(r['chrom37'],int(r['pos37'])) for r in table_rows(d['out']/'universes'/'U2_all_strong.tsv')}
        content='chrom\tpos\tbuild\n'+''.join(f'{c}\t{p}\tGRCh37\n' for c,p in sorted(target_rows))
        if pca_targets.exists():
            if pca_targets.read_text()!=content: raise ValueError('PCA strong-target exclusions changed')
        else: pca_targets.write_text(content)
    return prepare_pca_plan(prefix, d["metadata"], d["hq"], d["out"] / "pca", basis=a.pca_basis,
                            maf=a.pca_maf, geno=a.pca_geno, max_per_group=a.pca_max_per_group,
                            seed=a.seed, exclude_bed=a.pca_exclude_bed, target_variants=pca_targets, plink2=a.plink2,
                            n_components=a.pca_components, target_flank_bp=a.pca_target_flank_bp,
                            threads=a.pca_threads, memory_mb=a.pca_memory_mb)


def _pca_time(a, d, eigenvec=None):
    pass
    source = eigenvec or a.eigenvec
    if not source:
        raise ValueError("pca-time needs --eigenvec from one fixed PCA basis for all epochs")
    label = a.pca_basis if eigenvec else (a.pca_basis_label or "用户提供的共享 PCA 基底；请核对其训练样本与标记")
    kwargs = dict(basis_label=label, max_variants=a.pca_max_variants, simulated=a.simulated, count_kind=a.count_kind, candidate_path=d.get("candidates"),
                  layer_years=a.pca_layer_years, time_view=a.pca_time_view, color_by=a.pca_color_by)
    if d["selected"].exists() and d["pgen"].exists() and d["psam"].exists():
        from allele_common import tsv_writer
        rows = list(table_rows(d["selected"]))
        by_id = {r["variant_id"]: r for r in rows}
        links = d["out"] / "panel" / "site_variant_links.tsv"
        ordered = []
        if links.exists() and not a.selected_variants:
            for link in table_rows(links):
                r = by_id.get(link["variant_id"])
                if r and r not in ordered:
                    ordered.append(r)
                if len(ordered) >= a.pca_max_variants:
                    break
        else:
            ordered = rows[:a.pca_max_variants]
        if ordered:
            bounded = d["out"] / "pca" / "overlay_targets.tsv"
            h, w = tsv_writer(bounded, list(ordered[0]))
            with h:
                w.writerows(ordered)
            dump_json(d["out"] / "pca" / "overlay_selection.json",
                      dict(rule="first candidate sites in pre-functional selection order, anchor then extra; bounded visualization only",
                           total_available=len(rows), displayed=len(ordered), max_variants=a.pca_max_variants,
                           source=fingerprint(d["selected"])))
            kwargs.update(pgen_path=d["pgen"], psam_path=d["psam"], selected_variants_path=bounded)
    return render_pca_time(source, d["metadata"], a.html or d["out"] / "pca" / "pca_time.html", **kwargs)


def run_evolution(a, d):
    from report import compare_geometry, write_report
    started = time.monotonic()
    result = dict(version=VERSION, status="running", stages={}, api_requested=bool(a.with_atlas),
                  pca_execution_requested=bool(a.with_pca), modern_out=str(Path(a.modern_out).resolve()))
    path = d["out"] / "evolution.json"

    def stage(name, fn):
        print(f"[evolution] {name}", flush=True)
        value = fn()
        result["stages"][name] = value
        dump_json(path, result)
        return value

    try:
        if a.counts_tsv:
            stage("counts_audit", lambda: _counts(a, d))
        if a.metadata:
            result["stages"]["metadata"] = dict(source=fingerprint(d["metadata"]), adapter="user_supplied_canonical")
        else:
            stage("metadata", lambda: _prepare(a, d))
        stage("join", lambda: _join(a, d))
        stage("candidates", lambda: _select(a, d))
        if a.with_ukb:
            if not a.with_universes: raise ValueError('--with-ukb requires --with-universes so the full U0 is scanned first')
            from scan import scan_mfi_dir
            stage('ukb_mfi',lambda: scan_mfi_dir(a.ukb_mfi_dir,d['out']/'universes'/'U0.tsv.gz',d['out']/'ukb'/'U0_mfi.tsv'))
        stage("panel", lambda: _panel(a, d))
        if not a.skip_trajectories:
            stage("trajectories", lambda: _trajectories(a, d))
        if not a.with_universes: stage("atlas_plan", lambda: _atlas_plan(a, d))
        pca = stage("pca_plan", lambda: _pca_plan(a, d))
        if a.with_pca:
            # Executable paths/arguments are written with shlex.quote by planner.
            # No shell interpolation of secrets or uncontrolled command strings.
            script = pca["script_path"]
            stage("pca_execution", lambda: _execute_pca(script))
            stage("pca_visualization", lambda: _pca_time(a, d, pca["projected_path"]))
            from pca3d.variant_centroid import export_dosages, run as centroid
            from types import SimpleNamespace
            dosage_path=d['out']/'pca'/'selected_variant_dosages.tsv.gz'
            selected=d['out']/'pca'/'overlay_targets.tsv'
            if selected.exists():
                stage('dosage_export',lambda:export_dosages(d['pgen'],d['psam'],selected,dosage_path,a.count_kind))
                stage('pca_centroid',lambda:centroid(SimpleNamespace(coordinates=pca['projected_path'],metadata=d['metadata'],dosages=dosage_path,out=d['out']/'pca'/'centroids',bin_years=2000)))
        elif a.eigenvec:
            stage("pca_visualization", lambda: _pca_time(a, d))
        if a.with_universes:
            from report import write_local_review
            stage('local_review',lambda: write_local_review(a.modern_out,d['out'],a.metag_census,make_plots=not a.no_plots))
            from scan import select_u3
            stage('U3_selection',lambda: select_u3(d['out']/'universes',d['out']/'universes'/'U3.tsv',a.u3_sites,a.seed))
            stage('atlas_plan',lambda: _atlas_plan(a,d))
        if a.with_atlas:
            stage("atlas_query", lambda: _atlas_query(a, d))
            geometry = stage("atlas_geometry", lambda: _atlas_summarize(a, d))
            stage("functional_comparison", lambda: compare_geometry(geometry["outputs"]["geometry"], d["out"] / "comparison", seed=a.seed))
            if a.with_universes:
                stage('final_figures',lambda: write_local_review(a.modern_out,d['out'],a.metag_census,make_plots=not a.no_plots))
        result.update(status="complete", elapsed_seconds=time.monotonic() - started)
    except BaseException as exc:
        result.update(status="incomplete", error_type=type(exc).__name__, elapsed_seconds=time.monotonic() - started)
        # Do not persist potentially credential-bearing exception text.
        dump_json(path, result)
        write_report(d["out"])
        raise
    result["report"] = write_report(d["out"])
    dump_json(path, result)
    return result


def _execute_pca(script):
    process = subprocess.run(["bash", str(script)], check=False)
    if process.returncode:
        raise RuntimeError(f"PLINK2 PCA script failed with exit {process.returncode}; inspect its local logs")
    return dict(status="complete", script=str(script))


def audit(a, d):
    result = dict(files={})
    for key in ("psam", "pvar", "pgen", "selection", "hq", "supplement"):
        p = d[key]
        result["files"][key] = fingerprint(p) if p.exists() else dict(path=str(p), missing=True)
    if d["psam"].exists():
        result["psam_data_rows"] = sum(1 for _ in table_rows(d["psam"], ("IID",)))
    for key in ("pvar", "selection", "hq"):
        if d[key].exists():
            rows = table_rows(d[key])
            first = next(rows, None)
            rows.close()
            result[key + "_first_record"] = first
    result["note"] = "PVAR full record count is obtained by panel-index; wc -l includes headers. PGEN fields are not inferred from filename."
    dump_json(d["out"] / "ancient_audit.json", result)
    return result


def main_evolution(argv=None):
    a = parser().parse_args(argv)
    d = paths(a)
    d["out"].mkdir(parents=True, exist_ok=True)
    try:
        with (d["out"] / ".evolution.lock").open("a") as lock:
            fcntl.flock(lock, fcntl.LOCK_EX)
            actions = {"evolve": run_evolution, "prepare-ancient": _prepare, "join-selection": _join,
                       "select-candidates": _select, "panel-index": _panel, "trajectories": _trajectories,
                       "atlas-plan": _atlas_plan, "atlas-query": _atlas_query, "atlas-summarize": _atlas_summarize,
                       "pca-plan": _pca_plan, "pca-time": _pca_time, "audit-ancient": audit,
                       "audit-counts": _counts}
            if a.command == "report-evolution":
                from report import write_report
                result = dict(report=write_report(d["out"]))
            elif a.command == "compare-functions":
                from report import compare_geometry
                if not a.geometry:
                    raise ValueError("compare-functions requires --geometry")
                result = compare_geometry(a.geometry, d["out"] / "comparison", seed=a.seed)
            else:
                result = actions[a.command](a, d)
        # Large manifests stay in files; avoid printing secrets or 6,000 variants.
        print(json.dumps({k: result[k] for k in ("status", "manifest_path", "metadata_path", "candidate_path",
                                                 "plan_path", "report", "report_path", "qc", "pvar_records") if k in result},
                         ensure_ascii=False, indent=2))
        return 0
    except (ValueError, OSError, RuntimeError) as exc:
        print(f"allele3 evolution: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main_evolution())
