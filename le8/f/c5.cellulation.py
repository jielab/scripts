#!/usr/bin/env python3
"""C5: expression-context annotation and native CIGMA integration.

Drop-in CLI replacement for le8/f/c5.cellulation.py. No C1-C4 code is changed.
Cell labels are external expression annotations, NOT measured cell abundance,
secretion tracing, cell age, or proof of a causal cell of action.

Commands: annotate, import-cellage, import-cigma, inspect-table, cigma,
          prepare-pseudobulk, kinship; no subcommand = existing LE8 driver.
"""
from __future__ import annotations
import argparse
import csv
import gzip
import hashlib
import importlib.metadata
import importlib.util
import inspect
import itertools
import json
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import warnings

import numpy as np
import pandas as pd
from scipy import sparse
from scipy.stats import fisher_exact, hypergeom

VERSION = "2026-10-04.cell-context-v2"
HERE = Path(__file__).resolve().parent
REPOSITORY = "https://github.com/Minhui-Chen/CIGMA"
VERIFIED_COMMIT = "5813e4ae84d7b3733dfcd938fe42d12c6b30a8aa"
VERIFIED_FIT_BLOB = "5ad0b510c193543f633152e097fe29fe72cc73cf"
CELLAGE_COMMIT = "a0de4a124e0a41986185a090fb142c6c8195399e"
# Gene-level alias only. Assay identity is never collapsed for assay budgets.
ALIASES = {"NTPROBNP": "NPPB"}
NOTE = "External expression context; not secretion tracing, cell abundance, cell aging or causal direction"


def sha(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def canonical_json(value):
    def convert(x):
        if isinstance(x, np.ndarray): return convert(x.tolist())
        if isinstance(x, np.generic): return convert(x.item())
        if isinstance(x, Path): return str(x)
        if isinstance(x, dict): return {str(k): convert(v) for k, v in x.items()}
        if isinstance(x, (list, tuple)): return [convert(v) for v in x]
        if isinstance(x, float) and not math.isfinite(x): return None
        return x
    return json.dumps(convert(value), ensure_ascii=False, indent=2, sort_keys=True, allow_nan=False) + "\n"


def atomic_text(text, path):
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix="." + path.name + ".", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as f: f.write(text)
        os.replace(name, path)
    finally:
        if os.path.exists(name): os.unlink(name)


def write_json(value, path): atomic_text(canonical_json(value), path)


def write_csv(d, path):
    atomic_text(d.to_csv(index=False, na_rep=""), path)


def provenance_path(path):
    """Stable published identity while reading the disposable table workspace."""
    p = Path(path).resolve()
    work, published = os.getenv("LE8_ANALYSIS_ROOT"), os.getenv("LE8_PUBLISHED_ROOT")
    if work and published:
        try: return str(Path(published).resolve() / p.relative_to(Path(work).resolve()))
        except ValueError: pass
    return str(p)


def stamp(path):
    if path is None: return None
    p = Path(path).resolve()
    return dict(path=provenance_path(p), sha256=sha(p)) if p.is_file() else dict(path=provenance_path(p), missing=True)


def read_table(path, sheet=None, header_row=0):
    """Read data without altering identifiers or silently accepting duplicate headers."""
    p = Path(path)
    if not p.is_file(): raise FileNotFoundError(p)
    if p.suffix.lower() in {".xlsx", ".xlsm"}:
        # Optional user-side importer; no spreadsheet is created by this module.
        import openpyxl
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        try:
            if sheet is None:
                if len(wb.sheetnames) != 1:
                    raise ValueError("Select an explicit sheet: " + ", ".join(wb.sheetnames))
                sheet = wb.sheetnames[0]
            ws = wb[sheet]; rows = ws.iter_rows(min_row=header_row + 1, values_only=True)
            hdr = next(rows, ())
            # Remove only completely empty trailing columns, not interior headers.
            hdr = list(hdr)
            while hdr and hdr[-1] is None: hdr.pop()
            names = [str(x).strip() if x is not None else "" for x in hdr]
            data = [list(x[:len(names)]) for x in rows]
            d = pd.DataFrame(data, columns=names)
        finally: wb.close()
    else:
        op = gzip.open if str(p).endswith(".gz") else open
        with op(p, "rt", encoding="utf-8-sig", newline="") as f:
            for _ in range(header_row): next(f)
            line = next(f, "")
        sep = "\t" if line.count("\t") > line.count(",") else ","
        names = next(csv.reader([line], delimiter=sep), [])
        d = pd.read_csv(p, sep=sep, header=header_row, dtype=str, keep_default_na=False)
    if not names or len(names) != len(set(names)) or any(not x for x in names):
        raise ValueError(f"Missing/duplicate column labels: {p}")
    if list(d.columns) != names: raise ValueError(f"Parser changed headers: {p}")
    return d


def require(d, columns, label):
    missing = set(columns) - set(d)
    if missing: raise ValueError(f"{label}: missing columns {sorted(missing)}")


def textcol(s): return s.fillna("").astype(str).str.strip()


def gene_name(x):
    if pd.isna(x): return ""
    s = str(x).strip().upper()
    if s in {"", "NAN", "NONE", "NULL", "NA", "."}: return ""
    if re.search(r"[;,|/\s]", s):
        raise ValueError(f"Ambiguous multi-gene mapping {x!r}; supply one verified gene or blank")
    if re.fullmatch(r"ENSG\d+\.\d+", s): s = s.split(".")[0]
    return ALIASES.get(s, s)


def bh(values, family_size=None):
    p = np.asarray(values, dtype=float)
    finite = np.isfinite(p)
    if ((p[finite] < 0) | (p[finite] > 1)).any(): raise ValueError("P outside [0,1]")
    m = len(p) if family_size is None else int(family_size)
    if m < len(p): raise ValueError("Declared family is smaller than the supplied test table")
    out = np.full(len(p), np.nan); ix = np.flatnonzero(finite)
    ix = ix[np.argsort(p[ix], kind="stable")]
    if len(ix):
        out[ix] = np.minimum(1, np.minimum.accumulate((p[ix] * m / np.arange(1, len(ix)+1))[::-1])[::-1])
    return out


def pcolumn(d, column, optional=False):
    if column not in d:
        if optional: return np.full(len(d), np.nan)
        raise ValueError("Missing " + column)
    raw = textcol(d[column]); v = pd.to_numeric(raw.replace("", np.nan), errors="raise").to_numpy(float)
    if np.isinf(v).any() or ((v[np.isfinite(v)] < 0) | (v[np.isfinite(v)] > 1)).any():
        raise ValueError("Invalid probability in " + column)
    return v


def normalized_panels(p):
    """Separate folds before deduplication. A fold is never pooled into a panel."""
    p=p.copy(); require(p,["model","feature"],"panels")
    p["model"]=textcol(p.model); p["feature"]=textcol(p.feature)
    p=p[p.feature.ne("")].copy()
    if p.model.eq("").any(): raise ValueError("Panel has no model identifier")
    if "fold" in p:
        p["fold"]=textcol(p.fold)
        if p.fold.eq("").any(): raise ValueError("Explicit fold column must not have missing labels")
        if p.model.str.contains("|fold=",regex=False).any():
            raise ValueError("Use either explicit fold column or encoded |fold= model, not both")
        p["model_base"]=p.model
        p["model"]=p.model+"|fold="+p.fold
    return p.drop_duplicates(["model","feature"])


def prepare_annotation(universe, atlas, panels):
    u, a, p = (read_table(x) for x in (universe, atlas, panels))
    require(u, ["assay", "gene"], "universe")
    require(a, ["gene", "cell_type", "source", "source_version"], "atlas")
    require(p, ["model", "feature"], "panels")
    u = u.copy(); a = a.copy(); p = p.copy()
    u["assay"] = textcol(u.assay); u["gene_original"] = textcol(u.gene)
    u["gene"] = u.gene.map(gene_name)
    if u.assay.eq("").any() or u.assay.duplicated().any(): raise ValueError("Assay universe must be unique/nonmissing")
    for col in ["cell_type", "source", "source_version"]:
        a[col] = textcol(a[col])
        if a[col].eq("").any(): raise ValueError("Missing atlas " + col)
    a["gene"] = a.gene.map(gene_name)
    if a.gene.eq("").any(): raise ValueError("Atlas has missing gene identifiers")
    if len(a[["source", "source_version"]].drop_duplicates()) != 1:
        raise ValueError("Run one prespecified atlas/version at a time; never pool label systems")
    a = a.drop_duplicates(["gene", "cell_type"])
    p = normalized_panels(p)
    if not set(p.feature) <= set(u.assay): raise ValueError("Panel assay not in full measured assay universe")
    if not set(u.gene) - {""}: raise ValueError("No mapped genes in full assay universe")
    return u, a, p


def conditional_panel_test(A, B, labelled):
    """Exact conditional gene-label exchangeability test; shared genes stay fixed.

    This is NOT an independent validation of the panel learning algorithm.
    """
    if not A or not B: return dict(delta=np.nan, null_center=np.nan, p=np.nan, exclusive_genes=0)
    common = A & B; onlya = A-B; onlyb = B-A; pool = onlya | onlyb
    observed = len(A & labelled)/len(A) - len(B & labelled)/len(B)
    if not pool: return dict(delta=observed, null_center=0., p=1., exclusive_genes=0)
    N, successes, draws = len(pool), len(pool & labelled), len(onlya)
    low, high = max(0, draws - (N-successes)), min(draws, successes)
    x = np.arange(low, high+1); prob = hypergeom.pmf(x, N, successes, draws)
    shared_hits = len(common & labelled)
    vals = (shared_hits+x)/len(A) - (shared_hits+successes-x)/len(B)
    center = float(np.dot(prob, vals))
    pval = float(prob[np.abs(vals-center) >= abs(observed-center)-1e-12].sum())
    return dict(delta=observed, null_center=center, p=min(1., pval), exclusive_genes=N)


def matched_enrichment(genes, background, cellsets, strata, draws, rng):
    """Sample gene sets matched on prespecified technical/annotation strata."""
    pools = {}; need = {}
    for g in sorted(background): pools.setdefault(strata[g], []).append(g)
    for g in genes: need[strata[g]] = need.get(strata[g], 0) + 1
    cells = sorted(cellsets); observed = np.array([len(genes & cellsets[c]) for c in cells])
    exceed = np.zeros(len(cells), int); means = np.zeros(len(cells), float)
    for _ in range(draws):
        pick = set()
        for key, n in sorted(need.items()): pick.update(rng.choice(pools[key], n, replace=False).tolist())
        vals = np.array([len(pick & cellsets[c]) for c in cells])
        exceed += vals >= observed; means += vals
    return {c: ((exceed[i]+1)/(draws+1), means[i]/draws) for i, c in enumerate(cells)}


def annotate(universe, atlas, panels, outdir, prefix="cell", contrasts=None,
             matched_draws=0, seed=2026, plots=True):
    if not re.fullmatch(r"[A-Za-z0-9_.-]+",prefix): raise ValueError("Invalid file prefix")
    if matched_draws < 0: raise ValueError("Negative matched draws")
    u, a, p = prepare_annotation(universe, atlas, panels)
    outdir = Path(outdir); outdir.mkdir(parents=True, exist_ok=True)
    background = set(u.gene) - {""}; labelled = set(a.gene) & background
    cellsets = {c: set(t.gene) & background for c, t in a.groupby("cell_type", sort=True)}
    all_annotations = u.merge(a, on="gene", how="left", validate="many_to_many")
    all_annotations["annotation_status"] = np.where(all_annotations.gene.eq(""), "gene_unmapped",
        np.where(all_annotations.cell_type.notna(), "expression_label_present", "no_label_in_this_atlas"))
    all_annotations["interpretation"] = NOTE
    ann = p.merge(all_annotations, left_on="feature", right_on="assay", how="left")
    msets = {}; coverage = []; rows = []; strata = None
    if matched_draws:
        require(u, ["match_stratum"], "Matched-null universe")
        tmp = u[u.gene.ne("")][["gene", "match_stratum"]].copy(); tmp.match_stratum = textcol(tmp.match_stratum)
        if tmp.match_stratum.eq("").any() or tmp.groupby("gene").match_stratum.nunique().max() > 1:
            raise ValueError("A gene must have one nonmissing, prespecified match_stratum")
        strata = dict(tmp.drop_duplicates("gene").values)
    for model, group in p.groupby("model", sort=True):
        m = group.merge(u, left_on="feature", right_on="assay", validate="many_to_one")
        genes = set(m.gene)-{""}; msets[model] = genes
        coverage.append(dict(model=model, assays=len(group), mapped_assays=int(m.gene.ne("").sum()),
            unique_genes=len(genes), cell_labelled_genes=len(genes & labelled),
            labelled_gene_fraction=len(genes & labelled)/len(genes) if genes else np.nan,
            labelled_assay_fraction=float(m.gene.isin(labelled).mean()), background_genes=len(background),
            interpretation=NOTE))
        for kind, bg in [("all_measured_genes", background), ("atlas_labelled_measured_genes", labelled)]:
            g = genes & bg
            for cell, target in cellsets.items():
                cg = target & bg; hit = len(g & cg)
                table = [[hit, len(g)-hit], [len(cg-g), len(bg-g-cg)]]
                valid = bool(g and cg and len(cg) < len(bg) and len(g) < len(bg))
                odds, pv = fisher_exact(table, alternative="greater") if valid else (np.nan, 1. if g else np.nan)
                rows.append(dict(model=model, cell_type=cell, background_type=kind, hits=hit,
                    panel_genes=len(g), atlas_genes_in_assay_background=len(cg), background_genes=len(bg),
                    odds_ratio=odds, p=pv, status="tested" if valid else "uninformative",
                    hit_genes=";".join(sorted(g & cg))))
    res = pd.DataFrame(rows, columns=["model","cell_type","background_type","hits","panel_genes",
        "atlas_genes_in_assay_background","background_genes","odds_ratio","p","status","hit_genes"])
    # Primary family preserves the old name and uses all measured genes. Sensitivity has its own declared family.
    primary = res[res.background_type.eq("all_measured_genes")].copy()
    secondary = res[res.background_type.eq("atlas_labelled_measured_genes")].copy()
    primary["FDR_all_panel_cell_tests"] = bh(primary.p)
    secondary["FDR_all_panel_cell_tests"] = bh(secondary.p)
    matched = []
    if matched_draws:
        for model, genes in msets.items():
            rng = np.random.default_rng(seed + int(hashlib.sha256(model.encode()).hexdigest()[:8], 16))
            z = matched_enrichment(genes, background, cellsets, strata, matched_draws, rng)
            for cell, (pv, mean) in z.items():
                matched.append(dict(model=model, cell_type=cell, hits=len(genes & cellsets[cell]),
                    expected_matched_hits=mean, p=pv, draws=matched_draws,
                    null="fixed gene-stratum counts; conditional random gene-set null, not participant bootstrap"))
    mt = pd.DataFrame(matched, columns=["model","cell_type","hits","expected_matched_hits","p","draws","null"])
    mt["FDR_matched_family"] = bh(mt.p)
    ct = []
    if contrasts is not None:
        pairs = read_table(contrasts); require(pairs, ["model", "reference"], "contrasts")
        pairs = pairs.drop_duplicates(["model", "reference"])
        budgets = dict(zip([x["model"] for x in coverage], [x["assays"] for x in coverage]))
        for _, pair in pairs.iterrows():
            left, right = pair["model"], pair["reference"]
            if left not in msets or right not in msets: raise ValueError("Unknown contrast model/reference")
            if budgets[left] != budgets[right]: raise ValueError("Panel contrast requires equal actual assay counts")
            if "fold" in p:
                fa=set(p.loc[p.model.eq(left),"fold"].astype(str)); fb=set(p.loc[p.model.eq(right),"fold"].astype(str))
                if fa != fb: raise ValueError("Panel contrast requires the same fold")
            A, B = msets[left], msets[right]
            for cell, target in {"__ANY_LABEL__": labelled, **cellsets}.items():
                z = conditional_panel_test(A, B, target)
                ct.append(dict(model=left, reference=right, cell_type=cell, **z,
                    assay_budget=budgets[left], genes_a=len(A), genes_b=len(B), shared_genes=len(A&B),
                    inference="Conditional exclusive-gene allocation; not algorithm superiority or intervention response"))
    con = pd.DataFrame(ct, columns=["model","reference","cell_type","delta","null_center","p","exclusive_genes",
        "assay_budget","genes_a","genes_b","shared_genes","inference"])
    con["FDR_all_contrasts"] = bh(con.p)
    write_csv(pd.DataFrame([dict(analysis="same_budget_contrasts",status="completed" if contrasts is not None and len(con) else "no_estimable_contrasts" if contrasts is not None else "not_requested",
        comparisons=len(con),reason="Explicit, prespecified same-budget/same-fold contrast table required")]),outdir/f"{prefix}.contrast_status.csv")
    cov = pd.DataFrame(coverage, columns=["model","assays","mapped_assays","unique_genes","cell_labelled_genes",
        "labelled_gene_fraction","labelled_assay_fraction","background_genes","interpretation"])
    for name, table in [("all_assay_annotation", all_annotations), ("panel_annotation",ann), ("coverage",cov),
                         ("enrichment",primary), ("enrichment_annotated_background",secondary),
                         ("matched_enrichment",mt), ("panel_contrasts",con)]:
        write_csv(table, outdir / f"{prefix}.{name}.csv")
    # Stability is label inclusion frequency, NOT an average of transformed P values over folds.
    stability = []
    if "fold" in p:
        for model, g in p.groupby("model_base"):
            nfold = g.fold.nunique()
            for feature, gg in g.groupby("feature"):
                stability.append(dict(model=model, feature=feature, included_folds=gg.fold.nunique(),
                    observed_folds=nfold, frequency=gg.fold.nunique()/nfold))
    write_csv(pd.DataFrame(stability, columns=["model","feature","included_folds","observed_folds","frequency"]),
              outdir / f"{prefix}.fold_coverage.csv")
    prov = dict(version=VERSION, inputs=[stamp(x) for x in [universe,atlas,panels,contrasts]],
        atlas=a[["source","source_version"]].drop_duplicates().to_dict("records"),
        aliases=ALIASES, universe_assays=len(u), universe_genes=len(background),
        atlas_labelled_measured_genes=len(labelled), atlas_cells=len(cellsets),
        primary_family=len(primary), sensitivity_family=len(secondary),
        matched_draws=matched_draws, seed=seed, interpretation=NOTE)
    write_json(prov, outdir / f"{prefix}.provenance.json")
    write_csv(pd.DataFrame([dict(source=a.source.iloc[0],source_version=a.source_version.iloc[0],
        interpretation=NOTE,denominator="All measured unique mapped genes; assay identities retained")]),
        outdir/f"{prefix}.provenance.csv")
    if plots: annotation_plots(primary, cov, outdir, prefix)
    return dict(coverage=cov, enrichment=primary, contrasts=con, annotation=all_annotations)


def annotation_plots(primary, coverage, outdir, prefix):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    for old in list(outdir.glob(f"{prefix}.enrichment*.png")) + list(outdir.glob(f"{prefix}.enrichment*.panel_data.csv")):
        old.unlink()
    # Each figure has one axis. Models/folds are never silently averaged or recoded.
    page_size = 16
    models = list(coverage.model)
    for page, start in enumerate(range(0,len(models),page_size), 1):
        selected = models[start:start+page_size]
        rows = primary[primary.model.isin(selected)]
        cells = rows.groupby("cell_type").hits.sum().sort_values(ascending=False, kind="stable").head(20).index
        grid = rows.pivot(index="cell_type", columns="model", values="hits").reindex(index=cells, columns=selected)
        q = rows.pivot(index="cell_type", columns="model", values="FDR_all_panel_cell_tests").reindex_like(grid)
        fig, ax = plt.subplots(figsize=(max(9,len(selected)*.6), max(5,len(cells)*.32+2)))
        im = ax.imshow(grid.fillna(0), aspect="auto")
        ax.set_xticks(range(len(selected)), selected, rotation=65, ha="right",fontsize=8)
        ax.set_yticks(range(len(cells)), cells,fontsize=9)
        for i in range(len(cells)):
            for j in range(len(selected)):
                if pd.notna(q.iloc[i,j]) and q.iloc[i,j]<.05: ax.text(j,i,"*",ha="center",va="center")
        ax.set_title("Cell-expression annotation: hit counts; * primary-family FDR < 0.05")
        fig.colorbar(im, ax=ax,label="Unique labelled genes (not secretion evidence)")
        fig.tight_layout(); name=f"{prefix}.enrichment" + ("" if page==1 else f".page{page}") + ".png"
        write_csv(rows[rows.cell_type.isin(cells)], outdir / name.replace(".png", ".panel_data.csv"))
        fig.savefig(outdir/name,dpi=180); plt.close(fig)
    if not models:
        fig, ax=plt.subplots(figsize=(9,3)); ax.text(.5,.5,"No frozen panel supplied; see all-assay annotation",ha="center")
        write_csv(pd.DataFrame([{"analysis":"frozen panel annotation","status":"unavailable: no frozen panel supplied"}]), outdir/f"{prefix}.enrichment.panel_data.csv")
        ax.set_axis_off(); fig.savefig(outdir/f"{prefix}.enrichment.png",dpi=150); plt.close(fig)


# ---------------- Native CIGMA: labelled donors, real kinship, no fake single-cell input ----------------
def matrix(path):
    d = read_table(path)
    if d.shape[1] < 2: raise ValueError(f"Expected labelled matrix: {path}")
    ids = textcol(d.iloc[:,0])
    if ids.eq("").any() or ids.duplicated().any(): raise ValueError(f"Missing/duplicate matrix row IDs: {path}")
    d = d.iloc[:,1:].copy(); d.index = ids; d.index.name = "donor"
    d = d.apply(pd.to_numeric, errors="raise")
    if not np.isfinite(d.to_numpy()).all(): raise ValueError(f"Nonfinite matrix values: {path}")
    return d


def aligned(path, donors, columns=None):
    d = matrix(path)
    if set(d.index) != set(donors): raise ValueError(f"Donor set mismatch: {path}")
    if columns is not None:
        if set(d.columns) != set(columns): raise ValueError(f"Cell/column set mismatch: {path}")
        d = d.loc[:,columns]
    return d.loc[donors]


def valid_kinship(d, label):
    x = d.to_numpy(float)
    if x.shape[0] != x.shape[1] or not np.allclose(x,x.T,rtol=0,atol=1e-7):
        raise ValueError(label+": nonsymmetric kinship")
    if np.min(np.linalg.eigvalsh(x)) < -1e-6 or np.min(np.diag(x)) <= 0:
        raise ValueError(label+": kinship not positive semidefinite with positive diagonal")
    if np.allclose(x,np.eye(len(x))*np.trace(x)/len(x)):
        raise ValueError(label+": identity kinship cannot distinguish genetic covariance from residual covariance")
    return x


def read_inputs(row, base):
    if row.get("ctnu_definition") != "variance_of_pseudobulk_mean":
        raise ValueError("ctnu_definition must be variance_of_pseudobulk_mean (SEM squared, not SD or cell variance)")
    paths = {k:(base/row[k]).resolve() for k in ["ctp","ctnu","P","K"]}
    y = matrix(paths["ctp"]); donors, cells = y.index, list(y.columns)
    if len(y) < 20 or len(cells) < 2: raise ValueError("Need >=20 donors and >=2 cell types; this is not a power guarantee")
    if row.get("cell_types") and set(row["cell_types"].split(";")) != set(cells):
        raise ValueError("Manifest cell_types does not match CTP")
    nu = aligned(paths["ctnu"],donors,cells); prop = aligned(paths["P"],donors,cells)
    kin = aligned(paths["K"],donors,donors)
    if (nu.to_numpy()<0).any(): raise ValueError("Negative pseudobulk-mean noise variance")
    if (prop.to_numpy()<=0).any() or (prop.to_numpy()>1).any() or (prop.sum(axis=1)>1+1e-6).any():
        raise ValueError("Represented donor/cell pairs need positive proportions; row totals must be <=1")
    if (np.var(y.to_numpy(),axis=0)<=0).any(): raise ValueError("Constant/unexpressed cell type")
    args=dict(Y=y.to_numpy(float),ctnu=nu.to_numpy(float),P=prop.to_numpy(float),K=valid_kinship(kin,"K"))
    q=0
    for key in ["fixed_covars","random_covars"]:
        if row.get(key):
            paths[key]=(base/row[key]).resolve(); cov=aligned(paths[key],donors)
            arr=cov.to_numpy(float)
            if key=="fixed_covars":
                # CIGMA already adds cell intercepts. Do not add another intercept.
                if np.linalg.matrix_rank(np.column_stack([np.ones(len(y)),arr])) != arr.shape[1]+1:
                    raise ValueError("Fixed covariates contain a constant/intercept or collinear columns")
                q=arr.shape[1]
            args[key]={"design":arr}
    if row.get("Kt"):
        paths["Kt"]=(base/row["Kt"]).resolve()
        args["Kt"]=valid_kinship(aligned(paths["Kt"],donors,donors),"Kt")
        if np.allclose(args["Kt"],args["K"]): raise ValueError("K and Kt are identical; separate variance components unidentified")
    # Conservative safeguard for the native donor-jackknife Wald denominator.
    npar=2+3*len(cells)+q+len(args.get("random_covars",{}))+(len(cells)+1 if "Kt" in args else 0)
    if len(y)<=npar+2: raise ValueError(f"Insufficient donors for CIGMA parameter count: N={len(y)}, parameters~{npar}")
    if row.get("n"):
        paths["n"]=(base/row["n"]).resolve(); n=aligned(paths["n"],donors,cells).to_numpy(float)
        if ((n<2)|(n!=np.floor(n))).any(): raise ValueError("Need >=2 cells per represented donor/cell pair to estimate SEM squared")
    return args,cells,paths


def load_cigma(allow_untested=False):
    from cigma import fit
    function=fit.free_HE
    expected={"Y","K","ctnu","P","fixed_covars","random_covars","Kt","jk","verbose"}
    if not expected <= set(inspect.signature(function).parameters):
        raise ValueError("Unsupported CIGMA free_HE API")
    src=Path(inspect.getfile(fit)); data=src.read_bytes()
    blob=hashlib.sha1(b"blob "+str(len(data)).encode()+b"\0"+data).hexdigest()
    if blob!=VERIFIED_FIT_BLOB and not allow_untested:
        raise ValueError("Installed CIGMA source differs from inspected commit. Install the pinned release, or explicitly pass --allow-untested-cigma after checking its API.")
    info=dict(repository=REPOSITORY, inspected_commit=VERIFIED_COMMIT,
        fit_source=str(src), fit_blob=blob, fit_sha256=sha(src), source_matches_inspected=blob==VERIFIED_FIT_BLOB,
        installed_version=importlib.metadata.version("cigma"))
    return function, info


def scalar(x):
    a=np.asarray(x,dtype=float)
    if a.size!=1: raise ValueError("Expected one scalar from CIGMA")
    return float(a.reshape(-1)[0])


def variance_summary(est,p,cells,base,component="cis"):
    second=component=="trans"; suffix="_b" if second else ""
    shared=scalar(est["hom_g2"+suffix]); V=np.asarray(est["V"+suffix],float)
    if V.shape!=(len(cells),len(cells)): raise ValueError("Native V dimension mismatch")
    v=np.diag(V); mean=float(v.mean()); denom=shared+mean
    if not np.isfinite(shared) or not np.isfinite(v).all(): raise ValueError("Nonfinite native variance estimates")
    raw=mean/denom if denom>0 else np.nan
    admissible=shared>=0 and (v>=0).all() and denom>0
    var_s=scalar(p.get("var_hom_g2"+suffix,np.nan))
    Vs=np.asarray(p.get("var_V"+suffix,np.full_like(V,np.nan)),float)
    se=math.sqrt(var_s) if var_s>=0 else np.nan
    semean=math.sqrt(float(Vs.sum())/len(cells)**2) if Vs.shape==V.shape and np.isfinite(Vs).all() and Vs.sum()>=0 else np.nan
    pvc=np.asarray(p.get("vc"+suffix,np.full(len(cells),np.nan)),float).reshape(-1)
    if pvc.size!=len(cells): raise ValueError("Native cell-specific P dimension mismatch")
    specific_p=scalar(p.get("V"+suffix,np.nan)); shared_p=scalar(p.get("hom_g2"+suffix,np.nan))
    probs=np.r_[specific_p,shared_p,pvc]
    if ((probs[np.isfinite(probs)]<0)|(probs[np.isfinite(probs)]>1)).any(): raise ValueError("Native P value outside [0,1]")
    summary=dict(**base,component=component,shared_variance=shared,mean_specific_variance=mean,
        shared_se=se,mean_specific_se=semean,specificity_raw=raw,
        specificity=float(raw) if admissible else np.nan,variance_admissible=bool(admissible),
        specific_p=specific_p,shared_p=shared_p,method="CIGMA free_HE; full donor jackknife",
        interpretation="Genetic regulation of expression; no disease mediation or cell origin inferred")
    cellrows=[]
    for i,c in enumerate(cells):
        vv=Vs[i,i] if Vs.shape==V.shape else np.nan
        cellrows.append(dict(**base,component=component,cell_type=c,specific_variance=v[i],
            specific_se=np.sqrt(vv) if vv>=0 else np.nan,total_genetic_variance=shared+v[i],p=pvc[i]))
    return summary,cellrows


def run(manifest,outdir,validate_only=False,seed=2026,allow_untested=False,_fit_function=None):
    manifest=Path(manifest).resolve(); outdir=Path(outdir); outdir.mkdir(parents=True,exist_ok=True)
    jobs=read_table(manifest); required=["gene","ctp","ctnu","P","K","tissue","build","kinship_scope","ctnu_definition"]
    require(jobs,required,"CIGMA manifest")
    if jobs.empty: raise ValueError("Empty CIGMA manifest")
    for c in required:
        jobs[c]=textcol(jobs[c])
        if jobs[c].eq("").any(): raise ValueError("Blank manifest field "+c)
    if not jobs.kinship_scope.isin(["cis","trans","genomewide"]).all(): raise ValueError("kinship_scope must be cis, trans or genomewide")
    jobs["gene"]=jobs.gene.map(gene_name)
    if jobs.gene.eq("").any(): raise ValueError("Empty CIGMA gene")
    if jobs.duplicated(["gene","tissue"]).any(): raise ValueError("Duplicate gene/tissue jobs; use separate manifests for cohorts")
    # Clear published result tables BEFORE loading native code, so a failed rerun cannot leave apparent success.
    result_cols=["gene","tissue","component","specific_p","specific_FDR_manifest","specificity"]
    write_csv(pd.DataFrame(columns=result_cols),outdir/"cigma.results.csv")
    write_csv(pd.DataFrame(columns=["gene","tissue","component","cell_type","specific_variance","p","FDR_cell_manifest"]),outdir/"cigma.cell_types.csv")
    status=[]; results=[]; cellrows=[]; provenance=[]; fits=[]; cell_family_complete=True; expected_cells=0
    pkg=dict(installed_version="validation-only", inspected_commit=VERIFIED_COMMIT)
    fn=None
    try:
        if not validate_only:
            if _fit_function is not None:
                fn=_fit_function; pkg={"installed_version":"test double; NOT a native fit"}
            else: fn,pkg=load_cigma(allow_untested)
        for index,row in enumerate(jobs.to_dict("records")):
            base={k:row[k] for k in ["gene","tissue","build","kinship_scope","ctnu_definition"]}
            base["cohort"]=row.get("cohort","")
            declared=[]
            try:
                declared=list(matrix(manifest.parent/row["ctp"]).columns)
                expected_cells+=len(declared)*(2 if row.get("Kt") else 1)
            except Exception:
                declared=row.get("cell_types","").split(";") if row.get("cell_types") else []
                expected_cells+=len(declared)*(2 if row.get("Kt") else 1)
                if not declared: cell_family_complete=False
            try:
                args,cells,paths=read_inputs(row,manifest.parent)
                for role,path in paths.items(): provenance.append(dict(**base,role=role,**stamp(path)))
                if not validate_only:
                    np.random.seed(seed+index)
                    with warnings.catch_warnings(record=True) as caught:
                        est,p=fn(**args,jk=True,verbose=True)
                    base.update(N=args["Y"].shape[0],cell_types=len(cells),package_version=pkg["installed_version"])
                    comps=["cis","trans"] if "Kt" in args else [row["kinship_scope"]]
                    # For a single K the component name can be 'cis', 'trans', or 'genomewide'.
                    for component in comps:
                        s,cr=variance_summary(est,p,cells,base,"trans" if "Kt" in args and component=="trans" else "cis")
                        s["component"]=component
                        for rr in cr: rr["component"]=component
                        results.append(s); cellrows.extend(cr)
                    # Stream large donor-jackknife arrays instead of retaining every gene in memory.
                    native_file=outdir/"native_fits"/(f"{index:06d}.json")
                    write_json(dict(**base,estimates=est,pvalues_and_uncertainty=p,
                        warnings=[str(w.message) for w in caught]),native_file)
                    fits.append(dict(gene=base["gene"],tissue=base["tissue"],
                        file=str(native_file.relative_to(outdir)),sha256=sha(native_file)))
                    finite=all(np.isfinite(x["specific_p"]) for x in results if x["gene"]==base["gene"] and x["tissue"]==base["tissue"])
                    status.append(dict(**base,status="ok" if finite else "numerically_incomplete",message="; ".join(str(w.message) for w in caught)))
                else: status.append(dict(**base,status="validated",message="Input validation only; no fit performed"))
            except Exception as exc:
                status.append(dict(**base,status="failed",message=f"{type(exc).__name__}: {exc}"))
    except Exception as exc:
        status=[dict(gene=r["gene"],tissue=r["tissue"],status="failed",message=str(exc)) for r in jobs.to_dict("records")]
    result=pd.DataFrame(results)
    family_size=len(jobs)+sum(bool(r.get("Kt")) for r in jobs.to_dict("records"))
    if len(result):
        result["specific_FDR_manifest"]=bh(result.specific_p, family_size)
        result["shared_FDR_manifest"]=bh(result.shared_p, family_size)
        write_csv(result,outdir/"cigma.results.csv")
    c=pd.DataFrame(cellrows)
    if len(c):
        c["FDR_cell_manifest"]=bh(c.p,expected_cells) if cell_family_complete else np.nan
        c["cell_family_complete"]=cell_family_complete
        write_csv(c,outdir/"cigma.cell_types.csv")
    write_csv(pd.DataFrame(status),outdir/"cigma.status.csv")
    write_csv(pd.DataFrame(provenance,columns=["gene","tissue","build","kinship_scope","ctnu_definition","cohort","role","path","sha256"]),outdir/"cigma.inputs.csv")
    # Store all native estimates/uncertainty, including negative HE estimates; never truncate to zero.
    write_json(fits,outdir/"cigma.native_fits.json")
    write_json(dict(version=VERSION,**pkg,manifest=stamp(manifest),tests=family_size,
        gene_family_size=family_size,cell_family_size=expected_cells,cell_family_complete=cell_family_complete,
        full_test_family=True,seed=seed,validate_only=validate_only,record_type="native" if _fit_function is None else "test_double"),
        outdir/"cigma.provenance.json")
    return 1 if any(s["status"] in {"failed","numerically_incomplete"} for s in status) else 0


# ---------------- External summaries: explicit column map, original inference family ----------------
def import_cigma(table, mapping, outdir):
    spec=json.loads(Path(mapping).read_text())
    needed={"source","source_version","cohort","tissue","component","full_test_family","columns"}
    if not needed<=set(spec): raise ValueError("CIGMA mapping JSON requires "+", ".join(sorted(needed)))
    for k in ["source","source_version","cohort","tissue","component"]:
        if not str(spec[k]).strip(): raise ValueError("Blank summary provenance "+k)
    if not isinstance(spec["full_test_family"],bool): raise ValueError("full_test_family must be JSON true/false")
    d=read_table(table,spec.get("sheet"),int(spec.get("header_row",0)))
    cols=spec["columns"]
    if not {"gene","specific_p"}<=set(cols): raise ValueError("Map gene and specific_p explicitly")
    allowed={"gene","specific_p","specific_FDR_manifest","shared_p","shared_variance","mean_specific_variance", "shared_se","mean_specific_se","specificity","specificity_raw"}
    if not set(cols)<=allowed: raise ValueError("Unknown canonical mapped column: "+", ".join(sorted(set(cols)-allowed)))
    require(d,list(cols.values()),"External CIGMA table")
    # Drop empty trailing workbook rows only when every mapped field is empty.
    d=d.loc[~d[list(cols.values())].apply(lambda x:textcol(x).eq("")).all(axis=1)].copy()
    out=pd.DataFrame({k:d[v].values for k,v in cols.items()})
    out["gene"]=out.gene.map(gene_name)
    if out.gene.eq("").any() or out.gene.duplicated().any():
        raise ValueError("Missing/duplicate canonical genes; select one cohort/model per import and resolve identifiers")
    out["specific_p"]=pcolumn(out,"specific_p")
    if "specific_FDR_manifest" in out:
        out["specific_FDR_manifest"]=pcolumn(out,"specific_FDR_manifest")
        family="Original supplied adjustment; "+str(spec.get("original_adjustment","unspecified"))
    elif spec["full_test_family"]:
        n=int(spec.get("n_tests",len(out)))
        out["specific_FDR_manifest"]=bh(out.specific_p,n)
        family=f"BH across declared complete source family, n={n}"
    else:
        out["specific_FDR_manifest"]=np.nan
        family="Not recomputed: imported rows are not the complete test family"
    for k in ["shared_variance","mean_specific_variance","shared_se","mean_specific_se","specificity_raw","specificity"]:
        if k in out: out[k]=pd.to_numeric(textcol(out[k]).replace("",np.nan),errors="raise")
    if "specificity" not in out: out["specificity"]=np.nan
    if "specificity_raw" not in out: out["specificity_raw"]=out.specificity
    out["specificity_in_unit_interval"]=np.isfinite(out.specificity)&out.specificity.between(0,1)
    out.loc[~out.specificity_in_unit_interval,"specificity"]=np.nan
    if {"shared_variance","mean_specific_variance"}<=set(out):
        denom=out.shared_variance+out.mean_specific_variance
        if "specificity_raw" not in out: out["specificity_raw"]=np.where(denom>0,out.mean_specific_variance/denom,np.nan)
    for key in ["source","source_version","cohort","tissue","component"]: out[key]=str(spec[key])
    out["kinship_scope"]=spec["component"]; out["adjustment_family"]=family
    out["evidence_origin"]="external_precomputed"; out["tested"]=np.isfinite(out.specific_p)
    cells=[]
    # Optional explicit wide-to-long cell columns; do not infer a specific cell from a joint gene P value.
    for cell in spec.get("cells",[]):
        require(pd.DataFrame([cell]),["cell_type","variance_col","p_col"],"Cell mapping")
        require(d,[cell["variance_col"],cell["p_col"]],"External CIGMA cell columns")
        cc=pd.DataFrame(dict(gene=out.gene,cell_type=cell["cell_type"],
            specific_variance=pd.to_numeric(d[cell["variance_col"]].values,errors="raise"),
            p=pd.to_numeric(d[cell["p_col"]].values,errors="raise")))
        cc["p"]=pcolumn(cc,"p")
        if cell.get("q_col"):
            require(d,[cell["q_col"]],"cell q column")
            cc["FDR_cell_manifest"]=pd.to_numeric(d[cell["q_col"]].values,errors="raise")
            cc["FDR_cell_manifest"]=pcolumn(cc,"FDR_cell_manifest")
        if cell.get("se_col"):
            cc["specific_se"]=pd.to_numeric(d[cell["se_col"]].values,errors="raise")
        cells.append(cc)
    celltab=pd.concat(cells,ignore_index=True) if cells else pd.DataFrame(columns=["gene","cell_type","specific_variance","p","FDR_cell_manifest"])
    if len(celltab):
        # An explicit complete gene x cell family is required, not implied by complete gene-level tests.
        if "FDR_cell_manifest" not in celltab:
            celltab["FDR_cell_manifest"]=bh(celltab.p,int(spec.get("n_cell_tests",len(celltab)))) if spec.get("complete_cell_test_family") is True else np.nan
        for k in ["source","source_version","cohort","tissue","component"]: celltab[k]=str(spec[k])
    outdir=Path(outdir); outdir.mkdir(parents=True,exist_ok=True)
    write_csv(out,outdir/"cigma.results.csv"); write_csv(celltab,outdir/"cigma.cell_types.csv")
    write_json(dict(version=VERSION,record_type="external_precomputed",source_table=stamp(table),mapping=stamp(mapping),
        full_test_family=spec["full_test_family"],gene_family_size=spec.get("n_tests",len(out)),
        adjustment_family=family,source=spec["source"],source_version=spec["source_version"],
        interpretation="Reused published CIGMA statistics; not rerun in UKB or independent of their source cohort"),outdir/"cigma.provenance.json")
    return out


def integrate_cigma(results, resources, out, cellfile=None):
    d=read_table(results); require(d,["gene","tissue","specific_p","specific_FDR_manifest","specificity"],"CIGMA results")
    d["gene"]=d.gene.map(gene_name)
    if "component" not in d: d["component"]="unspecified"
    if "cohort" not in d: d["cohort"]="unspecified"
    keys=["gene","tissue","cohort","component"]
    if d[keys].apply(lambda s:textcol(s).eq("")).any().any() or d.duplicated(keys).any(): raise ValueError("Nonunique/missing CIGMA identifiers")
    for k in ["specific_p","specific_FDR_manifest"]: d[k]=pcolumn(d,k)
    # External metadata is necessary to distinguish reuse from a local native fit.
    meta=Path(results).parent/"cigma.provenance.json"
    if not meta.is_file(): raise ValueError("CIGMA results need adjacent cigma.provenance.json; use import-cigma for a published table")
    provenance=json.loads(meta.read_text())
    if provenance.get("record_type") not in {"native","external_precomputed"}:
        raise ValueError("CIGMA provenance must describe native or external_precomputed results, not a test double")
    if provenance.get("validate_only"): raise ValueError("Validation-only CIGMA files are not inference results")
    d["interpretation"]="Expression-regulatory specificity, not protein source or disease-specific mediation"
    write_csv(d,out/"c5.CIGMA_results.csv")
    write_json(provenance,out/"c5.CIGMA_source.json")
    if cellfile is None:
        trial=Path(results).parent/"cigma.cell_types.csv"
        cellfile=trial if trial.is_file() else None
    c=None
    if cellfile:
        c=read_table(cellfile)
        require(c,["gene","tissue","cell_type","p","FDR_cell_manifest"],"CIGMA cell results") if len(c) else None
        if len(c):
            c["gene"]=c.gene.map(gene_name)
            for k in ["p","FDR_cell_manifest"]: c[k]=pcolumn(c,k)
        write_csv(c,out/"c5.CIGMA_cell_types.csv")
    if resources[1] is None: return d
    u=read_table(resources[1]); require(u,["assay","gene"],"universe"); u["gene"]=u.gene.map(gene_name)
    # Left join retains untested genes as missing, not CIGMA negatives.
    aa=u.merge(d,on="gene",how="left"); aa["CIGMA_available"]=aa.specific_p.notna()
    write_csv(aa,out/"c5.CIGMA_annotation.csv")
    if resources[2] is not None and provenance.get("full_test_family") is True:
        p=normalized_panels(read_table(resources[2])).merge(u,left_on="feature",right_on="assay",how="left")
        rows=[]
        for context,dt in d.groupby(["cohort","tissue","component"],dropna=False):
            # Only truly tested, assayed genes define the enrichment background.
            bg=(set(dt.loc[dt.specific_p.notna(),"gene"]) & set(u.gene)) - {""}
            sig=set(dt.loc[dt.specific_FDR_manifest<.05,"gene"]) & bg
            for model,pg in p.groupby("model"):
                g=set(pg.gene)&bg; hit=len(g&sig)
                valid=bool(g and sig and len(g)<len(bg) and len(sig)<len(bg))
                odds,pv=fisher_exact([[hit,len(g)-hit],[len(sig-g),len(bg-g-sig)]],alternative="greater") if valid else (np.nan,np.nan)
                rows.append(dict(model=model,cohort=context[0],tissue=context[1],component=context[2],
                    panel_tested_genes=len(g),assayed_tested_background=len(bg),hits=hit,
                    cs_egene_background=len(sig),odds_ratio=odds,p=pv,
                    interpretation="Enrichment of source-defined cs-eGenes; a gene-level joint test is not a cell label"))
        e=pd.DataFrame(rows)
        if len(e): e["FDR_panel_context"]=bh(e.p)
        write_csv(e,out/"c5.CIGMA_enrichment.csv")
    else:
        write_csv(pd.DataFrame([dict(status="not_tested",
            reason="A complete source test family and panel file are required; significant-only published rows are not a valid background")]),
            out/"c5.CIGMA_enrichment_status.csv")
    return d


# ---------------- Input helpers for real single-cell data ----------------
def pseudobulk_moments(X, donors, celltypes, gene_names, min_cells=11, min_expression_fraction=.1):
    """Small-matrix reference implementation; X must already be normalized.

    ctnu is unbiased sample variance / n_cells, matching SEM squared.
    No gene or donor is imputed. The same retained donors/cells are used for all genes.
    """
    X=sparse.csr_matrix(X,dtype=float)
    donors=np.asarray(donors,dtype=str); celltypes=np.asarray(celltypes,dtype=str)
    if len(donors)!=X.shape[0] or len(celltypes)!=X.shape[0]: raise ValueError("Cell metadata length mismatch")
    ds=sorted(set(donors)); cs=sorted(set(celltypes)); pairs=list(itertools.product(ds,cs)); index={x:i for i,x in enumerate(pairs)}
    groups=np.array([index[(d,c)] for d,c in zip(donors,celltypes)])
    S=sparse.csr_matrix((np.ones(len(groups)),(groups,np.arange(len(groups)))),shape=(len(pairs),len(groups)))
    n=np.asarray(S.sum(axis=1)).ravel(); sums=(S@X).toarray(); sq=(S@X.power(2)).toarray()
    return moments_from_sums(n,sums,sq,ds,cs,gene_names,min_cells,min_expression_fraction)


def moments_from_sums(n,sums,sq,donors,cells,genes,min_cells,expression_fraction):
    D,C=len(donors),len(cells); counts=np.asarray(n).reshape(D,C)
    keep=(counts>=min_cells).all(axis=1)
    if not keep.any(): raise ValueError("No donor has enough cells in every requested cell type")
    nn=np.asarray(n)[:,None]
    with np.errstate(divide="ignore",invalid="ignore"):
        mu=sums/nn; varmean=(sq-sums*sums/nn)/((nn-1)*nn)
    # Only repair floating point cancellation at machine precision, never negative biological estimates.
    tolerance=1e-9*np.maximum(1.,np.abs(sq))
    if ((varmean< -tolerance)&np.isfinite(varmean)).any(): raise ValueError("Invalid negative noise estimate")
    varmean=np.maximum(varmean,0)
    mu=mu.reshape(D,C,-1)[keep]; varmean=varmean.reshape(D,C,-1)[keep]
    counts=counts[keep]; ds=np.asarray(donors)[keep]
    usable=((mu>0).mean(axis=0)>expression_fraction).all(axis=0) & (np.var(mu,axis=0)>0).all(axis=0)
    return dict(ctp=mu,ctnu=varmean,n=counts,donors=ds,cells=list(cells),genes=list(genes),usable=usable,retained=keep)


def prepare_pseudobulk(args):
    import anndata
    ann=anndata.read_h5ad(args.h5ad,backed="r")
    try:
        require(ann.obs,[args.donor_col,args.cell_type_col],"h5ad.obs")
        rawdonor=ann.obs[args.donor_col]; rawcell=ann.obs[args.cell_type_col]
        if rawdonor.isna().any() or rawcell.isna().any(): raise ValueError("Missing donor or cell type; perform QC before C5")
        donors=textcol(rawdonor).to_numpy(); types=textcol(rawcell).to_numpy()
        if (donors=="").any() or (types=="").any(): raise ValueError("Blank donor/cell IDs")
        ds=sorted(set(donors)); cs=args.cell_types.split(",")
        if len(cs)<2 or len(set(cs))!=len(cs) or not set(cs)<=set(types): raise ValueError("Invalid requested cell types")
        genes=[x.strip() for x in Path(args.genes).read_text().splitlines() if x.strip()]
        if not genes or len(genes)!=len(set(genes)) or not ann.var_names.is_unique: raise ValueError("Nonunique gene IDs")
        missing=set(genes)-set(ann.var_names)
        if missing: raise ValueError("Requested genes absent from h5ad (use matching gene IDs): "+",".join(sorted(missing)[:10]))
        gi=ann.var_names.get_indexer(genes); di={x:i for i,x in enumerate(ds)}; ci={x:i for i,x in enumerate(cs)}
        G=len(ds)*len(cs); n=np.zeros(G); sums=np.zeros((G,len(genes))); sq=sums.copy()
        total_counts=pd.Series(donors).value_counts()
        data=ann.X if args.layer is None else ann.layers[args.layer]
        for start in range(0,ann.n_obs,args.chunk_cells):
            end=min(start+args.chunk_cells,ann.n_obs)
            xx=sparse.csr_matrix(data[start:end,:],dtype=float)
            if not np.isfinite(xx.data).all() or (xx.data<0).any(): raise ValueError("Nonfinite/negative expression values")
            if args.normalization=="log10-cp10k":
                if not np.allclose(xx.data,np.round(xx.data),atol=1e-6,rtol=0): raise ValueError("Input is not raw counts; choose pretransformed only with a documented scale")
                totals=np.asarray(xx.sum(axis=1)).ravel()
                if (totals<=0).any(): raise ValueError("Zero-library cells must be removed upstream")
                xx=xx[:,gi].multiply(1e4/totals[:,None]).tocsr(); xx.data=np.log10(xx.data+1)
            else: xx=xx[:,gi]
            ids=np.array([di[d]*len(cs)+ci[c] if c in ci else -1 for d,c in zip(donors[start:end],types[start:end])])
            ok=np.flatnonzero(ids>=0)
            H=sparse.csr_matrix((np.ones(len(ok)),(ids[ok],ok)),shape=(G,end-start))
            n+=np.asarray(H.sum(axis=1)).ravel(); sums+=(H@xx).toarray(); sq+=(H@xx.power(2)).toarray()
        pb=moments_from_sums(n,sums,sq,ds,cs,genes,args.min_cells,args.min_expression_fraction)
        out=Path(args.outdir); out.mkdir(parents=True,exist_ok=True)
        count=pd.DataFrame(pb["n"],index=pb["donors"],columns=cs); count.index.name="donor"
        P=count.div(total_counts.loc[pb["donors"]].to_numpy(),axis=0)
        atomic_text(count.to_csv(),out/"n.csv"); atomic_text(P.to_csv(),out/"P.csv")
        write_csv(pd.DataFrame(dict(donor=ds,retained=pb["retained"])),out/"donor_qc.csv")
        output=[]
        for i,gene in enumerate(genes):
            if not pb["usable"][i]:
                output.append(dict(gene=gene,status="excluded_expression_or_variance",ctp="",ctnu="")); continue
            tag=hashlib.sha256(gene.encode()).hexdigest()[:16]
            y=pd.DataFrame(pb["ctp"][:,:,i],index=pb["donors"],columns=cs); y.index.name="donor"
            v=pd.DataFrame(pb["ctnu"][:,:,i],index=pb["donors"],columns=cs); v.index.name="donor"
            atomic_text(y.to_csv(),out/f"{tag}.ctp.csv"); atomic_text(v.to_csv(),out/f"{tag}.ctnu.csv")
            output.append(dict(gene=gene,status="prepared",ctp=f"{tag}.ctp.csv",ctnu=f"{tag}.ctnu.csv"))
        tab=pd.DataFrame(output)
        for k,val in dict(P="P.csv",n="n.csv",K="",tissue=args.tissue,build=args.build,cohort=args.cohort,
                          kinship_scope="cis",ctnu_definition="variance_of_pseudobulk_mean",cell_types=";".join(cs)).items(): tab[k]=val
        # Explicitly incomplete until paired donor genotype-derived K is supplied. Never write a runnable fake K.
        write_csv(tab,out/"pseudobulk_manifest_needs_K.csv")
        write_json(dict(version=VERSION,source=stamp(args.h5ad),genes=stamp(args.genes),normalization=args.normalization,
            expression_scale=args.expression_scale,proportions="selected cell count / all QC cells for donor",min_cells=args.min_cells,
            min_expression_fraction=args.min_expression_fraction,donors_in=len(ds),donors_retained=len(pb["donors"]),
            note="Only pseudobulk preparation; supply gene-specific K from genotypes of these same donors; not a CIGMA result"),out/"pseudobulk.provenance.json")
    finally: ann.file.close()


def kinship_from_dosage(dosage,min_maf=.05,max_missing=.02):
    x=np.asarray(dosage,float)
    if x.ndim!=2 or x.shape[0]<3: raise ValueError("Expected donor x variant dosage matrix")
    valid=np.isfinite(x)
    if ((x[valid]<0)|(x[valid]>2)).any(): raise ValueError("Autosomal diploid dosages must lie in [0,2]")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore",RuntimeWarning); freq=np.nanmean(x,axis=0)/2
    keep=np.isfinite(freq)&(freq>0)&(freq<1)&(np.minimum(freq,1-freq)>=min_maf)&((~valid).mean(axis=0)<=max_missing)
    if keep.sum()<2: raise ValueError("Fewer than two usable polymorphic variants; no kinship generated")
    x=x[:,keep]; f=freq[keep]; x=np.where(np.isfinite(x),x,2*f)
    z=(x-2*f)/np.sqrt(2*f*(1-f)); K=z@z.T/z.shape[1]
    return K,keep,f


# ---------------- LE8 integration; original public interfaces and output names retained ----------------
def default_cell_atlas():
    return Path(os.getenv("C5_CELL_ATLAS","F:/annot/cellage/atlas.csv" if os.name=="nt" else "/mnt/f/annot/cellage/atlas.csv"))


def import_cellage(repo,output):
    repo=Path(repo); source=repo/"preprocessing/cell_type_mapping_update.csv"
    table=read_table(source); require(table,["Marker Genes","Original Cell types"],"CellAge author table")
    version=subprocess.check_output(["git","-C",str(repo),"rev-parse","HEAD"],text=True).strip()
    dirty=subprocess.check_output(["git","-C",str(repo),"status","--porcelain","--",str(source.relative_to(repo))],text=True).strip()
    if dirty: raise ValueError("CellAge marker file is locally modified; commit/provenance would be misleading")
    rows=[]
    for _,row in table.iterrows():
        for gene in str(row["Marker Genes"]).split(","):
            gene=gene_name(gene)
            if gene: rows.append(dict(gene=gene,cell_type=str(row["Original Cell types"]).strip(),
                source="Ding et al. CellAge author mapping; HPA-derived",source_version=version))
    out=pd.DataFrame(rows).drop_duplicates(["gene","cell_type"])
    if out.empty or out.cell_type.eq("").any(): raise ValueError("Empty/malformed CellAge atlas")
    output=Path(output); output.mkdir(parents=True,exist_ok=True)
    write_csv(out,output/"atlas.csv")
    if (repo/"LICENSE").is_file(): shutil.copy2(repo/"LICENSE",output/"LICENSE")
    write_json(dict(repository="https://github.com/dingdaisy/cellage",commit=version,source=stamp(source),
        paper="https://doi.org/10.1038/s41591-026-04446-y",rule="Author Marker Genes; not SomaScan-only assay IDs",
        genes=out.gene.nunique(),cell_types=out.cell_type.nunique(),interpretation=NOTE),output/"provenance.json")
    print(f"Imported {out.gene.nunique()} genes and {out.cell_type.nunique()} labels: {output/'atlas.csv'}")


def default_resources(root,atlas):
    """Use the repository's existing native-table resolver before older Final exports."""
    scratch=Path(tempfile.gettempdir())/"le8-c5-inputs"/hashlib.sha256(provenance_path(root).encode()).hexdigest()[:16]
    scratch.mkdir(parents=True,exist_ok=True)
    messages=[]
    try:
        from final import cell_annotation_inputs
        paths=cell_annotation_inputs(root,scratch)
        if paths is not None and all(Path(x).is_file() for x in paths): return [atlas,*map(Path,paths)],messages
    except Exception as exc: messages.append("Native input resolver: "+str(exc))
    for directory,universe,panels in [
        (root.parent/"final"/root.name/"systematic","c5.systematic.cell_universe.csv","c5.systematic.cell_panels.csv"),
        (root.parent/"final"/root.name/"joint","c5.cell_assay_universe.csv","c5.cell_panels.csv")]:
        if (directory/universe).is_file() and (directory/panels).is_file():
            messages.append("Using existing Final export; inspect provenance to ensure it matches the desired frozen run")
            return [atlas,directory/universe,directory/panels],messages
    return [atlas,None,None],messages


def evidence_long(root,universe,out):
    """Attach C1-C4 rows without best-P/H4 selection, gene vote counts or inferred direction."""
    u=read_table(universe); require(u,["assay","gene"],"universe"); u=u[["assay","gene"]].copy(); u.gene=u.gene.map(gene_name)
    specs=[
        ("C1_association","c1_correlate/pwas_incident_adj2.csv","term","beta","p.value","FDR",""),
        ("C2_MR","c2_cause/c2.MR_all.csv","exposure","b","pval","FDR_all","analysis"),
        ("C3_coloc","c3_coloc/c3.coloc_summary.csv","feature","PP.H4","","","locus"),
        ("C4_proxy","c4_connect/c4.primary_pillar_assignment.csv","feature","r_disc","p_disc","FDR_disc","primary_component"),
        ("C4_path","c4_connect/c4.mediation_all.csv","feature","indirect_beta","indirect_p","FDR_indirect","component")]
    records=[]; inventory=[]
    for domain,relative,idcol,value,pcol,qcol,context in specs:
        path=root/"prot"/relative
        if not path.is_file(): inventory.append(dict(domain=domain,path=provenance_path(path),status="not_published_at_canonical_path")); continue
        d=read_table(path)
        if idcol not in d: inventory.append(dict(domain=domain,path=provenance_path(path),status="schema_mismatch")); continue
        filehash=sha(path)
        for i,row in enumerate(d.to_dict("records")):
            records.append(dict(feature=row[idcol],evidence_domain=domain,
                context=row.get(context,""),value=row.get(value,""),p=row.get(pcol,""),q=row.get(qcol,""),
                locus=row.get("locus",""),locus_class=row.get("locus_class",row.get("analysis","")),
                signal_id=row.get("signal_id",""),row_status=row.get("status",""),
                source_path=provenance_path(path),source_row=i+2,source_sha256=filehash))
        inventory.append(dict(domain=domain,path=provenance_path(path),status="read",rows=len(d),sha256=filehash))
    cols=["feature","evidence_domain","context","value","p","q","locus","locus_class","signal_id","row_status","source_path","source_row","source_sha256"]
    d=pd.DataFrame(records,columns=cols).merge(u,left_on="feature",right_on="assay",how="left")
    write_csv(d,out/"c5.evidence_long.csv")
    # Annotate each original evidence row without combining its P values or borrowing another locus.
    labels=out/"c5.cell.all_assay_annotation.csv"
    if labels.is_file():
        la=read_table(labels)
        require(la,["assay","cell_type","annotation_status"],"All-assay context")
        select=[x for x in ["assay","cell_type","source","source_version","annotation_status"] if x in la]
        context=d.merge(la[select].drop_duplicates(),on="assay",how="left")
        context["context_interpretation"]=NOTE
        write_csv(context,out/"c5.evidence_cell_context.csv")
    write_csv(pd.DataFrame(inventory),out/"c5.evidence_sources.csv")


def main_driver(argv):
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--Y",default="cvd_cad,ra"); ap.add_argument("--biom",default="prot,met")
    ap.add_argument("--analysis-root",type=Path,required=True)
    for name in ["atlas","universe","panels","contrasts","cigma-manifest","cigma-results","cigma-cells"]:
        ap.add_argument("--"+name,type=Path)
    ap.add_argument("--replace",action="store_true"); ap.add_argument("--no-plots",action="store_true")
    ap.add_argument("--matched-draws",type=int,default=0); ap.add_argument("--seed",type=int,default=int(os.getenv("SEED","2026")))
    ap.add_argument("--python",default=os.getenv("C5_CIGMA_PYTHON",sys.executable))
    ap.add_argument("--allow-untested-cigma",action="store_true")
    ap.add_argument("--preflight",action="store_true")
    a=ap.parse_args(argv)
    if (a.universe is None)!=(a.panels is None): ap.error("Provide --universe and --panels together")
    if a.cigma_manifest and a.cigma_results: ap.error("Choose a native manifest OR precomputed CIGMA results")
    if a.cigma_cells and not a.cigma_results: ap.error("--cigma-cells requires --cigma-results")
    if a.matched_draws and a.matched_draws<100: ap.error("Use >=100 matched draws, or 0 to disable")
    if a.panels and ("," in a.Y or a.biom!="prot"): ap.error("Explicit panels require one --Y and --biom prot")
    if (a.cigma_manifest or a.cigma_results) and "," in a.Y: ap.error("Attach CIGMA resources to one outcome per invocation")
    for name in ["atlas","universe","panels","contrasts","cigma_manifest","cigma_results","cigma_cells"]:
        value=getattr(a,name)
        if value is not None and not value.is_file(): ap.error(f"Missing --{name.replace('_','-')}: {value}")
    if a.preflight:
        if a.contrasts:
            require(read_table(a.contrasts),["model","reference"],"contrasts")
        print(canonical_json(dict(status="preflight_ok",seed=a.seed,matched_draws=a.matched_draws,
            contrasts="requested" if a.contrasts else "not_requested",cigma="input_supplied" if a.cigma_manifest or a.cigma_results else "not_requested")))
        return 0
    explicit_atlas=a.atlas is not None; a.atlas=a.atlas or default_cell_atlas()
    rc=0
    for trait in a.Y.split(","):
        for layer in a.biom.split(","):
            if layer not in {"prot","met"} or not re.fullmatch(r"[A-Za-z0-9_-]+",trait): raise ValueError("Unsafe/invalid trait or layer")
            root=a.analysis_root.resolve()/trait; out=root/layer/"c5_cellulation"; out.mkdir(parents=True,exist_ok=True)
            resources=[a.atlas,a.universe,a.panels]; notes=[]
            if layer=="prot" and a.universe is None: resources,notes=default_resources(root,a.atlas)
            signature=dict(version=VERSION,trait=trait,layer=layer,code=sha(Path(__file__)),
                inputs=[stamp(x) for x in resources+[a.contrasts,a.cigma_results,a.cigma_cells]],
                seed=a.seed,matched_draws=a.matched_draws,plots=not a.no_plots,allow_untested_cigma=a.allow_untested_cigma)
            # Include result provenance and cell table; changes invalidate reuse.
            if a.cigma_results:
                signature["cigma_meta"]=stamp(a.cigma_results.parent/"cigma.provenance.json")
                signature["auto_cells"]=stamp(a.cigma_results.parent/"cigma.cell_types.csv")
            # Include actual C1-C4 source signatures used by the evidence ledger.
            signature["evidence"]=[stamp(root/"prot"/p) for p in [
                "c1_correlate/pwas_incident_adj2.csv","c2_cause/c2.MR_all.csv",
                "c3_coloc/c3.coloc_summary.csv","c4_connect/c4.primary_pillar_assignment.csv",
                "c4_connect/c4.mediation_all.csv"]] if layer=="prot" else []
            completed=out/"c5.completed.json"
            old={}
            if completed.is_file():
                try: old=json.loads(completed.read_text())
                except (ValueError,OSError): pass
            if not a.replace and not a.cigma_manifest and old.get("signature")==signature and old.get("outputs"):
                if all((out/n).is_file() and sha(out/n)==h for n,h in old["outputs"].items()):
                    print(f"C5 {trait}/{layer}: unchanged completed outputs reused"); continue
            # Stage the full new output set; old results are not published beside failed reruns.
            with tempfile.TemporaryDirectory(prefix=".c5-stage-",dir=out) as staging:
                stage=Path(staging); status=[]
                write_csv(pd.DataFrame(columns=["gene","tissue","specific_p","specific_FDR_manifest","specificity"]),stage/"c5.CIGMA_results.csv")
                write_csv(pd.DataFrame(columns=["assay","gene","tissue","specific_p","specific_FDR_manifest","specificity"]),stage/"c5.CIGMA_annotation.csv")
                if layer=="met":
                    status.append(dict(analysis="cell_expression",status="not_applicable",detail="No unique metabolite encoding gene; no direct gene/cell conversion"))
                elif all(p is not None and p.is_file() for p in resources):
                    try:
                        annotate(resources[1],resources[0],resources[2],stage,"c5.cell",a.contrasts,a.matched_draws,a.seed,not a.no_plots)
                        evidence_long(root,resources[1],stage)
                        status.append(dict(analysis="cell_expression",status="completed",detail="All measured genes; explicit sensitivity background; no averaging of fold P values"))
                    except Exception as exc:
                        for part in list(stage.glob("c5.cell.*"))+list(stage.glob("c5.evidence_*")): part.unlink()
                        rc=1; status.append(dict(analysis="cell_expression",status="failed",detail=str(exc)))
                else:
                    failed=explicit_atlas or a.universe is not None
                    rc=max(rc,int(failed))
                    status.append(dict(analysis="cell_expression",status="failed" if failed else "unavailable",
                        detail="Need existing atlas.csv, full assay/gene universe and frozen panels; not a negative result"))
                results=a.cigma_results
                native_failed=False
                if a.cigma_manifest and layer=="prot":
                    native_out=out/"cigma_native"; native_out.mkdir(exist_ok=True)
                    command=[a.python,str(Path(__file__).resolve()),"cigma","--manifest",str(a.cigma_manifest.resolve()),
                        "--outdir",str(native_out),"--seed",str(a.seed)]
                    if a.allow_untested_cigma: command.append("--allow-untested-cigma")
                    with (native_out/"runner.log").open("w") as log:
                        proc=subprocess.run(command,stdout=log,stderr=subprocess.STDOUT)
                    results=native_out/"cigma.results.csv"; native_failed=proc.returncode!=0
                    if native_failed:
                        rc=1; results=None
                        status.append(dict(analysis="CIGMA",status="failed",detail="See cigma_native status and log; partial fits are not promoted to completed"))
                if results and layer=="prot":
                    try:
                        d=integrate_cigma(results,resources,stage,a.cigma_cells)
                        status.append(dict(analysis="CIGMA",status="completed" if len(d) else "unavailable",
                            detail=f"{len(d)} gene/context rows; read c5.CIGMA_source.json for external vs native provenance"))
                    except Exception as exc:
                        for part in stage.glob("c5.CIGMA*"): part.unlink()
                        write_csv(pd.DataFrame(columns=["gene","tissue","specific_p","specific_FDR_manifest","specificity"]),stage/"c5.CIGMA_results.csv")
                        write_csv(pd.DataFrame(columns=["assay","gene","tissue","specific_p","specific_FDR_manifest","specificity"]),stage/"c5.CIGMA_annotation.csv")
                        rc=1; status.append(dict(analysis="CIGMA",status="failed",detail=str(exc)))
                elif not native_failed:
                    status.append(dict(analysis="CIGMA",status="unavailable" if layer=="prot" else "not_applicable",
                        detail="Needs published CIGMA summaries or real paired-donor single-cell/genotype inputs; no inference from UKB plasma alone"))
                write_csv(pd.DataFrame(status),stage/"c5.cellulation_status.csv")
                write_csv(pd.DataFrame([dict(role=r,**stamp(p)) for r,p in zip(["atlas","universe","panels"],resources) if p is not None]),stage/"c5.cellulation_inputs.csv")
                write_json(dict(version=VERSION,notes=notes,interpretation=NOTE),stage/"c5.run_notes.json")
                newfiles=[p for p in stage.iterdir() if p.is_file()]
                names={p.name for p in newfiles}
                for oldfile in out.glob("c5.*"):
                    owned=oldfile.name.startswith(("c5.cell.","c5.CIGMA","c5.evidence_","c5.cellulation_","c5.run_notes"))
                    if oldfile.is_file() and owned and oldfile.name not in names and oldfile.suffix!=".rds": oldfile.unlink()
                for p in newfiles: os.replace(p,out/p.name)
                outputs={n:sha(out/n) for n in names}
                record=dict(signature=signature,status=status,outputs=outputs)
                # Failure is never a reusable success marker.
                if any(s["status"]=="failed" for s in status):
                    if completed.exists(): completed.unlink()
                    write_json(record,out/"c5.failed.json")
                else:
                    write_json(record,completed)
                    if (out/"c5.failed.json").exists(): (out/"c5.failed.json").unlink()
                print(json.dumps(dict(trait=trait,layer=layer,status=status),ensure_ascii=False))
    return rc


def command_line(argv=None):
    argv=list(sys.argv[1:] if argv is None else argv)
    if not argv or argv[0].startswith("-"): return main_driver(argv)
    cmd=argv.pop(0); ap=argparse.ArgumentParser(prog=f"c5.cellulation.py {cmd}")
    if cmd=="annotate":
        for k in ["universe","atlas","panels","outdir"]: ap.add_argument("--"+k,type=Path,required=True)
        ap.add_argument("--prefix",default="c5.cell"); ap.add_argument("--contrasts",type=Path)
        ap.add_argument("--matched-draws",type=int,default=0); ap.add_argument("--seed",type=int,default=int(os.getenv("SEED","2026")))
        ap.add_argument("--no-plots",action="store_true"); a=ap.parse_args(argv)
        if not re.fullmatch(r"[A-Za-z0-9_.-]+",a.prefix): ap.error("Invalid file prefix")
        if a.matched_draws and a.matched_draws<100: ap.error("Use >=100 matched draws or 0")
        annotate(a.universe,a.atlas,a.panels,a.outdir,a.prefix,a.contrasts,a.matched_draws,a.seed,not a.no_plots)
    elif cmd=="import-cellage":
        ap.add_argument("--repo",type=Path,required=True); ap.add_argument("--outdir",type=Path,default=default_cell_atlas().parent)
        a=ap.parse_args(argv); import_cellage(a.repo,a.outdir)
    elif cmd=="import-cigma":
        for k in ["table","mapping","outdir"]: ap.add_argument("--"+k,type=Path,required=True)
        a=ap.parse_args(argv); import_cigma(a.table,a.mapping,a.outdir)
    elif cmd=="inspect-table":
        ap.add_argument("--table",type=Path,required=True); ap.add_argument("--sheet"); ap.add_argument("--header-row",type=int,default=0)
        a=ap.parse_args(argv)
        if a.table.suffix.lower() in {".xlsx",".xlsm"} and a.sheet is None:
            import openpyxl
            w=openpyxl.load_workbook(a.table,read_only=True,data_only=True)
            try: print(json.dumps(w.sheetnames,ensure_ascii=False))
            finally: w.close()
        else:
            d=read_table(a.table,a.sheet,a.header_row); print(canonical_json(dict(columns=list(d),rows=len(d),preview=d.head(3).to_dict("records"))))
    elif cmd=="cigma":
        for k in ["manifest","outdir"]: ap.add_argument("--"+k,type=Path,required=True)
        ap.add_argument("--validate-only",action="store_true"); ap.add_argument("--seed",type=int,default=int(os.getenv("SEED","2026")))
        ap.add_argument("--allow-untested-cigma",action="store_true")
        a=ap.parse_args(argv); return run(a.manifest,a.outdir,a.validate_only,a.seed,a.allow_untested_cigma)
    elif cmd=="prepare-pseudobulk":
        for k in ["h5ad","genes","outdir"]: ap.add_argument("--"+k,type=Path,required=True)
        for k in ["donor-col","cell-type-col","cell-types","tissue","build","cohort"]: ap.add_argument("--"+k,required=True)
        ap.add_argument("--layer"); ap.add_argument("--normalization",choices=["log10-cp10k","pretransformed"],default="log10-cp10k")
        ap.add_argument("--expression-scale",default="log10(CP10K+1)")
        ap.add_argument("--min-cells",type=int,default=11); ap.add_argument("--min-expression-fraction",type=float,default=.1)
        ap.add_argument("--chunk-cells",type=int,default=2048)
        a=ap.parse_args(argv)
        if a.min_cells<2 or a.chunk_cells<1 or not 0<=a.min_expression_fraction<1: ap.error("Invalid pseudobulk QC settings")
        if a.normalization=="pretransformed" and a.expression_scale=="log10(CP10K+1)": ap.error("Specify the actual --expression-scale for pretransformed input")
        prepare_pseudobulk(a)
    elif cmd=="kinship":
        ap.add_argument("--dosage",type=Path,required=True); ap.add_argument("--out",type=Path,required=True)
        for k in ["gene","tissue","build","cohort","scope","variant-region"]: ap.add_argument("--"+k,required=True)
        ap.add_argument("--min-maf",type=float,default=.05); ap.add_argument("--max-missing",type=float,default=.02)
        a=ap.parse_args(argv)
        if not 0<=a.min_maf<.5 or not 0<=a.max_missing<1: ap.error("Invalid genotype QC thresholds")
        # Missing dosages are allowed here, unlike model matrices.
        d=read_table(a.dosage); ids=textcol(d.iloc[:,0])
        if ids.eq("").any() or ids.duplicated().any(): raise ValueError("Invalid dosage donor IDs")
        x=d.iloc[:,1:].replace({"":np.nan,"NA":np.nan}).apply(pd.to_numeric,errors="raise")
        K,keep,freq=kinship_from_dosage(x,a.min_maf,a.max_missing)
        z=pd.DataFrame(K,index=ids,columns=ids); z.index.name="donor"; atomic_text(z.to_csv(),a.out)
        write_json(dict(version=VERSION,source=stamp(a.dosage),gene=a.gene,tissue=a.tissue,build=a.build,cohort=a.cohort,
            scope=a.scope,variant_region=a.variant_region,variants_in=x.shape[1],variants_used=int(keep.sum()),donors=len(ids),
            formula="Z=(G-2p)/sqrt(2p(1-p)); K=Z Z'/M; genotype mean imputation only",min_maf=a.min_maf,max_missing=a.max_missing),
            Path(str(a.out)+".provenance.json"))
        write_csv(pd.DataFrame(dict(variant=np.array(x.columns)[keep],effect_allele_frequency=freq)),Path(str(a.out)+".variants.csv"))
    else: raise ValueError("Unknown command "+cmd)
    return 0


if __name__=="__main__":
    try: raise SystemExit(command_line())
    except (ValueError,FileNotFoundError,KeyError,ImportError) as exc:
        print(f"C5 ERROR: {exc}",file=sys.stderr); raise SystemExit(2)
