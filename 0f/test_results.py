"""Storage regressions: Excel interoperability and recoverable GU native data."""
import csv
import importlib.util
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
import unittest

import openpyxl


def module(name,path):
 spec=importlib.util.spec_from_file_location(name,path);value=importlib.util.module_from_spec(spec);spec.loader.exec_module(value);return value

HERE=Path(__file__).resolve().parent
io=module('result_io',HERE/'results.py')
gu=module('gu_io',HERE.parent/'gu/f/0.common.py')


class StorageTests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory(prefix='result-test-',dir='/tmp');self.root=Path(self.tmp.name)
 def tearDown(self):self.tmp.cleanup()
 def workbook(self,rows,types,name='results'):
  source=self.root/'source.tsv'
  with source.open('w',newline='') as f:csv.writer(f,delimiter='\t').writerows(rows)
  spec=self.root/'spec.json';spec.write_text(json.dumps({'tables':[{'name':name,'path':str(source),'types':types}]}))
  dest=self.root/'result.xlsx';io.stream_workbook(spec,dest);return dest
 def test_ids_precision_and_literal_text(self):
  dest=self.workbook([['eid','estimate','label'],['00001','1.234567890123456','=1+1'],['00002','','NA']],['character','numeric','character'])
  wb=openpyxl.load_workbook(dest,read_only=True)
  self.assertEqual(list(wb['results'].values)[1],('00001',1.234567890123456,'=1+1'))
  self.assertEqual(wb['results']['C2'].data_type,'s')
  self.assertEqual(list(wb['results'].values)[2],('00002',None,'NA'));wb.close()
  restored=io.materialize_workbook(dest)
  self.assertEqual(restored.read_bytes(),(self.root/'source.tsv').read_bytes())
  # Older software using the former .rds name can resolve the replacement.
  self.assertEqual(io.resolve_table(dest.with_suffix('.rds')),restored)
 def test_r_reader_interoperability(self):
  dest=self.workbook([['eid','estimate'],['00001','1.234567890123456']],['character','numeric'])
  script='x<-openxlsx::read.xlsx(commandArgs(TRUE)[1]);stopifnot(identical(x$eid,"00001"),identical(x$estimate,1.234567890123456))'
  subprocess.run(['/usr/bin/Rscript','-e',script,str(dest)],check=True,capture_output=True,text=True)
 def test_long_text_is_split_without_loss(self):
  value='AGCT'*20000
  dest=self.workbook([['id','sequence'],['00001',value]],['character','character'])
  wb=openpyxl.load_workbook(dest,read_only=True)
  parts=list(wb['_long_text'].values)[1:]
  self.assertEqual(''.join(r[4] for r in parts),value)
  self.assertIn('_long_text',list(wb['results'].values)[1][1]);wb.close()
 def test_excel_row_boundary(self):
  def rows():
   yield ['index']
   for i in range(1048576):yield [str(i)]
  dest=self.workbook(rows(),['numeric'])
  wb=openpyxl.load_workbook(dest,read_only=True)
  self.assertEqual(wb.sheetnames,['results','results_2'])
  self.assertEqual(list(wb['results_2'].values),[('index',),(1048575,)]);wb.close()
 def test_failed_export_preserves_prior_destination(self):
  source=self.root/'bad.tsv';source.write_text('a\tb\n1\n')
  spec=self.root/'bad.json';spec.write_text(json.dumps({'tables':[{'name':'data','path':str(source),'types':['numeric','numeric']}]}))
  dest=self.root/'old.xlsx';dest.write_bytes(b'prior verified output')
  with self.assertRaises(ValueError):io.stream_workbook(spec,dest)
  self.assertEqual(dest.read_bytes(),b'prior verified output')
 def test_persistent_native_archive_survives_tmp_cleanup(self):
  with tempfile.TemporaryDirectory(prefix='gu-storage-test-',dir=HERE) as directory:
   root=Path(directory);run=root/'ibdmix/test/chr1';(run/'final').mkdir(parents=True);(run/'samples/C1').mkdir(parents=True)
   (run/'run.meta.tsv').write_text('refs\tAltai\n');(run/'samples/C1/ALL.txt').write_text('00001\n')
   (run/'linked_final').symlink_to('final',target_is_directory=True)
   (run/'final/segments.tsv').write_text('sample_id\tstart\tend\n00001\t1\t20\n')
   gu.gu_publish_results(root,root,['ibdmix'])
   self.assertTrue((run/'ibdmix.tracts.xlsx').is_file())
   gu.gu_compact_results(root,['ibdmix'])
   self.assertFalse((run/'final').exists())
   self.assertTrue((run/'ibdmix.raw.tar.gz').is_file())
   view=gu.gu_native_view_root(root)
   shutil.rmtree(view);gu.gu_prepare_read_view(root)
   self.assertIn('00001',(view/'ibdmix/test/chr1/final/segments.tsv').read_text())
   self.assertTrue((view/'ibdmix/test/chr1/linked_final/segments.tsv').is_file())
   gu.gu_restore_results(root,'ibdmix',run)
   self.assertTrue((run/'final/segments.tsv').is_file())
   # Partial newer work must not be overwritten by a previously archived run.
   (run/'final/segments.tsv').write_text('newer native work')
   gu.gu_restore_results(root,'ibdmix',run)
   self.assertEqual((run/'final/segments.tsv').read_text(),'newer native work')
   shutil.rmtree(view)


if __name__=='__main__':unittest.main(verbosity=2)
